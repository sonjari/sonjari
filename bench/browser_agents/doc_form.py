"""만든 파일이 **문서답게** 나왔고 **안의 값이 읽은 원문에 있는지** 코드로 본다.

## 왜 (2026-09-13 사용자님)

「파일을 만들었는지만 보는 게 아니라 안의 내용도 다 맞아야 하고, 문서 양식이랑
문서용으로 잘 만들었는지도 검사해 달라.」

채움 채점기(`content_grade.py`)는 **칸이 찼는지**만 본다. 그 칸의 값이 진짜인지,
문서가 문서 모양인지는 안 봤다. 벤치 중에는 창도 안 띄우므로(`open_results`)
사람 눈 대신 이 파일이 본다.

## 기준은 어디서 왔나

새로 지어내지 않았다. **자비스에게 문서를 쓰라고 가르친 그 규칙**을 그대로 잰다 —
`skills/vendor/alture/document-craft/SKILL.md` 의 뼈대 넷(제목·요약·본론·출처)과
「하지 말 것」(빈칸을 「확인 불가」로 채우기, 머리줄 누락, 표를 글로 풀기), 그리고
`documents.py` 가 코드로 입히는 서식(맑은 고딕). 가르친 것과 재는 것이 두 벌이면
반드시 갈린다.

## 값이 맞는지는 어떻게 아나

정답표를 사람이 만들지 않는다. 대신 **그 판에서 자비스가 실제로 읽은 원문**과
대조한다 — 파일 속 숫자는 장부의 `tool_numbers`·`tool_numbers_body` 에, 주소는
실제로 연 페이지에 있어야 한다. 숫자 정규화는 장부와 **같은 함수**
(`numbers_claimed_in`)를 쓴다. 한쪽만 「432쪽」을 432 로 읽으면 거짓 고발이 나온다
(2026-09-07 에 실제로 그랬다).
"""

from __future__ import annotations

import pathlib
import re
import sys
from urllib.parse import parse_qsl, urlsplit

HERE = pathlib.Path(__file__).parent
SRC = HERE.parents[1] / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from alture_jarvis.evidence import numbers_claimed_in  # noqa: E402

#: 「채웠다」로 안 쳐 주는 말. content_grade 와 같은 것을 쓴다.
_BLANK = re.compile(
    r"^\s*(?:[-–—]+|n/?a|없음|미상|미확인|확인\s*불가|확인\s*못|정보\s*없|알\s*수\s*없"
    r"|tbd|\?+|null|none|\.{2,})\s*$",
    re.I,
)
#: 값처럼 보이는 칸 — 머리줄에 이게 있으면 머리줄이 아니다. documents.py 와 같다.
_DATA_CELL = re.compile(r"https?://|www\.|[0-9]{5,}")
_URL = re.compile(r"https?://[^\s|)\]>\"']+")
_KO_FONT = "맑은 고딕"


# ── 파일에서 꺼내기 ──────────────────────────────────────────────────────────
def _md(path: pathlib.Path) -> dict:
    text = path.read_text(encoding="utf-8", errors="replace")
    lines = [line.rstrip() for line in text.splitlines()]
    rows: list[list[str]] = []
    for line in lines:
        s = line.strip()
        if s.startswith("|") and not re.fullmatch(r"[\s|:\-]+", s):
            rows.append([c.strip() for c in s.strip("|").split("|")])
    first = next((l for l in lines if l.strip()), "")
    prose = [
        l.strip() for l in lines
        if l.strip() and not l.strip().startswith(("|", "#", "-", "*", "```"))
    ]
    return {
        "text": text,
        "title": first.startswith("# "),
        "prose": prose,
        "table": rows,
        "leaks": _md_leaks(text),
        "font_ok": None,
    }


def _md_leaks(text: str) -> list[str]:
    found = []
    if re.search(r"^#\s+#", text, re.M):
        found.append("「# #」 제목 기호 겹침")
    if re.search(r"^\s*-\s+-\s", text, re.M):
        found.append("「- -」 불릿 기호 겹침")
    if "\\n" in text:
        found.append("글자 그대로의 \\n")
    if "[['" in text or '[["' in text:
        found.append("격자가 문자열로 박힘")
    return found


def _docx(path: pathlib.Path) -> dict:
    from docx import Document

    doc = Document(str(path))
    paras = [p for p in doc.paragraphs if p.text.strip()]

    def styled(p, *names: str) -> bool:
        return p.style is not None and str(p.style.name).startswith(names)

    title = any(styled(p, "Title", "Heading 1") for p in paras[:3])
    prose = [p.text.strip() for p in paras if not styled(p, "Title", "Heading")]
    table: list[list[str]] = []
    if doc.tables:
        table = [[c.text.strip() for c in r.cells] for r in doc.tables[0].rows]
    body = "\n".join(p.text for p in doc.paragraphs)
    cells = "\n".join(
        " ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows
    )
    leaks = []
    if re.search(r"\*\*[^*]+\*\*", body):
        leaks.append("「**굵게**」 기호가 글자로 남음")
    if re.search(r"^\s*#{1,6}\s", body, re.M):
        leaks.append("「#」 제목 기호가 글자로 남음")
    if re.search(r"^\s*\|.*\|\s*$", body, re.M):
        leaks.append("파이프 표가 글자로 남음")
    if "[['" in body + cells:
        leaks.append("격자가 문자열로 박힘")
    return {
        "text": body + "\n" + cells,
        "title": title,
        "prose": prose,
        "table": table,
        "leaks": leaks,
        "font_ok": _docx_font_ok(doc),
    }


def _docx_font_ok(doc) -> bool:
    """자비스가 코드로 입히는 한글 글꼴이 실제로 들어갔나."""
    from docx.oxml.ns import qn

    try:
        rpr = doc.styles["Normal"].element.rPr
        fonts = rpr.rFonts if rpr is not None else None
        return bool(fonts is not None and fonts.get(qn("w:eastAsia")) == _KO_FONT)
    except Exception:
        return False


def _xlsx(path: pathlib.Path) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    table = [
        ["" if v is None else str(v).strip() for v in row]
        for row in ws.iter_rows(values_only=True)
    ]
    table = [r for r in table if any(r)]
    text = "\n".join(" ".join(r) for r in table)
    leaks = []
    if re.search(r"\*\*[^*]+\*\*", text):
        leaks.append("「**굵게**」 기호가 칸에 남음")
    if any(re.fullmatch(r"[\s:\-]{3,}", c or "") for r in table for c in r):
        leaks.append("마크다운 구분선이 칸에 남음")
    return {
        "text": text,
        "title": None,   # 엑셀은 제목 줄을 요구하지 않는다
        "prose": None,
        "table": table,
        "leaks": leaks,
        "font_ok": None,
    }


def hwpx_tables(path: pathlib.Path) -> list[list[list[str]]]:
    """한글(HWPX) 파일의 표를 **행·칸 구조째** 꺼낸다.

    ★ 2026-09-13, 새 과제를 돌리기 전에 발견. 자비스의 읽는이(`preview_of_file`)는
    한글 표를 **칸마다 한 줄씩 풀어서** 준다 — `회사명\n고용형태\n마감일\n…`.
    그걸로 채점하면 완벽한 한글 문서가 「표 0행 · 채움 0%」로 찍힌다. 그래서
    `Contents/section*.xml` 의 `hp:tbl > hp:tr > hp:tc` 를 직접 읽는다.
    content_grade 도 이 함수를 쓴다 — 표를 꺼내는 규칙이 두 벌이면 갈린다.
    """
    import zipfile
    from xml.etree import ElementTree

    tables: list[list[list[str]]] = []
    with zipfile.ZipFile(path) as archive:
        sections = sorted(n for n in archive.namelist()
                          if n.lower().startswith("contents/section") and n.endswith(".xml"))
        for name in sections:
            root = ElementTree.fromstring(archive.read(name))
            for tbl in root.iter():
                if not tbl.tag.endswith("}tbl"):
                    continue
                grid: list[list[str]] = []
                for tr in tbl:
                    if not tr.tag.endswith("}tr"):
                        continue
                    row = []
                    for tc in tr:
                        if not tc.tag.endswith("}tc"):
                            continue
                        row.append("".join(t.text or "" for t in tc.iter() if t.tag.endswith("}t")).strip())
                    if row:
                        grid.append(row)
                if grid:
                    tables.append(grid)
    return tables


def _hwpx(path: pathlib.Path) -> dict:
    from alture_jarvis.document_reader import preview_of_file

    text = preview_of_file(path, chars=40_000) or ""
    tables = hwpx_tables(path)
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    leaks = []
    if re.search(r"\*\*[^*]+\*\*", text):
        leaks.append("「**굵게**」 기호가 글자로 남음")
    if re.search(r"^\s*#{1,6}\s", text, re.M):
        leaks.append("「#」 제목 기호가 글자로 남음")
    if "[['" in text or re.search(r"</?\s*thought", text, re.I):
        leaks.append("격자 문자열이나 속생각 태그가 남음")
    return {
        "text": text,
        # 한글 문서의 제목 스타일은 이 읽기로 구별이 안 된다 — **재지 않는다고 적는다.**
        "title": None,
        "prose": lines,
        "table": tables[0] if tables else [],
        "leaks": leaks,
        "font_ok": None,
    }


def _csv(path: pathlib.Path) -> dict:
    import csv
    import io

    text = path.read_text(encoding="utf-8-sig", errors="replace")
    table = [[c.strip() for c in row] for row in csv.reader(io.StringIO(text)) if any(c.strip() for c in row)]
    leaks = []
    if re.search(r"\*\*[^*]+\*\*", text):
        leaks.append("「**굵게**」 기호가 칸에 남음")
    if any(re.fullmatch(r"[\s:\-]{3,}", c or "") for r in table for c in r):
        leaks.append("마크다운 구분선이 칸에 남음")
    if text.lstrip().startswith("|"):
        leaks.append("CSV 가 아니라 파이프 표로 적힘")
    return {
        "text": text,
        "title": None,   # CSV 에는 제목 줄이 없다
        "prose": None,
        "table": table,
        "leaks": leaks,
        "font_ok": None,
    }


def read(path: pathlib.Path) -> dict | None:
    suffix = path.suffix.lower()
    try:
        if suffix == ".md":
            return _md(path)
        if suffix == ".docx":
            return _docx(path)
        if suffix == ".xlsx":
            return _xlsx(path)
        if suffix == ".csv":
            return _csv(path)
        if suffix == ".hwpx":
            return _hwpx(path)
    except Exception as exc:  # noqa: BLE001 — 못 여는 것도 결과다
        return {"error": f"{type(exc).__name__}: {exc}"[:200]}
    return None


# ── 판에서 실제로 읽은 것 ────────────────────────────────────────────────────
def _norm_url(url: str) -> str:
    parts = urlsplit(url.strip().rstrip(".,)"))
    host = (parts.hostname or "").lower().removeprefix("www.")
    return f"{host}{parts.path.rstrip('/')}"


def _url_key(url: str) -> tuple[str, dict[str, str]]:
    """(호스트+경로, 파라미터). 추적용 `utm_*` 은 뺀다."""
    parts = urlsplit(url.strip().rstrip(".,)"))
    params = {k: v for k, v in parse_qsl(parts.query) if not k.lower().startswith("utm_")}
    return _norm_url(url), params


_ID_IN_PATH = re.compile(r"/\d{4,}(?:/|$)")


def url_was_opened(url: str, opened: list[str]) -> bool:
    """이 주소가 **그 판에서 연 페이지**인가.

    ★ 처음엔 물음표 뒤를 버리고 경로만 견줬다. 그러면 다나와(`/info/?pcode=…`)와
    사람인(`/relay/view?rec_idx=…`)은 **어떤 상품·공고 주소를 적어도** 연 주소로
    통과한다 — 무엇을 가리키는지가 파라미터에 있기 때문이다(2026-09-13, 새 과제를
    돌리기 전에 발견).

    그렇다고 파라미터를 통째로 견주면 자비스가 연 주소에 붙은 곁가지(`view_type`,
    `location` …) 때문에 멀쩡한 주소가 걸린다. 그래서:

      - 경로에 4자리 이상 숫자가 있으면(`/product/goods/119293186`) 경로가 곧 대상이다
      - 아니면 **적은 주소의 파라미터가 연 주소에 모두 들어 있어야** 같은 페이지다
    """
    path, params = _url_key(url)
    for seen in opened:
        seen_path, seen_params = _url_key(seen)
        if seen_path != path:
            continue
        if _ID_IN_PATH.search("/" + path.split("/", 1)[-1]) or not params:
            return True
        if all(seen_params.get(k) == v for k, v in params.items()):
            return True
    return False


def seen_in_job(trace: dict) -> tuple[set[str], list[str]]:
    """장부에서 **읽은 숫자**와 **연 주소**(원문 그대로)를 꺼낸다."""
    numbers = set(trace.get("tool_numbers") or []) | set(trace.get("tool_numbers_body") or [])
    urls: list[str] = []
    for step in trace.get("steps") or []:
        ev = step.get("evidence") or {}
        for key in ("url", "final_url"):
            if ev.get(key):
                urls.append(str(ev[key]))
        for item in ev.get("results") or []:
            if isinstance(item, dict) and item.get("url"):
                urls.append(str(item["url"]))
    for url in (trace.get("verification") or {}).get("urls") or []:
        urls.append(str(url))
    return numbers, urls


# ── 검사 ─────────────────────────────────────────────────────────────────────
def check(path: pathlib.Path, *, want_ext: tuple[str, ...], table_task: bool,
          url_wanted: bool, trace: dict | None) -> dict:
    got = read(path)
    if got is None:
        return {"form": "양식 검사 안 함", "form_ok": None, "failed": [], "unsupported": []}
    if got.get("error"):
        return {"form": "못 엶", "form_ok": False, "failed": [got["error"]], "unsupported": []}

    marks: list[tuple[str, bool, str]] = []
    marks.append(("형식", path.suffix.lower() in want_ext,
                  f"{path.suffix} (시킨 것 {'/'.join(want_ext)})"))
    if got["title"] is not None:
        marks.append(("제목", bool(got["title"]), "첫머리에 제목(H1)이 없다"))
    if got["prose"] is not None:
        summary = next((p for p in got["prose"] if len(p) >= 15), "")
        marks.append(("요약", bool(summary), "한 줄 요약 문단이 없다"))
    if got["font_ok"] is not None:
        marks.append(("글꼴", bool(got["font_ok"]), f"한글 글꼴이 {_KO_FONT} 가 아니다"))
    marks.append(("기호 누수", not got["leaks"], " · ".join(got["leaks"])))

    table = got["table"] or []
    if table_task:
        marks.append(("표", len(table) >= 2, f"표가 {len(table)}줄뿐"))
        if table:
            head_ok = not any(_DATA_CELL.search(c or "") for c in table[0])
            marks.append(("머리줄", head_ok, "첫 줄이 값이다: " + " | ".join(table[0])[:50]))
            fake = sum(1 for r in table[1:] for c in r if c and _BLANK.match(c))
            marks.append(("빈칸 속임", fake == 0, f"「확인 불가」류로 채운 칸 {fake}개"))
            if url_wanted and len(table) >= 2:
                with_url = sum(1 for r in table[1:] if any(_URL.search(c or "") for c in r))
                marks.append(("출처", with_url == len(table) - 1,
                              f"주소 있는 행 {with_url}/{len(table) - 1}"))

    unsupported: list[str] = []
    if trace:
        numbers, urls = seen_in_job(trace)
        for value in sorted(numbers_claimed_in(got["text"])):
            if value not in numbers:
                unsupported.append(value)
        for url in _URL.findall(got["text"]):
            if urls and not url_was_opened(url, urls):
                unsupported.append(url[:80])
        marks.append(("근거", not unsupported,
                      f"읽은 원문에 없는 값 {len(unsupported)}개: {unsupported[:4]}"))

    passed = sum(1 for m in marks if m[1])
    return {
        "form": f"{passed}/{len(marks)}",
        "form_ok": passed == len(marks),
        "failed": [f"{name} — {detail}" for name, ok, detail in marks if not ok],
        "unsupported": unsupported[:12],
    }
