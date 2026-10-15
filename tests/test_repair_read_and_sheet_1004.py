"""검사가 「먼저 열어서 읽어라」고 한 주소는 열 수 있어야 한다 · 엑셀 고치기는 엉뚱한 시트에 조용히 쓰지 않는다 (2026-10-04).

레딧 조사 두 판(30f003f6 · aa741c81)이 「…확인하겠습니다」 예고로 끝났다. 걸음 기록:

    26  local_write_document          파일을 씀 → 대조: 「열지 않은 주소가 줄에 붙어 있다 — 빼거나, 먼저 열어서 읽은 뒤에 적어라」
    27  browser_open  ✗ not_this_stage   그 주소를 열려 함 → 차례 문이 거절
    28  browser_open  ✗ not_this_stage
    29  local_edit_spreadsheet        「확인 불가」로 고치려 함 → 시트 이름이 없어 **첫 시트(공통점)** 에 들어감

검사는 열라 하고 문은 열지 말라 했다. 그리고 고치는 손은 틀린 자리에 조용히 썼다.
"""

from __future__ import annotations

import inspect
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from alture_jarvis.agent import BrowserAgent, _REPAIR_READ_OPENS, _repair_read_after_write_feedback
from alture_jarvis.documents import edit_spreadsheet

A = "https://www.reddit.com/r/LocalLLaMA/comments/aaa/context_rot"
B = "https://www.reddit.com/r/LocalLLM/comments/bbb/strata"
C = "https://www.reddit.com/r/LocalLLM/comments/ccc/other"


def _open(url: str, ok: bool = True) -> dict:
    return {"ok": ok, "tool": "browser_open", "arguments": {"url": url}}


WRITE = {"ok": True, "tool": "local_write_document", "arguments": {"path": "x.xlsx"}}
WRITE_FED = {**WRITE, "write_grounding_feedback": True}
EXTRACT = {"ok": True, "tool": "browser_extract", "arguments": {}}


# ── 산출물 차례의 「그 주소 열기」 ─────────────────────────────────────────────
def test_before_any_feedback_nothing_is_let_through() -> None:
    """대조가 되돌리기 전에는 예전 그대로다 — 산출물 차례에 더 찾으러 나가는 것은 막는다(2026-10-02 밤 판)."""
    steps = [_open(A), EXTRACT, WRITE]
    assert _repair_read_after_write_feedback("browser_open", {"url": B}, steps) is False


def test_after_feedback_a_never_opened_address_may_be_opened_and_read_once() -> None:
    steps = [_open(A), EXTRACT, WRITE_FED]
    assert _repair_read_after_write_feedback("browser_open", {"url": B}, steps) is True
    steps.append(_open(B))
    assert _repair_read_after_write_feedback("browser_extract", {}, steps) is True


def test_reopening_searching_and_rereading_stay_blocked() -> None:
    steps = [_open(A), EXTRACT, WRITE_FED]
    assert _repair_read_after_write_feedback("browser_open", {"url": A}, steps) is False, "이미 연 주소 되열기"
    assert _repair_read_after_write_feedback("browser_search", {"query": "more"}, steps) is False, "새 검색"
    assert _repair_read_after_write_feedback("browser_extract", {}, steps) is False, "연 직후가 아닌 읽기"
    assert _repair_read_after_write_feedback("browser_site_search", {"query": "Strata"}, steps) is False


def test_the_number_of_new_addresses_is_capped() -> None:
    steps = [_open(A), EXTRACT, WRITE_FED]
    for index in range(_REPAIR_READ_OPENS):
        url = f"https://www.reddit.com/r/x/comments/new{index}"
        assert _repair_read_after_write_feedback("browser_open", {"url": url}, steps) is True
        steps += [_open(url), EXTRACT]
    assert _repair_read_after_write_feedback("browser_open", {"url": C}, steps) is False


def test_a_refused_open_does_not_use_up_the_cap() -> None:
    steps = [_open(A), EXTRACT, WRITE_FED, _open(B, ok=False), _open(B, ok=False)]
    assert _repair_read_after_write_feedback("browser_open", {"url": B}, steps) is True


def test_the_stage_gate_uses_it_and_gives_the_steps_back() -> None:
    source = inspect.getsource(BrowserAgent)
    assert "or repair_read" in source
    assert "_repair_read_after_write_feedback(name, arguments, steps)" in source
    assert "repair_read_grace" in source and "step_limit += extra" in source
    # 대조 문구가 「열어서 읽어라」고 말하는 한, 문도 그 길을 열어 둬야 한다.
    assert "먼저 열어서 읽은 뒤에 적어라" in source


# ── 엑셀 고치기: 시트를 조용히 바꿔치지 않는다 ─────────────────────────────────
def _book(path: Path, names: list[str]) -> Path:
    book = Workbook()
    book.active.title = names[0]
    book.active["A1"] = names[0]
    for name in names[1:]:
        book.create_sheet(name)["A1"] = name
    book.save(path)
    return path


def test_a_multi_sheet_file_needs_the_sheet_name(tmp_path: Path) -> None:
    target = _book(tmp_path / "t.xlsx", ["공통점", "글 목록"])
    with pytest.raises(ValueError) as caught:
        edit_spreadsheet(target, [{"cell": "B7", "value": "확인 불가"}])
    assert "sheet_required" in str(caught.value) and "공통점" in str(caught.value) and "글 목록" in str(caught.value)
    assert load_workbook(target)["공통점"]["B7"].value is None, "거절했으면 아무 시트도 고치지 않는다"


def test_a_wrong_sheet_name_is_refused_with_the_real_names(tmp_path: Path) -> None:
    target = _book(tmp_path / "t.xlsx", ["공통점", "글 목록"])
    with pytest.raises(ValueError) as caught:
        edit_spreadsheet(target, [{"cell": "B7", "value": "x"}], "후기 목록")
    assert "sheet_not_found" in str(caught.value) and "글 목록" in str(caught.value)


def test_the_named_sheet_is_the_one_edited(tmp_path: Path) -> None:
    target = _book(tmp_path / "t.xlsx", ["공통점", "글 목록"])
    result = edit_spreadsheet(target, [{"cell": "B7", "value": "확인 불가"}], "글 목록")
    book = load_workbook(target)
    assert result["sheet"] == "글 목록" and book["글 목록"]["B7"].value == "확인 불가" and book["공통점"]["B7"].value is None


def test_case_and_spaces_around_the_name_are_forgiven(tmp_path: Path) -> None:
    target = _book(tmp_path / "t.xlsx", ["Summary", "Posts"])
    assert edit_spreadsheet(target, {"A2": "ok"}, " posts ")["sheet"] == "Posts"


def test_a_single_sheet_file_still_needs_no_name(tmp_path: Path) -> None:
    target = _book(tmp_path / "t.xlsx", ["Sheet1"])
    assert edit_spreadsheet(target, {"A2": "ok"})["changed"] == 1


# ── 대조표 오독 둘 (같은 판 aa741c81) ──────────────────────────────────────────
def test_sheets_in_the_output_file_are_not_sources_to_read() -> None:
    """「엑셀 파일 하나에 시트 두 개로」의 2가 「로컬 자료 2개 이상 직접 열어 읽기」가 됐고 「글 5개 이상」은 빠졌다."""
    from alture_jarvis import plan

    prompt = ("레딧(reddit.com)의 r/LocalLLaMA 에서 그래픽카드 메모리 16GB 이하에서 로컬 모델로 에이전트를 돌릴 때 컨텍스트가 모자라는 문제를 "
              "사람들이 어떻게 풀었는지 찾아 읽어 줘. 서로 다른 글을 5개 이상 열어 댓글까지 읽고, 엑셀(.xlsx) 파일 하나에 시트 두 개로 정리해 줘.")
    assert all(scope != "local" for scope, _count in plan._scoped_count_targets(prompt)), "시트 수는 읽을 로컬 자료 수가 아니다"
    counts = [req for req in plan.extract(prompt) if req.kind == "count"]
    assert [req.target for req in counts] == [5], [(req.kind, req.label, req.target) for req in plan.extract(prompt)]
    assert "로컬" not in counts[0].label and counts[0].source_scope != "local"


def test_a_short_latin_site_name_is_not_found_inside_another_word() -> None:
    from alture_jarvis import plan

    known = {"GB": ("gb.example",), "레딧": ("reddit.com",)}
    sites = [req.label for req in plan.extract("레딧에서 16GB 이하 카드 후기 글을 3개 읽어 줘", known_sites=known) if req.kind == "site"]
    assert sites == ["레딧에서 하기"], sites
    # 낱말로 나오면 예전처럼 건다 — 한글 조사가 붙어도, 마침표가 붙어도.
    assert plan._ascii_name_as_word("yes24에서 책을 찾아줘", "yes24")
    assert plan._ascii_name_as_word("go to gb.example and read", "gb")
    assert not plan._ascii_name_as_word("16gb 이하에서", "gb")
