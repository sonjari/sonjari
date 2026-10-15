"""`type` 없이 `rows` 만 보낸 블록도 **표**다.

★ 2026-09-28 헤르메스 대결 T2 온텔릭 판(ef11749c). 모델이 회의록 표 다섯을 `{"rows": [...]}` 로 보냈고, 종류 기본값이 문단이라
표가 전부 말없이 사라졌다 — 도구는 `verified: true`, 답은 「결정 내용·담당자·기한 표」를 설명했다. 같은 시각 헤르메스 판은 표가 있었다.
"""

from __future__ import annotations

from pathlib import Path

import docx

from alture_jarvis.documents import write_document

BLOCKS = [
    {"level": 1, "text": "2026년 9월 28일 회의록", "type": "heading"},
    {"level": 2, "text": "결정 사항", "type": "heading"},
    {"rows": [["순서", "결정 사항", "담당자", "기한", "비고"], ["1", "", "", "", ""]]},
    {"level": 2, "text": "다음 회의"},
]


def test_rows_without_a_type_become_a_table(tmp_path: Path) -> None:
    target = tmp_path / "회의록.docx"

    write_document(target, "docx", {"blocks": BLOCKS})

    document = docx.Document(str(target))
    assert len(document.tables) == 1, "표가 말없이 사라졌다"
    assert [cell.text for cell in document.tables[0].rows[0].cells] == ["순서", "결정 사항", "담당자", "기한", "비고"]
    assert any(p.text == "다음 회의" and p.style.name.startswith("Heading") for p in document.paragraphs)


def test_sheets_in_a_word_file_become_titled_tables(tmp_path: Path) -> None:
    """대결 T2 O3: 워드에 `sheets` 로 표 여럿을 보냈는데 아무것도 안 들어간 워드가 `verified: true` 로 나갔다."""
    target = tmp_path / "회의록.docx"

    write_document(target, "docx", {"sheets": {
        "기본 정보": [["항목", "내용"], ["날짜", "2026-09-28"]],
        "결정 사항": [["번호", "결정 사항", "담당자", "기한"], ["1", "", "", ""]],
    }})

    document = docx.Document(str(target))
    assert len(document.tables) == 2
    assert [p.text for p in document.paragraphs if p.style.name.startswith("Heading")] == ["기본 정보", "결정 사항"]


def test_a_table_whose_header_row_was_closed_early_is_repaired(tmp_path: Path) -> None:
    """T2 재측정 B1: `{"rows": [["순서","결정 사항","담당자","기한"]], ["1","","",""]]}` — 머리 줄을 일찍 닫은 표를 통째로 버렸다."""
    broken = (
        '[{"type": "heading", "level": 2, "text": "결정 사항"}, '
        '{"rows": [["순서", "결정 사항", "담당자", "기한"]], ["1", "", "", ""], ["2", "", "", ""]]}]'
    )
    target = tmp_path / "회의록.docx"

    shape = write_document(target, "docx", {"blocks": broken})

    document = docx.Document(str(target))
    assert len(document.tables) == 1
    assert [cell.text for cell in document.tables[0].rows[0].cells] == ["순서", "결정 사항", "담당자", "기한"]
    assert "unreadable_blocks" not in shape


def test_a_block_that_cannot_be_read_is_reported(tmp_path: Path) -> None:
    broken = '[{"type": "heading", "level": 1, "text": "제목"}, {"rows": [["a" "b"]]}]'

    shape = write_document(tmp_path / "x.docx", "docx", {"blocks": broken})

    assert "1개" in shape["unreadable_blocks"]
