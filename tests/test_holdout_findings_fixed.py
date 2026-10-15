"""손대지 않은 과제 5개(2026-09-14)가 드러낸 넷을 막는다.

    ① 「YES24에서 …」 판에 사이트 요구가 한 번도 안 뽑혔다 — YES24 가 아는 사이트 목록에 없었다
    ② 두 사이트 비교에서 모델이 브라우저 대신 인터넷 코드 실행으로 긁으려 했다 → 승인창에서 900초
    ③ 문서 칸마다 「확인 불가」 — 차 보이지만 비어 있는 표
    ④ 「6개 읽음 · 답에 2개」인데 되받는 말이 「본문을 열어라」였다 — 이미 여섯 곳을 열었다

**이 다섯 과제로 다시 재서 점수를 올리지 않는다**(715d9bf5 에 결과 보기 전에 박아 둔 규칙).
고친 것은 이 시험과, 다른 새 문장으로 확인한다.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document

from alture_jarvis.agent import (
    BrowserAgent,
    _FILL_FROM_OTHER_CANDIDATES_HINT,
    _OPEN_THE_BODY_HINT,
    _next_plan_step_hint,
    _read_enough_but_listed_short,
)
from alture_jarvis.config import JarvisConfig
from alture_jarvis.documents import write_document
from alture_jarvis.events import EventBus
from alture_jarvis.jobs import JobStore
from alture_jarvis.models import ToolResult
from alture_jarvis.plan import Requirement, extract as extract_plan

from test_agent import FakeBrowser


# ── ① YES24 ──────────────────────────────────────────────────────────────────
def _sites(prompt: str) -> list[tuple]:
    plan = extract_plan(prompt, has_named_source=True, known_sites={"교보문고": ("kyobobook.co.kr",)})
    return [req.target for req in plan if req.kind == "site"]


def test_both_bookstores_in_a_comparison_become_site_requirements() -> None:
    sites = _sites("교보문고와 YES24에서 '혼자 공부하는 파이썬' 책의 판매가를 각각 확인해서 비교해줘")
    assert ("kyobobook.co.kr",) in sites and ("yes24.com",) in sites


def test_latin_site_names_match_whatever_the_case() -> None:
    for spelled in ("YES24", "Yes24", "yes24", "예스24"):
        assert ("yes24.com",) in _sites(f"{spelled}에서 파이썬 책 3권의 쪽수를 조사해줘"), spelled


def test_one_site_is_not_counted_twice_under_two_spellings() -> None:
    sites = _sites("YES24에서 파이썬 책 3권 조사해줘")
    assert sites.count(("yes24.com",)) == 1


# ── ② 웹 조사 판의 인터넷 코드 실행 ──────────────────────────────────────────
class _Tools:
    names = {"sandbox_run", "browser_read"}
    ollama_tools = [
        {"type": "function", "function": {"name": name, "description": name,
                                          "parameters": {"type": "object", "properties": {}}}}
        for name in ("sandbox_run", "browser_read")
    ]


class _Gateway:
    def __init__(self) -> None:
        self.executed: list[str] = []

    def execute(self, name, arguments, *, job_id=""):
        self.executed.append(name)
        return ToolResult(True, name, result={"action": "ran"})


class _ScrapesThenAnswers:
    def __init__(self, *, network: bool) -> None:
        self.network = network
        self.calls = 0
        self.seen_instruction = ""

    def chat_messages(self, messages, **_kwargs) -> dict[str, Any]:
        self.calls += 1
        if self.calls == 1:
            return {"role": "assistant", "content": "", "tool_calls": [{"function": {
                "name": "sandbox_run",
                "arguments": {"code": "print(1)", "allow_network": self.network},
            }}]}
        self.seen_instruction = str(messages[-1].get("content") or "")
        return {"role": "assistant", "content": "확인한 내용을 정리했습니다."}


def _run(tmp_path: Path, *, network: bool):
    config = JarvisConfig(runtime_root=tmp_path, data_dir=tmp_path / "data",
                          output_dir=tmp_path / "out", browser_enabled=True, browser_task_timeout_sec=60)
    store = JobStore(tmp_path / "jarvis.db")
    gateway, model = _Gateway(), _ScrapesThenAnswers(network=network)
    agent = BrowserAgent(config, model, gateway, FakeBrowser(), _Tools(), store, EventBus())
    prompt = "사람인에서 여수 회계 채용 공고 2개를 열어 읽고 정리해줘"
    job = store.create_job("chat", "agent", {"prompt": prompt})
    store.claim_next()
    result = agent.run(prompt, job_id=job["id"], origin="chat")
    return gateway, model, result


def test_scraping_the_web_with_code_is_sent_back_to_the_browser(tmp_path: Path) -> None:
    gateway, model, result = _run(tmp_path, network=True)
    assert "sandbox_run" not in gateway.executed          # 승인창까지 가지 않는다
    first = result.steps[0]
    assert first["ok"] is False
    assert "web_pages_are_read_with_the_browser" in str(first.get("error") or first)
    assert "browser_site_search" in model.seen_instruction or "browser" in model.seen_instruction


def test_code_without_internet_still_runs_in_a_web_job(tmp_path: Path) -> None:
    """표 계산·정리는 막지 않는다 — 막기만 하면 제대로 된 길까지 막힌다."""
    gateway, _model, _result = _run(tmp_path, network=False)
    assert "sandbox_run" in gateway.executed


# ── ③ 「확인 불가」 칸 ──────────────────────────────────────────────────────────
ROWS = [
    ["회사명", "고용형태", "마감일", "경력"],
    ["가나상사", "정규직", "확인 불가", "경력 무관"],
    ["다라물산", "(확인불가)", "2026-10-05", "없음"],
]


def test_unverified_cells_are_left_blank_and_said_once_in_word(tmp_path: Path) -> None:
    result = write_document(tmp_path / "t.docx", "docx", {"blocks": [
        {"type": "heading", "level": 1, "text": "공고"},
        {"type": "table", "rows": ROWS},
    ]})
    assert result["blanked_unverified_cells"] == 2
    doc = Document(str(tmp_path / "t.docx"))
    cells = [[c.text for c in row.cells] for row in doc.tables[0].rows]
    assert cells[1] == ["가나상사", "정규직", "", "경력 무관"]
    assert cells[2] == ["다라물산", "", "2026-10-05", "없음"]      # 「없음」은 진짜 값일 수 있다
    notes = [p.text for p in doc.paragraphs if "비워 두었습니다" in p.text]
    assert notes == ["빈 칸 2개는 페이지에서 값을 확인하지 못해 비워 두었습니다."]


def test_header_and_ordinary_values_are_never_touched(tmp_path: Path) -> None:
    rows = [["확인 불가 사유", "비고"], ["지원자 없음", "N/A 표기 금지"]]
    result = write_document(tmp_path / "t.md", "md", {"blocks": [{"type": "table", "rows": rows}]})
    assert "blanked_unverified_cells" not in result
    text = (tmp_path / "t.md").read_text(encoding="utf-8")
    assert "확인 불가 사유" in text and "지원자 없음" in text and "N/A 표기 금지" in text


def test_spreadsheet_cells_are_blanked_without_adding_a_note_row(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    result = write_document(tmp_path / "t.xlsx", "xlsx", {"rows": ROWS})
    assert result["blanked_unverified_cells"] == 2
    sheet = load_workbook(tmp_path / "t.xlsx").worksheets[0]
    values = [[c.value for c in row] for row in sheet.iter_rows()]
    assert len(values) == 3                              # 표 밑에 메모 줄을 끼우지 않는다
    assert values[1][2] in (None, "")


# ── ④ 읽기는 충분한데 답에 모자람 ────────────────────────────────────────────
def _count(saw: str, target: int = 3):
    return [(Requirement("count", f"{target}개 이상 직접 열어 읽기", target), False, saw)]


def test_enough_pages_read_but_too_few_items_gets_the_other_candidates_hint() -> None:
    results = _count("6개 읽음 · 답에 2개")
    assert _read_enough_but_listed_short(results)
    assert _next_plan_step_hint(results) == _FILL_FROM_OTHER_CANDIDATES_HINT


def test_not_enough_pages_read_still_gets_the_open_the_body_hint() -> None:
    results = _count("2개 읽음 · 답에 2개")
    assert not _read_enough_but_listed_short(results)
    assert _next_plan_step_hint(results) == _OPEN_THE_BODY_HINT


# ── ⑤ 따옴표 안의 책 제목은 코딩 요청이 아니다 ─────────────────────────────────
def test_a_quoted_book_title_is_not_a_coding_request() -> None:
    """★ 새 과제 검사의 두 사이트 비교가 코드로 긁으러 간 **진짜 출발점**이었다.

    「교보문고와 YES24에서 '혼자 공부하는 파이썬' 책의 …」·「'클린 코드' 책의 …」 가 코딩
    과제로 분류됐다. 따옴표 검사는 대상 낱말 **바로 앞**만 봤고(앞은 「혼자 공부하는 」),
    소재 검사는 **바로 뒤**가 「책」이어야 했다(뒤에 닫는 따옴표가 낌). 과거 프롬프트 905개에
    대 보니 판정이 바뀐 4개가 전부 이런 책 비교였고, 진짜 코딩 요청은 하나도 안 바뀌었다.
    """
    from alture_jarvis.plan import wants_code

    for prompt in (
        "교보문고와 YES24에서 '혼자 공부하는 파이썬' 책의 판매가와 쪽수를 비교하는 마크다운 문서로 만들어줘",
        "알라딘과 YES24에서 '클린 코드' 책의 판매가와 쪽수를 비교하는 워드 문서로 만들어줘",
        "「혼자 공부하는 파이썬」 책 가격 비교해서 표로 만들어줘",
    ):
        assert wants_code(prompt) is False, prompt


def test_real_coding_requests_with_quotes_stay_coding() -> None:
    from alture_jarvis.plan import wants_code

    for prompt in (
        "'정렬' 함수 코드를 짜줘",
        "파이썬으로 '가계부' 프로그램 만들어줘",
        '"test.py" 파일의 코드를 고쳐줘',
    ):
        assert wants_code(prompt) is True, prompt
