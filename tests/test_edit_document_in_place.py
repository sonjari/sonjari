"""받은 문서를 **제자리에서** 고친다 — 통째로 다시 쓰지 않는다.

★ 2026-09-01. 엑셀만 셀 단위 수정이 있었고(`edit_spreadsheet`), 워드·한글은
**새로 만들기밖에** 없었다. 그래서 「받은 파일에서 이 표현만 바꿔줘」가 통째로
다시 쓰기가 됐고 서식·표·머리말·글꼴이 전부 날아갔다. 사람이 쓰던 파일에 몇 줄
고치는 것이 실제 사무의 모양이다.

가짜 바이트로는 「서식이 살았는가」를 못 잰다. 여기서는 **진짜 문서를 만들어
고치고 다시 열어** 본다.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from alture_jarvis.documents import edit_document


def _resume(path: Path) -> None:
    """굵은 글씨·표·머리말이 든 진짜 워드 문서."""
    from docx import Document

    document = Document()
    document.add_heading("나래 이력서", 0)
    paragraph = document.add_paragraph("")
    paragraph.add_run("가람대학교").bold = True
    paragraph.add_run(" 산학협력단 — 지출 서류 검토 담당")
    document.add_paragraph("42억 원 규모 세부과제 관리")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "기간"
    table.cell(0, 1).text = "기관"
    document.sections[0].header.paragraphs[0].text = "이력서 · 최나래"
    document.save(str(path))


def _read(path: Path):
    from docx import Document

    return Document(str(path))


def test_changing_a_phrase_keeps_the_table_and_header(tmp_path: Path) -> None:
    target = tmp_path / "이력서.docx"
    _resume(target)

    result = edit_document(target, [{"find": "지출 서류 검토", "replace": "증빙 적정성 검토"}])

    assert result["changed"] == 1
    document = _read(target)
    body = "\n".join(p.text for p in document.paragraphs)
    assert "증빙 적정성 검토" in body
    assert "지출 서류 검토" not in body
    # 통째로 다시 썼으면 이 셋이 전부 사라진다.
    assert len(document.tables) == 1
    assert document.tables[0].cell(0, 1).text == "기관"
    assert document.sections[0].header.paragraphs[0].text == "이력서 · 최나래"


def test_a_phrase_inside_one_run_keeps_its_neighbours_bold(tmp_path: Path) -> None:
    """부분서식이 사는 길을 **먼저** 시도한다 — 그게 제자리 수정의 이유다."""
    target = tmp_path / "이력서.docx"
    _resume(target)

    edit_document(target, [{"find": "산학협력단", "replace": "산학협력단(연구비팀)"}])

    document = _read(target)
    paragraph = next(p for p in document.paragraphs if "산학협력단" in p.text)
    bold_runs = [run.text for run in paragraph.runs if run.bold]
    assert "가람대학교" in bold_runs, "옆 글자의 굵기가 죽었다"
    assert "산학협력단(연구비팀)" in paragraph.text


def test_it_says_so_when_it_could_not_find_the_phrase(tmp_path: Path) -> None:
    """못 찾았는데 성공으로 돌려주면 「고쳤다」가 거짓이 된다."""
    target = tmp_path / "이력서.docx"
    _resume(target)

    result = edit_document(target, [{"find": "없는 표현", "replace": "새 글"}])

    assert result["changed"] == 0
    assert any(item.get("hits") == 0 for item in result["edits"])


def test_replacing_a_paragraph_reports_lost_run_formatting(tmp_path: Path) -> None:
    """서식 손실은 **침묵으로 오면 안 된다.**"""
    target = tmp_path / "이력서.docx"
    _resume(target)
    document = _read(target)
    index = next(i for i, p in enumerate(document.paragraphs) if "산학협력단" in p.text)

    result = edit_document(target, [{"paragraph": index, "text": "가람대 산학협력단 연구비 관리"}])

    assert result["changed"] == 1
    assert result["formatting_lost_paragraphs"] == 1
    assert "부분서식" in result["note"]
    assert _read(target).paragraphs[index].text == "가람대 산학협력단 연구비 관리"


def test_inserting_lands_right_after_the_anchor_not_at_the_end(tmp_path: Path) -> None:
    target = tmp_path / "이력서.docx"
    _resume(target)

    edit_document(target, [{"after": 0, "insert": "핵심 역량 요약"}])

    texts = [p.text for p in _read(target).paragraphs]
    assert texts[1] == "핵심 역량 요약", f"엉뚱한 자리에 들어갔다: {texts[:4]}"


def test_deleting_removes_only_that_paragraph(tmp_path: Path) -> None:
    target = tmp_path / "이력서.docx"
    _resume(target)
    before = [p.text for p in _read(target).paragraphs]

    edit_document(target, [{"delete": len(before) - 1}])

    after = [p.text for p in _read(target).paragraphs]
    assert len(after) == len(before) - 1
    assert after == before[:-1]


def test_the_result_carries_a_preview_of_what_it_now_says(tmp_path: Path) -> None:
    """고친 뒤에도 **열지 않고** 확인이 돼야 한다."""
    target = tmp_path / "이력서.docx"
    _resume(target)

    result = edit_document(target, [{"find": "지출 서류 검토", "replace": "증빙 적정성 검토"}])

    assert "증빙 적정성 검토" in result["preview"]


def test_hwpx_is_edited_in_place_too(tmp_path: Path) -> None:
    """국내 문서가 이 기능의 본론이다. 한글이 안 깔려 있어도 된다."""
    from alture_jarvis.documents import write_document

    target = tmp_path / "제안서.hwpx"
    write_document(
        target,
        "hwpx",
        {"blocks": [{"type": "paragraph", "text": "지출 서류 검토 업무"}]},
    )

    result = edit_document(target, [{"find": "지출 서류 검토", "replace": "증빙 적정성 검토"}])

    assert result["format"] == "hwpx"
    assert result["changed"] == 1
    assert "증빙 적정성 검토" in result["preview"]


def test_an_unsupported_format_is_refused_not_silently_rewritten(tmp_path: Path) -> None:
    target = tmp_path / "표.xlsx"
    target.write_bytes(b"x")

    with pytest.raises(ValueError, match="unsupported_edit_format"):
        edit_document(target, [{"find": "a", "replace": "b"}])
