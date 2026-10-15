"""발표자료도 **제자리에서** 고친다.

★ 2026-09-01 형식별 실측에서 이 칸만 비어 있었다. docx·hwpx 는 제자리 수정이
되는데 pptx 는 새로 만드는 수밖에 없어서, 「이 PPT 3번 슬라이드만 고쳐줘」가
통째로 다시 쓰기가 됐다 — 디자인·레이아웃·다른 장이 전부 날아간다.

python-pptx 에는 장을 지우거나 옮기는 공개 API 가 없다. `sldIdLst` 를 직접
다루므로, **정말 그렇게 되는지** 여기서 재 본다.
"""

from __future__ import annotations

import warnings
from pathlib import Path

import pytest

from alture_jarvis.document_reader import preview_of_file
from alture_jarvis.documents import EDITABLE, edit_document, write_document

warnings.filterwarnings("ignore")

DECK = [
    {"title": "현황", "bullets": ["42억 지출 서류 검토", "자금 관리"]},
    {"title": "제안", "bullets": ["표현 변경"]},
    {"title": "맺음", "bullets": ["끝"]},
]


def _deck(tmp_path: Path) -> Path:
    target = tmp_path / "발표.pptx"
    write_document(target, "pptx", {"blocks": DECK})
    return target


def _titles(path: Path) -> list[str]:
    from pptx import Presentation

    return [
        slide.shapes.title.text if slide.shapes.title is not None else ""
        for slide in Presentation(str(path)).slides
    ]


def test_pptx_is_in_the_editable_list() -> None:
    assert "pptx" in EDITABLE


def test_changing_a_phrase_keeps_every_other_slide(tmp_path: Path) -> None:
    target = _deck(tmp_path)

    result = edit_document(target, [{"find": "지출 서류 검토", "replace": "증빙 적정성 검토"}])

    assert result["changed"] == 1
    assert result["slides"] == 3, "통째로 다시 썼다"
    assert _titles(target) == ["현황", "제안", "맺음"]
    text = preview_of_file(target, chars=400)
    assert "증빙 적정성 검토" in text and "지출 서류 검토" not in text


def test_a_single_slide_can_be_rewritten(tmp_path: Path) -> None:
    target = _deck(tmp_path)

    result = edit_document(target, [{"slide": 1, "title": "수정 제안", "bullets": ["전문성 강조", "성과 수치화"]}])

    assert result["changed"] == 2
    assert _titles(target) == ["현황", "수정 제안", "맺음"]
    text = preview_of_file(target, chars=400)
    assert "성과 수치화" in text
    # 다른 장은 그대로다.
    assert "42억" in text and "끝" in text


def test_an_inserted_slide_lands_right_after_the_anchor(tmp_path: Path) -> None:
    """맨 뒤에 만들어 제자리로 옮긴다 — 공개 API 가 없어 목록을 직접 다룬다."""
    target = _deck(tmp_path)

    result = edit_document(target, [{"after": 0, "insert": {"title": "요약", "bullets": ["핵심 3줄"]}}])

    assert result["slides"] == 4
    assert _titles(target) == ["현황", "요약", "제안", "맺음"], "엉뚱한 자리에 들어갔다"
    assert "핵심 3줄" in preview_of_file(target, chars=400)


def test_deleting_removes_only_that_slide(tmp_path: Path) -> None:
    target = _deck(tmp_path)

    result = edit_document(target, [{"delete": 1}])

    assert result["slides"] == 2
    assert _titles(target) == ["현황", "맺음"]


def test_it_says_so_when_it_could_not_find_the_phrase(tmp_path: Path) -> None:
    """못 찾았는데 성공으로 돌려주면 「고쳤다」가 거짓이 된다."""
    target = _deck(tmp_path)

    result = edit_document(target, [{"find": "없는 말", "replace": "새 말"}])

    assert result["changed"] == 0
    assert any(item.get("hits") == 0 for item in result["edits"])


def test_a_slide_out_of_range_is_refused(tmp_path: Path) -> None:
    target = _deck(tmp_path)

    with pytest.raises(ValueError, match="slide_out_of_range"):
        edit_document(target, [{"slide": 99, "title": "x"}])


def test_the_result_carries_a_preview_of_what_it_now_shows(tmp_path: Path) -> None:
    target = _deck(tmp_path)

    result = edit_document(target, [{"slide": 2, "bullets": ["감사합니다"]}])

    assert "감사합니다" in result["preview"]
