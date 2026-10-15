"""헤르메스 대결 T1 O2~O4(2026-09-28 저녁)에서 잡은 셋 — 줄 값이 그 페이지 주 항목과 다름 · 답에 있는 값이 파일에 없음 · 머리줄이 세로로.

O3: 「Do it! 점프 투 파이썬 27,000원 · 2,700P」 — 그 책 상세를 열고도 27,000 을 페이지 1,350번째 글자의 추천 띠에서 가져왔다(책 값 19,800원은 260번째).
    표에 주소 칸이 없어 줄↔페이지 검사가 못 봤고, 경고 0.
O4: 답은 책 3권 표 + 「엑셀도 동일하게 생성」. 파일에는 머리줄 5칸이 **세로로**만 있고 값은 0줄.
O2: 머리줄만 평평한 목록으로 보내 세로로 쪼개졌다(과거 기록 2/2 가 이 모양, 둘 다 머리줄).
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl

from alture_jarvis.agent import BrowserAgent
from alture_jarvis.documents import write_document

DOIT = "https://product.kyobobook.co.kr/detail/S000202532365"
PYTORCH = "https://ebook-product.kyobobook.co.kr/dig/epd/ebook/E000003122079"
ML = "https://product.kyobobook.co.kr/detail/S000209068452"

# 09-28 17:1x 교보 상세에서 직접 본 첫머리(요약). 추천 띠는 1,200자 뒤.
DOIT_BODY = (
    "국내도서 컴퓨터/IT Do it! 점프 투 파이썬 중학생도 첫날부터 실습하는 초고속 입문서 2판 박응용 이지스퍼블리싱 "
    "10% 19,800 원 22,000 원 통합포인트 1,100 P (5% 적립) eBook 13,500원 "
    + "책 소개 " * 200  # 추천 띠를 값 대조 창(900자) 뒤에 둔다 — 실제 페이지는 1,350자째
    + " 함께 구매한 책 하네스 엔지니어링 with 클로드 코드 황민호 27,000원"
)
PYTORCH_BODY = (
    "eBook IT/프로그래밍 파이썬 딥러닝 파이토치(Python Deep Learning PyTorch) 국내도서 22,500원 eBook 15,300원 "
    "이경택 지음 정보문화사 소장 17,000원 10% 15,300 원 통합포인트 5%적립"
)
ML_BODY = "국내도서 파이썬(Python) 3학년 머신러닝의 구조 10% 16,200 원 18,000 원 통합포인트 900 P (5% 적립)"


def _opened(url: str, title: str) -> dict[str, Any]:
    return {"tool": "browser_extract", "ok": True, "error": "", "evidence": {"url": url, "title": title}}


READ = [
    _opened(ML, "파이썬(Python) 3학년 머신러닝의 구조 - 교보문고"),
    _opened(PYTORCH, "파이썬 딥러닝 파이토치(Python Deep Learning PyTorch) | 이경택 | 정보문화사 - 교보ebook"),
    _opened(DOIT, "Do it! 점프 투 파이썬 - 교보문고"),
]
BODIES = {BrowserAgent._url_key(ML): ML_BODY, BrowserAgent._url_key(PYTORCH): PYTORCH_BODY, BrowserAgent._url_key(DOIT): DOIT_BODY}


def _wrote_table(rows: list[list[Any]]) -> dict[str, Any]:
    return {"tool": "local_write_document", "ok": True, "arguments": {"format": "xlsx", "blocks": [{"type": "table", "rows": rows}]}}


def test_a_value_taken_from_the_recommendation_strip_is_flagged_even_without_urls() -> None:
    steps = [*READ, _wrote_table([
        ["순위", "제목", "판매가", "적립 포인트"],
        [1, "파이썬(Python) 3학년 머신러닝의 구조", "16,200원", "900P (5%)"],
        [2, "파이썬 딥러닝 파이토치(Python Deep Learning PyTorch)", "35,000원", "3,500P (10%)"],
        [3, "Do it! 점프 투 파이썬", "27,000원", "2,700P (10%)"],
    ])]

    note = BrowserAgent._row_page_binding_check("", steps, BODIES)

    assert "「Do it! 점프 투 파이썬」 27,000원 · 2,700P" in note and "19,800원" in note
    assert "35,000원" in note, "그 페이지에 아예 없는 가격"
    assert "3학년 머신러닝" not in note, "맞게 적은 줄까지 몰았다"


def test_right_values_stay_quiet() -> None:
    steps = [*READ, _wrote_table([
        ["순위", "제목", "판매가", "적립 포인트"],
        [1, "파이썬(Python) 3학년 머신러닝의 구조", "16,200원", "900P"],
        [2, "파이썬 딥러닝 파이토치", "15,300원", "5%"],
        [3, "Do it! 점프 투 파이썬 (개정판)", "19,800원", "1,100P"],
    ])]

    assert BrowserAgent._row_page_binding_check("", steps, BODIES) == ""


def test_a_page_whose_head_states_no_price_is_not_judged() -> None:
    """급여가 본문 뒤쪽에 있는 공고처럼 첫머리에 값이 없으면 판정하지 않는다 — 헛경고를 안 낸다."""
    url = "https://www.saramin.co.kr/zf_user/jobs/relay/view?rec_idx=9"
    steps = [_opened(url, "[(주)솔바람뜰여수] 재무/회계 경력직 채용 - 사람인"), _wrote_table([
        ["회사", "직무", "급여"], ["(주)솔바람뜰여수", "회계", "3,200만원"],
    ])]
    body = "[(주)솔바람뜰여수] 재무/회계 경력직 채용 모집부문 회계 자격요건 경력 3년 " + "상세 " * 400 + " 급여 3,200만원"

    assert BrowserAgent._row_page_binding_check("", steps, {BrowserAgent._url_key(url): body}) == ""


def test_an_answer_table_whose_values_never_reached_the_file_is_flagged(tmp_path: Path) -> None:
    empty = tmp_path / "파이썬 책 3권.xlsx"
    book = openpyxl.Workbook()
    for label in ("순위", "제목", "판매가", "적립 포인트", "출처 URL"):
        book.active.append([label])
    book.save(empty)
    answer = (
        "| 순위 | 제목 | 판매가 | 적립 포인트 |\n|---|---|---|---|\n"
        "| 1 | Do it! 점프 투 파이썬 | 19,800원 | 1,100P |\n"
        "| 2 | 초보자를 위한 파이썬 200제 | 18,000원 | 1,000P |\n"
        "| 3 | 혼자 공부하는 파이썬 | 19,800원 | 1,100P |\n\n엑셀 파일도 동일하게 생성했습니다."
    )
    steps = [{"tool": "local_write_document", "ok": True, "evidence": {"path": str(empty)}}]

    assert "만든 파일(파이썬 책 3권.xlsx)에 답의 표 값이 하나도 없습니다" in BrowserAgent._answer_missing_from_artifact(answer, steps)

    full = tmp_path / "채운 것.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["순위", "제목", "판매가"])
    book.active.append([1, "Do it! 점프 투 파이썬", "19,800원"])
    book.save(full)
    steps = [{"tool": "local_write_document", "ok": True, "evidence": {"path": str(full)}}]
    assert BrowserAgent._answer_missing_from_artifact(answer, steps) == ""


def test_a_flat_header_list_becomes_one_row_and_is_reported(tmp_path: Path) -> None:
    target = tmp_path / "책.xlsx"

    shape = write_document(target, "xlsx", {"sheets": {"파이썬 책": ["순위", "제목", "판매가", "적립 포인트", "URL"]}})

    sheet = openpyxl.load_workbook(target).active
    assert [cell.value for cell in sheet[1]] == ["순위", "제목", "판매가", "적립 포인트", "URL"], "머리줄이 세로로 쪼개졌다"
    assert sheet.max_row == 1
    assert "값 줄이 0개" in shape["header_only_sheets"]


def test_a_full_sheet_is_not_reported(tmp_path: Path) -> None:
    shape = write_document(tmp_path / "책.xlsx", "xlsx", {"sheets": {"책": [["제목", "가격"], ["Do it!", "19,800원"]]}})

    assert "header_only_sheets" not in shape


def test_only_the_last_write_of_a_file_is_read() -> None:
    """O6: 첫 쓰기의 오타(「파이썰」)와 고친 쓰기(「파이썬」)를 두 줄로 세어 「한 주소를 여러 줄이 가리킨다」로 건너뛰었다."""
    deep = "https://product.kyobobook.co.kr/detail/S000215599933"
    steps = [_opened(deep, "밑바닥부터 시작하는 딥러닝 1(리마스터판) - 교보문고")]
    for title in ("Do it! 점프 투 파이썰", "Do it! 점프 투 파이썬"):
        steps.append({"tool": "local_write_document", "ok": True, "arguments": {
            "format": "xlsx", "path": "책.xlsx", "sheets": {"책": [["순위", "제목", "판매가", "URL"], [1, title, "19,800원", deep]]}}})

    note = BrowserAgent._row_page_binding_check("", steps)

    assert "「Do it! 점프 투 파이썬」 줄의 주소는 「밑바닥부터 시작하는 딥러닝 1(리마스터판)」 페이지" in note


def test_a_row_citing_a_page_never_opened_is_flagged() -> None:
    """O6: 「혼자 공부하는 파이썬 27,000원」에 한 번도 안 연 `…216488888` 을 달았다."""
    steps = [*READ, _wrote_table([
        ["순위", "제목", "판매가", "URL"],
        [1, "혼자 공부하는 파이썬", "27,000원", "https://product.kyobobook.co.kr/detail/S000216488888"],
    ])]

    note = BrowserAgent._row_page_binding_check("", steps)

    assert "이번에 열지 않은 주소" in note and "S000216488888" in note


def test_the_same_posting_opened_under_another_url_shape_counts_as_opened() -> None:
    """사람인: 목록의 긴 주소와 실제로 연 주소는 꼴이 다르다 — 번호가 같으면 연 것이다(역검증 헛울림 57 → 7)."""
    opened = "https://www.saramin.co.kr/zf_user/jobs/relay/view?rec_idx=51000002"
    listed = "https://www.saramin.co.kr/zf_user/jobs/relay/view?view_type=search&rec_idx=51000002&search_uuid=35a12d42"
    steps = [_opened(opened, "(주)푸른빛에너지 회계사무직 채용 - 사람인"), _wrote_table([
        ["회사", "직무", "주소"], ["(주)푸른빛에너지", "회계사무직", listed],
    ])]

    assert "열지 않은 주소" not in BrowserAgent._row_page_binding_check("", steps)


class _Store:
    def __init__(self, bodies: dict[int, str]) -> None:
        self.bodies = bodies

    def get(self, job_id: str, step: int) -> dict[str, Any]:
        return {"payload": {"text": self.bodies.get(step, "")}}


def _agent_with(bodies: dict[int, str]) -> BrowserAgent:
    agent = BrowserAgent.__new__(BrowserAgent)
    agent.observations = _Store(bodies)
    agent.config = None
    agent.events = type("E", (), {"publish": lambda *a, **k: None})()
    return agent


def _read_step(step: int, url: str, title: str) -> dict[str, Any]:
    return {**_opened(url, title), "step": step, "evidence_record": {
        "type": "body_read", "ok": True, "source_id": url, "job_id": "j", "step": step}}


def test_a_wrong_row_is_sent_back_at_the_moment_of_writing() -> None:
    """쓰는 순간 되돌린다 — 끝나고 경고만 붙이던 것(O7)보다 한 단계 앞(O6·O7: 열고도 딴 값을 씀)."""
    agent = _agent_with({1: DOIT_BODY})
    steps = [_read_step(1, DOIT, "Do it! 점프 투 파이썬 - 교보문고"), {**_wrote_table([
        ["순위", "제목", "판매가", "적립 포인트"], [1, "Do it! 점프 투 파이썬", "27,000원", "2,700P"],
    ]), "step": 2}]

    feedback = agent._write_grounding_feedback("local_write_document", {"ok": True}, steps, "j")

    assert feedback.startswith("★") and "19,800원" in feedback and "다시 써라" in feedback
    assert steps[-1]["write_grounding_feedback"] is True


def test_the_write_feedback_stops_after_two_rounds() -> None:
    agent = _agent_with({1: DOIT_BODY})
    wrong = _wrote_table([["순위", "제목", "판매가"], [1, "Do it! 점프 투 파이썬", "27,000원"]])
    steps = [_read_step(1, DOIT, "Do it! 점프 투 파이썬 - 교보문고"),
             {**wrong, "write_grounding_feedback": True}, {**wrong, "write_grounding_feedback": True}, {**wrong}]

    assert agent._write_grounding_feedback("local_write_document", {"ok": True}, steps, "j") == ""


def test_a_right_write_gets_no_feedback() -> None:
    agent = _agent_with({1: DOIT_BODY})
    steps = [_read_step(1, DOIT, "Do it! 점프 투 파이썬 - 교보문고"), _wrote_table([
        ["순위", "제목", "판매가", "적립 포인트"], [1, "Do it! 점프 투 파이썬", "19,800원", "1,100P"],
    ])]

    assert agent._write_grounding_feedback("local_write_document", {"ok": True}, steps, "j") == ""


def test_a_name_missing_a_piece_of_the_title_still_finds_its_page() -> None:
    """★ 2026-09-28 T1 O9(`t1_grounding_O8-O10_O9`): 표에 주소 칸이 없고 이름이 제목과 글자 그대로 안 겹쳤다 —
    「파이썬 머신러닝의 구조」(제목은 「파이썬(Python) 3학년 …」) · 「[eBook] 파이썬 딥러닝 파이토치(…)」. 값 25,000원·2,500P · 36,000원 이
    경고 없이 지나갔다(실제 16,200원·900P · 15,300원)."""
    steps = [*READ, _wrote_table([
        ["순위", "제목", "판매가", "적립 포인트"],
        ["1", "파이썬 머신러닝의 구조", "25,000원", "2,500P (10%)"],
        ["2", "[eBook] 파이썬 딥러닝 파이토치(Python Deep Learning PyTorch)", "36,000원", "3,600P (10%)"],
        ["3", "Do it! 점프 투 파이썬", "19,800원", "1,100P (5%)"],
    ])]

    note = BrowserAgent._row_page_binding_check("", steps, BODIES)

    assert "25,000원" in note and "16,200원" in note
    assert "36,000원" in note
    assert "「Do it! 점프 투 파이썬」" not in note, "맞게 적은 줄까지 몰았다"


def test_a_sheet_given_as_header_and_rows_is_written(tmp_path: Path) -> None:
    """★ T1 O10: `{"header": [...], "rows": [...]}` 를 두 번 보냈는데 빈 시트 + 「rows 로 채우세요」(이미 rows 였다)."""
    target = tmp_path / "책.xlsx"

    shape = write_document(target, "xlsx", {"sheets": {"파이썬 책 정보": {
        "header": ["책 제목", "판매가", "적립 포인트"],
        "rows": [["파이썬(Python) 3학년 머신러닝의 구조", "16,200원", "900P"], ["Do it! 점프 투 파이썬", "19,800원", "1,100P"]],
    }}})

    rows = [[cell.value for cell in row] for row in openpyxl.load_workbook(target).active.iter_rows()]
    assert rows[0] == ["책 제목", "판매가", "적립 포인트"]
    assert rows[2] == ["Do it! 점프 투 파이썬", "19,800원", "1,100P"]
    assert "header_only_sheets" not in shape and "unreadable_sheets" not in shape


def test_a_sheet_we_cannot_read_says_so_instead_of_asking_for_rows(tmp_path: Path) -> None:
    shape = write_document(tmp_path / "책.xlsx", "xlsx", {"sheets": {"책": {"뭔가": "모르는 모양"}}})

    assert "빈 시트" in shape["unreadable_sheets"]
    assert "header_only_sheets" not in shape, "틀린 안내 — 머리줄조차 없다"


def test_an_answer_of_names_only_over_an_empty_file_is_flagged(tmp_path: Path) -> None:
    """★ T1 O10: 답 표는 책 셋, 값 칸은 전부 「확인 필요」(숫자 0개) — 파일은 완전히 빈 시트였는데 경고 0 · verified."""
    empty = tmp_path / "python_books (2).xlsx"
    book = openpyxl.Workbook()
    book.active.title = "파이썬 책 정보"
    book.save(empty)
    answer = (
        "교보문고에서 '파이썬' 검색 결과 중 상위 3권의 정보를 정리한 엑셀 파일입니다.\n\n"
        "| 책 제목 | 판매가 | 적립 포인트 |\n|---|---|---|\n"
        "| 파이썬(Python) 3학년 머신러닝의 구조 | 확인 필요 | 확인 필요 |\n"
        "| 초보자를 위한 파이썬(Python) 200제 | 확인 필요 | 확인 필요 |\n"
        "| 혼자 공부하는 파이썬 | 확인 필요 | 확인 필요 |\n\n실제로 읽은 페이지에서 값을 추출해 다시 정리하겠습니다."
    )
    steps = [{"tool": "local_write_document", "ok": True, "evidence": {"path": str(empty)}}]

    assert "python_books (2).xlsx" in BrowserAgent._answer_missing_from_artifact(answer, steps)


def test_a_short_names_only_answer_stays_quiet(tmp_path: Path) -> None:
    """이름이 셋 미만이고 숫자도 없으면 판정하지 않는다(좁게 문다)."""
    empty = tmp_path / "빈.xlsx"
    openpyxl.Workbook().save(empty)
    answer = "| 항목 | 상태 |\n|---|---|\n| 파이썬 책 목록 | 확인 필요 |\n| 적립 포인트 표 | 확인 필요 |\n| 정리 | 끝 |"
    steps = [{"tool": "local_write_document", "ok": True, "evidence": {"path": str(empty)}}]

    assert BrowserAgent._answer_missing_from_artifact(answer, steps) == ""


def test_points_are_matched_against_points_not_any_number_on_the_page() -> None:
    """★ T1 O13: 답에 적립 「12,960P」·「12,600P」 — 그 페이지의 **eBook 가격(원)** 이었다. 숫자가 페이지에 있어서 통과, 경고 0."""
    head = "국내도서\n18,000원\neBook\n12,600원\n장삼용\n10%\n18,000\n원\n20,000\n원\n통합포인트\n1,000\nP\n(5% 적립)"
    name = "초보자를 위한 파이썬(Python) 200제"

    wrong = BrowserAgent._row_value_mismatch(["2", name, "18,000원", "12,600P (5% 적립 기준)"], name, head)

    assert wrong is not None and "12,600P" in wrong[0] and "1,000P" in wrong[1]
    assert BrowserAgent._row_value_mismatch(["2", name, "18,000원", "1,000P"], name, head) is None


def _write_step(path: str, rows: int, sheets: list[str], **extra: Any) -> dict[str, Any]:
    return {"tool": "local_write_document", "ok": True,
            "evidence": {"action": "local_write_document", "format": "xlsx", "output_path": path, "rows": rows, "sheets": sheets, **extra}}


def test_a_header_only_table_file_is_flagged_at_the_end() -> None:
    """★ T1 O23(`032dfa64`): 머리줄만 쓰고 걸음 상한까지 책만 열다 끝남 — 「만듭니다」 예고 · goal_met · 경고 0."""
    steps = [_write_step(r"C:\Users\x\Documents\JARVIS\파이썬 책 3권 (9).xlsx", 1, ["파이썬 책"],
                         header_only_sheets="시트 파이썬 책 에는 머리줄만 있고 값 줄이 0개입니다")]

    assert "파이썬 책 3권 (9).xlsx" in BrowserAgent._empty_table_artifact_check(steps)


def test_a_filled_rewrite_of_the_same_file_clears_it() -> None:
    path = r"C:\x\책.xlsx"
    steps = [_write_step(path, 1, ["책"]), _write_step(path, 4, ["책"])]

    assert BrowserAgent._empty_table_artifact_check(steps) == ""


def test_a_separate_source_sheet_does_not_make_it_empty() -> None:
    """O19 모양: 「출처」 시트 4줄 + 책 시트 4줄 → 조용. 시트 둘 다 머리줄뿐일 때만 울린다."""
    assert BrowserAgent._empty_table_artifact_check([_write_step(r"C:\x\책.xlsx", 8, ["출처", "파이썬 책"])]) == ""
    assert BrowserAgent._empty_table_artifact_check([_write_step(r"C:\x\책.xlsx", 2, ["출처", "파이썬 책"])]) != ""


def test_a_header_only_file_filled_later_by_cell_edits_stays_quiet() -> None:
    """역검증 헛울림 O21(`4f556d8f`): 머리줄만 쓴 `(8)` 을 걸음 10 의 칸 고치기가 두 줄 채웠다."""
    path = r"C:\x\파이썬 책 3권 (8).xlsx"
    steps = [_write_step(path, 1, ["파이썬 책"]),
             {"tool": "local_edit_spreadsheet", "ok": True, "evidence": {"output_path": path}}]

    assert BrowserAgent._empty_table_artifact_check(steps) == ""
