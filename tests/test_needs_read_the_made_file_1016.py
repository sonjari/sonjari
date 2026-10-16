"""「돌려받을 항목(needs)」 검사가 **만든 파일의 머리글**도 본다 · 칸 이름을 늘어놓은 요청은 셈이 아니다 (2026-10-16).

10-16 저녁 바깥 AI 가 맡긴 판(job 759ba894): 「r/LLMDevs 이번 달 인기 글 25개의 제목 · 점수(업보트) · 댓글 수 · 올린 지
얼마나 됐는지를 엑셀로」. 엑셀에는 머리글(순위 · 제목 · 점수(업보트) · 댓글 수 · 올린 지 · 올린 시각)과 25줄이 다 있었는데
needs 「제목 · 점수 · 댓글 수 · 올린 시각」이 넷 다 met:false(「필요한 것 0/4」)였다 — 검사가 **답 글만** 봤고, 답은 파일 경로만
적었다. 같은 판의 「시킨 것 대조」에는 시키지 않은 「직접 계산해 보기」가 ❌ 로 들어가 completion_state 가 partial 이 됐다 —
셈 대상 낱말표의 「점수」와 「정리해 줘」가 짝지어졌다. 그 「점수」는 화면에서 읽어 옮길 **칸 이름**이었다.

양방향으로 지킨다: 칸이 있는 파일 → 채움 · 칸이 없는 파일 → 못 채움 · 칸 이름 목록 → 셈 아님 · 진짜 셈 → 그대로 셈.
표본 글과 값은 지어낸 것이다.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest

from alture_jarvis.delegation import summarize_job
from alture_jarvis.outbound import check_needs
from alture_jarvis.plan import extract, wants_a_computation

NEEDS = ["제목", "점수", "댓글 수", "올린 시각"]
HEAD = ["순위", "제목", "점수(업보트)", "댓글 수", "올린 지", "올린 시각"]
ROWS = [[1, "가상 글 하나", 440, 57, "16일 전", "2026-09-29 22:45"],
        [2, "가상 글 둘", 310, 12, "20일 전", "2026-09-25 08:10"]]
ANSWER = "엑셀 파일을 만들었습니다: `{path}` (머리글 1 + 글 2)"


def _xlsx(path: Path, head: list[Any], rows: list[list[Any]]) -> Path:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.append(head)
    for row in rows:
        sheet.append(row)
    book.save(path)
    return path


def _met(needs: list[str], answer: str, files: list[str]) -> dict[str, bool]:
    return {item["label"]: item["met"] for item in check_needs(needs, answer, files=files)}


# ── 만든 파일의 머리글 ─────────────────────────────────────────────────────────────


def test_a_made_xlsx_with_the_columns_fills_the_needs(tmp_path: Path) -> None:
    made = _xlsx(tmp_path / "인기글.xlsx", HEAD, ROWS)

    assert _met(NEEDS, ANSWER.format(path=made), [str(made)]) == dict.fromkeys(NEEDS, True)


def test_a_made_xlsx_without_the_columns_does_not(tmp_path: Path) -> None:
    made = _xlsx(tmp_path / "다른표.xlsx", ["순위", "주소"], [[1, "https://example.com/a"]])

    assert _met(NEEDS, ANSWER.format(path=made), [str(made)]) == dict.fromkeys(NEEDS, False)


def test_a_column_with_no_values_does_not_count(tmp_path: Path) -> None:
    made = _xlsx(tmp_path / "빈칸.xlsx", HEAD, [[1, "가상 글 하나", 440, 57, "16일 전", None]])

    assert _met(["올린 시각", "제목"], "", [str(made)]) == {"올린 시각": False, "제목": True}


def test_a_csv_and_a_docx_table_count_too(tmp_path: Path) -> None:
    csv_path = tmp_path / "인기글.csv"
    csv_path.write_text(",".join(HEAD) + "\n" + "\n".join(",".join(map(str, row)) for row in ROWS) + "\n", encoding="utf-8-sig")
    assert _met(NEEDS, "", [str(csv_path)]) == dict.fromkeys(NEEDS, True)

    from docx import Document

    document = Document()
    table = document.add_table(rows=1 + len(ROWS), cols=len(HEAD))
    for index, name in enumerate(HEAD):
        table.rows[0].cells[index].text = name
    for r, row in enumerate(ROWS, start=1):
        for index, value in enumerate(row):
            table.rows[r].cells[index].text = str(value)
    docx_path = tmp_path / "인기글.docx"
    document.save(docx_path)
    assert _met(NEEDS, "", [str(docx_path)]) == dict.fromkeys(NEEDS, True)


def test_a_missing_or_broken_file_is_not_counted(tmp_path: Path) -> None:
    broken = tmp_path / "깨짐.xlsx"
    broken.write_bytes(b"not a workbook")

    assert _met(["제목"], "", [str(tmp_path / "없음.xlsx"), str(broken)]) == {"제목": False}


def test_the_job_summary_reads_the_file_it_made(tmp_path: Path) -> None:
    """★ 재현 그대로 — 답은 경로만, 표는 파일에만 있다."""
    made = _xlsx(tmp_path / "rLLMDevs_이번달_인기글_25.xlsx", HEAD, ROWS)
    job: dict[str, Any] = {
        "id": "j1", "status": "completed",
        "payload": {"prompt": "인기 글을 엑셀로", "origin": "mcp", "needs": NEEDS},
        "result": {"answer": ANSWER.format(path=made), "agent_trace": {
            "steps": [{"tool": "local_write_document", "ok": True,
                       "evidence": {"action": "local_write_document", "format": "xlsx", "output_path": str(made)},
                       "evidence_record": {"effect": "write_local", "source_id": str(made).lower()}}],
            "tool_numbers": ["1", "2", "440", "57", "310", "12"]}},
    }

    summary = summarize_job(job)

    assert [item["met"] for item in summary["needs"]] == [True] * 4
    assert summary["receipt"]["needs_met"] == 4
    assert "비어 있는 항목" not in str(summary)


# ── 칸 이름을 늘어놓은 요청은 셈이 아니다 ──────────────────────────────────────────

REDDIT = ("크롬에서 https://www.reddit.com/r/LLMDevs/top/?t=month 를 열고, 이번 달 인기 글 위에서부터 25개의 제목 · "
          "점수(업보트) · 댓글 수 · 올린 지 얼마나 됐는지를 엑셀 파일 하나로 정리해 줘. 글은 열지 말고 목록 화면에서 읽으면 돼.")


@pytest.mark.parametrize("prompt", [
    REDDIT,
    "학원 후기 10개의 작성자, 점수, 날짜를 엑셀로 만들어줘",
    "다나와 무선마우스 5개의 이름과 점수를 표로 정리해줘",
    "Open r/LLMDevs top of the month and put the title, score, comment count and age of the top 25 posts into an Excel file.",
])
def test_listed_column_names_are_not_a_computation(prompt: str) -> None:
    assert not wants_a_computation(prompt), prompt
    assert "sandbox_run" not in [r.target for r in extract(prompt)]


@pytest.mark.parametrize("prompt", [
    # 칸을 늘어놓았어도 **셈 동사**가 있으면 셈이다.
    "학생 다섯 명의 이름 · 점수를 보고 평균 점수를 계산해줘",
    "제목 · 점수 · 댓글 수를 모으고 점수의 합계를 구해줘",
    # 목록이 아닌 셈 대상 + 만드는 동사는 예전처럼 셈이다.
    "세 반의 평균을 엑셀로 만들어줘",
    "Calculate the average score of these five students",
])
def test_a_real_computation_still_is_one(prompt: str) -> None:
    assert wants_a_computation(prompt), prompt
