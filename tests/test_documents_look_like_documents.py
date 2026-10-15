"""**내용이 맞아도 문서 꼴이 아니면 못 쓴다.**

## 라이브에서 나온 것 (2026-09-11)

    여수_회계_채용_공고_조사_결과.docx
      [Normal] 여수 지역 회계 채용 공고 조사 결과 (2026-09-10 기준)   ← 제목인데 Normal
      문서 기본 글꼴: None / None                                     ← 한글 글꼴 미지정

    python_books_info.md   (134바이트)
      제목⇥쪽수⇥ISBN⇥상세페이지 URL
      Do it! 점프 투 파이썬⇥432쪽⇥9791163034735⇥https://…

앞의 것은 **모델이 heading 을 안 쓴 것**이고(코드엔 있었다), 뒤의 것은
**우리 코드에 마크다운 쓰기가 아예 없어서** 평문 갈래로 떨어진 것이다.

서식은 모델에게 부탁할 일이 아니다 — 결정론으로 박을 수 있으면 코드로 간다.
"""
from __future__ import annotations

import pathlib
import re
import tempfile
import zipfile


def _tmp(name: str) -> pathlib.Path:
    return pathlib.Path(tempfile.mkdtemp()) / name


def test_a_korean_document_pins_its_korean_font() -> None:
    """`font.name` 은 `w:ascii` 만 건드린다. 한글은 `w:eastAsia` 를 본다 —
    그것만 안 정하면 여는 PC 마다 굴림·바탕·맑은고딕이 제각각이다."""
    from alture_jarvis.documents import _write_docx

    target = _tmp("a.docx")
    _write_docx(target, [{"type": "paragraph", "text": "한글 본문"}], [], "")
    styles = zipfile.ZipFile(target).read("word/styles.xml").decode("utf-8", "replace")
    assert "맑은 고딕" in styles, "동아시아 글꼴이 안 박혔다"


def test_a_heading_block_becomes_a_real_heading() -> None:
    import docx

    from alture_jarvis.documents import _write_docx

    target = _tmp("b.docx")
    _write_docx(target, [
        {"type": "heading", "level": 1, "text": "조사 결과"},
        {"type": "paragraph", "text": "본문"},
    ], [], "")
    styles = [p.style.name for p in docx.Document(str(target)).paragraphs if p.text.strip()]
    assert styles[0] == "Heading 1", styles


def test_a_table_header_is_bold_and_repeats_across_pages() -> None:
    """긴 표는 다음 쪽에서 머리줄이 사라져 어느 칸이 무엇인지 알 수 없게 된다."""
    import docx

    from alture_jarvis.documents import _write_docx

    target = _tmp("c.docx")
    _write_docx(target, [{"type": "table", "rows": [["회사명", "마감일"],
                                                    ["가나", "2026.10.09"]]}], [], "")
    table = docx.Document(str(target)).tables[0]
    assert table.rows[0].cells[0].paragraphs[0].runs[0].bold is True
    body = zipfile.ZipFile(target).read("word/document.xml").decode("utf-8", "replace")
    assert "w:tblHeader" in body


def test_markdown_is_written_as_markdown_not_tabs() -> None:
    """★ 이게 라이브에서 나온 134바이트짜리 파일의 원인이었다."""
    from alture_jarvis.documents import _write_markdown

    target = _tmp("d.md")
    _write_markdown(target, [
        {"type": "heading", "level": 1, "text": "파이썬 도서 조사"},
        {"type": "table", "rows": [["도서명", "쪽수"], ["점프 투 파이썬", "432"]]},
    ], [], "")
    body = target.read_text(encoding="utf-8")
    assert body.startswith("# 파이썬 도서 조사"), body
    assert "| 도서명 | 쪽수 |" in body, body
    assert "| --- | --- |" in body, "구분선이 없으면 표로 안 그려진다"
    assert "\t" not in body, "탭 구분으로 떨어졌다"


def test_rows_alone_still_make_a_markdown_table() -> None:
    """모델이 `blocks` 를 안 주고 `rows` 만 줘도 표가 나와야 한다 —
    형식을 못 맞춰 실패하는 것이 제일 흔한 고장이다(이 모듈의 설계 규율)."""
    from alture_jarvis.documents import _write_markdown

    target = _tmp("e.md")
    _write_markdown(target, None, [["제목", "쪽수"], ["점프 투 파이썬", "432"]], "")
    body = target.read_text(encoding="utf-8")
    assert "| 제목 | 쪽수 |" in body
    assert "| --- | --- |" in body


def test_a_pipe_inside_a_cell_does_not_break_the_table() -> None:
    from alture_jarvis.documents import _md_table

    body = _md_table([["이름"], ["가|나"]])
    assert r"가\|나" in body, body
    assert len(body.splitlines()) == 3


def test_plain_text_still_uses_tabs() -> None:
    """txt 는 마크다운이 아니다 — 파이프 표를 쓸 이유가 없다."""
    from alture_jarvis.documents import write_document

    target = _tmp("f.txt")
    write_document(target, "txt", {"rows": [["가", "나"], ["1", "2"]]})
    assert "	" in target.read_text(encoding="utf-8")


def test_markdown_goes_to_its_own_writer_before_the_plain_fallback() -> None:
    """갈래 순서가 뒤집히면 `md` 가 다시 평문으로 떨어진다."""
    import inspect

    from alture_jarvis import documents

    source = inspect.getsource(documents.write_document)
    assert 'elif fmt == "md":' in source
    assert source.index('elif fmt == "md":') < source.index("body = text or")


def test_the_tool_tells_the_model_how_to_send_a_heading() -> None:
    """코드에 `heading` 이 있는데 **쓰는 법을 안 알려 줘서** 안 쓰였다.

    라이브 문서의 큰 제목이 `Normal` 로 나온 진짜 원인이 여기였다.
    """
    body = pathlib.Path(
        "src/alture_jarvis/local_computer.py"
    ).read_text(encoding="utf-8")
    spot = body[body.index('"blocks": {'):][:2000]
    assert "heading" in spot and "table" in spot
    assert re.search(r"제목은 heading", spot), "제목 주는 법이 설명에 없다"


def test_a_heading_does_not_get_two_hashes() -> None:
    """모델이 `#` 를 붙여 보낸다. 우리가 또 붙이면 `# # 제목` 이 된다 (2026-09-11 라이브)."""
    from alture_jarvis.documents import _write_markdown

    target = _tmp("g.md")
    _write_markdown(target, [{"type": "heading", "level": 1,
                              "text": "# YES24 파이썬 도서 조사 결과"}], [], "")
    body = target.read_text(encoding="utf-8")
    assert body.startswith("# YES24"), body
    assert "# #" not in body


def test_a_word_heading_does_not_keep_the_hash() -> None:
    """워드는 스타일로 제목을 준다 — 글에 `#` 가 남으면 그대로 보인다."""
    import docx

    from alture_jarvis.documents import _write_docx

    target = _tmp("h.docx")
    _write_docx(target, [{"type": "heading", "level": 1, "text": "## 조사 결과"}], [], "")
    text = [p.text for p in docx.Document(str(target)).paragraphs if p.text.strip()]
    assert text[0] == "조사 결과", text


def test_a_newline_sent_as_two_characters_becomes_a_newline() -> None:
    """라이브에서 「**부족한 정보**\n* 도서 2권 미확보」가 한 줄로 나왔다.

    모델이 줄바꿈을 **글자 두 개**(역슬래시+n)로 보낸 것이다. 받는 쪽에서 푼다 —
    이렇게 코드로 옮긴 규칙만큼 작은 모델도 같은 결과를 낸다.
    """
    from alture_jarvis.documents import _write_markdown

    two_chars = chr(92) + "n"          # 역슬래시 + n. 진짜 줄바꿈이 아니다.
    target = _tmp("i.md")
    _write_markdown(target, [{"type": "paragraph",
                              "text": "**부족한 정보**" + two_chars + "* 도서 2권 미확보"}],
                    [], "")
    body = target.read_text(encoding="utf-8")
    assert two_chars not in body, body
    assert "**부족한 정보**" + chr(10) + "* 도서 2권 미확보" in body, body


def test_a_real_newline_is_left_alone() -> None:
    """이미 진짜 줄바꿈이 있으면 손대지 않는다 — 코드 블록의 역슬래시를 망치지 않게."""
    from alture_jarvis.documents import _write_markdown

    target = _tmp("j.md")
    _write_markdown(target, [{"type": "code", "language": "python",
                              "text": "print('a\\nb')\nprint('c')"}], [], "")
    body = target.read_text(encoding="utf-8")
    assert r"print('a\nb')" in body, body


def test_a_table_with_only_a_header_is_called_out() -> None:
    """머리줄만 있는 표는 **빈 표다** (2026-09-11).

    사슬 벤치에서 yes24 판이 두 모델·두 라운드 모두 이렇게 나왔다 — 표 산출물
    8개 중 **4개**다. 내용은 다 조사해 놓고 담는 그릇만 틀렸다.

    기계가 합쳐 넣을 수도 있지만 열이 한 칸씩 밀려 있어(제목이 절 머리로 빠졌다)
    저자 칸에 출판사가 들어간다. **데이터를 망가뜨리느니 짚어 주고 모델이 고치게
    한다** — 이 저장소의 「막지 말고 알려주기」와 같은 규율이다.
    """
    from alture_jarvis.documents import _shape_notes

    notes = _shape_notes("md", [
        {"type": "heading", "level": 1, "text": "YES24 파이썬 책 3권"},
        {"type": "table", "rows": [["제목", "저자", "쪽수", "ISBN", "URL"]]},
        {"type": "heading", "level": 2, "text": "1. Do it! 첫 파이썬"},
        {"type": "paragraph", "text": "엘리스 코딩 | 이지스퍼블리싱 | 268쪽"},
    ], [], "")
    assert any("머리줄만" in note for note in notes), notes


def test_a_filled_table_says_nothing() -> None:
    """넓힌 만큼 반대쪽도 붙든다 — 제대로 준 판에 잔소리가 붙으면 안 된다."""
    from alture_jarvis.documents import _shape_notes

    assert _shape_notes("md", [
        {"type": "heading", "level": 1, "text": "제목"},
        {"type": "table", "rows": [["제목", "쪽수"], ["책1", "268"], ["책2", "300"]]},
    ], [], "") == []


def test_a_table_without_a_header_is_called_out() -> None:
    """머리줄만 있던 것을 고쳤더니 이번엔 **머리줄이 없었다** (2026-09-11).

        | Do it! 점프 투 파이썬 | 432쪽 | 9791163034735 | https://… |
        | --- | --- | --- | --- |

    표의 첫 줄은 열 이름이다. 값이 거기 앉으면 그 항목이 통째로 사라지고, 받는 쪽은
    어느 칸이 무엇인지 모른다.
    """
    from alture_jarvis.documents import _shape_notes

    notes = _shape_notes("md", [
        {"type": "table", "rows": [
            ["Do it! 점프 투 파이썬", "432쪽", "9791163034735", "https://a"],
            ["혼자 공부하는 파이썬", "552쪽", "9791162245651", "https://b"],
        ]},
    ], [], "")
    assert any("머리줄이 없다" in note for note in notes), notes


def test_a_year_column_name_is_not_mistaken_for_data() -> None:
    """★ 넓힌 만큼 반대쪽. 「2026년 매출」은 열 이름이지 값이 아니다.

    그래서 판별은 **주소**와 **다섯 자리 넘는 숫자**만 본다 — 연도(네 자리)는 살린다.
    """
    from alture_jarvis.documents import _shape_notes

    notes = _shape_notes("md", [
        {"type": "heading", "level": 1, "text": "매출"},
        {"type": "table", "rows": [["항목", "2026년 매출", "증감"], ["A", "1200", "+3%"]]},
    ], [], "")
    assert notes == [], notes


def test_a_bullet_does_not_get_two_markers() -> None:
    """모델이 「- 」를 이미 붙여 보낸다 (`# #` 과 같은 병)."""
    from alture_jarvis.documents import _write_markdown

    target = _tmp("k.md")
    _write_markdown(target, [
        {"type": "bullet", "text": "- **조사 완료:** 2건"},
        {"type": "bullet", "text": "1. 정상 항목"},
    ], [], "")
    body = target.read_text(encoding="utf-8")
    assert "- - " not in body, body
    assert "- **조사 완료:** 2건" in body
    assert "- 정상 항목" in body


def test_a_spreadsheet_accepts_blocks_too() -> None:
    """★ 내가 낸 회귀 (2026-09-11).

    스킬에 「표는 table 블록으로 줘라」를 넣었더니 모델이 **엑셀에도** 그렇게 보냈다.
    그런데 xlsx 는 `rows`/`sheets` 만 받아서 **네 번 연속 거부**했다:

        오류: rows_or_sheets_required
        인자: blocks=[heading, paragraph, table…]   ← 내용은 다 있었다

    조사도 표 작성도 다 해 놓고 그릇 이름이 안 맞아 파일이 안 나왔다. 사슬 9/10 →
    7/10, 채움 67% → 38%.

    안내문을 되돌리는 대신 **받는 쪽을 넓힌다.** 부르는 쪽이 형식마다 다른 모양을
    외워야 하면 그건 언젠가 틀린다.
    """
    import openpyxl

    from alture_jarvis.documents import write_document

    blocks = [
        {"type": "heading", "level": 1, "text": "프로젝트"},
        {"type": "paragraph", "text": "직접 확인한 결과입니다."},
        {"type": "table", "rows": [["프로젝트명", "기간"], ["가", "90일"], ["나", "60일"]]},
    ]
    target = _tmp("l.xlsx")
    got = write_document(target, "xlsx", {"blocks": blocks})
    assert got["rows"] == 3, got
    cells = list(openpyxl.load_workbook(target).active.iter_rows(values_only=True))
    assert cells[0] == ("프로젝트명", "기간")
    assert cells[1] == ("가", "90일")


def test_a_csv_accepts_blocks_too() -> None:
    from alture_jarvis.documents import write_document

    target = _tmp("m.csv")
    got = write_document(target, "csv", {"blocks": [
        {"type": "table", "rows": [["가", "나"], ["1", "2"]]},
    ]})
    assert got["rows"] == 2
    assert "가,나" in target.read_text(encoding="utf-8-sig")


def test_a_spreadsheet_with_no_table_still_says_why() -> None:
    """넓혔다고 조용히 빈 파일을 만들면 안 된다 — 못 하는 것은 못 한다고 말한다."""
    import pytest

    from alture_jarvis.documents import write_document

    with pytest.raises(ValueError) as caught:
        write_document(_tmp("n.xlsx"), "xlsx", {"blocks": [
            {"type": "paragraph", "text": "표가 없는 글"},
        ]})
    assert "표가 있어야" in str(caught.value)
