"""디자인된 장표 손(`local_make_slide`) — 규칙 · 저장 · 화면 알림 · PPTX 굽기 · 숫자 대조 (2026-10-25 대표)."""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest

from alture_jarvis import slides
from alture_jarvis.agent import ungrounded_write_numbers
from alture_jarvis.config import JarvisConfig
from alture_jarvis.local_computer import LocalComputerTools

SLIDE = """<!DOCTYPE html><html><head><style>
html,body{width:1280px;height:720px;margin:0;font-family:'Malgun Gothic',sans-serif;background:linear-gradient(135deg,#0b1020,#26215c)}
h1{position:absolute;left:80px;top:120px;font-size:56px;color:#fff;margin:0}
p{position:absolute;left:80px;top:240px;font-size:24px;color:#cecbf6;margin:0}
.ghost{position:absolute;right:40px;top:300px;font-size:200px;color:rgba(255,255,255,0.06)}
</style><script>alert(1)</script></head><body onload="x()">
<h1>판 첫 걸음 30.6초 → 1.08초</h1><p>도구 목록 순서를 고정한 뒤 · 한 판 기준</p><div class="ghost">END</div>
</body></html>"""


@pytest.fixture(autouse=True)
def _no_browser_fit(request, monkeypatch):
    if "real_fit" not in request.node.name:
        monkeypatch.setattr(slides, "fit_in_subprocess", lambda path, **_: {"overflow": 0, "base_px": 20})


@pytest.fixture
def html_path(monkeypatch):
    """틀째 html 길 — 2026-10-26 부터 부품 결(지금은 모든 결)은 틀째 html 을 받지 않는다(표지 · 마무리까지 부품).
    이 길은 부품이 없는 결에만 남았다. 그 길의 동작(결 CSS 입히기 · 색 면 리듬 · 틀 글 남김)을 보는 시험은 「부품 없는 결」로 흉내 낸다."""
    from alture_jarvis import slide_parts

    monkeypatch.setattr(slide_parts, "supports", lambda theme: False)


def _tools(tmp_path: Path) -> tuple[LocalComputerTools, list]:
    config = JarvisConfig(runtime_root=tmp_path, data_dir=tmp_path / "data", output_dir=tmp_path / "out",
                          documents_dir=str(tmp_path / "docs"))
    tools = LocalComputerTools(config)
    sent: list = []
    tools.publish = lambda event, payload: sent.append((event, payload))
    return tools, sent


def _put(tools, args):
    """틀째 본문 장 — 10-25 저녁부터 모든 결이 부품 결이라 첫 번은 parts_required 로 돌아온다. 이 시험들은 틀 동작을 보므로 두 번째로 넣는다."""
    try:
        return tools.make_slide(args)
    except ValueError as exc:
        if "parts_required" not in str(exc):
            raise
        return tools.make_slide(args)


def test_without_html_it_hands_over_the_rules_themes_and_layouts(tmp_path: Path) -> None:
    """4차: 규칙 + 결 목록 + 배치 틀 12개. 영어 꼬리표를 강제하던 Akxan 플레이북은 더 주지 않는다(GUIDE 와 부딪쳤다)."""
    tools, _ = _tools(tmp_path)
    out = tools.make_slide({"title": "손자리 소개"})
    assert "<section" in out["guide"] and out["theme"] == "paper"
    assert "forest" in out["themes"] and "ink" in out["themes"]
    # 10-26 부터 모든 결이 틀(html) 없이 부품 — 표지 · 마무리도 frames 조립표(대표 「템플릿은 참고만」).
    assert "<section" not in out["layouts"] and "parts" in out and "assembly" in out
    assert '"frame":"cover"' in out["frames"] and '"frame":"closing"' in out["frames"]
    assert "design_playbook" not in out and "style_example_html" not in out
    picked = tools.make_slide({"title": "손자리 소개", "theme": "forest", "__job_id": "j"})
    assert picked["theme"] == "forest"
    assert (tmp_path / "out/slides/j/손자리 소개/theme.txt").read_text(encoding="utf-8") == "forest"
    assert "local_make_slide" in tools.names


def test_each_slide_is_cleaned_saved_and_shown(tmp_path: Path, monkeypatch) -> None:
    tools, sent = _tools(tmp_path)
    built = []
    def fake_render(deck, out):
        built.append((deck, out))
        out.write_bytes(b"pptx")
        return {"slides": 2, "text_boxes": 4}

    monkeypatch.setattr(slides, "render_in_subprocess", fake_render)
    first = tools.make_slide({"title": "손자리 소개", "index": 1, "total": 2, "html": SLIDE, "__job_id": "job1"})
    assert first["next"].startswith("2번") and "path" not in first
    saved = (tmp_path / "out/slides/job1/손자리 소개/slide_01.html").read_text(encoding="utf-8")
    assert "<script" not in saved and "onload" not in saved and "data-sonjari-size" in saved
    assert sent[0][0] == "deck_slide" and sent[0][1]["index"] == 1 and "<script" not in sent[0][1]["html"]
    last = tools.make_slide({"title": "손자리 소개", "index": 2, "total": 2, "html": SLIDE, "__job_id": "job1"})
    assert last["path"].endswith("손자리 소개.pptx") and last["verified"] and len(built) == 1
    assert sent[-1][0] == "deck_built"
    # 같은 제목으로 다시 구우면 덮지 않고 새 이름으로 둔다.
    again = tools.make_slide({"title": "손자리 소개", "index": 2, "total": 2, "html": SLIDE, "__job_id": "job1"})
    assert again["path"].endswith("손자리 소개 (2).pptx")


def test_a_missing_slide_is_named_before_baking(tmp_path: Path, monkeypatch) -> None:
    tools, _ = _tools(tmp_path)
    monkeypatch.setattr(slides, "render_in_subprocess", lambda deck, out: pytest.fail("빠진 장이 있는데 구웠다"))
    out = tools.make_slide({"title": "x", "index": 3, "total": 3, "html": SLIDE, "__job_id": "job2"})
    assert out["missing"] == [1, 2]


def test_bad_numbers_are_refused(tmp_path: Path) -> None:
    tools, _ = _tools(tmp_path)
    with pytest.raises(ValueError):
        tools.make_slide({"title": "x", "index": 3, "total": 2, "html": SLIDE})


def test_number_check_reads_only_the_visible_words() -> None:
    """CSS 의 1280px · 56px · 135deg 는 쓴 숫자가 아니다. 보이는 30.6 · 1.08 만 대조한다."""
    seen = {"30.6", "1.08"}
    assert ungrounded_write_numbers({"title": "t", "index": 1, "total": 5, "html": SLIDE}, seen) == []
    loose = SLIDE.replace("1.08초", "1.08초 · 28.1초 아낌 · 96.5% · 4,200만 원")
    assert len(ungrounded_write_numbers({"html": loose}, seen)) >= 3


@pytest.mark.skipif(not (shutil.which("chrome") or Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe").exists()
                         or Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe").exists()),
                    reason="굽는 데 쓸 크롬 · 엣지가 없다")
def test_real_bake_keeps_the_background_and_editable_text(tmp_path: Path) -> None:
    """배경이 빠지던 문제(Akxan 변환기)의 재발 방지: 배경은 그림 한 장, 글은 고칠 수 있는 글 상자. 거의 투명한 물결 글자는 그림에 남는다."""
    from pptx import Presentation

    deck = tmp_path / "deck"
    deck.mkdir()
    (deck / "slide_01.html").write_text(slides.clean_slide_html(SLIDE), encoding="utf-8")
    out = slides.render_in_subprocess(deck, tmp_path / "out.pptx")
    # PPTX 옆에 같은 장의 PDF 도 굽는다(2026-10-26 앱 글꼴 · `test_slide_fonts_1026`).
    assert {k: out[k] for k in ("slides", "text_boxes")} == {"slides": 1, "text_boxes": 2}
    assert out.get("pdf") == str(tmp_path / "out.pdf") and "pdf_error" not in out
    slide = Presentation(str(tmp_path / "out.pptx")).slides[0]
    pictures = [s for s in slide.shapes if s.shape_type == 13]
    texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
    assert len(pictures) == 1
    assert any("30.6초" in text for text in texts) and not any("END" in text for text in texts)


def test_jobs_handed_over_by_claude_can_see_the_slide_hand() -> None:
    """바깥 클로드가 맡긴 판은 코드 쪽 목록으로 돈다 — 거기 없으면 새 손이 안 보여 흰 바탕 PPT 로 돌아간다(10-25 첫 판)."""
    from alture_jarvis.code_runtime import CODE_CORE_TOOL_NAMES

    assert "local_make_slide" in CODE_CORE_TOOL_NAMES


def test_the_gateway_hands_the_job_id_to_the_slide_hand(tmp_path: Path) -> None:
    """판 번호가 없으면 장표가 「nojob」으로 저장 · 알림돼 컴퓨터 창이 이 판의 장표를 못 찾는다(10-25 첫 실사용)."""
    from alture_jarvis.tool_gateway import ToolGateway

    tools, _ = _tools(tmp_path)
    gateway = ToolGateway.__new__(ToolGateway)
    gateway.__dict__.update({"_specs": {}, "approval_describers": {}, "pre_approval_checks": {}})
    registered = {}
    gateway.register = lambda spec: registered.__setitem__(spec.name, spec)
    tools.register(gateway)
    assert registered["local_make_slide"].needs_job_context


def test_a_built_deck_counts_as_made_and_checked_for_code_jobs() -> None:
    """맡긴 PPT 판이 코드 산출물로 읽혀도, PPTX 까지 구웠으면 「파일을 고치지 않았다」로 접지 않는다(10-25 둘째 판)."""
    from alture_jarvis.agent import _deck_built

    slide = {"ok": True, "tool": "local_make_slide", "evidence": {"saved": True}}
    built = {"ok": True, "tool": "local_make_slide", "evidence": {"verified": True, "output_path": "C:/x/손자리 소개.pptx"}}
    assert not _deck_built([slide])
    assert _deck_built([slide, built])


def test_outbound_number_check_reads_only_the_visible_words_of_slides() -> None:
    """바깥 클로드에게 돌려주는 숫자 검사도 CSS · 그림 좌표(0.92 · 880)를 세지 않는다(10-25 둘째 판 9개 헛경고)."""
    from alture_jarvis.delegation import _strings

    out: list[str] = []
    _strings({"title": "t", "index": 2, "html": SLIDE}, out)
    joined = " ".join(out)
    assert "30.6" in joined and "1280" not in joined and "rgba" not in joined


def test_the_research_route_lets_the_slide_hand_through() -> None:
    """「근거 읽기」 요구가 붙은 PPT 판은 웹 조사 갈래로 간다 — 거기서 새 손이 막히면 python-pptx 로 돌아간다(10-25 셋째 판)."""
    from alture_jarvis.tool_roster import _in_web_route

    assert _in_web_route("local_make_slide") and _in_web_route("local_write_document")


TELL_SLIDE = """<!DOCTYPE html><html><head><style>
.card{border-radius:14px;background:#f5f7fa;box-shadow:0 2px 8px rgba(0,0,0,.08)}
.title{background:linear-gradient(90deg,#7c3aed,#2563eb);-webkit-background-clip:text;color:transparent}
</style></head><body>
<div class="eyebrow">PART 01 · WHAT</div><h1 class="title">손자리가 뭔가</h1>
<div class="card"><b>01</b> 로컬 AI</div><div class="card"><b>02</b> 실제 크롬</div><div class="card"><b>03</b> 내 파일</div>
<footer>손자리 소개 · 02 / 07 · SONJARI</footer></body></html>"""


def test_the_tell_check_names_what_the_first_real_deck_did() -> None:
    """첫 실사용 「손자리 소개」에 있던 꼴들(대문자 꼬리표 · 01/02/03 · 가운데 점 줄 · 같은 카드)을 기계로 가른다."""
    ids = {tell["id"]: tell for tell in slides.design_tells(TELL_SLIDE)}
    assert {"caps_label", "number_marks", "dot_chain", "gradient_text", "same_cards"} <= set(ids)
    # 돌려보내는 것은 대표가 싫다고 한 「강한 느낌」뿐이다 — 나머지 티는 알리기만(디자인 맛을 죽이지 않게 · 10-25 3차).
    assert {key for key, tell in ids.items() if tell["hard"]} == {"gradient_text"}
    # 숫자 · 고유명사 하나짜리 대문자(RTX 3090)와 본문 속 가운데 점(긴 문장)은 티가 아니다.
    plain = "<p>RTX 3090</p><p>결정 · 할 일 · 담당을 첫 칸에 두고, 나머지 이야기는 그 아래에 둡니다. 다음 회의는 할 일부터.</p>"
    assert slides.design_tells(plain) == []


def test_a_slide_with_tells_is_sent_back_once_then_taken(tmp_path: Path) -> None:
    tools, sent = _tools(tmp_path)
    with pytest.raises(ValueError, match="design_tells"):
        tools.make_slide({"title": "x", "index": 1, "total": 3, "html": TELL_SLIDE, "__job_id": "j"})
    assert not (tmp_path / "out/slides/j/x/slide_01.html").exists() and sent == []
    again = tools.make_slide({"title": "x", "index": 1, "total": 3, "html": TELL_SLIDE, "__job_id": "j"})
    assert again["saved"] and "카드" in again.get("design_note", "") and "꼬리표" in again.get("design_note", "")


def test_money_the_person_gave_is_not_called_unvisited_store_prices() -> None:
    """과제에 직접 준 금액을 옮긴 답에 「스토어에 실제로 가지 않았다」가 붙지 않는다(10-25 PPT 판 셋)."""
    from alture_jarvis.agent import BrowserAgent

    prompt = "손자리 소개 PPT 를 만들어 줘. 비용: 클로드 소네트 5.5 $40.70 · DeepSeek V4 Flash $0.77. 마이크로소프트 스토어에서 받을 수 있다."
    answer = "7장을 만들었습니다. 비용 장: 소네트 $40.70 · DeepSeek $0.77. 마이크로소프트 스토어에서 받기."
    assert BrowserAgent._cited_source_check(prompt, answer, []) == ""
    made_up = answer + " 스토어 가격은 $12.99 입니다."
    assert BrowserAgent._cited_source_check(prompt, made_up, []) != ""


def test_a_cut_off_slide_is_sent_back(tmp_path: Path) -> None:
    """출력이 CSS 중간에서 끊긴 장은 빈 화면으로 그려진다 — 받지 않고 다시 쓰게 한다(10-25 시험대 4~7장)."""
    tools, _ = _tools(tmp_path)
    cut = SLIDE.split("</style>")[0]
    with pytest.raises(ValueError, match="slide_html_truncated"):
        tools.make_slide({"title": "x", "index": 1, "total": 2, "html": cut, "__job_id": "j"})


COVER = """<section class="cover"><p class="kicker">처음 쓰는 분께</p><h1>내 PC에서 일하는 AI</h1>
<p class="lede">본문은 이 PC에서 읽는다.</p></section>"""
BIG = """<section class="bignum"><p class="kicker">첫 걸음</p><h2>첫 걸음이 빨라졌다</h2>
<div class="hero"><p class="was">30.6초</p><p class="num">1.08<small>초</small></p><p class="desc">한 판 기준.</p></div></section>"""


@pytest.mark.usefixtures("html_path")
def test_a_section_fragment_is_dressed_in_the_chosen_theme(tmp_path: Path) -> None:
    """모델은 <section> 조각만 넣는다 — 결 CSS · 쪽 표시 · 색 면은 손자리가 끼운다. 결은 처음 고른 것을 기억한다."""
    tools, sent = _tools(tmp_path)
    tools.make_slide({"title": "소개", "theme": "forest", "__job_id": "j"})
    first = tools.make_slide({"title": "소개", "index": 1, "total": 3, "html": COVER, "__job_id": "j"})
    assert first["layout"] == "cover" and first["theme"] == "forest"
    saved = (tmp_path / "out/slides/j/소개/slide_01.html").read_text(encoding="utf-8")
    assert saved.lower().startswith("<!doctype html>") and 'data-sonjari-theme="forest"' in saved
    assert '<section class="cover field">' in saved  # forest 는 표지가 색 면
    second = _put(tools, {"title": "소개", "index": 2, "total": 3, "html": BIG.replace('class="bignum"', 'class="bignum alt"'),
                               "__job_id": "j"})
    page = (tmp_path / "out/slides/j/소개/slide_02.html").read_text(encoding="utf-8")
    assert '<section class="bignum">' in page  # forest 의 bignum 은 면인데 alt 로 뒤집었다
    assert '<div class="folio"><span>소개</span><span>2 / 3</span></div>' in page
    assert sent[-1][0] == "deck_slide" and "data-sonjari-theme" in sent[-1][1]["html"]
    assert "design_note" not in second


@pytest.mark.usefixtures("html_path")
def test_an_unknown_layout_is_named_back(tmp_path: Path) -> None:
    tools, _ = _tools(tmp_path)
    with pytest.raises(ValueError, match="slide_layout_unknown"):
        tools.make_slide({"title": "x", "index": 1, "total": 2, "html": '<section class="hero-card"><h2>x</h2></section>', "__job_id": "j"})
    with pytest.raises(ValueError, match="slide_html_truncated"):
        tools.make_slide({"title": "x", "index": 1, "total": 2, "html": '<section class="cover"><h1>잘린', "__job_id": "j"})


@pytest.mark.usefixtures("html_path")
def test_neighbours_with_the_same_layout_get_a_note(tmp_path: Path) -> None:
    tools, _ = _tools(tmp_path)
    tools.make_slide({"title": "x", "index": 1, "total": 3, "html": BIG, "__job_id": "j"})
    again = _put(tools, {"title": "x", "index": 2, "total": 3, "html": BIG, "__job_id": "j"})
    assert "같은 짜임(bignum)" in again["design_note"]


@pytest.mark.usefixtures("html_path")
def test_overflow_that_shrinking_cannot_fix_is_sent_back_once(tmp_path: Path, monkeypatch) -> None:
    tools, sent = _tools(tmp_path)
    monkeypatch.setattr(slides, "fit_in_subprocess", lambda path, **_: {"overflow": 120, "where": "div.grid", "base_px": 14})
    with pytest.raises(ValueError, match="slide_overflow: 1번 장.*120px"):
        tools.make_slide({"title": "x", "index": 1, "total": 2, "html": BIG, "__job_id": "j"})
    assert not (tmp_path / "out/slides/j/x/slide_01.html").exists() and sent == []
    taken = tools.make_slide({"title": "x", "index": 1, "total": 2, "html": BIG, "__job_id": "j"})
    assert taken["saved"] and "14px 로 줄여" in taken["design_note"]


def test_real_fit_shrinks_a_crowded_slide_and_bakes_the_size(tmp_path: Path) -> None:
    """진짜 브라우저: 글이 넘치는 칸 묶음은 장 기본 글자(20px)를 줄여 맞추고, 그 크기를 파일에 박는다(미리보기 = PPTX)."""
    from alture_jarvis import slides_style

    crowded = ('<section class="tiles"><p class="kicker">묶음</p><h2>칸이 많은 장</h2><div class="grid">'
               + '<div class="tile"><b>첫 칸</b><span>' + "아주 긴 설명 문장이 계속 이어진다. " * 40 + "</span></div>"
               + '<div class="tile"><b>둘</b><span>짧게</span></div><div class="tile"><b>셋</b><span>짧게</span></div></div></section>')
    html, _ = slides_style.wrap_fragment(crowded, "paper", title="t", index=2, total=3)
    page = tmp_path / "slide_02.html"
    page.write_text(html, encoding="utf-8")
    shape = slides.fit_in_subprocess(page)
    assert shape["overflow_before"] > 0 and shape["base_px"] < 20
    assert "data-sonjari-fit" in page.read_text(encoding="utf-8")


def test_font_pair_splits_latin_and_korean_like_the_browser() -> None:
    assert slides.font_pair("Georgia, 'Batang', serif") == ("Georgia", "Batang")
    assert slides.font_pair("'Malgun Gothic', sans-serif") == ("Malgun Gothic", "Malgun Gothic")
    assert slides.font_pair("Bahnschrift, 'Segoe UI', 'Malgun Gothic', sans-serif") == ("Bahnschrift", "Malgun Gothic")
    assert slides.font_pair("") == ("Malgun Gothic", "Malgun Gothic")


def test_real_bake_keeps_the_browsers_line_breaks(tmp_path: Path) -> None:
    """★ 두 줄 제목이 PPTX 에서 한 줄로 이어져 잘리던 것(10-25 파워포인트로 열어 봄): 줄은 브라우저가 바꾼 자리 그대로, 자동 줄바꿈은 끈다.
    작은 글자(1.08<small>초</small>)는 같은 줄로 남는다."""
    from pptx import Presentation
    from alture_jarvis import slides_style

    deck = tmp_path / "deck"
    deck.mkdir()
    cover = COVER.replace("내 PC에서 일하는 AI", "내 PC에서 일하는 AI, 손자리를 처음 쓰는 분께 드리는 안내")
    for i, frag in enumerate((cover, BIG), 1):
        html, _ = slides_style.wrap_fragment(frag, "forest", title="t", index=i, total=2)
        (deck / f"slide_{i:02d}.html").write_text(html, encoding="utf-8")
    slides.render_in_subprocess(deck, tmp_path / "out.pptx")
    prs = Presentation(str(tmp_path / "out.pptx"))
    frames = {s.text_frame.text: s.text_frame for s in prs.slides[0].shapes if s.has_text_frame}
    title = next(frame for text, frame in frames.items() if text.startswith("내 PC"))
    assert "\v" in title.text or len(title._txBody.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}br")) >= 1
    assert title.word_wrap is False
    latin = title._txBody.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}latin").get("typeface")
    east = title._txBody.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}ea").get("typeface")
    # forest 의 제목 글꼴(고운바탕)은 PPTX 에 심으므로 그 이름 그대로(10-26 저녁 · pptx_fonts). 못 심는 PC 면 옛 윈도우 글꼴.
    from alture_jarvis import pptx_fonts
    assert (latin, east) == (("Georgia", "Batang") if pptx_fonts.unavailable() else ("Gowun Batang", "Gowun Batang"))
    number = next(s.text_frame for s in prs.slides[1].shapes if s.has_text_frame and s.text_frame.text.startswith("1.08"))
    assert number.text == "1.08초"


@pytest.mark.usefixtures("html_path")
def test_two_light_slides_in_a_row_turn_the_next_big_slide_into_a_colour_field(tmp_path: Path) -> None:
    """리듬(B 조사): 흰 장이 둘 이어지면 다음 큰 숫자 장은 색 면으로. 마지막 장 · alt 를 단 장은 건드리지 않는다."""
    tools, _ = _tools(tmp_path)
    steps = '<section class="steps"><h2>순서</h2><ol><li><b>하나</b><span>설명</span></li></ol></section>'
    compare = '<section class="compare"><h2>비교</h2><div class="cols"><div class="col"><h3>가</h3></div><div class="col on"><h3>나</h3></div></div></section>'
    for i, frag in enumerate((COVER, steps, compare, BIG), 1):
        _put(tools, {"title": "r", "index": i, "total": 6, "html": frag, "theme": "paper", "__job_id": "j"})
    page = (tmp_path / "out/slides/j/r/slide_04.html").read_text(encoding="utf-8")
    assert '<section class="bignum field">' in page


@pytest.mark.usefixtures("html_path")
def test_a_doubly_escaped_fragment_is_unwrapped(tmp_path: Path) -> None:
    """10-25 시험대 N2: 모델이 html 을 한 번 더 문자열로 감싸 class=\\"cover\\" · 글자 그대로의 \n 이 들어왔다 — 배치를 못 읽어 맨글자 장."""
    tools, _ = _tools(tmp_path)
    raw = COVER.replace('"', '\\"').replace("\n", "\\n")
    out = tools.make_slide({"title": "e", "index": 1, "total": 2, "html": raw, "__job_id": "j"})
    assert out["layout"] == "cover"
    page = (tmp_path / "out/slides/j/e/slide_01.html").read_text(encoding="utf-8")
    assert "\\n" not in page.split("<body>")[1] and '\\"' not in page


def test_a_slide_list_table_does_not_bring_back_the_unvisited_store_warning() -> None:
    """4차 실제 판 2b81aa61: 답에 장 목록 표(7줄)가 있어 「표 셋 줄 미만」 조건을 빠져나가 같은 헛경고가 붙었다."""
    from alture_jarvis.agent import BrowserAgent

    prompt = "손자리 소개 PPT 를 만들어 줘. 비용: 클로드 소네트 5.5 $40.70 · Qwen3.8 27B $10.70. 마이크로소프트 스토어에서 받을 수 있다."
    rows = "\n".join(f"| {i} | 장 {i} | cover |" for i in range(1, 8))
    answer = "7장을 만들었습니다.\n| 장 | 내용 | 배치 |\n|---|---|---|\n" + rows + "\n비용 장: $40.70 · $10.70."
    assert BrowserAgent._cited_source_check(prompt, answer, []) == ""


def test_supplied_facts_with_no_address_do_not_ask_for_reading() -> None:
    """「아래 사실만 쓰고 다른 숫자는 지어내지 마」 + 사실 목록 → 읽을 대상이 없다. 실제 주소가 있으면 예전대로 읽기를 요구한다."""
    from alture_jarvis import plan

    facts = ("손자리 소개 장표(PPTX)를 7장으로 만들어 줘. 아래 사실만 쓰고 다른 숫자는 지어내지 마. "
             "- 클로드가 맡기면 손자리가 이 PC 에서 실제 크롬과 내 파일을 읽고 근거만 돌려준다. - 마이크로소프트 스토어에서 받을 수 있다.")
    kinds = [item.kind for item in plan.extract(facts, has_named_source=True, must_observe=True)]
    assert "observe" not in kinds and "artifact" in kinds
    research = "reddit.com 의 r/ollama 에서 글을 찾아 아래 내용으로 엑셀 만들어. 다른 숫자는 지어내지 마."
    assert any(item.kind in {"observe", "count"} for item in plan.extract(research, has_named_source=True, must_observe=True))


# ── 6차: 옮긴 원본 템플릿 · 추천 · 큐레이션 화면 ───────────────────────────────────────────────


def test_eighteen_ported_templates_ship_with_their_own_layouts_and_license() -> None:
    """원본 템플릿 18개(MIT) — 결마다 자기 CSS 와 장 틀, 고지 파일. 5차의 「색만 옮긴 35개」는 걷어 냈다(대표 「템플릿이 별로」)."""
    from alture_jarvis import slides_style

    # 옮긴 것 18개 + 우리가 쓴 것(own · IR 브리프 등)은 따로 센다.
    assert len([k for k, m in slides_style.TEMPLATES.items() if not m.get("own")]) == 18
    assert (slides_style.TEMPLATE_DIR / "LICENSE").read_text(encoding="utf-8").startswith("MIT License")
    for key, meta in slides_style.TEMPLATES.items():
        _, css = slides_style.theme_css(key)
        assert css.startswith("/* sonjari-template */") and "fonts.googleapis" not in css and "1920px" not in css, key
        sections = slides_style.template_sections(key)
        assert len(sections) >= 8 and [sec["class"] for sec in sections] == meta["layouts"], key
    assert not ({"forest", "emerald", "cobalt"} & {line.split(":")[0][2:] for line in slides_style.theme_menu().splitlines()})


def test_no_template_borrows_a_real_company_name() -> None:
    from alture_jarvis import slides_style

    banned = ("mckinsey", "bcg", "bain", "deloitte", "airbnb", "sequoia", "netflix", "맥킨지", "구글", "애플")
    for key, meta in slides_style.TEMPLATES.items():
        text = f"{key} {meta['name']} {meta['when']}".lower()
        assert not any(word in text for word in banned), key


def test_recommend_reads_the_occasion_and_mood() -> None:
    from alture_jarvis.slide_catalog import recommend

    assert recommend("동네 카페 3분기 매출 보고")[0]["theme"] in {"editorial-forest", "emerald-editorial"}
    assert recommend("대학원 학위 논문 심사 발표")[0]["theme"] == "cobalt-grid"
    assert recommend("신입사원 온보딩 교육 자료")[0]["theme"] == "daisy-days"
    assert recommend("개발팀 서버 성능 보고 어둡게")[0]["theme"] == "ink"
    assert recommend("아이디어 워크숍")[0]["theme"] == "scatterbrain"
    plain = recommend("")
    assert plain[0]["theme"] == "paper" and "단서" in str(plain[0]["why"])


def test_the_first_call_shows_three_title_covers_for_the_suggested_styles(tmp_path: Path) -> None:
    tools, sent = _tools(tmp_path)
    out = tools.make_slide({"title": "카페 3분기 보고", "about": "동네 카페 분기 매출 보고", "__job_id": "j"})
    assert out["theme"] in {"editorial-forest", "emerald-editorial"} and len(out["recommended"]) == 3
    assert "<section" not in out["layouts"] and "frames" in out
    event, payload = next(item for item in sent if item[0] == "deck_themes")
    assert [p["theme"] for p in payload["picks"]] == [p["theme"] for p in out["recommended"]]
    assert all("카페 3분기 보고" in p["html"] and "data-sonjari-theme" in p["html"] for p in payload["picks"])
    # 미리보기 표지는 실제로 만들 표지 부품이다(10-26) — 템플릿 표지에 제목만 넣던 것과 실제 장이 달랐다.
    assert all("pz-frame" in p["html"] and "fr-title" in p["html"] for p in payload["picks"])


def _forest_cover(text: str) -> str:
    from alture_jarvis import slides_style

    cover = slides_style.template_sections("editorial-forest")[0]["html"]
    import re as _re

    cover = _re.sub(r"(<h1\b[^>]*>).*?(</h1>)", lambda m: m.group(1) + text + m.group(2), cover, flags=_re.S)
    cover = _re.sub(r'(<span class="label">).*?(</span>)', lambda m: m.group(1) + "분기 보고" + m.group(2), cover, flags=_re.S)
    return _re.sub(r"(<div class=\"footline\">).*?(</div>)", lambda m: m.group(1) + "<span>동네 카페</span><span>운영팀</span>" + m.group(2),
                   cover, flags=_re.S)


@pytest.mark.usefixtures("html_path")
def test_a_template_slide_keeps_the_template_and_is_shown_as_the_visible_slide(tmp_path: Path) -> None:
    tools, sent = _tools(tmp_path)
    tools.make_slide({"title": "t", "theme": "editorial-forest", "__job_id": "j"})
    out = tools.make_slide({"title": "t", "index": 1, "total": 3, "html": _forest_cover("동네 카페 3분기"), "__job_id": "j"})
    assert out["layout"] == "cover" and out["theme"] == "editorial-forest"
    page = (tmp_path / "out/slides/j/t/slide_01.html").read_text(encoding="utf-8")
    assert "/* sonjari-template */" in page and "active is-active" in page and 'class="folio"' not in page
    with pytest.raises(ValueError, match="slide_layout_unknown.*agenda"):
        tools.make_slide({"title": "t", "index": 2, "total": 3, "html": '<section class="bignum"><h2>x</h2></section>', "__job_id": "j"})


@pytest.mark.usefixtures("html_path")
def test_a_slide_that_keeps_the_template_words_is_sent_back_once(tmp_path: Path) -> None:
    """시험대 T1: 4번 장을 원본 영어(Fewer bets, stronger commitments.) 그대로 베꼈다."""
    from alture_jarvis import slides_style

    tools, _ = _tools(tmp_path)
    tools.make_slide({"title": "t", "theme": "editorial-forest", "__job_id": "j"})
    copied = slides_style.template_sections("editorial-forest")[3]["html"]
    with pytest.raises(ValueError, match="template_text_left: 2번 장"):
        _put(tools, {"title": "t", "index": 2, "total": 3, "html": copied, "__job_id": "j"})
    again = _put(tools, {"title": "t", "index": 2, "total": 3, "html": copied, "__job_id": "j"})
    assert again["saved"]


def test_real_fit_squeezes_a_long_korean_title_in_a_template(tmp_path: Path) -> None:
    """진짜 브라우저: 옮긴 템플릿은 글자가 px 라 장 기본 글자로 안 준다 — 큰 글자를 줄여 문서째 저장한다(T1 8번 장이 넘쳤다)."""
    from alture_jarvis import slides_style

    long_title = "동네 카페의 3분기를 숫자와 손님 이야기로 돌아보고 4분기에 무엇을 할지 정리한 보고"
    html, _ = slides_style.wrap_fragment(_forest_cover(long_title), "editorial-forest", title="t", index=1, total=3)
    page = tmp_path / "slide_01.html"
    page.write_text(html, encoding="utf-8")
    shape = slides.fit_in_subprocess(page)
    assert shape["overflow_before"] > 0 and shape["overflow"] == 0 and shape["squeezed"] > 0
    saved = page.read_text(encoding="utf-8")
    assert "font-size" in saved.split("<body", 1)[1] and "data-sonjari-static" not in saved
