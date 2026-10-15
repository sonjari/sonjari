"""문서가 **문서답게** 나온다 — 양식 검사기가 첫 판에 잡은 셋.

## 왜 (2026-09-13)

벤치 창을 끄고 대신 코드로 내용·양식을 보게 하자(사용자님 요청) 검사기가 처음 돌자마자
만든 파일 열두 개 중 넷을 걸었다. 셋이 진짜였다:

    1. 워드에 「1. **Python 중심의 에이전트 생태계:**」 — 별표가 글자 그대로
    2. 사람인 문서 첫 줄이 「<thought」, 둘째 줄이 「>사람인에서 …」
    3. 이력서를 시켰는데 파일 속이 「이력서 초안을 작성하겠습니다」 — 계획문

파일은 나왔으니 전부 성공으로 셀 뻔했다. 사람이 열면 셋 다 못 쓰는 문서다.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from alture_jarvis.agent import (
    _artifact_still_owed,
    _forced_artifact_decision,
)
from alture_jarvis.documents import write_document
from alture_jarvis.local_computer import _unused_sibling
from alture_jarvis.plan import check as check_plan, extract as extract_plan


def _docx(tmp_path: Path, blocks: list[dict]) -> Document:
    write_document(tmp_path / "t.docx", "docx", {"blocks": blocks})
    return Document(str(tmp_path / "t.docx"))


# ── 1. 굵게는 기호가 아니라 서식 ────────────────────────────────────────────
def test_bold_markers_become_real_bold_in_word(tmp_path: Path) -> None:
    doc = _docx(tmp_path, [
        {"type": "paragraph", "text": "1. **Python 중심:** 수요가 많다"},
        {"type": "bullet", "text": "- **경력직:** 3년 이상"},
    ])
    texts = [p.text for p in doc.paragraphs]
    assert all("**" not in t for t in texts)
    bold = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert "Python 중심:" in bold and "경력직:" in bold
    # 불릿은 「- 」를 이미 붙여 온다. 워드 불릿 위에 또 「-」가 서면 안 된다.
    assert texts[1] == "경력직: 3년 이상"


def test_table_cells_and_headings_carry_no_markers(tmp_path: Path) -> None:
    doc = _docx(tmp_path, [
        {"type": "heading", "level": 1, "text": "# **여수 회계** 공고"},
        {"type": "table", "rows": [["회사", "형태"], ["**비케이**", "정규직"]]},
    ])
    assert doc.paragraphs[0].text == "여수 회계 공고"
    assert doc.tables[0].rows[1].cells[0].text == "비케이"


def test_a_markdown_rule_is_not_printed_as_dashes(tmp_path: Path) -> None:
    doc = _docx(tmp_path, [
        {"type": "paragraph", "text": "위"},
        {"type": "paragraph", "text": "---"},
        {"type": "paragraph", "text": "아래"},
    ])
    assert [p.text for p in doc.paragraphs] == ["위", "아래"]


# ── 2. 속생각 태그가 본문에 새지 않는다 ──────────────────────────────────────
def test_a_thought_tag_split_over_two_blocks_is_removed(tmp_path: Path) -> None:
    doc = _docx(tmp_path, [
        {"type": "paragraph", "text": "<thought"},
        {"type": "paragraph", "text": ">사람인에서 공고 3개를 정리했다."},
        {"type": "heading", "level": 1, "text": "여수 회계 공고"},
    ])
    assert [p.text for p in doc.paragraphs] == ["사람인에서 공고 3개를 정리했다.", "여수 회계 공고"]


def test_a_real_quote_without_any_tag_is_left_alone(tmp_path: Path) -> None:
    """태그가 없으면 한 글자도 안 바꾼다 — 「a > b」 같은 글을 망가뜨리지 않는다."""
    write_document(tmp_path / "q.md", "md", {"blocks": [
        {"type": "paragraph", "text": "매출 > 비용 이면 흑자다"},
    ]})
    assert "매출 > 비용 이면 흑자다" in (tmp_path / "q.md").read_text(encoding="utf-8")


# ── 3. 계획문은 문서가 아니다 ────────────────────────────────────────────────
RESUME_PROMPT = "사람인에서 여수 회계 채용 공고 2개를 보고, 거기 지원할 이력서 초안을 문서로 만들어줘."

PLAN_ANSWER = """사람인에서 검색된 여수 지역 회계 채용 공고 2곳의 요건을 분석했습니다.

### 1. 분석한 채용 공고 요건

- **경력직 지원 시:** 경력 3년 이상 요건에 맞춰 과거 경력을 수치 중심으로 기술합니다.
- **신입 지원 시:** 자격증과 꼼꼼한 문서 관리 능력을 중심으로 작성합니다.

이 내용을 바탕으로 두 곳 모두에 활용할 수 있는 이력서 초안을 작성하겠습니다.
"""

REAL_RESUME = """# 이력서 — 회계 사무직 지원

## 자기소개
전산회계 1급을 보유하고 3년간 거래처 장부와 월 결산을 맡았습니다.
입사 후에는 정확한 결산으로 신뢰받는 직원이 되도록 성실히 근무하겠습니다.

## 경력
- 2022~2025 ○○상사 회계팀: 월 결산, 부가세 신고 보조, 거래처 원장 관리
- 2020~2022 △△세무회계사무소: 기장 대리, 급여 대장 작성, 원천세 신고

## 자격
- 전산회계 1급 (2020) · 전산세무 2급 (2021) · 컴퓨터활용능력 2급

## 지원 동기
공고에서 요구한 경력 3년과 전산세무 자격을 모두 갖췄고, 여수 지역에서 오래 일할 수 있습니다.
"""


def _resume_results() -> list:
    plan = extract_plan(RESUME_PROMPT, has_named_source=True)
    urls = ["https://www.saramin.co.kr/zf_user/jobs/1", "https://www.saramin.co.kr/zf_user/jobs/2"]
    steps = [{"tool": "browser_extract", "ok": True, "evidence": {"url": u}} for u in urls]
    return check_plan(plan, sources=urls, steps=steps), steps


def test_a_promise_to_write_the_document_is_not_written_as_the_document() -> None:
    results, steps = _resume_results()
    call, reason = _forced_artifact_decision(results, steps, PLAN_ANSWER, RESUME_PROMPT)
    assert call is None
    assert reason == "answer_is_a_plan"


def test_a_resume_that_pledges_to_work_hard_is_still_a_resume() -> None:
    """「성실히 근무하겠습니다」는 다짐이지 예고가 아니다 — 거르면 진짜 이력서를 막는다."""
    results, steps = _resume_results()
    call, reason = _forced_artifact_decision(results, steps, REAL_RESUME, RESUME_PROMPT)
    assert reason == ""
    assert call is not None


# ── 잠긴 파일을 비켜 갈 이름 ────────────────────────────────────────────────
def test_the_spare_name_is_never_one_that_already_exists(tmp_path: Path) -> None:
    """상한 9 가 박혀 있어서, (2)~(9) 가 다 차면 열려 있던 (9) 에 그대로 썼다."""
    base = tmp_path / "이력서.docx"
    base.write_bytes(b"x")
    for n in range(2, 13):
        (tmp_path / f"이력서 ({n}).docx").write_bytes(b"x")
    spare = _unused_sibling(base)
    assert not spare.exists()
    assert spare.name == "이력서 (13).docx"


# ── 접기 전에 산출물 차례를 줄지 ──────────────────────────────────────────────
def test_an_owed_artifact_gets_a_turn_only_if_nobody_tried_to_write() -> None:
    results, steps = _resume_results()
    assert _artifact_still_owed(results, steps) is True
    tried = steps + [{"tool": "local_write_document", "ok": False, "evidence": {}}]
    assert _artifact_still_owed(results, tried) is False


# ── 안전망이 만든 문서에도 제목이 선다 ──────────────────────────────────────
def test_a_fallback_document_starts_with_a_real_title() -> None:
    """안전망이 만든 파일 넷이 `정리.docx`·`1_채용_공고_요약.docx` 로 제목 없이 나왔다."""
    from alture_jarvis.agent import _with_document_title

    numbered = [{"type": "heading", "level": 3, "text": "1. 채용 공고 요약"},
                {"type": "paragraph", "text": "본문"}]
    titled = _with_document_title(numbered, "")
    assert titled[0] == {"type": "heading", "level": 1, "text": "채용 공고 요약"}

    intro_first = [{"type": "paragraph", "text": "조사한 결과입니다."},
                   {"type": "heading", "level": 1, "text": "YES24 파이썬 도서"},
                   {"type": "table", "rows": [["a"], ["b"]]}]
    moved = _with_document_title(intro_first, "")
    assert moved[0]["text"] == "YES24 파이썬 도서" and len(moved) == 3

    no_heading = [{"type": "table", "rows": [["회사", "형태"], ["가나", "정규직"]]}]
    made = _with_document_title(no_heading, "여수 회계 공고 3건을 정리했습니다. 표를 보세요.")
    assert made[0]["type"] == "heading" and made[0]["level"] == 1
    assert made[0]["text"].startswith("여수 회계 공고 3건")
