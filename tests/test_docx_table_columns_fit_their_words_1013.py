"""워드 표의 칸 너비 — 긴 주소 한 칸이 폭을 다 차지해 나머지가 한두 글자씩 꺾이지 않게 한다.

★ 2026-10-13 영상 촬영 판: 「이걸 보고서로 만들어 줘」로 만든 워드 보고서의 표가 다섯 칸(가게 · 상품 이름 · 판매 가격 · 배송비 ·
  읽은 페이지 주소)이었는데, 칸 너비를 정하지 않아 워드의 자동 맞춤이 끊을 수 없는 주소에 폭을 몰아 주었다.
  「상품 이름」은 두세 글자씩, 「101,890원」은 「101,8 / 90 원」으로 꺾였다. 값은 맞는데 읽을 수가 없는 표였다.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from alture_jarvis.documents import write_document

ROWS = [
    ["가게", "상품 이름", "판매 가격", "배송비", "읽은 페이지 주소"],
    ["가람상회 (온새몰, 해외직구)", "솔빛 무선 블루투스 마우스 M3", "101,890원", "무료배송 (제주/도서산간 6,000원)",
     "https://item.example.co.kr/Item?goodscode=4389353840&utparam=abcdefghijklmnop"],
    ["한울마트", "[정품] 솔빛 마우스 M3 블랙 수신기 포함", "139,700원", "무료배송 (제주/도서산간 5,000원)",
     "https://www.example.co.kr/products/8852665861"],
]


def _table(tmp_path: Path):
    path = tmp_path / "보고서.docx"
    write_document(path, "docx", {"blocks": [{"type": "heading", "text": "가격 비교 보고서"}, {"type": "table", "rows": ROWS}]})
    document = Document(str(path))
    section = document.sections[0]
    return document.tables[0], section.page_width - section.left_margin - section.right_margin


def _widths(table) -> list[int]:
    return [int(cell.width) for cell in table.rows[1].cells]


def test_the_table_has_a_fixed_layout_that_fills_the_text_width(tmp_path: Path) -> None:
    table, available = _table(tmp_path)

    layout = table._tbl.tblPr.find(qn("w:tblLayout"))
    assert layout is not None and layout.get(qn("w:type")) == "fixed", "자동 맞춤에 맡기면 주소 칸이 폭을 다 가져간다"
    assert abs(sum(_widths(table)) - available) < available * 0.02


def test_no_single_column_takes_most_of_the_width(tmp_path: Path) -> None:
    table, available = _table(tmp_path)
    widths = _widths(table)

    assert max(widths) <= available * 0.36, "주소 칸은 정해 준 폭 안에서 줄을 바꾼다"
    assert min(widths) >= available * 0.11, "어느 칸도 한두 글자 폭으로 찌그러지지 않는다"


def test_a_value_is_not_broken_in_the_middle(tmp_path: Path) -> None:
    """「101,890원」 한 낱말이 한 줄에 들어갈 폭 — 글자 하나를 약 0.085인치(본문 크기)로 어림한다."""
    table, _ = _table(tmp_path)
    price = _widths(table)[2]

    assert price >= int(914400 * 0.85), "가격 칸이 값 하나보다 좁다"


def test_a_short_table_keeps_its_natural_widths(tmp_path: Path) -> None:
    """두 칸짜리 짧은 표는 폭을 억지로 늘리지 않아도 된다 — 다만 넘치지는 않는다."""
    path = tmp_path / "짧은표.docx"
    write_document(path, "docx", {"blocks": [{"type": "table", "rows": [["항목", "값"], ["기간", "270일"], ["인원", "4명"]]}]})
    document = Document(str(path))
    section = document.sections[0]
    available = section.page_width - section.left_margin - section.right_margin

    assert sum(_widths(document.tables[0])) <= available * 1.02
