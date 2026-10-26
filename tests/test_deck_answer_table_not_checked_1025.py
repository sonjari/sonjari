"""장표 손이 구운 PPTX 에는 「답의 표 값이 파일에 없다」를 걸지 않는다 (2026-10-25 밤 scatterbrain 판 1dbcb16e — 답 표가 장 목록)."""

from pathlib import Path

from alture_jarvis.agent import BrowserAgent

ANSWER = """PPT 6장 완성했습니다.

| 장 | 내용 |
|---|---|
| 1 | 표지 — 봄 신메뉴 알리기 (포스트잇 3개: 9개·3개·3월) |
| 3 | 투표 결과 TOP 3 — 막대 5·4·3표 |
| 5 | 다음 할 일 3건 — 민지(3/2)·준호(3/4)·점장님(3/3) |
"""


def _pptx(tmp_path: Path, text: str) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text_frame.text = text
    path = tmp_path / "회의 정리.pptx"
    prs.save(str(path))
    return path


def test_a_baked_deck_is_not_held_to_the_answer_table(tmp_path: Path) -> None:
    path = _pptx(tmp_path, "투표 결과 _ 3개 선정")
    steps = [{"tool": "local_make_slide", "ok": True, "evidence": {"index": 6, "total": 6, "saved": True, "output_path": str(path)}}]
    assert BrowserAgent._answer_missing_from_artifact(ANSWER, steps) == ""


def test_the_same_file_written_as_a_document_is_still_checked(tmp_path: Path) -> None:
    """같은 파일을 문서 쓰기 손으로 만들었으면 예전처럼 본다 — 이 규칙은 장표 손이 구운 파일에만."""
    path = _pptx(tmp_path, "아무 관계 없는 글")
    steps = [{"tool": "local_write_document", "ok": True, "evidence": {"output_path": str(path)}}]
    assert "답의 표 값이 하나도 없습니다" in BrowserAgent._answer_missing_from_artifact(ANSWER, steps)
