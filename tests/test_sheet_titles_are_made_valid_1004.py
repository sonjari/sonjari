"""시트 이름에 엑셀이 못 받는 글자가 있어도 쓰기는 된다 (2026-10-04 허깅페이스 판 3ad09518).

시트 이름의 `/` 하나로 쓰기가 통째로 튕겼다(`Invalid character / found in sheet title`). 모델은 이름을 바꿔 다섯 줄을 한 번 더 썼다.
글자 하나 바꾸면 되는 일은 도구가 한다.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from alture_jarvis.documents import _sheet_titles, write_document

ROWS = [["모델", "수치"], ["A", "74.0"]]


def test_a_sheet_named_with_a_slash_is_written(tmp_path: Path) -> None:
    target = tmp_path / "models.xlsx"

    write_document(target, "xlsx", {"sheets": {"벤치마크/출처": ROWS}})

    book = load_workbook(target)
    assert book.sheetnames == ["벤치마크 출처"]
    assert [cell.value for cell in book["벤치마크 출처"][1]] == ["모델", "수치"]


def test_every_character_excel_refuses_is_replaced() -> None:
    assert _sheet_titles(["a\\b/c*d?e:f[g]h"]) == ["a b c d e f g h"]


def test_long_names_are_cut_and_do_not_overwrite_each_other(tmp_path: Path) -> None:
    """31자로 자르다 같은 이름이 되면 뒤 시트가 앞 시트를 덮는다 — 번호를 붙여 둘 다 남긴다."""
    long_a = "허깅페이스에서 찾은 벤치마크 근거가 있는 모델 목록 — 첫째 묶음"
    long_b = "허깅페이스에서 찾은 벤치마크 근거가 있는 모델 목록 — 둘째 묶음"
    titles = _sheet_titles([long_a, long_b])
    target = tmp_path / "two.xlsx"

    write_document(target, "xlsx", {"sheets": {long_a: ROWS, long_b: [["모델", "수치"], ["B", "16.0"]]}})

    assert len(set(titles)) == 2 and all(len(title) <= 31 for title in titles)
    book = load_workbook(target)
    assert len(book.sheetnames) == 2
    assert book[book.sheetnames[0]]["A2"].value == "A" and book[book.sheetnames[1]]["A2"].value == "B"


def test_an_empty_name_gets_a_default_and_a_plain_name_is_untouched() -> None:
    assert _sheet_titles(["", "모델"]) == ["Sheet1", "모델"]
