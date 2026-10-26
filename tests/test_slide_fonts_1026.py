"""장표 결의 앱 글꼴 (2026-10-26) — 파일 · 라이선스 · CSS · 실제로 그리는 글꼴 · PPTX 글꼴 이름 · PDF.

결 다섯(paper · cobalt · emerald · forest · ink)이 윈도우 기본 글꼴 대신 앱에 든 오픈 글꼴(전부 SIL OFL 1.1)로 그린다.
- 파일은 만든 곳이 배포한 그대로(`slide_themes/fonts/fonts.json` 의 해시) · 옆에 OFL.txt(저작권 줄 포함).
- 장은 주소 없는 문서라 글꼴 주소를 아무 데도 없는 주소로 적고, 그리는 브라우저가 그 요청만 이 폴더에서 돌려준다.
- PPTX 에는 쓴 글자만 담은 글꼴을 심고(`pptx_fonts` · 10-26 저녁) 글 상자에 그 이름을 그대로 적는다. 심을 수 없으면
  예전처럼 윈도우 기본 글꼴을 적고 글자를 그 폭에 맞춰 줄인다. PDF 도 같이 굽는다.
"""

from __future__ import annotations

import hashlib
import re
import shutil
from pathlib import Path

import pytest

from alture_jarvis import slides, slides_style

OWN = ("paper", "cobalt", "emerald", "forest", "ink")
#: 결마다 제목 · 본문 글꼴(조사 문서의 「결마다 짝」 · generated_reports 쪽 기록).
PAIRS = {"paper": ("Pretendard", "Pretendard"), "cobalt": ("IBM Plex Sans KR", "Pretendard"),
         "emerald": ("Paperlogy", "Pretendard"), "forest": ("Gowun Batang", "Gowun Dodum"), "ink": ("SUIT", "Pretendard")}
#: 옛 결이 쓰던 윈도우 글꼴 — 새 글꼴 뒤에 남아 있어야 한다(글꼴을 못 부르면 옛 모양으로).
OLD_FALLBACK = {"paper": "'Segoe UI'", "cobalt": "Georgia", "emerald": "Georgia", "forest": "Georgia", "ink": "'Segoe UI'"}

HAS_BROWSER = bool(shutil.which("chrome") or Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe").exists()
                   or Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe").exists())
needs_browser = pytest.mark.skipif(not HAS_BROWSER, reason="그리는 데 쓸 크롬 · 엣지가 없다")


def _var(css: str, name: str) -> str:
    match = re.search(rf"{name}:\s*([^;]+);", css)
    assert match, name
    return match.group(1)


def test_every_font_file_is_the_makers_file_with_its_license_beside_it() -> None:
    fams = slides_style.bundled_fonts()
    assert {fam["family"] for fam in fams} == {"Pretendard", "IBM Plex Sans KR", "Paperlogy", "Gowun Dodum", "Gowun Batang", "SUIT"}
    listed = {"fonts.json"}
    for fam in fams:
        license_path = slides_style.FONT_DIR / fam["license"]
        text = license_path.read_text(encoding="utf-8")
        assert "SIL Open Font License, Version 1.1" in text and "Copyright" in text.splitlines()[0], fam["family"]
        assert license_path.name == "OFL.txt"
        listed.add(fam["license"])
        assert 1 <= len(fam["faces"]) <= 2, fam["family"]   # 굵기는 둘까지(보통 · 굵게)
        for face in fam["faces"]:
            path = slides_style.FONT_DIR / face["file"]
            data = path.read_bytes()
            assert len(data) == face["bytes"] and hashlib.sha256(data).hexdigest() == face["sha256"], face["file"]
            assert path.parent == license_path.parent  # 라이선스는 글꼴 바로 옆
            assert data[:4] in (b"wOF2", b"\x00\x01\x00\x00"), face["file"]   # 글꼴 파일만(실행 파일 아님)
            listed.add(face["file"])
    on_disk = {p.relative_to(slides_style.FONT_DIR).as_posix() for p in slides_style.FONT_DIR.rglob("*") if p.is_file()}
    assert on_disk == listed, on_disk ^ listed


def test_each_hand_theme_names_its_bundled_font_first_and_keeps_the_old_windows_font_behind() -> None:
    for key in OWN:
        css = (slides_style.THEME_DIR / f"{key}.css").read_text(encoding="utf-8")
        title, body = PAIRS[key]
        assert _var(css, "--sans").startswith(f"'{body}'") and "'Malgun Gothic'" in _var(css, "--sans"), key
        for name in ("--display", "--quote", "--figure"):
            value = _var(css, name)
            assert value.startswith(f"'{title}'") and OLD_FALLBACK[key] in value, (key, name, value)


def test_theme_css_declares_local_fonts_only_for_hand_themes() -> None:
    for key in OWN:
        _, css = slides_style.theme_css(key)
        faces = re.findall(r"@font-face\{font-family:'([^']+)';src:url\('([^']+)'\)", css)
        assert {fam for fam, _ in faces} >= set(PAIRS[key]), key
        for _, url in faces:
            assert url.startswith(slides_style.FONT_ORIGIN + "/") and slides_style.font_file(url) is not None, url
        assert "http://" not in css and "fonts.googleapis" not in css
    for key, meta in slides_style.TEMPLATES.items():   # 옮긴 템플릿 · IR 은 그대로
        assert slides_style.FONT_ORIGIN not in slides_style.theme_css(key)[1], key


def test_the_font_server_only_hands_out_listed_files() -> None:
    origin = slides_style.FONT_ORIGIN
    assert slides_style.font_file(f"{origin}/pretendard/Pretendard-Regular.woff2").name == "Pretendard-Regular.woff2"
    assert slides_style.font_file(f"{origin}/pretendard/Pretendard-Regular.woff2?v=1") is not None
    for bad in (f"{origin}/../slides.py", f"{origin}/pretendard/OFL.txt", f"{origin}/fonts.json",
                f"{origin}/pretendard/..%2F..%2Fslides.py", "https://example.com/pretendard/Pretendard-Regular.woff2",
                f"{origin}/pretendard/Pretendard-Black.woff2"):
        assert slides_style.font_file(bad) is None, bad


def test_pptx_names_the_bundled_font_only_when_it_was_embedded() -> None:
    """심은 가족만 그 이름을 라틴 · 한글 둘 다에. 못 심었으면 목록 뒤의 윈도우 글꼴(옛 결과 같은 이름)."""
    assert slides.font_pair("'Pretendard', 'Malgun Gothic', sans-serif", {"pretendard"}) == ("Pretendard", "Pretendard")
    assert slides.font_pair("'SUIT', Bahnschrift, 'Malgun Gothic', sans-serif", {"pretendard"}) == ("Bahnschrift", "Malgun Gothic")
    assert slides.font_pair("Georgia, 'Batang', serif", {"pretendard"}) == ("Georgia", "Batang")


def test_pptx_text_never_names_a_font_the_reader_does_not_have() -> None:
    """심지 않았을 때: 앱 글꼴은 윈도우에 깔리지 않는다 — PPTX 는 목록 뒤의 윈도우 글꼴(옛 결과 같은 이름)을 적는다."""
    assert slides.font_pair("'Pretendard', 'Malgun Gothic', sans-serif") == ("Malgun Gothic", "Malgun Gothic")
    assert slides.font_pair("'Pretendard', 'Segoe UI', 'Malgun Gothic', sans-serif") == ("Segoe UI", "Malgun Gothic")
    assert slides.font_pair("'Gowun Batang', Georgia, 'Batang', serif") == ("Georgia", "Batang")
    assert slides.font_pair("'SUIT', Bahnschrift, 'Segoe UI', 'Malgun Gothic', sans-serif") == ("Bahnschrift", "Malgun Gothic")
    assert slides.font_pair("Pretendard") == ("Malgun Gothic", "Malgun Gothic")
    for key in OWN:
        css = (slides_style.THEME_DIR / f"{key}.css").read_text(encoding="utf-8")
        for name in ("--sans", "--display", "--quote", "--figure"):
            latin, east = slides.font_pair(_var(css, name))
            assert latin.lower() not in slides_style.bundled_family_names() and east.lower() not in slides_style.bundled_family_names()


def test_pptx_text_is_shrunk_only_where_the_browser_drew_a_bundled_font() -> None:
    """옮긴 템플릿 · IR(앱 글꼴을 안 쓰는 상자)은 브라우저와 PPTX 가 같은 글꼴이라 손대지 않는다 — 브라우저도 부르지 않는다."""
    blocks = [{"x": 0, "y": 0, "w": 300, "h": 30, "runs": [{"text": "가나다", "size": 20, "weight": 400, "font": "Georgia, 'Batang', serif"}]},
              {"x": 0, "y": 0, "w": 300, "h": 30, "runs": [{"text": "가나다", "size": 20, "weight": 400, "font": "'Malgun Gothic', sans-serif"}]}]
    assert slides.fit_pptx_fonts(None, blocks) == 0 and not any("pptx_scale" in b for b in blocks)


class _FakePage:
    """줄 폭을 돌려주는 가짜 — 글자 크기에 비례(첫 판 1.15배 넓게)."""

    def __init__(self) -> None:
        self.calls = 0

    def evaluate(self, _js: str, lines: list) -> list[float]:
        self.calls += 1
        return [sum(len(r["text"]) * r["size"] * 1.15 for r in line) for line in lines]


def test_pptx_text_shrinks_until_every_line_fits_the_width_the_browser_drew() -> None:
    page = _FakePage()
    blocks = [{"x": 0, "y": 0, "w": 100, "h": 60, "runs": [{"text": "가나다라마", "size": 20, "weight": 700, "font": "'Pretendard', 'Malgun Gothic'"},
                                                        {"br": True}, {"text": "가나", "size": 20, "weight": 700, "font": "'Pretendard', 'Malgun Gothic'"}]}]
    assert slides.fit_pptx_fonts(page, blocks) == 1
    scale = blocks[0]["pptx_scale"]
    assert slides.PPTX_SCALE_FLOOR <= scale < 1 and 5 * 20 * 1.15 * scale <= 101
    wide = [{"x": 0, "y": 0, "w": 10, "h": 30, "runs": [{"text": "가나다라마", "size": 20, "weight": 400, "font": "'SUIT', 'Malgun Gothic'"}]}]
    slides.fit_pptx_fonts(_FakePage(), wide)
    assert wide[0]["pptx_scale"] == slides.PPTX_SCALE_FLOOR   # 아무리 넓어도 이 밑으로는 안 줄인다


def test_the_pdf_never_overwrites_a_file_that_is_already_there(tmp_path: Path) -> None:
    out = tmp_path / "보고.pptx"
    assert slides._pdf_target(out) == tmp_path / "보고.pdf"
    (tmp_path / "보고.pdf").write_bytes(b"mine")
    assert slides._pdf_target(out) == tmp_path / "보고 (2).pdf"


COVER = ('<section class="cover"><p class="kicker">분기 보고</p><h1>3분기 매출 두 배, 다음은 해외</h1>'
         '<p class="lede">국내 1,240곳 · 재구매 62%</p></section>')


@needs_browser
def test_the_renderer_really_draws_with_the_bundled_fonts_and_asks_nobody_outside() -> None:
    """`document.fonts` 로 확인 — 글꼴 파일이 실제로 들어왔고(loaded) 제목 · 본문이 그 글꼴이다. 밖으로 나간 요청은 0."""
    from playwright.sync_api import sync_playwright

    asked: list[str] = []
    with sync_playwright() as pw:
        browser = slides._launch(pw)
        try:
            page = browser.new_page(viewport={"width": slides.WIDTH, "height": slides.HEIGHT})
            page.on("request", lambda req: asked.append(req.url))
            slides._serve_fonts(page)
            for key in OWN:
                html, _ = slides_style.wrap_fragment(COVER, key, title="t", index=1, total=3)
                slides._open_slide(page, html)
                got = page.evaluate("""() => ({status: document.fonts.status,
                    loaded: [...document.fonts].filter((f) => f.status === 'loaded').map((f) => f.family.replace(/["']/g, '')),
                    failed: [...document.fonts].filter((f) => f.status === 'error').map((f) => f.family)})""")
                title, body = PAIRS[key]
                assert got["status"] == "loaded" and not got["failed"], (key, got)
                assert title in got["loaded"] and body in got["loaded"], (key, got)
                # 글꼴이 정말 바뀌었나 — 같은 제목을 옛 윈도우 글꼴로만 그린 폭과 다르다.
                widths = page.evaluate("""(old) => { const h = document.querySelector('h1'); const a = h.getBoundingClientRect().width;
                    const r = document.createRange(); r.selectNodeContents(h); const w1 = r.getBoundingClientRect().width;
                    h.style.fontFamily = old; const w2 = r.getBoundingClientRect().width; return [w1, w2]; }""", OLD_FALLBACK[key] + ", 'Malgun Gothic'")
                assert abs(widths[0] - widths[1]) > 1, (key, widths)
        finally:
            browser.close()
    outside = [url for url in asked if not url.startswith(("about:", "data:", slides_style.FONT_ORIGIN))]
    assert not outside, outside
    assert any(url.startswith(slides_style.FONT_ORIGIN) for url in asked)


@needs_browser
def test_a_real_bake_writes_a_pdf_with_the_new_font_beside_the_pptx(tmp_path: Path) -> None:
    from pptx import Presentation
    from pypdf import PdfReader

    deck = tmp_path / "deck"
    deck.mkdir()
    big = '<section class="bignum"><p class="kicker">재구매</p><h2>열 곳 중 여섯이 다시 샀다</h2><p class="num">62<small>%</small></p></section>'
    for i, frag in enumerate((COVER, big), 1):
        html, _ = slides_style.wrap_fragment(frag, "paper", title="t", index=i, total=2)
        (deck / f"slide_{i:02d}.html").write_text(slides.clean_slide_html(html), encoding="utf-8")
    out = slides.render_in_subprocess(deck, tmp_path / "보고.pptx")
    assert out["slides"] == 2 and "pdf_error" not in out, out
    # 쓴 글자만 담은 프리텐다드를 심었고(보통 · 굵게), 그래서 글자를 줄이지 않았다.
    assert {"Pretendard/regular", "Pretendard/bold"} <= set(out.get("embedded_fonts") or []), out
    assert "pptx_scaled_boxes" not in out and "font_embed_errors" not in out, out
    assert 0 < out["embedded_font_bytes"] < 400_000, out
    pdf = PdfReader(out["pdf"])
    assert len(pdf.pages) == 2
    names = set()
    for page in pdf.pages:
        for font in (page["/Resources"].get("/Font") or {}).values():
            names.add(str(font.get_object().get("/BaseFont")))
    assert any("Pretendard" in name for name in names), names
    prs = Presentation(str(tmp_path / "보고.pptx"))
    xml = "".join(slide.shapes._spTree.xml for slide in prs.slides)
    faces = set(re.findall(r'typeface="([^"]+)"', xml))
    assert "Pretendard" in faces, faces
    root = prs.part._element
    assert root.get("embedTrueTypeFonts") == "1"
    listed = re.findall(r'<p:font typeface="([^"]+)"', prs.part.blob.decode("utf-8"))
    assert "Pretendard" in listed, listed


def test_the_store_package_and_the_public_export_carry_the_fonts() -> None:
    """포장(MSIX)은 `src/alture_jarvis` 를 통째로 옮기고, 공개판은 깃이 추적하는 src/ 를 옮긴다 — 글꼴 폴더가 빠지는 규칙이 없다."""
    import importlib.util

    root = Path(__file__).resolve().parents[1]
    build = (root / "packaging" / "msix" / "build_app_package.ps1").read_text(encoding="utf-8")
    assert re.search(r'robocopy \(Join-Path \$Root "src\\alture_jarvis"\) \(Join-Path \$site "alture_jarvis"\) /E', build)
    sync = (root / "scripts" / "sync_to_runtime.ps1").read_text(encoding="utf-8-sig")
    assert '"src"' in sync and "-Recurse" in sync
    spec = importlib.util.spec_from_file_location("export_public", root / "scripts" / "export_public.py")
    if spec is None or not (root / "scripts" / "export_public.py").exists():
        pytest.skip("공개판에는 내보내기 도구가 없다")
    export = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(export)  # type: ignore[union-attr]
    for path in slides_style.FONT_DIR.rglob("*"):
        if path.is_file():
            rel = path.relative_to(root).as_posix()
            assert export.is_public(rel), rel
            if path.suffix == ".txt":   # 라이선스 글은 이름 바꾸기에 걸리지 않는다(저작권 줄이 그대로 나간다)
                text = path.read_text(encoding="utf-8")
                assert export.transform(text) == text, rel


def test_embedded_font_header_matches_the_font_it_carries() -> None:
    """EOT 머리의 검사합 · 이름은 **실은 줄인 글꼴**의 값이어야 한다(원본 값을 쓰면 데이터와 어긋난다 · 10-26 첫 시도).
    심기 허용 값(fsType)은 0 — SUIT 는 원래 4(보기 · 인쇄만)라 받는 사람이 고칠 수 없게 열릴 수 있었다."""
    import io
    import struct

    from fontTools.ttLib import TTFont

    from alture_jarvis import pptx_fonts

    if pptx_fonts.unavailable():
        pytest.skip(pptx_fonts.unavailable())
    path, real_bold = pptx_fonts.face_file("SUIT", True)
    assert path is not None and real_bold
    data = pptx_fonts.subset_font(path, "가나다 ABC 123")
    blob = pptx_fonts.eot(data)
    total, size, version, flags = struct.unpack("<IIII", blob[:16])
    assert total == len(blob) and size == len(data) and blob[-size:] == data
    assert version == 0x00020002 and flags == 0x1
    font = TTFont(io.BytesIO(data))
    assert struct.unpack("<I", blob[60:64])[0] == font["head"].checkSumAdjustment
    assert struct.unpack("<H", blob[32:34])[0] == 0 == font["OS/2"].fsType
    family_len = struct.unpack("<H", blob[82:84])[0]
    assert blob[84:84 + family_len].decode("utf-16-le") == "SUIT"
    cmap = font.getBestCmap()
    assert {ord(c) for c in "가나다ABC123"} <= set(cmap) and ord("라") not in cmap   # 쓴 글자만


def test_a_family_without_a_bold_face_carries_bold_text_in_its_regular_face() -> None:
    from pptx import Presentation

    from alture_jarvis import pptx_fonts

    if pptx_fonts.unavailable():
        pytest.skip(pptx_fonts.unavailable())
    deck = Presentation()
    report = pptx_fonts.embed(deck, {"Gowun Dodum": {False: "가나", True: "다라"}})
    assert report["embedded"] == ["Gowun Dodum/regular"] and not report["errors"], report
    blob = next(p.blob for p in deck.part.package.iter_parts() if str(p.partname).endswith(".fntdata"))
    from fontTools.ttLib import TTFont
    import io
    import struct
    size = struct.unpack("<I", blob[4:8])[0]
    cmap = TTFont(io.BytesIO(blob[-size:])).getBestCmap()
    assert {ord(c) for c in "가나다라"} <= set(cmap)


def test_without_embedding_the_old_windows_font_path_still_bakes() -> None:
    """글꼴을 못 심으면(부품 없음) 예전 길: 윈도우 글꼴 이름 + 줄인 글자."""
    import tempfile

    from pptx import Presentation

    blocks = [{"x": 10, "y": 10, "w": 300, "h": 40, "size": 20, "pptx_scale": 0.85,
               "runs": [{"text": "가나다", "size": 20, "weight": 400, "font": "'Pretendard', 'Malgun Gothic', sans-serif"}]}]
    with tempfile.TemporaryDirectory() as tmp:
        from PIL import Image
        bg = Path(tmp) / "bg.png"
        Image.new("RGB", (16, 9), "white").save(bg)
        out = slides.build_pptx([{"background": bg, "blocks": blocks}], Path(tmp) / "x.pptx", embed_fonts=False)
        assert out.get("pptx_scaled_boxes") == 1 and "embedded_fonts" not in out
        xml = Presentation(str(Path(tmp) / "x.pptx")).slides[0].shapes._spTree.xml
        assert 'typeface="Malgun Gothic"' in xml and "Pretendard" not in xml


def test_the_preview_font_route_hands_out_only_listed_fonts(tmp_path: Path) -> None:
    """미리보기 칸은 글꼴을 본체 `/v1/slides/fonts/…` 에서 받는다(토큰 없이 · 출처 없는 칸이라 CORS 열림). 목록 밖은 404."""
    from fastapi.testclient import TestClient

    from alture_jarvis.api import create_app
    from alture_jarvis.config import JarvisConfig

    config = JarvisConfig(runtime_root=tmp_path, data_dir=tmp_path / "data", output_dir=tmp_path / "outputs",
                          control_token="test-token", camera_enabled=False, vision_enabled=False,
                          voice_enabled=False, tts_enabled=False, profile_path=tmp_path / "data" / "JARVIS_PROFILE.md")
    with TestClient(create_app(config, start_services=False)) as client:
        ok = client.get("/v1/slides/fonts/pretendard/Pretendard-Regular.woff2")
        assert ok.status_code == 200 and ok.content[:4] == b"wOF2"
        assert ok.headers["content-type"] == "font/woff2" and ok.headers["access-control-allow-origin"] == "*"
        for bad in ("pretendard/OFL.txt", "fonts.json", "../slides.py", "pretendard/..%2F..%2Fslides.py", "pretendard/Nope.woff2"):
            assert client.get(f"/v1/slides/fonts/{bad}").status_code == 404, bad


def test_the_packaged_app_lets_the_preview_load_fonts_from_the_backend() -> None:
    conf = (Path(__file__).resolve().parents[1] / "ui" / "src-tauri" / "tauri.conf.json").read_text(encoding="utf-8")
    csp = re.search(r'"csp":\s*"([^"]+)"', conf).group(1)
    font_src = next(part for part in csp.split(";") if part.strip().startswith("font-src"))
    assert "'self'" in font_src and "http://127.0.0.1:8091" in font_src
