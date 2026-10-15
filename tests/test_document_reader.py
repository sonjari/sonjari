from __future__ import annotations

from pathlib import Path

from alture_jarvis.document_reader import DocumentReader


def test_document_reader_decodes_cp949(tmp_path: Path) -> None:
    path = tmp_path / "회의록.txt"
    path.write_bytes("알튜어 회의 결과입니다.".encode("cp949"))
    result = DocumentReader().read_path(path)
    assert "알튜어 회의" in result.text
    assert result.encoding in {"cp949", "euc-kr"}


def test_document_reader_reads_docx(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "기획서.docx"
    document = Document()
    document.add_heading("JARVIS 기획", 1)
    document.add_paragraph("문서 형식과 한글을 정상적으로 읽습니다.")
    document.save(path)
    result = DocumentReader().read_path(path)
    assert result.format == "docx"
    assert "한글을 정상적으로" in result.text


def test_scanned_pdf_uses_vision_ocr_callback(tmp_path: Path) -> None:
    from pypdf import PdfWriter

    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    with path.open("wb") as stream:
        writer.write(stream)

    reader = DocumentReader(pdf_ocr=lambda raw, pages: ("[PDF OCR 1쪽]\n판독된 한글", 1))
    result = reader.read_path(path)
    assert "판독된 한글" in result.text
    assert any("Gemma 비전 OCR" in warning for warning in result.warnings)
