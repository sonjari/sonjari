"""저장한 문서에서 글을 다시 꺼내, 쓰려던 내용과 견준다 (2026-10-15)."""

from __future__ import annotations

import re
import unicodedata
from pathlib import Path

_READABLE = {"docx", "xlsx", "md", "txt"}


def _normalize(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    for curved, straight in (("\u201c", '"'), ("\u201d", '"'), ("\u2018", "'"), ("\u2019", "'")):
        text = text.replace(curved, straight)
    text = text.replace("#", "").replace("**", "")
    text = re.sub(r"\s+", " ", text)
    return text


def _extract_file_text(path: Path, fmt: str) -> str:
    if fmt == "docx":
        from docx import Document

        document = Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)
    if fmt == "xlsx":
        import openpyxl

        book = openpyxl.load_workbook(str(path), data_only=True)
        parts: list[str] = []
        for sheet in book.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        parts.append(str(cell.value))
        return "\n".join(parts)
    if fmt in ("md", "txt"):
        return path.read_text(encoding="utf-8")
    raise ValueError(f"읽을 수 없는 형식: {fmt}")


def _wanted_units(payload: dict) -> list[str]:
    units: list[str] = []
    for row in payload.get("rows") or []:
        for cell in row:
            if cell is not None:
                units.append(str(cell))
    for block in payload.get("blocks") or []:
        kind = block.get("type")
        if kind == "table":
            for row in block.get("rows") or []:
                for cell in row:
                    if cell is not None:
                        units.append(str(cell))
        elif kind in ("heading", "paragraph"):
            units.append(str(block.get("text") or ""))
    text = payload.get("text")
    if text:
        for line in str(text).splitlines():
            units.append(line)
    return units


def readback_gaps(path: Path, fmt: str, payload: dict) -> list[str] | None:
    fmt = (fmt or "").lower()
    if fmt not in _READABLE:
        return None
    try:
        stored = _extract_file_text(Path(path), fmt)
    except Exception:
        return None
    hay = _normalize(stored)
    gaps: list[str] = []
    for unit in _wanted_units(payload):
        needle = _normalize(unit).strip()
        if len(needle) < 2:
            continue
        if needle not in hay:
            gaps.append(unit)
    return gaps


def readback_report(path: Path, fmt: str, payload: dict, *, skip=None) -> dict:
    """그림자 기록용 — 견준 단위 수와 빠진 것. `skip(단위)` 가 참이면 그 단위는 세지 않는다(일부러 비운 칸 등).

    못 읽는 형식이거나 다시 읽기가 실패하면 `checked` 거짓 · `gaps` None — 「확인 못 함」이지 「통과」가 아니다.
    """
    wanted = [unit for unit in _wanted_units(payload)
              if len(_normalize(unit).strip()) >= 2 and not (skip and skip(unit))]
    gaps = readback_gaps(path, fmt, payload)
    if gaps is None:
        return {"checked": False, "units": len(wanted), "gaps": None}
    kept = [gap for gap in gaps if not (skip and skip(gap))]
    return {"checked": True, "units": len(wanted), "gaps": kept}
