"""장표 만들기 — 모델이 장마다 HTML 로 디자인하고, 다 되면 PPTX 로 굽는다 (2026-10-25 대표).

## 왜

`documents._write_pptx` 는 python-pptx 기본 틀에 「제목 + 글머리표」 한 모양만 썼다 — 흰 바탕에 글만 있는 장표가 나왔다.
같은 날 같은 주제 · 같은 숫자 · 같은 Swift 27B 로 공개 스킬 셋을 견줬다(scratchpad `pptx_three_way.png` · 한 판씩):

    Akxan(모델이 장마다 HTML)      보기 가장 좋음 · 약 6분 · PPTX 로 옮기면 배경이 빠져 흰 글이 안 보임
    MiniMax(모델이 PPT 코드)        70초 · 깔끔하지만 수수함 · PPTX 멀쩡
    siril9(설계도만, 도구가 디자인) 22초 · 회사 보고서 느낌 · 짐이 큼(파일 758개)

대표: 「첫 번째의 강한 느낌만 빼면 낫겠다」. 그래서 모델은 HTML 로 디자인하고(차분한 스타일 예시 · `slides_style`),
PPTX 는 우리가 굽는다 — **배경은 그림 한 장, 글은 그 위에 같은 자리 · 같은 크기의 고칠 수 있는 글 상자.**
배경이 빠지던 까닭(그러데이션 배경을 PPT 도형으로 옮기지 못함)을 아예 피한다.

## 손 하나 (`local_make_slide`)

    html 없이 부르면   디자인 규칙 + 스타일 예시를 돌려준다(규칙은 PPT 를 만들 때만 대화에 들어간다)
    장마다 html 과     그 장을 저장하고 화면(컴퓨터 창)에 바로 띄운다(`deck_slide` 알림)
    마지막 장이면      PPTX 를 굽는다(`deck_built` 알림)

## 지키는 것

- 모델이 쓴 HTML 의 `<script>` 는 지운다 — 굽는 브라우저 · 화면 미리보기 어디서도 모델 코드를 돌리지 않는다.
- 굽기는 **따로 된 프로세스**에서 한다(`python -m alture_jarvis.slides render`) — 손자리가 쥔 사용자 크롬과 섞이지 않게,
  설치된 크롬(없으면 윈도우 기본 엣지)을 화면 없이 띄운다. 새로 받는 브라우저는 없다.
- 차트 글자(SVG 안 글)는 배경 그림에 남는다 — 고칠 수 있는 진짜 차트는 다음 판(MiniMax 방식)이다.
"""

from __future__ import annotations

import html as _html
import json
import math
import re
import subprocess
import sys
from pathlib import Path
from typing import Any

WIDTH, HEIGHT = 1280, 720
#: 1px = 9525 EMU (96 dpi). 1280 × 720 px → 13.333 × 7.5 인치(16:9).
EMU_PER_PX = 9525
MAX_HTML_CHARS = 120_000
MAX_SLIDES = 40

GUIDE = """장표는 **디자이너가 만든 결(템플릿)의 장 틀을 골라 내용만 채운다.** 디자인 · CSS 는 손자리가 끼운다. 순서:
1) 결 — `recommended`(요청과 꼬리표로 고른 셋) 중 하나를 고른다. 사람이 바란 분위기 · 색이 있으면 전체 목록(`themes`)에서
   그에 맞는 것을. 다른 결을 보고 싶으면 html 없이 theme 만 넣어 다시 부르면 **그 결의 장 틀(layouts)**이 온다.
   모든 장에 같은 theme 을 넣는다. 고른 결과 까닭을 답에 한 줄로 밝힌다(다른 후보 둘의 이름도).
2) 계획 — 장 수(보통 6~10장)와 장마다 「쓸 틀(layouts 의 ### 이름 · 부품 결은 cover · parts · closing) · 한 줄 주장」을 먼저 정한다.
   `order` 가 오면 그 장표 유형의 흐름을 출발점으로 짠다.
   - 1번은 표지 틀, 마지막은 마무리 틀. **이웃한 두 장은 다른 틀.** 틀을 고루 쓴다.
   - 내용이 틀을 고른다: 숫자 하나 · 여럿 → 숫자 틀, 값 비교 → 차트 · 표 틀, 순서 → 단계 틀, 주장 하나 → 한 문장 틀.
3) 1번 장부터 차례로 `local_make_slide`: title(모든 장에 같게) · index · total · theme · parts.
   guide 에 parts · frames 가 오는 결은 **모든 장을 조립표(parts)로** 넣는다 — 표지 · 마무리는 frames, 본문은 parts 메뉴.
   글자 크기 · 간격 · 줄 · 장식(도형)은 손자리가 결에 맞게 그린다 — 너는 내용과 부품만 고른다. 사진 자리는 photo 부품.
   (부품이 없는 결만: layouts 의 틀 하나를 복사한 `<section>` 조각을 html 로 넣고 글만 바꾼다. 틀의 영어 · 자리 표시가 남으면 돌려보낸다.)
   부를 때마다 그 장이 사용자 화면에 그려지고, 마지막 장(index == total)을 넣으면 PPTX 가 구워진다.

글 — 장표는 문서가 아니다:
- 제목(h2 · h1)은 **주장 문장**(「3분기 매출 두 배」 ○, 「3분기 매출」 ×), 한 줄~두 줄. kicker 는 한국어 2~8자.
- 칸 하나에 1~2문장. 한 장에 글이 많으면 장을 나눈다(넘치면 손자리가 글자를 줄이고, 그래도 넘치면 돌려보낸다).
- 쓰지 않는 것: 이모지 · 「감사합니다 · 질문?」만 있는 마지막 장(대신 다음 걸음 한 문장).
- 글은 모두 한국어(고유명사 · 단위만 원문).

숫자 규칙: 장표의 숫자는 **이 판에서 읽은 자료 · 요청에 있는 값만** 쓴다. 막대 길이(%)는 가장 큰 값을 100% 로 둔 비율이다.
계산한 값(차이 · 배수 · 퍼센트)은 계산이 맞는지 다시 보고 쓴다. 근거 없는 숫자가 여럿이면 그 장은 한 번 돌려보내진다."""


def guide_payload(theme: str | None = None, style_example: str | None = None, about: str = "",
                  deck_type: str | None = None) -> dict[str, Any]:
    """html 없이 불렀을 때 돌려주는 것 — 규칙 + 결 목록 + 배치 틀 12개(`slides_style`).

    ★ 3차까지 주던 Akxan 플레이북은 더 주지 않는다 — 내용 장마다 영어 꼬리표(`span.overline`)를 넣으라는 지시가 있어
    GUIDE 의 「대문자 꼬리표 빼기」와 부딪쳤다(B 조사, `page-html-playbook.md:121`).
    """
    from .slide_catalog import recommend
    from .slides_style import THEMES, layout_menu, theme_menu

    from .slides_style import TEMPLATES

    picks = recommend(about)
    chosen = theme if theme in THEMES else str(picks[0]["theme"])
    payload = {"ok": True, "action": "local_make_slide", "guide": GUIDE,
               "recommended": [{"theme": p["theme"], "name": p["name"], "when": p["when"], "why": p["why"]} for p in picks],
               "theme": chosen, "themes": theme_menu(), "layouts": layout_menu(chosen),
               "next": "계획을 세운 뒤 1번 장(cover)부터 theme 과 html(<section> 조각)을 넣어 부른다."}
    # 장 순서가 정해진 결(IR 등) — 이야기판을 이 순서로 짠다.
    flow = str(TEMPLATES.get(chosen, {}).get("flow") or "")
    if flow:
        payload["order"] = flow
    # 부품 결(2026-10-25 대표 「부품으로 증류 · 모델은 골라 배치」) — 틀은 표지 · 마무리만, 본문은 부품 메뉴.
    from . import slide_parts

    if slide_parts.supports(chosen):
        # 틀(html)을 하나도 주지 않는다 — 표지 · 마무리도 부품이다(2026-10-26 대표 「템플릿은 참고만」).
        #   예전엔 표지 · 마무리 틀을 줬고 모델은 장식 SVG 까지 베꼈다(daisy 표지 한 장 약 8,000토큰 · 10-26 결과.md).
        #   장식은 손자리가 결 값으로 그린다(`slide_frames.DECOS`) — 모델은 보지도 쓰지도 않는다.
        from . import slide_frames

        payload["parts"] = slide_parts.menu(chosen)
        payload["frames"] = slide_frames.menu()
        payload["assembly"] = slide_parts.GUIDE
        payload["layouts"] = "이 결은 틀(html)이 없다 — 표지 · 마무리는 frames, 본문은 parts 의 조립표를 parts 칸에 넣는다."
        payload["next"] = ("1번 장은 parts 에 {frame:\"cover\", …}, 본문 장은 parts(조립표), 마지막 장은 {frame:\"closing\", …}. "
                           "이야기판 카드의 틀 이름은 cover · parts · closing. 계획(장마다 부품 · 한 줄 주장)을 세운 뒤 1번 장부터.")
        # 장표 유형(결 × 유형) — 결은 느낌, 유형은 장 흐름. 결에 정해진 흐름(IR)이 있으면 그것이 먼저다.
        kind = slide_frames.deck_type(about, deck_type)
        payload["deck_type"] = kind
        if not flow:
            payload["order"] = slide_frames.flow_text(kind)
    return payload


# ── AI 티 기계 검사 (2026-10-25 · 사고 R: 안내문보다 결정론) ──────────────────────────────────
#
# 출처: 앤트로픽 anthropics/skills 의 `frontend-design` 스킬(Apache-2.0)이 적은 「생성된 페이지의 흔한 티」와
# f42gh/anti-ai-slop-design(MIT)의 안티패턴 목록에서, **HTML 만 보고 확실히 가를 수 있는 것**만 골라 장표에 맞게 다시 썼다.
# (팔레트 무게 · 리듬 · 위계 같은 것은 기계로 못 가린다 — 그건 규칙 글이 맡는다.) 고지: packaging/THIRD_PARTY_NOTICES.md.
# 첫 실사용(「손자리 소개」 7장 · 10-25)에 대어 보니 7장 모두에 대문자 꼬리표 · 가운데 점 쪽 번호가 있었다.

_CAPS_SEGMENT = re.compile(r"^(?=.*[A-Z]{2,}.*[A-Z]{2,})[A-Z0-9][A-Z0-9 .·/&:+\-—–]{3,40}$")
_CAPS_WORDS = re.compile(r"[A-Z]{2,}")
_NUMBER_MARK = re.compile(r"^(?:0\d|STEP\s*\d+|PART\s*\d+)$", re.I)
_DOT_CHAIN = re.compile(r"\S\s+[·•]\s+\S.*\s[·•]\s+\S")
_GRADIENT_TEXT = re.compile(r"(?:-webkit-)?background-clip\s*:\s*text", re.I)
_GLOW = re.compile(r"(?:text-shadow|box-shadow)\s*:[^;{}]*?\b0(?:px)?\s+0(?:px)?\s+([2-9]\d|\d{3,})px", re.I)
_CLASS_ATTR = re.compile(r"class\s*=\s*[\"']([^\"']+)[\"']", re.I)
_CSS_RULE = re.compile(r"([^{}]+)\{([^{}]*)\}")


def _text_segments(raw: str) -> list[str]:
    """태그 사이 글 토막 — 스타일 · 주석은 빼고, 토막마다 공백을 접는다."""
    text = _STYLE_BLOCK.sub(" ", str(raw or ""))
    text = re.sub(r"<!--.*?-->", " ", text, flags=re.S)
    return [seg for seg in (re.sub(r"\s+", " ", _html.unescape(part)).strip() for part in re.split(r"<[^>]+>", text)) if seg]


def design_tells(raw: str) -> list[dict[str, Any]]:
    """장 HTML 에서 **기계로 확실히 보이는** AI 티. 각 항목 {id, what, fix, seen}. `hard` 가 참인 것만 돌려보낸다."""
    found: list[dict[str, Any]] = []
    segments = _text_segments(raw)
    caps = [seg for seg in segments
            if (_CAPS_SEGMENT.match(seg) and len(_CAPS_WORDS.findall(seg)) >= 2) or re.fullmatch(r"[A-Z]{7,}", seg)]
    if caps:
        found.append({"id": "caps_label", "hard": False, "seen": caps[:4],
                      "what": "영어 대문자 꼬리표", "fix": "꼬리표를 지우거나, 꼭 필요하면 한국어 보통 글씨로."})
    marks = [seg for seg in segments if _NUMBER_MARK.match(seg)]
    if len(marks) >= 2:
        found.append({"id": "number_marks", "hard": False, "seen": marks[:4],
                      "what": "번호 표지(01 / 02 · STEP 1)", "fix": "순서가 아니면 번호를 지운다. 순서면 1 · 2 · 3 처럼 내용 안에서."})
    dots = [seg for seg in segments if len(seg) <= 50 and _DOT_CHAIN.search(seg)]
    if dots:
        found.append({"id": "dot_chain", "hard": False, "seen": dots[:4],
                      "what": "가운데 점으로 이은 짧은 줄", "fix": "한 가지만 남기거나 문장으로. 쪽 번호는 「3 / 8」만."})
    css = " ".join(re.findall(r"<style\b[^>]*>(.*?)</style\s*>", str(raw or ""), flags=re.I | re.S))
    style_attrs = " ".join(re.findall(r"style\s*=\s*[\"']([^\"']*)[\"']", str(raw or ""), flags=re.I))
    if _GRADIENT_TEXT.search(css + " " + style_attrs):
        found.append({"id": "gradient_text", "hard": True, "seen": [], "what": "그러데이션 글자", "fix": "글자는 단색으로."})
    if _GLOW.search(css + " " + style_attrs):
        found.append({"id": "glow", "hard": True, "seen": [], "what": "빛 번짐(glow) 그림자", "fix": "번짐 그림자를 지운다."})
    # 같은 꼴 카드 셋 이상 — 둥근 모서리에 바탕 · 그림자 · 테두리를 입힌 클래스가 셋 이상 쓰였나(돌려보내지는 않고 알린다).
    boxed = {sel.strip().lstrip(".") for sels, body in _CSS_RULE.findall(css) for sel in sels.split(",")
             if re.match(r"^\s*\.[\w-]+\s*$", sel) and "border-radius" in body and re.search(r"background|box-shadow|border\s*:", body)}
    counts: dict[str, int] = {}
    for attr in _CLASS_ATTR.findall(str(raw or "")):
        for name in attr.split():
            if name in boxed:
                counts[name] = counts.get(name, 0) + 1
    cards = {name: count for name, count in counts.items() if count >= 3}
    if cards:
        found.append({"id": "same_cards", "hard": False, "seen": [f"{name} ×{count}" for name, count in cards.items()][:3],
                      "what": "똑같은 둥근 카드 반복", "fix": "다음 장부터는 글줄 · 표 · 선으로 나눈 칸처럼 내용에 맞는 모양을 고른다."})
    return found


_SCRIPT = re.compile(r"<script\b[^>]*>.*?</script\s*>|<script\b[^>]*/>", re.I | re.S)
_EXTERNAL_LINK = re.compile(r"<link\b[^>]*>", re.I)
_EVENT_ATTR = re.compile(r"\son[a-z]+\s*=\s*(\"[^\"]*\"|'[^']*'|[^\s>]+)", re.I)
_SIZE_LOCK = ("<style data-sonjari-size>html,body{width:1280px!important;height:720px!important;"
              "overflow:hidden!important;margin:0!important}</style>")


def clean_slide_html(raw: str) -> str:
    """모델이 쓴 장 HTML 을 굽기 · 미리보기 전에 다듬는다 — 스크립트 · 바깥 링크 · 이벤트 속성을 지우고 크기를 고정한다."""
    text = str(raw or "").strip()
    text = re.sub(r"^```(?:html)?\s*|\s*```$", "", text)
    text = _SCRIPT.sub("", text)
    text = _EXTERNAL_LINK.sub("", text)
    text = _EVENT_ATTR.sub("", text)
    if "<head" in text.lower():
        text = re.sub(r"(</head\s*>)", _SIZE_LOCK + r"\1", text, count=1, flags=re.I)
    else:
        text = _SIZE_LOCK + text
    return text


_STYLE_BLOCK = re.compile(r"<style\b[^>]*>.*?</style\s*>", re.I | re.S)
_TAG = re.compile(r"<[^>]+>")


def visible_text(raw: str) -> str:
    """장 HTML 에서 **눈에 보이는 글만** — 쓰기 전 숫자 대조가 CSS 값(1280px · 48px)을 숫자로 세지 않게."""
    text = _STYLE_BLOCK.sub(" ", str(raw or ""))
    text = re.sub(r"<!--.*?-->", " ", text, flags=re.S)
    text = _TAG.sub(" ", text)
    return re.sub(r"\s+", " ", _html.unescape(text)).strip()


def safe_title(title: str) -> str:
    name = re.sub(r'[\\/:*?"<>|\r\n\t]+', " ", str(title or "").strip())
    name = re.sub(r"\s+", " ", name).strip(" .")[:60]
    return name or "장표"


# ── 굽기 (따로 된 프로세스) ─────────────────────────────────────────────────────────────

_EXTRACT_JS = r"""
() => {
  const INLINE = new Set(['inline', 'contents']);
  const SKIP = new Set(['SCRIPT', 'STYLE', 'NOSCRIPT', 'TEMPLATE']);
  const blocks = [];
  const taken = new Set();
  const inSvg = (el) => !!el.closest('svg');
  const hasText = (el) => (el.textContent || '').trim().length > 0;
  const allInline = (el) => Array.from(el.children).every((c) => c.tagName === 'BR' || (INLINE.has(getComputedStyle(c).display) && allInline(c)));
  const color = (cs) => {
    let c = cs.webkitTextFillColor && !/rgba\(0, 0, 0, 0\)|transparent/.test(cs.webkitTextFillColor) ? cs.webkitTextFillColor : cs.color;
    if (/rgba\(\d+, \d+, \d+, 0\)|transparent/.test(c)) {
      const m = (cs.backgroundImage || '').match(/rgba?\([^)]*\)|#[0-9a-fA-F]{3,8}/);
      c = m ? m[0] : 'rgb(30, 41, 59)';
    }
    return c;
  };
  // 거의 투명한 글자(물결 글자 · 장식 번호)는 뽑지 않는다 — 배경 그림에 그대로 남긴다.
  const alphaOf = (node, c) => {
    const m = String(c).match(/rgba\([^)]*,\s*([\d.]+)\)/);
    let a = m ? parseFloat(m[1]) : 1;
    for (let e = node; e && e.nodeType === 1; e = e.parentElement) a *= parseFloat(getComputedStyle(e).opacity || '1');
    return a;
  };
  const charRange = document.createRange();
  const walk = (el) => {
    if (SKIP.has(el.tagName) || inSvg(el)) return;
    const cs = getComputedStyle(el);
    if (cs.display === 'none' || cs.visibility === 'hidden' || parseFloat(cs.opacity) === 0) return;
    if (hasText(el) && allInline(el) && !INLINE.has(cs.display)) {
      const fs = parseFloat(cs.fontSize) || 16;
      const lhRaw = parseFloat(cs.lineHeight);
      const lh = isNaN(lhRaw) ? fs * 1.25 : lhRaw;
      const runs = [];
      let lineTop = null;   // 지금 줄의 아래쪽(이름은 옛것) — 글자 하나하나의 위치로 **브라우저가 실제로 바꾼 줄**을 찾는다
      let lastSpace = true;
      const pushText = (p, text, node) => {
        const last = runs[runs.length - 1];
        if (last && last.node === node && !last.br) { last.text += text; return; }
        runs.push({ text, size: parseFloat(p.fontSize), weight: parseInt(p.fontWeight, 10) || 400, italic: p.fontStyle === 'italic',
                    color: color(p), font: p.fontFamily, spacing: parseFloat(p.letterSpacing) || 0, alpha: alphaOf(node.parentElement, p.color), node });
      };
      const visit = (node) => {
        if (node.nodeType === 3) {
          const p = getComputedStyle(node.parentElement);
          const raw = node.textContent;
          const upper = p.textTransform === 'uppercase';
          for (let i = 0; i < raw.length; i += 1) {
            let ch = raw[i];
            const space = /\s/.test(ch);
            if (space) { if (lastSpace) continue; ch = ' '; }
            else {
              charRange.setStart(node, i); charRange.setEnd(node, i + 1);
              const rect = Array.from(charRange.getClientRects()).find((x) => x.width > 0 || x.height > 0);
              if (rect) {
                // 줄은 글자 **아래쪽**으로 가른다 — 같은 줄의 작은 글자(1.08<small>초</small>)는 위쪽이 한참 낮아도 아래쪽은 거의 같다.
                if (lineTop !== null && rect.bottom > lineTop + lh * 0.5) {
                  // 줄이 바뀌었다 — 앞의 빈칸을 지우고 줄바꿈을 넣는다
                  const last = runs[runs.length - 1];
                  if (last && !last.br) last.text = last.text.replace(/ +$/, '');
                  runs.push({ br: true, text: '\n' });
                }
                if (lineTop === null || rect.bottom > lineTop + lh * 0.5) lineTop = rect.bottom;
              }
            }
            lastSpace = space;
            pushText(p, upper ? ch.toUpperCase() : ch, node);
          }
        } else if (node.nodeType === 1) {
          if (node.tagName === 'BR') {
            const last = runs[runs.length - 1];
            if (last && !last.br) last.text = last.text.replace(/ +$/, '');
            runs.push({ br: true, text: '\n' }); lastSpace = true; lineTop = null; return;
          }
          const ncs = getComputedStyle(node);
          if (ncs.display === 'none' || ncs.visibility === 'hidden') return;
          node.childNodes.forEach(visit);
        }
      };
      el.childNodes.forEach(visit);
      while (runs.length && runs[runs.length - 1].br) runs.pop();
      if (runs.length && !runs[runs.length - 1].br) runs[runs.length - 1].text = runs[runs.length - 1].text.replace(/ +$/, '');
      const range = document.createRange();
      range.selectNodeContents(el);
      const r = range.getBoundingClientRect();
      const solid = runs.some((run) => !run.br && run.text.trim() && run.alpha >= 0.35);
      if (solid && r.width > 0 && r.height > 0) {
        runs.forEach((run) => { if (run.node) taken.add(run.node); delete run.node; delete run.alpha; });
        blocks.push({ x: r.left, y: r.top, w: r.width, h: r.height, align: cs.textAlign, line_px: lh, size: fs,
                      lines: 1 + runs.filter((run) => run.br).length, runs: runs.filter((run) => run.br || run.text) });
      }
      return;
    }
    Array.from(el.children).forEach(walk);
  };
  walk(document.body);
  // 뽑은 글자만 투명하게 — 테두리 · 아이콘(currentColor)은 그대로 남긴다.
  taken.forEach((node) => {
    const span = document.createElement('span');
    span.setAttribute('style', 'color:transparent!important;-webkit-text-fill-color:transparent!important;text-shadow:none!important');
    span.textContent = node.textContent;
    node.parentNode.replaceChild(span, node);
  });
  return blocks;
}
"""


def _launch(playwright: Any) -> Any:
    last: Exception | None = None
    for channel in ("chrome", "msedge", None):
        try:
            return playwright.chromium.launch(channel=channel, headless=True) if channel else playwright.chromium.launch(headless=True)
        except Exception as exc:  # noqa: BLE001 — 다음 브라우저로
            last = exc
    raise RuntimeError(f"no_browser_for_slides: {last}")


def _rgb(text: str) -> tuple[int, int, int]:
    text = str(text or "").strip()
    m = re.match(r"rgba?\((\d+),\s*(\d+),\s*(\d+)", text)
    if m:
        return int(m.group(1)), int(m.group(2)), int(m.group(3))
    m = re.match(r"#([0-9a-fA-F]{6})", text)
    if m:
        value = m.group(1)
        return int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16)
    return 30, 41, 59


#: 받는 사람 PC 에 있는 한글 글꼴 — CSS 글꼴 목록에서 이 중 처음 것을 한글(`a:ea`)에 쓴다.
_KOREAN_FONTS = {"malgun gothic": "Malgun Gothic", "맑은 고딕": "Malgun Gothic", "batang": "Batang", "바탕": "Batang",
                 "gulim": "Gulim", "굴림": "Gulim", "dotum": "Dotum", "돋움": "Dotum", "gungsuh": "Gungsuh", "궁서": "Gungsuh"}
_GENERIC = {"serif", "sans-serif", "monospace", "cursive", "fantasy", "system-ui", "ui-monospace", "ui-sans-serif", "ui-serif"}


def font_pair(family: str, embedded: frozenset[str] | set[str] = frozenset()) -> tuple[str, str]:
    """CSS font-family → (라틴 글꼴, 한글 글꼴). 브라우저처럼 라틴은 목록 첫 글꼴, 한글은 목록에 있는 첫 한글 글꼴.

    ★ 앱에 든 글꼴(프리텐다드 등)은 **PPTX 에 심었을 때만**(`embedded` · 소문자 이름) 그 이름을 라틴 · 한글 둘 다에 쓴다
    (2026-10-26 · `pptx_fonts`). 못 심었으면 건너뛴다 — 앱 안에만 있어 받는 사람 PC 에 없고, 없는 이름을 적으면 PowerPoint 가
    아무 글꼴로 바꿔 그린다. 그때 글 상자는 목록 뒤의 윈도우 기본 글꼴(옛 결과 같은 글꼴)로 간다.
    """
    from .slides_style import bundled_family_names

    bundled = bundled_family_names()
    names = [part.strip().strip("'\"") for part in str(family or "").split(",") if part.strip()]
    if names and names[0].lower() in embedded:
        return names[0], names[0]
    names = [n for n in names if n.lower() not in bundled]
    korean = next((_KOREAN_FONTS[n.lower()] for n in names if n.lower() in _KOREAN_FONTS), "Malgun Gothic")
    first = next((n for n in names if n.lower() not in _GENERIC), korean)
    latin = _KOREAN_FONTS.get(first.lower(), first)
    return latin, korean


#: 글 상자 줄마다 **PPTX 에 적을 글꼴**로 그린 폭(px) — 브라우저가 그 글꼴로 대신 잰다.
_PPTX_WIDTH_JS = r"""
(lines) => {
  const box = document.createElement('div');
  box.style.cssText = 'position:absolute;left:0;top:0;white-space:pre;visibility:hidden';
  document.body.appendChild(box);
  const out = lines.map((runs) => {
    box.textContent = '';
    runs.forEach((r) => {
      const s = document.createElement('span');
      s.textContent = r.text;
      s.style.fontFamily = `'${r.latin}', '${r.ea}'`;
      s.style.fontSize = r.size + 'px';
      s.style.fontWeight = r.bold ? '700' : '400';
      s.style.fontStyle = r.italic ? 'italic' : 'normal';
      s.style.letterSpacing = (r.spacing || 0) + 'px';
      box.appendChild(s);
    });
    return box.getBoundingClientRect().width;
  });
  box.remove();
  return out;
}
"""

#: PPTX 글자를 이보다 더 줄이지는 않는다(맑은 고딕이 프리텐다드보다 한글 문장에서 약 17% 넓다 · 10-26 잼).
PPTX_SCALE_FLOOR = 0.8


def _uses_bundled_font(family: str) -> bool:
    from .slides_style import bundled_family_names

    names = [part.strip().strip("'\"").lower() for part in str(family or "").split(",") if part.strip()]
    return bool(names) and names[0] in bundled_family_names()


def fit_pptx_fonts(page: Any, blocks: list[dict[str, Any]]) -> int:
    """앱 글꼴로 그린 글 상자의 PPTX 글자 크기를 맞춘다(2026-10-26 · 사다리 0단). 고친 상자 수.

    줄은 브라우저가 앱 글꼴(프리텐다드 등)의 폭으로 바꿨는데, PPTX 에는 받는 PC 에 있는 윈도우 글꼴을 적는다(`font_pair`).
    윈도우 글꼴이 한글에서 10~17% 넓어(10-26 잼) 줄이 상자 밖으로 나가 옆 글에 닿았다(운영 장 대조: 손 결 장 거의 전부).
    그래서 그 상자만 **PPTX 글꼴로 줄 폭을 다시 재서** 브라우저가 그린 폭에 들게 글자를 같은 비율로 줄인다(`pptx_scale`).
    앱 글꼴을 쓰지 않는 상자(옮긴 템플릿 · IR)는 손대지 않는다 — 브라우저와 PPTX 가 같은 글꼴이다.
    """
    lines: list[list[dict[str, Any]]] = []
    owner: list[int] = []
    for i, block in enumerate(blocks):
        if not any(_uses_bundled_font(str(run.get("font") or "")) for run in block.get("runs") or [] if not run.get("br")):
            continue
        current: list[dict[str, Any]] = []
        for run in list(block.get("runs") or []) + [{"br": True}]:
            if run.get("br"):
                if current:
                    lines.append(current)
                    owner.append(i)
                current = []
                continue
            latin, east = font_pair(str(run.get("font") or ""))
            current.append({"text": str(run.get("text") or ""), "latin": latin, "ea": east, "size": float(run.get("size") or 16),
                            "bold": int(run.get("weight") or 400) >= 600, "italic": bool(run.get("italic")),
                            "spacing": float(run.get("spacing") or 0)})
    if not lines:
        return 0
    scale: dict[int, float] = {}
    # 글자 크기에 폭이 딱 비례하지는 않는다(픽셀 맞춤 · 힌팅) — 줄인 크기로 다시 재서 아직 넘치면 더 줄인다(세 번까지).
    for _ in range(3):
        sized = [[{**run, "size": run["size"] * scale.get(i, 1.0), "spacing": run["spacing"] * scale.get(i, 1.0)} for run in line]
                 for i, line in zip(owner, lines)]
        widths = page.evaluate(_PPTX_WIDTH_JS, sized)
        widest: dict[int, float] = {}
        for i, width in zip(owner, widths):
            widest[i] = max(widest.get(i, 0.0), float(width or 0))
        changed = False
        for i, width in widest.items():
            room = float(blocks[i]["w"]) + 1
            now = scale.get(i, 1.0)
            if width > room and now > PPTX_SCALE_FLOOR:
                scale[i] = max(PPTX_SCALE_FLOOR, now * room / width * 0.995)
                changed = True
        if not changed:
            break
    for i, value in scale.items():
        blocks[i]["pptx_scale"] = math.floor(value * 1000) / 1000
    return len(scale)


def _set_fonts(run: Any, latin: str, east_asian: str) -> None:
    """라틴(`a:latin`)과 한글(`a:ea`)을 따로 정한다 — python-pptx 의 font.name 은 라틴만 바꾼다."""
    from pptx.oxml.ns import qn

    rpr = run._r.get_or_add_rPr()
    for tag, name in (("a:latin", latin), ("a:ea", east_asian), ("a:cs", latin)):
        node = rpr.find(qn(tag))
        if node is None:
            node = rpr.makeelement(qn(tag), {})
            rpr.append(node)
        node.set("typeface", name)


def _bundled_first(family: str) -> str:
    """CSS 글꼴 목록의 첫 이름이 앱에 든 글꼴이면 그 이름, 아니면 빈 글."""
    from .slides_style import bundled_family_names

    names = [part.strip().strip("'\"") for part in str(family or "").split(",") if part.strip()]
    return names[0] if names and names[0].lower() in bundled_family_names() else ""


def fonts_used(slides: list[dict[str, Any]]) -> dict[str, dict[bool, str]]:
    """장마다 앱 글꼴로 쓴 글자 — {가족: {굵게?: 글자들}} (PPTX 에 심을 글자)."""
    used: dict[str, dict[bool, set[str]]] = {}
    for item in slides:
        for block in item.get("blocks") or []:
            for run in block.get("runs") or []:
                family = "" if run.get("br") else _bundled_first(str(run.get("font") or ""))
                if family and run.get("text"):
                    bold = int(run.get("weight") or 400) >= 600
                    used.setdefault(family, {}).setdefault(bold, set()).update(str(run["text"]))
    return {family: {bold: "".join(sorted(chars)) for bold, chars in by.items()} for family, by in used.items()}


def build_pptx(slides: list[dict[str, Any]], out_path: Path, *, embed_fonts: bool = False) -> dict[str, Any]:
    """[{background: png 경로, blocks: [...]}] → PPTX. 배경 그림 한 장 + 고칠 수 있는 글 상자.

    ★ 줄은 **브라우저가 바꾼 자리 그대로** 박는다(2026-10-25). 예전엔 줄 수만 세어 PowerPoint 에 줄바꿈을 맡겼는데,
    큰 제목은 두 줄이 겹쳐 「한 줄」로 세어져 오른쪽으로 잘렸고(표지 · 마무리), PowerPoint 는 한국어 낱말 중간에서 줄을 바꿨다.
    이제 글 뽑기(_EXTRACT_JS)가 글자마다 위치를 재서 줄바꿈 자리를 넘기고, 여기서는 자동 줄바꿈을 끈다.
    """
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Pt

    deck = Presentation()
    deck.slide_width = Emu(WIDTH * EMU_PER_PX)
    deck.slide_height = Emu(HEIGHT * EMU_PER_PX)
    blank = deck.slide_layouts[6]
    # ★ 앱 글꼴을 쓴 글자만 PPTX 에 심는다(2026-10-26 · `pptx_fonts`). 심은 가족은 글 상자에 그 이름을 그대로 적고,
    #   글자를 줄이지 않는다(받는 PC 에서도 같은 글꼴 · 같은 폭). 못 심은 가족은 예전처럼 윈도우 글꼴로 적는다.
    embedded: set[str] = set()
    font_report: dict[str, Any] = {}
    if embed_fonts:
        from . import pptx_fonts

        font_report = pptx_fonts.embed(deck, fonts_used(slides))
        embedded = {entry.split("/", 1)[0].lower() for entry in font_report.get("embedded") or []}
    align_of = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT, "justify": PP_ALIGN.LEFT}
    boxes = 0
    scaled = 0
    for item in slides:
        slide = deck.slides.add_slide(blank)
        slide.shapes.add_picture(str(item["background"]), 0, 0, deck.slide_width, deck.slide_height)
        for block in item.get("blocks") or []:
            align = str(block.get("align") or "")
            # PowerPoint 글자 폭이 조금 넓을 수 있어 상자를 넓힌다 — 맞춤 쪽은 그대로 두고 반대쪽으로(가운데면 양쪽으로).
            extra = block["w"] * 0.08 + 8
            x = block["x"] - (extra / 2 if align == "center" else extra if align in {"right", "end"} else 1)
            x = max(0.0, x)
            w = min(WIDTH - x, block["w"] + extra + 1)
            line_px = float(block.get("line_px") or 0) or float(block.get("size") or 16) * 1.25
            h = max(block["h"], 4)
            box = slide.shapes.add_textbox(Emu(int(x * EMU_PER_PX)), Emu(int(block["y"] * EMU_PER_PX)),
                                           Emu(int(w * EMU_PER_PX)), Emu(int(h * EMU_PER_PX)))
            frame = box.text_frame
            frame.word_wrap = False
            frame.auto_size = None
            frame.vertical_anchor = MSO_ANCHOR.TOP
            frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
            paragraph = frame.paragraphs[0]
            paragraph.alignment = align_of.get(align, PP_ALIGN.LEFT)
            paragraph.line_spacing = Pt(line_px * 0.75)
            ppr = paragraph._p.get_or_add_pPr()
            ppr.set("eaLnBrk", "0")  # 한국어 낱말을 반으로 자르지 않는다(줄은 이미 정해 넣었다)
            # 앱 글꼴 상자를 윈도우 글꼴 폭에 맞춘 비율(fit_pptx_fonts) — 그 글꼴을 심었으면 줄이지 않는다.
            scale = float(block.get("pptx_scale") or 1)
            if embedded and any(_bundled_first(str(r.get("font") or "")).lower() in embedded for r in block["runs"] if not r.get("br")):
                scale = 1.0
            for run_data in block["runs"]:
                if run_data.get("br"):
                    paragraph.add_line_break()
                    continue
                text = str(run_data.get("text") or "")
                if not text:
                    continue
                run = paragraph.add_run()
                run.text = text
                font = run.font
                font.size = Pt(max(6.0, float(run_data.get("size") or 16) * 0.75 * scale))
                font.bold = int(run_data.get("weight") or 400) >= 600
                font.italic = bool(run_data.get("italic"))
                font.color.rgb = RGBColor(*_rgb(run_data.get("color", "")))
                spacing = float(run_data.get("spacing") or 0) * scale
                if abs(spacing) >= 0.2:  # 자간(px) → 1/100 pt
                    run._r.get_or_add_rPr().set("spc", str(int(round(spacing * 0.75 * 100))))
                _set_fonts(run, *font_pair(str(run_data.get("font") or ""), embedded))
            boxes += 1
            scaled += scale < 1
        if item.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(item["notes"])[:4000]
    out_path.parent.mkdir(parents=True, exist_ok=True)
    deck.save(str(out_path))
    shape = {"slides": len(slides), "text_boxes": boxes}
    if scaled:
        shape["pptx_scaled_boxes"] = scaled
    if font_report.get("embedded"):
        shape["embedded_fonts"] = font_report["embedded"]
        shape["embedded_font_bytes"] = font_report["bytes"]
    if font_report.get("errors"):
        shape["font_embed_errors"] = font_report["errors"]
    return shape


#: 넘침 재기 + 맞추기. 결 조각(<section>)은 글자가 전부 em 이라 장 기본 크기 하나만 줄이면 다 같이 준다.
_FIT_JS = r"""
(fit) => {
  const sec = document.querySelector('section') || document.body;
  const W = 1280, H = 720;
  const hasOwnText = (el) => Array.from(el.childNodes).some((n) => n.nodeType === 3 && n.textContent.trim());
  // 글을 품은 「상자」 — 바탕색 · 바탕 그림 · 테두리가 있는 가장 가까운 조상(장 자신은 빼고 장 경계로 따로 잰다).
  const boxOf = (el) => {
    for (let e = el.parentElement; e && e !== sec; e = e.parentElement) {
      const cs = getComputedStyle(e);
      const bg = cs.backgroundColor && !/rgba\(\d+, \d+, \d+, 0\)|transparent/.test(cs.backgroundColor);
      if (bg || (cs.backgroundImage && cs.backgroundImage !== 'none') || parseFloat(cs.borderTopWidth) > 0 && parseFloat(cs.borderBottomWidth) > 0) return e;
    }
    return null;
  };
  const tag = (el) => el.tagName.toLowerCase() + (el.className ? '.' + String(el.className).split(' ')[0] : '');
  // ★ 글 상자(요소의 테두리 상자)로 잰다 — 글자 윤곽(큰 숫자는 줄 높이보다 크다)으로 재면 넘침이 아닌데 넘침으로 센다(10-25 시험대 N1).
  // 거의 투명한 장식 글(워터마크 번호 · 물결 글자)은 넘침으로 세지 않는다 — 굽기도 이 글은 뽑지 않는다(같은 문턱 0.35).
  //   grove 의 표지 · 마무리 장 「01」(투명도 6%, 원본부터 장 아래로 걸침)이 판마다 「59px 넘친다」로 돌려보냈다(10-25 밤).
  const faint = (el, cs) => {
    const m = String(cs.color).match(/rgba\([^)]*,\s*([\d.]+)\)/);
    let a = m ? parseFloat(m[1]) : 1;
    for (let e = el; e && e.nodeType === 1; e = e.parentElement) a *= parseFloat(getComputedStyle(e).opacity || '1');
    return a < 0.35;
  };
  // 옮긴 템플릿의 틀째 장은 쪽 번호 · 바닥 줄을 원래 장 아래 40px 안에 둔다 — 40px 여백은 우리 결의 규칙이다.
  //   원본 190장 중 93장이 손대지 않고도 「넘친다」였고(쪽 번호 44 · 바닥 글 34, 중앙 16px), 카페 판 7번 장이 이것으로 막혔다(10-25 밤).
  //   그 장은 실제 장 경계로 잰다. 칸 안 넘침(아래 상자 검사)은 그대로다. 부품 장(pz-slide)은 우리 규칙 그대로.
  const borrowed = !sec.classList.contains('pz-slide')
    && Array.from(document.querySelectorAll('style')).some((s) => s.textContent.includes('/* sonjari-template */'));
  const floorY = borrowed ? H : H - 40, floorX = borrowed ? W : W - 16;
  const measure = () => {
    let over = 0, where = '';
    sec.querySelectorAll('*').forEach((el) => {
      // 장식(.pz-deco · 표지 부품의 무늬 · 본문 무늬 한 점) 안의 글(쪽 번호 · No.)은 글이 아니다 — 넘침으로 재지 않는다(10-26).
      if (el.closest('.folio, .pz-deco') || !hasOwnText(el)) return;
      const cs = getComputedStyle(el);
      if (cs.display === 'none' || cs.visibility === 'hidden' || faint(el, cs)) return;
      const r = el.getBoundingClientRect();
      if (!r.width && !r.height) return;
      let o = Math.max(r.bottom - floorY, r.right - floorX, 0);
      const box = boxOf(el);
      if (box) {
        const b = box.getBoundingClientRect();
        o = Math.max(o, r.bottom - b.bottom - 2, r.right - b.right - 2);
      }
      // 부품 장의 칸 안 글은 칸 영역(.pz-grid) 바닥 밑으로 내려가면 넘침이다 — 그 밑은 결론 띠 · 바닥 줄(출처) 자리다.
      //   10-26 16:09 실제 판(5c0a3d8d 4번 장): 키우기가 오른쪽 칸의 「시사점」 큰 글을 1.5배로 키워 바닥 가로선 · 출처 줄까지 내려갔는데,
      //   장 바닥(720-40)만 보던 재기는 0 으로 셌고(글끼리 가로로 떨어져 겹침도 아님) 그대로 구워졌다.
      const grid = el.closest('.pz-grid');
      if (grid) {
        const gb = grid.getBoundingClientRect();
        o = Math.max(o, r.bottom - gb.bottom - 2);
      }
      if (el.scrollWidth - el.clientWidth > 2 && cs.whiteSpace === 'nowrap') o = Math.max(o, el.scrollWidth - el.clientWidth);
      if (o > over) { over = o; where = tag(el); }
    });
    return { over: Math.round(over), where };
  };
  // ★ 글 겹침 (2026-10-26 c81a10dd 표지) — 넘침 재기는 「장 · 칸 밖으로」만 봐서, 절대 위치로 놓인 글끼리 겹친 것을 못 잡았다
  //   (짧은 꼬리말 칸에 긴 문장이 들어가 아래 작은 바닥 글을 덮었다). 글 줄의 상자끼리, 그리고 글과 「글이 오면 안 되는 장식 자리」
  //   (표지 부품의 .pz-deco[data-zone])의 겹침을 잰다. 줄 상자는 글자 윤곽이 아니라 줄 높이로 — 줄 간격이 좁은 큰 제목의
  //   글자 윤곽은 아래 글에 몇 px 걸쳐도 겹친 게 아니다. 거의 투명한 장식 글(0.35 밑) · 숨긴 장식(aria-hidden)은 넘침처럼 뺀다.
  const lineBoxes = (el) => {
    const cs = getComputedStyle(el);
    const fs = parseFloat(cs.fontSize) || 16;
    const lhRaw = parseFloat(cs.lineHeight);
    const lh = isNaN(lhRaw) ? fs * 1.2 : lhRaw;
    const out = [];
    el.childNodes.forEach((n) => {
      if (n.nodeType !== 3 || !n.textContent.trim()) return;
      const range = document.createRange();
      range.selectNodeContents(n);
      Array.from(range.getClientRects()).forEach((x) => {
        if (x.width < 1 || x.height < 1) return;
        const h = Math.min(x.height, lh), cy = (x.top + x.bottom) / 2;
        out.push({ l: x.left, r: x.right, t: cy - h / 2, b: cy + h / 2 });
      });
    });
    return out;
  };
  const overlapMeasure = () => {
    const els = Array.from(sec.querySelectorAll('*')).filter((el) => {
      if (!hasOwnText(el) || el.closest('.pz-deco, [aria-hidden="true"], svg')) return false;
      const cs = getComputedStyle(el);
      return cs.display !== 'none' && cs.visibility !== 'hidden' && !faint(el, cs);
    });
    const boxes = els.map((el) => [el, lineBoxes(el)]);
    sec.querySelectorAll('.pz-deco[data-zone]').forEach((z) => {
      const r = z.getBoundingClientRect();
      if (r.width && r.height) boxes.push([z, [{ l: r.left, r: r.right, t: r.top, b: r.bottom }]]);
    });
    let over = 0, pair = null;
    for (let i = 0; i < boxes.length; i += 1) {
      for (let j = i + 1; j < boxes.length; j += 1) {
        const [a, ra] = boxes[i], [b, rb] = boxes[j];
        if (a.contains(b) || b.contains(a) || (a.dataset && 'zone' in a.dataset && b.dataset && 'zone' in b.dataset)) continue;
        for (const x of ra) for (const y of rb) {
          const w = Math.min(x.r, y.r) - Math.max(x.l, y.l), h = Math.min(x.b, y.b) - Math.max(x.t, y.t);
          if (w > 3 && h > 3 && Math.min(w, h) > over) { over = Math.min(w, h); pair = [a, b]; }
        }
      }
    }
    return { over: Math.round(over), pair, where: pair ? pair.map(tag).join(' × ') : '' };
  };
  const first = measure();
  const overlapFirst = overlapMeasure();
  // 1) 한 줄짜리 큰 숫자 · 낱말(.v · .num · .val · .big · .was)이 제 칸보다 넓으면 **그것만** 줄인다(장 전체가 아니라).
  if (fit) {
    // 부품 장의 숫자 칸(.n)도 — 빠져 있어 좁은 숫자 카드의 「4,820만 원」이 48px 넘쳐 돌려보내졌다(10-26 sakura 판 둘).
    sec.querySelectorAll('.v, .num, .val, .big, .was, .n').forEach((el) => {
      const start = parseFloat(getComputedStyle(el).fontSize);
      let size = start;
      const parent = el.parentElement.getBoundingClientRect();
      const wide = () => el.scrollWidth - el.clientWidth > 1 || el.getBoundingClientRect().right > Math.min(parent.right, W - 16) + 1;
      while (wide() && size > start * 0.4) { size *= 0.94; el.style.fontSize = size.toFixed(1) + 'px'; }
    });
    // 숫자 띠는 칸마다 크기가 다르면 어수선하다 — 가장 작게 줄어든 크기로 맞춘다.
    const vs = Array.from(sec.querySelectorAll('.stat .v'));
    if (vs.some((el) => el.style.fontSize)) {
      const least = Math.min(...vs.map((el) => parseFloat(getComputedStyle(el).fontSize)));
      vs.forEach((el) => { el.style.fontSize = least.toFixed(1) + 'px'; });
    }
    // 큰 숫자가 너무 넓어 옆 설명이 눌리면(240px 미만) 숫자를 줄인다.
    const num = sec.querySelector('.bignum .num'), desc = sec.querySelector('.bignum .desc');
    if (num && desc) {
      const start = parseFloat(getComputedStyle(num).fontSize);
      let size = start;
      while (desc.getBoundingClientRect().width < 240 && size > start * 0.5) { size *= 0.94; num.style.fontSize = size.toFixed(1) + 'px'; }
    }
  }
  // 2) 그래도 넘치면 장 기본 글자(em 의 바탕) 하나를 줄인다 — 글자가 전부 em 이라 같이 준다.
  let base = parseFloat(getComputedStyle(sec).fontSize) || 20;
  let now = measure();
  // 1b) 부품 장의 칸이 칸 영역 밑으로 넘치면 **그 칸의 글만** 큰 것부터 6% 씩, 처음의 70% · 11px 까지 줄인다(사다리 0단 · 10-26).
  //     부품 CSS 는 px 라 2) 장 기본 글자가 안 먹고, 3) 큰 글자 줄이기는 장 제목까지 줄인다 — 넘친 칸만 손댄다.
  let columnShrunk = 0;
  const pzGrid = sec.querySelector('.pz-grid');
  if (fit && pzGrid && now.over > 0) {
    const g = pzGrid.getBoundingClientRect();
    const colStart = new Map();
    for (let k = 0; k < 14 && now.over > 0; k += 1) {
      let touched = false;
      Array.from(pzGrid.querySelectorAll(':scope > .pz-col')).forEach((col) => {
        const els = Array.from(col.querySelectorAll('*')).filter((el) => hasOwnText(el) && !el.closest('.pz-deco'));
        if (!els.length || Math.max(...els.map((el) => el.getBoundingClientRect().bottom)) <= g.bottom + 2) return;
        const biggest = Math.max(...els.map((el) => parseFloat(getComputedStyle(el).fontSize)));
        els.forEach((el) => {
          const size = parseFloat(getComputedStyle(el).fontSize);
          if (!colStart.has(el)) colStart.set(el, size);
          // 큰 글(인용 · 시사점 · 큰 숫자)부터 — 그 칸에서 가장 큰 글의 85% 이상인 글을 먼저, 그다음 나머지.
          const floor = Math.max(11, colStart.get(el) * 0.7);
          if (size > floor && (size >= biggest * 0.85 || k >= 4)) { el.style.fontSize = Math.max(floor, size * 0.94).toFixed(1) + 'px'; touched = true; }
        });
      });
      if (!touched) break;
      columnShrunk += 1;
      now = measure();
    }
  }
  if (fit && sec.tagName === 'SECTION') {
    const em = parseFloat(getComputedStyle(sec).fontSize);
    while (now.over > 0 && base > 14) { base = Math.round((base - 0.5) * 10) / 10; sec.style.fontSize = base + 'px'; now = measure(); }
    // 3) 그래도 넘치면(옮긴 템플릿은 글자가 px 라 위가 안 먹는다) 큰 글자(26px 넘는 것)를 6% 씩, 처음의 55% 까지.
    let round = 0;
    const big = () => Array.from(sec.querySelectorAll('*')).filter((el) => hasOwnText(el) && !el.closest('.pz-deco') && parseFloat(getComputedStyle(el).fontSize) > 26);
    const start = new Map(big().map((el) => [el, parseFloat(getComputedStyle(el).fontSize)]));
    while (now.over > 0 && round < 12) {
      round += 1;
      start.forEach((first, el) => {
        const size = parseFloat(getComputedStyle(el).fontSize);
        if (size > first * 0.55) el.style.fontSize = (size * 0.94).toFixed(1) + 'px';
      });
      now = measure();
    }
    if (round) base = Math.min(base, em);
  }
  // 3b) 글이 겹치면 코드가 먼저 고친다(사다리 0~1단) — 겹친 둘 중 글이 긴 쪽(장식 자리와 겹치면 그 글)을 8% 씩, 처음의 55% · 11px 까지.
  //     그쪽이 바닥이면 다른 쪽을 줄인다. 그래도 겹치면 돌려보낸다(make_slide · slide_overlap).
  let overlapNow = overlapFirst, overlapFixed = 0;
  if (fit && overlapNow.over > 0) {
    const startSize = new Map();
    const textLen = (el) => (el.dataset && 'zone' in el.dataset) ? -1 : (el.textContent || '').trim().length;
    const canShrink = (el) => {
      if (textLen(el) < 0) return false;
      const size = parseFloat(getComputedStyle(el).fontSize);
      if (!startSize.has(el)) startSize.set(el, size);
      return size > Math.max(11, startSize.get(el) * 0.55);
    };
    for (let k = 0; k < 16 && overlapNow.over > 0; k += 1) {
      const [a, b] = overlapNow.pair;
      const order = textLen(a) >= textLen(b) ? [a, b] : [b, a];
      const victim = order.find(canShrink);
      if (!victim) break;
      victim.style.fontSize = (parseFloat(getComputedStyle(victim).fontSize) * 0.92).toFixed(1) + 'px';
      overlapFixed += 1;
      overlapNow = overlapMeasure();
    }
    now = measure();
  }
  // 4) 부품 장(ir-parts)이 비면 글자를 키운다 (10-25 대표 「여백이 너무 많이 남는다」) — 줄이기의 거울.
  //    부품 칸에서 글이 끝나는 자리가 칸 높이의 72% 에 못 미치면 6% 씩, 처음의 1.5배까지, 넘치기 직전에 멈춘다.
  let grown = 0;
  let emptyRatio = 0;
  const grid = fit ? sec.querySelector('.pz-grid') : null;
  if (grid && now.over === 0) {
    // 사진 · 증빙 칸의 자리 글은 빈 칸을 재는 데 넣지 않는다(칸이 바닥까지 차 있어 「빈 곳 0」으로 보였다 · 10-25 밤 grove 판 4번 장).
    const texts = Array.from(grid.querySelectorAll('*')).filter((el) => hasOwnText(el) && !el.closest('.pz-shot, .pz-photo, [data-slot]'));
    const start = new Map(texts.map((el) => [el, parseFloat(getComputedStyle(el).fontSize)]));
    const inlineBefore = new Map(texts.map((el) => [el, el.style.fontSize]));
    const g = grid.getBoundingClientRect();
    const empty = () => (g.bottom - Math.max(...texts.map((el) => el.getBoundingClientRect().bottom))) / Math.max(g.height, 1);
    // 키우기는 **가장 빈 칸**으로 잰다 — 한 칸이 차 있어도 옆 칸 글이 작게 남지 않게. 돌려보내기(empty)는 장 전체 그대로.
    const cols = Array.from(grid.querySelectorAll(':scope > .pz-col'));
    const emptiest = () => Math.max(0, ...cols.map((col) => {
      const own = texts.filter((el) => col.contains(el));
      return own.length ? (g.bottom - Math.max(...own.map((el) => el.getBoundingClientRect().bottom))) / Math.max(g.height, 1) : 0;
    }));
    const apply = (k) => start.forEach((size, el) => { el.style.fontSize = (size * Math.pow(1.06, k)).toFixed(1) + 'px'; });
    while (texts.length && grown < 7 && Math.max(empty(), emptiest()) > 0.28) {
      grown += 1;
      apply(grown);
      if (measure().over > 0 || overlapMeasure().over > 0) { grown -= 1; apply(grown); break; }
    }
    // 못 키웠으면 **키우기 직전의 크기로** 되돌린다 — 지우면 1단계에서 줄여 둔 숫자 크기까지 날아가 다시 넘쳤다(10-26 sakura 판 「4,820만 원」 48px).
    if (!grown) start.forEach((_, el) => { const kept = inlineBefore.get(el); if (kept) el.style.fontSize = kept; else el.style.removeProperty('font-size'); });
    now = measure();
    // 「비었다」(돌려보내기) 판정은 사진 · 증빙 칸을 **채운 자리**로 센다 — 사람이 사진을 넣을 칸이다. 키우기(위)는 글만 본다.
    //   10-26 3판씩 재기: 비었다 27번 중 9번이 사진 칸 장이었다(키우기를 고치며 같은 목록을 써서 사진 칸이 빈 곳으로 세어졌다).
    const slots = Array.from(grid.querySelectorAll('.pz-photo-box, .pz-shot, [data-slot]'));
    const filledBottom = Math.max(...texts.map((el) => el.getBoundingClientRect().bottom), ...slots.map((el) => el.getBoundingClientRect().bottom));
    emptyRatio = texts.length ? Math.max(0, (g.bottom - filledBottom) / Math.max(g.height, 1)) : 0;
  }
  let smallest = 99;
  sec.querySelectorAll('*').forEach((el) => {
    if (el.closest('.folio, .pz-deco') || !hasOwnText(el)) return;
    smallest = Math.min(smallest, parseFloat(getComputedStyle(el).fontSize) || 99);
  });
  const own = Array.from(sec.querySelectorAll('.v, .num, .val, .big, .was, .n')).filter((el) => el.style.fontSize).map((el, i) => [i, el.style.fontSize]);
  const squeezed = Array.from(sec.querySelectorAll('[style*="font-size"]')).length;
  if (fit && grown) overlapNow = overlapMeasure();
  // 5) 본문 무늬 한 점(.pz-mark[data-optional])은 글이 다 자리 잡은 뒤에 본다 — 글 · 사진 칸에 6px 안으로 닿으면 지운다(글이 먼저 · 10-26 대표).
  //    장식 때문에 글을 줄이거나 옮기지 않는다. 지운 장은 문서째 저장된다(fit_slide · deco_dropped).
  let decoDropped = 0, decoKept = 0;
  if (fit) {
    const marks = Array.from(sec.querySelectorAll('.pz-deco[data-optional]'));
    if (marks.length) {
      const texts = Array.from(sec.querySelectorAll('*')).filter((el) => hasOwnText(el) && !el.closest('.pz-deco, svg'));
      const solids = [...texts.flatMap((el) => lineBoxes(el)),
        ...Array.from(sec.querySelectorAll('.pz-photo-box, .pz-shot, [data-slot]')).map((el) => { const r = el.getBoundingClientRect(); return { l: r.left, r: r.right, t: r.top, b: r.bottom }; })];
      marks.forEach((mark) => {
        const m = mark.getBoundingClientRect(), pad = 6;
        const hit = solids.some((x) => Math.min(x.r, m.right + pad) - Math.max(x.l, m.left - pad) > 0 && Math.min(x.b, m.bottom + pad) - Math.max(x.t, m.top - pad) > 0);
        if (hit) { mark.remove(); decoDropped += 1; } else { decoKept += 1; }
      });
    }
  }
  return { column_shrunk: columnShrunk, overflow_before: first.over, overflow: now.over, where: now.where || first.where, base_px: base, smallest_px: smallest, shrunk: own.length, squeezed, grown, empty: Math.round(emptyRatio * 100) / 100,
           overlap_before: overlapFirst.over, overlap: overlapNow.over, overlap_where: overlapNow.where || overlapFirst.where, overlap_fixed: overlapFixed,
           deco_dropped: decoDropped, deco_kept: decoKept };
}
"""


#: 옮긴 템플릿의 등장 애니메이션을 끝난 상태로 — 사진 · PPTX · 재기용(화면 미리보기에서는 그대로 움직인다).
STATIC_CSS = ("*,*::before,*::after{animation-duration:0s!important;animation-delay:0s!important;"
              "transition-duration:0s!important;transition-delay:0s!important}")


def _settle(page: Any) -> None:
    page.evaluate("(css) => { const s = document.createElement('style'); s.setAttribute('data-sonjari-static', ''); "
                  "s.textContent = css; document.head.appendChild(s); }", STATIC_CSS)


def _serve_fonts(page: Any) -> None:
    """앱에 든 글꼴 요청(`slides_style.FONT_ORIGIN`)만 이 폴더의 파일로 돌려준다 — 밖으로 나가는 요청은 없다(2026-10-26).

    목록(`fonts.json`)에 없는 파일은 404. 장은 주소 없는 빈 문서(set_content)라 글꼴 요청이 다른 출처로 보여 CORS 머리를 단다.
    """
    from .slides_style import FONT_ORIGIN, font_file

    def handle(route: Any) -> None:
        path = font_file(route.request.url)
        if path is None:
            route.fulfill(status=404, body="")
            return
        kind = "font/woff2" if path.suffix == ".woff2" else "font/ttf"
        route.fulfill(status=200, body=path.read_bytes(), headers={"Content-Type": kind, "Access-Control-Allow-Origin": "*"})

    page.route(FONT_ORIGIN + "/**", handle)


#: 글꼴이 다 들어올 때까지 기다린다 — 「load」는 글꼴을 안 기다린다. 글꼴은 글을 배치할 때 불리므로 배치를 먼저 시킨다.
_FONTS_READY_JS = "async () => { document.body && document.body.getBoundingClientRect(); await document.fonts.ready; return document.fonts.status; }"


def _open_slide(page: Any, html: str, *, timeout: float = 20_000) -> None:
    """장 HTML 을 열고 애니메이션을 멈추고 글꼴을 기다린다 — 재기 · 굽기 · PDF 가 같은 글꼴로 본다."""
    page.set_content(html, wait_until="load", timeout=timeout)
    _settle(page)
    page.evaluate(_FONTS_READY_JS)


def fit_slide(path: Path) -> dict[str, Any]:
    """장 하나를 브라우저로 재고, 결 조각이면 넘치지 않게 장 기본 글자를 줄여 그 크기를 파일에 박는다. **따로 된 프로세스 안에서** 부른다."""
    from playwright.sync_api import sync_playwright

    html = path.read_text(encoding="utf-8")
    themed = "data-sonjari-theme" in html
    with sync_playwright() as playwright:
        browser = _launch(playwright)
        try:
            page = browser.new_page(viewport={"width": WIDTH, "height": HEIGHT})
            _serve_fonts(page)
            _open_slide(page, html)
            page.wait_for_timeout(120)
            shape = page.evaluate(_FIT_JS, themed)
            template = themed and 'data-sonjari-theme="' in html and "/* sonjari-template */" in html
            # ★ 부품 장에서 키운 글자(grown)도 요소마다 붙어 있다 — 우리 결(paper 등)은 이 길이 없어 키운 크기가 파일에 안 남았다(10-25 밤).
            # 겹침을 코드가 고친 장(overlap_fixed)도 — 줄인 글자가 요소마다 붙어 있다(2026-10-26).
            if ((template and shape.get("squeezed")) or (shape.get("grown") and "pz-slide" in html) or shape.get("overlap_fixed")
                    or ("pz-frame" in html and shape.get("squeezed")) or shape.get("deco_dropped") or shape.get("column_shrunk")):
                # 옮긴 템플릿은 줄인 글자가 요소마다 붙어 있다 — 브라우저가 그린 문서째 저장한다(정적 고정 표는 빼고).
                page.evaluate("() => document.querySelectorAll('style[data-sonjari-static]').forEach((s) => s.remove())")
                path.write_text(page.content(), encoding="utf-8")
                return shape
            if themed:
                sizes = page.evaluate("() => Array.from(document.querySelectorAll('section .v, section .num, section .val, section .big, section .was, section .n')).map((el) => el.style.fontSize || '')")
            else:
                sizes = []
        finally:
            browser.close()
    html = re.sub(r"<style data-sonjari-fit>.*?</style>", "", html, flags=re.S)
    if themed and (shape.get("base_px", 20) < 20 or any(sizes)):
        rules = [f"section{{font-size:{shape['base_px']}px!important}}"] if shape.get("base_px", 20) < 20 else []
        # 큰 숫자 하나만 줄인 것 — 같은 선택자 목록에서 몇 번째인지로 집는다(:is 목록 순서 = 문서 순서).
        marks = [i for i, size in enumerate(sizes) if size]
        if marks:
            html = _mark_shrunk(html, {i: sizes[i] for i in marks})
        lock = "<style data-sonjari-fit>" + "".join(rules) + "".join(
            f"[data-sonjari-fit-{i}]{{font-size:{sizes[i]}!important}}" for i in marks) + "</style>"
        html = re.sub(r"(</head\s*>)", lock + r"\1", html, count=1, flags=re.I)
        path.write_text(html, encoding="utf-8")
    return shape


_FIGURE_CLASS = re.compile(r"""<(\w+)\b([^>]*\bclass\s*=\s*["'][^"']*\b(?:v|num|val|big|was|n)\b[^"']*["'][^>]*)>""", re.I)


def _mark_shrunk(html: str, marks: dict[int, str]) -> str:
    """문서 순서로 n 번째 큰 숫자 칸에 표시를 단다 — 줄인 크기를 그 칸에만 박으려고."""
    body_at = html.lower().find("<section")
    head, body = html[:body_at], html[body_at:]
    count = -1

    def tagger(match: re.Match) -> str:
        nonlocal count
        count += 1
        return f"<{match.group(1)}{match.group(2)} data-sonjari-fit-{count}>" if count in marks else match.group(0)

    return head + _FIGURE_CLASS.sub(tagger, body)


def fit_in_subprocess(path: Path, *, timeout: float = 60.0) -> dict[str, Any]:
    proc = subprocess.run(
        [sys.executable, "-X", "utf8", "-m", "alture_jarvis.slides", "fit", str(Path(path).resolve())],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=timeout,
        cwd=str(Path(__file__).resolve().parent.parent),
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )
    lines = [line for line in (proc.stdout or "").splitlines() if line.startswith("{")]
    if proc.returncode != 0 or not lines:
        raise RuntimeError(f"slide_fit_failed: {(proc.stderr or proc.stdout or '')[-400:]}")
    return json.loads(lines[-1])


def render_deck(deck_dir: Path, out_path: Path) -> dict[str, Any]:
    """장 HTML 들을 배경 그림 + 글 상자 위치로 풀고 PPTX 를 굽는다. **따로 된 프로세스 안에서** 부른다."""
    from playwright.sync_api import sync_playwright

    pages = sorted(deck_dir.glob("slide_*.html"))
    if not pages:
        raise ValueError("no_slides")
    shots = deck_dir / "render"
    shots.mkdir(exist_ok=True)
    slides: list[dict[str, Any]] = []
    pdf_pages: list[bytes] = []
    pdf_error = ""
    from . import pptx_fonts

    # 글꼴을 심을 수 있으면 PPTX 글 상자는 앱 글꼴 그대로(줄이지 않음). 못 심으면 예전처럼 윈도우 글꼴 폭에 맞춰 줄인다.
    embed_fonts = not pptx_fonts.unavailable()
    with sync_playwright() as playwright:
        browser = _launch(playwright)
        try:
            page = browser.new_page(viewport={"width": WIDTH, "height": HEIGHT}, device_scale_factor=2)
            _serve_fonts(page)
            page.emulate_media(media="screen")  # PDF 도 화면과 같은 CSS 로(인쇄용 규칙을 따로 타지 않게)
            for path in pages:
                _open_slide(page, path.read_text(encoding="utf-8"))
                page.wait_for_timeout(250)
                stem = path.stem
                page.screenshot(path=str(shots / f"{stem}_preview.png"))
                # PDF 는 글을 뽑기(_EXTRACT_JS · 글자를 투명하게 바꾼다) **전에** — 글꼴이 박힌 그대로의 장.
                if not pdf_error:
                    try:
                        pdf_pages.append(page.pdf(width=f"{WIDTH}px", height=f"{HEIGHT}px", print_background=True, page_ranges="1",
                                                  margin={"top": "0", "right": "0", "bottom": "0", "left": "0"}))
                    except Exception as exc:  # noqa: BLE001 — PDF 가 안 돼도 PPTX 는 굽는다
                        pdf_error = f"{type(exc).__name__}: {exc}"[:300]
                blocks = page.evaluate(_EXTRACT_JS)
                if not embed_fonts:
                    fit_pptx_fonts(page, blocks)
                background = shots / f"{stem}_background.png"
                page.screenshot(path=str(background))
                slides.append({"background": background, "blocks": blocks, "notes": visible_text(path.read_text(encoding="utf-8"))[:1500]})
        finally:
            browser.close()
    shape = build_pptx(slides, out_path, embed_fonts=embed_fonts)
    # ★ PPTX 옆에 같은 장의 PDF (2026-10-26 · 앱에 든 글꼴). PPTX 에는 쓴 글자만 담은 글꼴을 심지만(`pptx_fonts`),
    #   PowerPoint 가 없는 곳(휴대폰 · 웹 보기)에서도 같은 모양으로 열리게 PDF 를 같이 둔다 — PDF 도 쓴 글자만 담긴다.
    if pdf_error or len(pdf_pages) != len(pages):
        shape["pdf_error"] = pdf_error or "pdf_pages_missing"
        return shape
    try:
        shape["pdf"] = str(_write_pdf(pdf_pages, out_path))
    except Exception as exc:  # noqa: BLE001
        shape["pdf_error"] = f"{type(exc).__name__}: {exc}"[:300]
    return shape


def _pdf_target(out_path: Path) -> Path:
    """PPTX 와 같은 이름의 PDF. 그 이름이 이미 있으면(사람 파일일 수 있다) 덮지 않고 번호를 붙인다."""
    target = out_path.with_suffix(".pdf")
    counter = 2
    while target.exists():
        target = out_path.with_name(f"{out_path.stem} ({counter}).pdf")
        counter += 1
    return target


def _write_pdf(pages: list[bytes], out_path: Path) -> Path:
    """장마다 뽑은 한 쪽짜리 PDF 들을 차례로 한 파일로."""
    import io

    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    for data in pages:
        for one in PdfReader(io.BytesIO(data)).pages[:1]:
            writer.add_page(one)
    target = _pdf_target(out_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("wb") as handle:
        writer.write(handle)
    return target


def render_in_subprocess(deck_dir: Path, out_path: Path, *, timeout: float = 300.0) -> dict[str, Any]:
    """굽기를 따로 된 파이썬 프로세스로 돌린다 — 손자리가 쥔 브라우저(이벤트 루프)와 섞이지 않게."""
    proc = subprocess.run(
        [sys.executable, "-X", "utf8", "-m", "alture_jarvis.slides", "render", str(deck_dir), str(out_path)],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=timeout,
        cwd=str(Path(__file__).resolve().parent.parent),
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )
    lines = [line for line in (proc.stdout or "").splitlines() if line.startswith("{")]
    if proc.returncode != 0 or not lines:
        raise RuntimeError(f"slide_render_failed: {(proc.stderr or proc.stdout or '')[-600:]}")
    return json.loads(lines[-1])


def main(argv: list[str]) -> int:
    if len(argv) == 2 and argv[0] == "fit":
        print(json.dumps(fit_slide(Path(argv[1])), ensure_ascii=False))
        return 0
    if len(argv) == 3 and argv[0] == "render":
        print(json.dumps(render_deck(Path(argv[1]), Path(argv[2])), ensure_ascii=False))
        return 0
    print("usage: python -m alture_jarvis.slides render <deck_dir> <out.pptx>", file=sys.stderr)
    return 2


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
