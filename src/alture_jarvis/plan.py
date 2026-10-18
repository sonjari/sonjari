"""시킨 것을 하나씩 세어 두고, 끝에 대조한다.

에이전트에는 "무엇이 남았는가"를 담는 자리가 없었다. 반응형 도구 루프라 매 걸음
그때그때 정할 뿐이다. 실측(2026-08-10)에서 그 결과는 26스텝 방황이었고, 17~30번째
걸음은 거의 전부 검색이었다 — 이미 가진 것과 아직 없는 것을 구별하지 못하니
가장 싼 행동(검색)을 되풀이했다.

★ 그런데 **LLM 에게 계획을 세우게 하지는 않는다.** 오늘 하루 종일 잡은 병이
"안 한 일을 했다고, 안 본 것을 봤다고" 하는 것이었는데, 계획을 지어내게 하면 그
병이 계획 층에서 다시 난다 — 그리고 계획은 아무도 대조하지 않으므로 더 조용하다.

대신 **프롬프트에서 결정론적으로 뽑을 수 있는 요구만** 뽑는다. 적게 뽑고 확실하게
대조한다. 못 뽑는 요구(문체·구성 같은 것)는 아예 다루지 않는다 — 반쯤 아는 체하는
체크리스트가 없느니만 못하다.

지금 뽑는 것 다섯:

    개수   "공고를 최소 4개는 직접 열어 읽어라"   -> 서로 다른 출처 수로 대조
    사이트 "위시켓과 크몽에서"                    -> 읽은 출처의 호스트로 대조
    행위   "임시저장까지 해줘"                    -> 그 커밋 도구 성공 여부로 대조
    산출물 "엑셀로 만들어줘"                      -> 파일이 실제로 존재하고 열리는지 대조
    검증   "테스트 통과시켜줘"                    -> 실행 도구의 성공 종료 증거로 대조

사이트 요구가 특히 중요하다. 어제 `_goal_shortfall_check` 는 출처를 4개로 세어
통과시켰는데, 그 넷 중 크몽은 하나도 없었고 위시켓 공고도 하나뿐이었다.
**개수만 세면 "아무거나 네 개"로 통과한다.**
"""

from __future__ import annotations

import json
import re

from . import english as _en
from .quantities import ROLE_CONSTRAINT, ROLE_INPUT, ROLE_OUTPUT, counted_object, quantity_role
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from .site_atlas import SEED as _SEED

# 사람이 부르는 이름 -> 실제 호스트 조각.
#
# ★ 2026-08-13: 여기 여덟 줄이 하드코딩이라, 목록에 없는 사이트(다나와·쿠팡·교보문고)는
# "그 사이트에서 읽었나" 대조가 **아예 안 돌았다.** 여덟 곳 밖에서는 범용이 아니었다는
# 뜻이다. 이제 씨앗은 `site_atlas.SEED` 하나로 두고, 자비스가 가 본 곳을 스스로
# 적립한다(`site_atlas.SiteAtlas`). 목록을 두 벌 두면 갈리고, 갈린 쪽은 조용히 틀린다.
#
# 여기 남는 것은 **계획 대조용 정적 뷰**다. 계획은 판이 시작될 때 뽑히므로 그 시점에
# 아는 것만 쓰면 된다 — 배운 것은 판이 끝날 때 검사가 쓴다.
SITE_ALIASES: dict[str, tuple[str, ...]] = {
    name: (host,) for name, host in _SEED.items()
}

# "N개/곳/건/군데". 단위 뒤 부정 후방탐색으로 "4개월"·"3개국"을 걸러낸다.
# ⚠ `agent._COUNT_REQUEST` 와 쌍둥이다 — 세는 말(권·명)을 한쪽만 넣으면 갈린다.
# 가지·종은 일부러 뺐다: 종류를 세는 말이라 "3가지 원칙"이 읽기 요구로 걸린다.
_COUNT = re.compile(
    r"(최소\s*|적어도\s*)?"
    # ★ 2026-08-24: 한글 수사를 넣었다. 전에는 `\d{1,2}` 뿐이라
    #   「세 군데」·「두 곳」·「다섯 개」가 **통째로 안 잡혔다.**
    #
    #   라이브에서 「여수 쪽 괜찮은 국밥집 조사해서 **세 군데만** 알려줘」가
    #   계획 `{}` 으로 끝났다. 개수 요구가 없으니 `_forced_read_call` 이 밀어 줄
    #   것도 없었고, 손자리는 검색 한 번(1걸음)으로 목록 요약만 옮겨 적고 끝냈다.
    #   도구 자신은 "안 읽은 결과 12건이 있다. 들어가라"라고 말하고 있었는데도.
    #
    #   사람은 「3군데」만큼이나 「세 군데」라고 말한다. 숫자만 보는 눈금은
    #   말하는 방식 절반을 못 본다.
    #
    # ★ 앞에 경계를 요구한다. 안 하면 낱말 **속**의 글자가 수사로 읽힌다 —
    #   「인터넷 개발」의 '넷', 「에서 명령」의 '서'. 실측에서 실제로 걸렸다.
    #   그리고 '서·너·석·넉'은 아예 뺐다: 「통신 **서명**」이 3으로 읽히는 것을
    #   막을 방법이 없고("서" 앞이 진짜 띄어쓰기다), 그 수사들은 개·곳·건·군데
    #   앞에 거의 안 온다. 못 잡아서 잃는 것보다 잘못 잡아서 잃는 것이 크다.
    r"(?:^|(?<=[\s,·.｢\)\]]))"
    r"(\d{1,2}|하나|둘|셋|넷|[한두세네]|다섯|여섯|일곱|여덟|아홉|열)"
    # ★ 「순위」를 넣었다 (2026-09-02). 「1순위부터 5순위까지 찾아줘」가 개수 0 이라
    #   첫 판이 **상세 하나만 열고 멈춰도 아무도 안 잡았다** — 사용자가 본
    #   「끈기가 없다」의 정체가 이것이다. 끈기가 아니라 목표가 없었다.
    #   실측: 프롬프트 382개에 새로 잡히는 것 2건(둘 다 진짜), 기존 개수와 충돌 0.
    #   「1순위」는 1 이라 `1 < value` 에서 걸러진다 — 순서를 말하는 것이지
    #   개수가 아니기 때문이다. 「5순위까지」만 5 로 선다.
    # ★ 「글」·「편」을 넣었다 (2026-10-17 운영 판 9f8645bd). 「목록에서 아래 다섯 글을 찾아 하나씩 열어 읽어 줘 ①…⑤ …
    #   글마다 … 댓글 위쪽 세 개」가 단위를 못 알아봐 「다섯 글」을 건너뛰고 「세 개」(글마다 볼 댓글 수)를 집었다 →
    #   「3개 이상 직접 열어 읽기」. 넷을 읽자 「목표 채움」으로 접혀, 마지막 「⑤번을 열겠습니다」 한 줄이 답이 됐다(읽은 넷의 정리 없음).
    #   「세 글자」·「두 글씨」는 개수가 아니다. 숫자와 **띄어 쓴** 「글」·「편」만 본다 — 붙이면 「한글 파일」이 「한(1) 글」,
    #   「한편」이 「한(1) 편」으로 읽힌다(처음엔 안 띄운 것도 봤다가 운영 요청 2,167개 대조에서 「한글 파일」 둘이 바뀌어 좁힘).
    r"\s*(?:개|곳|건|군데|권|명|순위|(?<=\s)편|(?<=\s)글(?![자씨]))(?![월년국인])(\s*이상)?"
)


# 뒤따라오는 **부정**. "발행은 하지 마" 의 '발행' 을 요구로 뽑으면, 시키지도 않은 것을
# 못 했다고 답에 적는다 — 2026-08-20 쓰기 벤치에서 실제로 그랬다. 임시저장까지
# 성공한 판이 "발행 미이행" 때문에 스스로 실패했다고 보고했다.
#
# 넓게 잡지 않는다. 뒤 20자 안의 금지 표현만 본다 — 문장 전체에서 '마' 를 찾으면
# "발행해 줘. 그리고 메일은 보내지 마" 같은 문장에서 발행까지 같이 죽는다.
# ★ 2026-09-23: 「보여주지는 마」·「발행하지는 마」·「보내지 마세요」를 못 봤다 — 「하지 마」 꼴만 있었다.
#   그래서 「임시저장만 해줘. 발행하지는 마」가 **발행을 요구로** 잡았다(안 시킨 일을 못 했다고 적을 자리).
#   `…지(는|도) 마` 를 더했다. 「마」 뒤에 한글이 오면 부정이 아니다(「지마켓」·「마우스」). 넓히기 전에 재 봤다:
#   최근 서로 다른 요청 1,181개에서 요구가 바뀐 것 0개.
_NEGATION = re.compile(r"(하지\s*마|하지\s*말|말고|말아|않(?:게|고|도록)|금지|빼고|제외"
                       r"|지\s*(?:는|도)?\s*(?:마(?![가-힣])|마세요|마라|말아|말고))")
_NEGATION_WINDOW = 20


#: 사람이 **적을 내용을 불러준** 표현 — 「…라고 적어」·「내용은 '…'」.
_DICTATED_CONTENT = re.compile(
    r"(?:이?라고|라는\s*(?:말|문장|글|내용)을?)\s*(?:적어|써|쓰고|쓴|입력|기록|남겨|넣어)"
    r"|내용(?:은|을|으로)?\s*[:：]?\s*['\"「“『]"
    r"|(?:주소|링크|URL|url|문구|텍스트|글자|번호)(?:를|을)?\s*(?:담은|담아|넣은|넣어|인코딩한|인코딩해)"
)
#: 가서 읽을 수 있는 자리의 모양 — 웹 주소·도메인·이 PC 경로(드라이브·UNC).
_READABLE_ADDRESS = re.compile(
    r"https?://|www\.|(?<![\w.])[a-z0-9-]+\.(?:com|net|org|io|co|kr|ai|dev|app|me|gg)\b|(?<![A-Za-z0-9])[A-Za-z]:[\\/]|\\\\",
    re.I,
)


def _names_a_readable_place(compact: str, known_sites: "dict[str, tuple[str, ...]] | None" = None) -> bool:
    """요청이 **실제로 가서 읽을 수 있는 자리**를 댔나 — 아는 사이트 이름(씨앗+배운 것) · 주소 · 이 PC 경로."""
    if _READABLE_ADDRESS.search(compact):
        return True
    names = {**SITE_ALIASES, **(known_sites or {})}
    folded = compact.casefold()
    return any(name.casefold() in folded for name in names)


#: 주소를 **산출물에 넣으라는** 말 — 「robometafit.com 주소를 담은 QR 코드」.
#
# ★ 2026-09-30 QR 판: 「근거가 되는 대상 실제로 읽기」가 붙었다. 주소가 있으니 「출처를 댔다」로 읽혔고, 출처를 댄
#   요청은 위 「불러 준 내용」 예외에서 빠진다. 그런데 이 주소는 **넣을 내용**이지 읽을 출처가 아니다. 넣으라는 말이
#   주소에 직접 붙어 있을 때만 출처 판정을 넘어선다 — 「yes24.com 에서 찾아 링크를 담은 문서」는 읽기 동사가 막는다.
_EMBEDDED_ADDRESS = re.compile(
    r"(?:주소|링크|URL|url)(?:를|을)?\s*(?:담은|담아|넣은|넣어|인코딩한|인코딩해)"
)


def _negated(compact: str, end: int) -> bool:
    """이 자리에서 끝난 표현이 **하지 말라는 말**인가.

    ★ 창을 문장에서 끊는다. 안 끊으면 다음 문장의 금지가 앞 문장까지 죽인다:

        "임시저장까지 해줘. 발행은 하지 마."
                     ↑ 여기서 끊지 않으면 '임시저장' 도 부정으로 읽힌다

    (2026-08-20 실측. 처음 20자만 보게 했더니 정확히 그렇게 됐다 — 시킨 일이
     통째로 사라져 "아무것도 안 시켰다"가 됐다.)
    """
    tail = compact[end : end + _NEGATION_WINDOW]
    for stop in (".", "!", "?", chr(10)):
        cut = tail.find(stop)
        if cut >= 0:
            tail = tail[:cut]
    return bool(_NEGATION.search(tail))


def _counts_something_to_produce(compact: str, match: "re.Match[str]") -> bool:
    """이 숫자가 **내놓을 것**을 세고 있는가.

    ★ 2026-09-03 스모크. 「이 문서를 직접 읽고 핵심 3개를 정리해줘」에서 `3개` 가
    읽기 개수로 잡혀, 지목한 문서 말고 **두 개를 더 읽었다.** 숫자는 자기가 세는
    낱말에 붙는다 — 판정은 `quantities` 한 곳에서만 한다.
    """
    # 매치는 「3개 이상」처럼 뒤 수식까지 먹는다. 단위는 숫자 바로 뒤 한 덩어리다.
    unit_end = match.start(3) if match.group(3) else match.end()
    unit = compact[match.end(2) : unit_end].strip()
    obj, _, _ = counted_object(compact, match.start(2), unit_end)
    return quantity_role(unit, obj, compact[unit_end : unit_end + 6]) == ROLE_OUTPUT


def _has_explicit_input_count(compact: str) -> bool:
    """자료·문서처럼 실제 입력 대상을 세는 수량이 하나라도 있는가."""
    if _en.english_dominant(compact):
        return any(role == _en.ROLE_INPUT for _v, _q, role, _s, _e in _en.counts(compact))
    for match in _COUNT.finditer(compact):
        unit_end = match.start(3) if match.group(3) else match.end()
        unit = compact[match.end(2) : unit_end].strip()
        obj, _, _ = counted_object(compact, match.start(2), unit_end)
        if quantity_role(unit, obj, compact[unit_end : unit_end + 6]) == ROLE_INPUT:
            return True
    return False


#: 개수가 **걸러 내는 조건**인가 — 「댓글이 7개 이상**인** 글」·「후기 10개 이상 **달린** 상품」. 요구가 아니라 대상을 꾸미는 말이다.
#:
#: ★ 2026-10-04 레딧 기준선 판(8faf3a05): 「…글 30개의 제목·댓글 수를 엑셀 파일로 정리해줘 … 댓글이 7개 이상인 글이 30개 중
#:   몇 개인지 알려줘」가 「로컬 자료 7개 이상 직접 열어 읽기」가 됐다. 손자리는 채울 수 없는 요구를 붙들고 447초를 헤매다
#:   파일 없이 예고로 끝났다. 틀린 자리가 둘이었다 — 이 조건의 7, 그리고 아래 「만들 파일」.
_FILTER_CONDITION = re.compile(r"^\s*(?:인|된|되는|넘는|달린|붙은|있는|없는|짜리)(?![가-힣])")

#: **만들 파일**을 가리키는 「파일」 — 「엑셀 파일 「이름.xlsx」로 정리」·「파일로 저장」. 읽을 로컬 출처가 아니다.
#: 「파일을 읽고」·「파일에서」는 그대로 출처다(뒤에 「로/으로」가 올 때만 만들 것으로 본다).
_OUTPUT_FILE = re.compile(r"파일\s*(?:[「\"'][^」\"']{0,80}[」\"'])?\s*(?:으로|로)(?=\s|$)")


def _is_filter_condition(compact: str, match: "re.Match[str]") -> bool:
    if _FILTER_CONDITION.search(compact[match.end() : match.end() + 8]):
        return True
    # 「인원은 2명」·「2명이서」 — 같이 가는 사람 수는 조건값이다(판정은 `quantities` 한 곳).
    unit_end = match.start(3) if match.group(3) else match.end()
    unit = compact[match.end(2) : unit_end].strip()
    obj, _, _ = counted_object(compact, match.start(2), unit_end)
    return quantity_role(unit, obj, compact[unit_end : unit_end + 6]) == ROLE_CONSTRAINT


#: 개수 바로 뒤에 **요청 동사**가 붙었는가 — 「5개를 찾아줘」·「3개만 골라」·「6개 이상 읽고」. 조건절의 개수
#: (「두 개 이상 작업을 걸면」)는 뒤에 다른 말이 끼어 안 걸린다.
_ASKED_DIRECTLY = re.compile(
    r"^\s*(?:를|을|씩|만|정도|쯤|까지)?\s*(?:만\s*)?(?:직접\s*|골라서?\s*|먼저\s*)?"
    r"(?:찾|뽑|골라|고르|알려|정리|비교|조사|읽|추려|보여|추천|열어|모아|가져|수집|확인|요약|분석)"
)


#: 사람이 「열 필요 없다」고 한 요구의 이름표 꼬리. 판정은 글자가 아니라 `Requirement.list_only` 로 한다.
LISTING_LABEL_TAIL = "목록으로 담기"


def _count_label_tail(compact: str) -> str:
    """개수 요구의 이름표 꼬리 — 사람이 「열 필요 없다」고 했으면 「목록으로 담기」, 아니면 「직접 열어 읽기」.

    ★ 2026-10-04 레딧 기준선 판(92b3776a): 사람이 「글을 하나하나 열 필요는 없고 목록에 보이는 숫자만 적어」라고 했는데
      대조표·되받기가 「25개 이상 **직접 열어 읽기**」였다. 충족 판정은 이미 「답·파일에 25줄이면 된다」였지만, 모델이
      듣는 것은 판정이 아니라 **이름표**다 — 판 내내 「열어 읽어라」는 말을 들었고, 답에 「시스템이 본문 읽기를 요구하는
      알림을 보냈습니다」라고 적었다. 이름표가 판정과 같은 말을 하게 한다.
    """
    return LISTING_LABEL_TAIL if says_list_only(compact) else "직접 열어 읽기"


#: 사람이 **직접** 「열지 않아도 된다 · 목록에 보이는 것만」이라고 한 말. 추측이 아니라 지시다.
_LIST_ONLY = re.compile(r"열\s*필요(?:는|가)?\s*없|목록에\s*(?:보이는|나온)\s*[^.!?\n]{0,8}만\s*(?:적|써|옮|정리)")
#: 상세를 **열지 말라**고 한 말 — 「기사 본문은 열지 말고」·「저장소는 열지 말고」·「하나하나 들어가지 않아도」.
#: ★ 2026-10-08 실기(네이버 뉴스 헤드라인 판 ad60e731): 「기사 본문은 열지 말고 목록에 보이는 제목 그대로」에서 「본문」이라는
#:   낱말이 `_WANTS_DETAIL`(상세를 원한다)에 걸려, 열지 말라는 말이 「5개 이상 직접 열어 읽기」가 됐다. 손자리는 기사를 네 번
#:   대신 열었다(13걸음 · 171초). 거절한 대상은 원하는 것이 아니다 — 이 꼴을 지우고 나서 상세 낱말을 본다(`_wants_detail`).
_DETAIL_DECLINED = re.compile(
    r"(?:본문|원문|상세(?:\s*페이지|\s*내용)?|기사|게시글|글|저장소|페이지|링크|항목|하나하나|일일이)"
    r"\s*(?:[은는을를도]|까지는?)?\s*(?:하나하나\s*|일일이\s*|따로\s*|굳이\s*)?"
    r"(?:열|들어가|읽|클릭하|누르)지\s*(?:는|도)?\s*(?:말|마(?![가-힣])|마세요|마라|않아도)"
)


def _wants_detail(compact: str) -> bool:
    """상세(본문·원문·상세 페이지)를 보라는 말이 있는가 — **열지 말라고 한 것**은 뺀다."""
    return bool(_WANTS_DETAIL.search(_DETAIL_DECLINED.sub(" ", compact)))


#: 목록 화면에 보이는 칸 — 이것만 달라고 했으면 하나하나 열 필요가 없다(2026-10-15 · acba5f24).
_LISTING_FIELD = re.compile(
    r"제목|점수|추천|좋아요|댓글\s*수|조회\s*수|순위|등수|게시일|올린\s*날|날짜|작성자|글쓴이|가수|채널"
    r"|주소|링크|URL|종류|분류|유형|플레어"
)
#: 목록 밖의 칸 — 칸 이름에 이것이 들어 있으면 열어야 한다(「자격 요건」 · 「본문 요약」 · 「댓글 요지」).
_DETAIL_FIELD = re.compile(r"요건|자격|쪽수|ISBN|배송비|요약|요지|내용|본문|원문|설명|스펙|사양|리뷰|후기|연봉|급여|복지|우대|작성자와")
#: 개수 — 「10개」 · 「50위」. 그 바로 뒤에 늘어놓은 칸만 본다.
_COUNTED = re.compile(r"\d+\s*(?:개|위|편|건|곡)")
#: 칸 목록이 끝나는 자리 — 요청 동사.
_FIELDS_END = re.compile(r"(?:을|를)?\s*(?:표로|엑셀|정리|알려|적어|적고|써\s*줘|뽑|모아|보여|읽|찾|확인|담아|기록|열어)")
_FIELD_SPLIT = re.compile(r"\s*(?:[·,、/]|와\s|과\s|및\s)\s*")


def _asks_only_listing_fields(compact: str) -> bool:
    """개수 **바로 뒤에** 칸을 둘 이상 늘어놓았고, 그 칸이 **하나도 빠짐없이** 목록 칸인가.

    ★ 2026-10-15 00:24 레딧 판(acba5f24): 「인기 글 10개 — 각각 제목 · 점수 · 댓글 수 · 글 종류 한 단어」가 「10개 이상 직접 열어
      읽기」가 되어 모델이 글 여덟 개를 하나씩 열었다(26걸음).
    ★ 처음엔 요청 **전체**에서 낱말을 찾았다 — 운영 요청 2,109개에 옛·새를 나란히 돌리니 8개가 바뀌었고, 「각 게시물을 열어 댓글
      작성자와 요지」가 「목록으로 담기」가 됐으며, 정작 그 판은 뒤쪽 「② 규칙 페이지를 열어 요약」의 「요약」에 걸려 안 바뀌었다.
      그래서 개수 바로 뒤 칸 목록만 본다. 모르는 칸(「회사명」 · 「가격」 · 「파라미터 크기」)이 하나라도 있으면 지금처럼 연다.
    """
    text = re.sub(r"[(（][^)）]{0,80}[)）]", " ", compact)
    for count in _COUNTED.finditer(text):
        rest = text[count.end() : count.end() + 160]
        end = _FIELDS_END.search(rest)
        stop = re.search(r"[.!?。]\s|[②③④⑤⑥⑦⑧⑨]|$", rest)
        segment = rest[: min(end.start() if end else len(rest), stop.start() if stop else len(rest))]
        fields = [field for field in _FIELD_SPLIT.split(segment) if field.strip(" —-:의씩각각")]
        if len(fields) < 2:
            continue
        if all(_LISTING_FIELD.search(field) and not _DETAIL_FIELD.search(field) for field in fields):
            return True
    return False


def says_list_only(prompt: str) -> bool:
    """사람이 「하나하나 열 필요 없다」고 **말했는가.** 이때만 이름표·안내가 「목록으로 담기」가 된다.

    ★ 2026-10-15: 달라는 칸이 **전부 목록 칸**이어도 참이다(`_asks_only_listing_fields`) — 말하지 않아도 열 필요가 없는 요청.
    ★ `wants_a_count`(목록이 답인 요청)로 가르면 안 된다 — 그쪽은 「표로 정리」만 있어도 참이다. 처음엔 그걸로 갈랐더니
      운영 요청 17개의 이름표가 바뀌었고, 그 안에 「사람인 공고 3개의 **자격 요건**을 비교표로」·「YES24 책 2권의
      **쪽수·ISBN**」이 있었다 — 상세를 열어야 나오는 값인데 「하나하나 열 필요 없다」는 안내가 나갈 뻔했다.
      (교보·사람인 3/3 은 상세를 여는 배선에서 나온 숫자다.)
    """
    compact = " ".join(str(prompt or "").split())
    if _wants_detail(compact):
        return False
    return bool(_LIST_ONLY.search(compact) or _DETAIL_DECLINED.search(compact) or _asks_only_listing_fields(compact))


def _count_target(compact: str) -> int:
    """프롬프트가 **몇 개를 요구하는가.**

    맨 앞 숫자를 그냥 집으면 안 된다. 2026-08-18 라이브에서 "리뷰 0개 신규 판매자가
    …최소 6개를 뽑아줘" 를 던졌더니 앞의 `0개` 가 먼저 잡혀 0 이 되고, 범위 밖이라
    **개수 요구가 통째로 사라졌다.** 요구가 없으면 미달 판정도 시작되지 않으므로
    6개 중 1개를 준 판이 `complete` 로 기록됐다.

    그래서 (1) 전부 훑고 (2) 범위 밖은 버리고 (3) `최소·적어도·이상` 이 붙은 것을
    먼저 본다. 조건을 서술하는 숫자("리뷰 0개", "10만원 이하")와 실제 요구를
    가르는 것은 그 수식어다. 수식어가 하나도 없으면 예전처럼 첫 값을 쓴다.

    ★ 2026-10-04 새벽 레딧 판(df35948f·33cac79a·93f08e95): 그 수식어가 **조건절**에 붙어 있었다 —

        「한 번에 **두 개 이상** 작업을 걸면 느려지는 문제를 … 실사용 글 **5개를 찾아줘**」

      `이상` 이 붙은 2 가 이겨 「2개 이상 직접 열어 읽기」가 됐고, 둘을 읽자 「더 진전 없음」으로 산출물 차례가
      일찍 와서 다섯째 글을 읽는 걸음이 거절됐다. 수식어보다 센 신호가 있다: **요청 동사에 바로 붙은 개수**
      (「5개를 찾아줘」·「6개를 뽑아줘」·「3개 이상 읽고」). 그런 것이 있으면 그것이 요구다.
      운영 요청 1,508개에 옛·새 판정을 나란히 돌려 달라지는 것은 위 세 판뿐이었다.
    """
    direct: list[int] = []
    qualified: list[int] = []
    plain: list[int] = []
    if _en.english_dominant(compact):
        # 영어는 단위 대신 명사로 센다 — 판정은 `english.counts` 한 곳(한국어 `_COUNT` 의 짝).
        for value, is_qualified, role, start, _end in _en.counts(compact):
            if not 1 < value <= 30 or role == _en.ROLE_OUTPUT:
                continue
            if _en.ASKED_DIRECTLY.search(compact[max(0, start - 40) : start]):
                direct.append(value)
            (qualified if is_qualified else plain).append(value)
        if direct:
            return max(direct)
        if qualified:
            return max(qualified)
        return plain[0] if plain else 0
    # 긴 글이거나, **시키는 문장이 따로 있는** 글 — 그때 진술 문장에 든 맨 숫자는 붙여넣은 내용이다.
    pasted_text = len(compact) > _PASTED_TEXT_CHARS or _has_a_request_sentence(compact)

    def direct_here(found: "re.Match[str]") -> bool:
        return bool(_ASKED_DIRECTLY.search(compact[found.end() : found.end() + 16]))

    where: dict[int, "re.Match[str]"] = {}
    for match in _COUNT.finditer(compact):
        value = _as_number(match.group(2))
        if value <= 0:
            continue
        if not 1 < value <= 30:
            continue
        if _counts_something_to_produce(compact, match) or _is_filter_condition(compact, match):
            continue
        where.setdefault(value, match)
        if _ASKED_DIRECTLY.search(compact[match.end() : match.end() + 16]):
            direct.append(value)
        if match.group(1) or match.group(3):
            qualified.append(value)
        elif pasted_text and not direct_here(match) and not _in_a_request_sentence(compact, match):
            # 긴 글의 **진술 문장**에 든 맨 숫자 — 붙여넣은 글의 내용이지 시킨 개수가 아니다(`_count_only_in_a_statement`).
            continue
        else:
            plain.append(value)
    chosen = max(direct) if direct else max(qualified) if qualified else plain[0] if plain else 0
    return min(30, chosen * _groups_for(compact, where[chosen])) if chosen else 0


#: 묶음을 세는 낱말 — 「다섯 주제」 · 「세 사이트」의 뒤쪽.
_GROUP_NOUN = r"주제|사이트|가게|곳|분야|모델|게시판|카테고리|키워드|검색어|나라|도시|회사|커뮤니티"
#: 개수 바로 앞의 「○○마다」 — 「주제마다 댓글이 가장 많은 글 2개」.
_PER_GROUP = re.compile(rf"({_GROUP_NOUN})\s*마다\s*[^\d마]{{0,20}}$")
_CIRCLED = re.compile(r"[①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳]")


def _groups_for(compact: str, match: "re.Match[str]") -> int:
    """이 개수가 **묶음마다**의 개수면 묶음 수, 아니면 1.

    ★ 2026-10-15 a92acbe8: 「다섯 주제를 각각 검색해 … ①~⑤ …, 주제마다 글 2개를 열어」가 「2개 이상 직접 열어 읽기」가 되어
      첫 주제 두 글로 「시킨 것 2/2」가 됐다. 묶음 수가 요청에 드러나 있을 때만 곱한다(①②③… 의 개수 · 「다섯 주제」) — 모르면 1.
    """
    per = _PER_GROUP.search(compact[: match.start()])
    if not per:
        return 1
    noun = per.group(1)
    spelled = re.search(rf"([0-9]+|두|세|네|다섯|여섯|일곱|여덟|아홉|열)\s*(?:개의?\s*|곳의?\s*)?(?:{noun})", compact)
    if spelled:
        value = _as_number(spelled.group(1))
        if 1 < value <= 10:
            return value
    circled = len(set(_CIRCLED.findall(compact)))
    return circled if 1 < circled <= 10 else 1


#: 시키는 말의 끝 — 「…줘」·「…주세요」·「…해라」·「…읽어라」(뒤에 글자가 이어지는 「…어라고」는 아니다). 개수가 든 문장이 요청인지 진술(「…했습니다」)인지 가른다.
_REQUEST_ENDING = re.compile(
    r"(?:줘|주라|주세요|줄래|세요|십시오|해봐|해 봐|부탁|바랍니다|바람|할 것|(?:[아어여]라|해라|하라)(?![가-힣]))"
)
#: 문장이 끊기는 자리. 줄바꿈은 여기 오기 전에 빈칸이 되므로 마침표·글머리로 본다.
# 마침표는 **뒤가 빈칸이거나 끝일 때만** 문장 끝이다 — 「기록.xlsx」·「127.0.0.2」·「3.5」의 점에서 끊으면 요청 문장이 두 동강 난다
# (재 보니 「글 25개의 제목 … 을 기록.xlsx 로 만들어 줘」가 그렇게 끊겨 개수를 잃었다).
_SENTENCE_BREAK = re.compile(r"\.(?=\s|$)|[!?。！？●•▪◦]|\s-\s")
#: 이보다 긴 글에서만 따진다 — 한두 문장짜리 요청은 예전 그대로다.
_PASTED_TEXT_CHARS = 300


def _in_a_request_sentence(compact: str, match: "re.Match[str]") -> bool:
    """이 개수가 든 **문장**이 시키는 말로 끝나는가(「…줘」·「…읽어라」)."""
    breaks = [found.end() for found in _SENTENCE_BREAK.finditer(compact, 0, match.start())]
    nxt = _SENTENCE_BREAK.search(compact, match.end())
    sentence = compact[(breaks[-1] if breaks else 0) : (nxt.end() if nxt else len(compact))]
    return bool(_REQUEST_ENDING.search(sentence))


def _has_a_request_sentence(compact: str) -> bool:
    """문장이 둘 이상이고 그중 하나가 시키는 말로 끝나는가 — 「…넣고 임시저장해줘. ● 9개 분관의 …」.

    한 문장짜리 요청(「삼성전자 뉴스 3개」)은 여기 안 걸린다 — 그 숫자가 곧 요청이다.
    """
    parts = [part for part in _SENTENCE_BREAK.split(compact) if part.strip()]
    return len(parts) >= 2 and any(_REQUEST_ENDING.search(part) for part in parts)


def _count_only_in_a_statement(compact: str) -> bool:
    """긴 글에서, 개수가 **시키는 문장이 아닌 곳**에만 있는가.

    ★ 2026-10-07 대표 실사용. 사람이 자기 자기소개서를 붙여넣고 「각 항목 700자 이상으로 써 줘」라고 했는데,
      본문의 「저는 … **N개 ○○**의 현황을 … 했습니다」 같은 진술에서 「N개 이상 직접 열어 읽기」가 뽑혔다(두 판).
      그 뒤 하네스가 `browser_open_result` 를 대신 부르고 「목록 말고 본문을 읽어라」고 되받았고, 모델은 그 말에
      대답하느라 시킨 일을 놓쳤다. 답 끝에는 「13개를 열어 보라고 하셨는데 하나도 안 열었습니다」가 붙었다.

    붙여넣은 글의 숫자는 **그 글의 내용**이다. 가르는 것은 그 숫자가 든 문장이 시키는 말인가다:
    요청 동사에 바로 붙었거나(`_ASKED_DIRECTLY`) · 「최소·이상」이 붙었거나 · 그 문장이 「…줘/주세요/해라」로 끝나면 요구다.

    ※ 문단 모양으로 「자료」를 잘라내는 규칙을 먼저 재 봤다가 버렸다 — 운영 요청 1,642개 중 39개의 계획이 바뀌었고
      그중 37개가 정상 지시(이어받기 인계문·코딩 과제·목록으로 조건을 적은 조사)의 진짜 요구를 잃었다.
      이 규칙(`_count_target` 이 진술 문장의 맨 숫자를 후보에서 뺀다 — 긴 글이거나 시키는 문장이 따로 있을 때)으로는
      운영 요청 1,644개 중 7개가 바뀐다: 위 두 판 · 붙여넣은 기록·긴 말의 숫자 셋(「실행됨 명령 2개」 등) ·
      「성인 2명 조건으로 봐야 해. 숙소 3곳의 …」가 2 에서 3 으로 바로잡힘 · 「두 가지가 신고됐어」의 행위 개수 하나를 잃음.
      「아래 3권에 대해 각각 반복해줘」·「문서 3개를 각각 직접 읽어라」·「글 25개의 제목을 … 만들어 줘」는 그대로다.
    """
    if _en.english_dominant(compact):
        return False
    if len(compact) <= _PASTED_TEXT_CHARS and not _has_a_request_sentence(compact):
        return False
    seen = False
    for match in _COUNT.finditer(compact):
        value = _as_number(match.group(2))
        if not 1 < value <= 30:
            continue
        if _counts_something_to_produce(compact, match) or _is_filter_condition(compact, match):
            continue
        seen = True
        if match.group(1) or match.group(3) or _ASKED_DIRECTLY.search(compact[match.end() : match.end() + 16]):
            return False
        if _in_a_request_sentence(compact, match):
            return False
    return seen


def _last_word_head(text: str, term: str) -> int:
    """`term` 이 **낱말 머리**에 선 마지막 자리. 바로 앞 글자가 한글이면 다른 낱말의 속이라 건너뛴다.

    ★ 2026-09-29 밤 AWS 첫 실행: 「YES24에서 파이썬 입문서 3권…」의 대조표가 「로컬 자료 3개」가 됐다 —
      「입**문서**」 안의 「문서」를 로컬 출처로 셌다(`rfind` 는 낱말 경계를 안 본다). 프로**파일**도 같다.
    """
    at = text.rfind(term)
    while at > 0 and "가" <= text[at - 1] <= "힣":
        at = text.rfind(term, 0, at)
    return at


#: 「로컬」이 **출처**인 꼴은 몇 개뿐이다 — 뒤에 파일을 담는 것(파일·문서·폴더·자료…)이 오거나 「로컬에 있는」. 그 밖의 「로컬 X」는 X 의 종류다.
#: ★ 2026-10-04: 처음엔 금지 목록이었다(로컬 AI·LLM·모델·에이전트·앱 — 09-30 e7dd3831 「로컬 AI 제품의 회고 글 5개」). 그 목록 밖의
#:   「로컬**용** GGUF 텍스트 생성 모델 … 5개」가 또 「로컬 자료 5개 이상 직접 열어 읽기」가 됐다(5749ff16 — 판 내내 「0개 읽음」, 끝은 부분 완료).
#:   「로컬」 오독만 세 번째다(09-29·09-30·10-04). 종류는 끝이 없고 출처는 몇 개뿐이라 **허용 목록으로 뒤집는다.**
#:   못 알아본 출처는 범위 없는 개수 요구가 된다 — 어디서 읽든 세니 손해가 없다(틀려도 싼 쪽으로 넘어진다).
_LOCAL_NOT_SOURCE = re.compile(
    r"로컬(?!\s*(?:파일|문서|폴더|자료|출처|PC|컴퓨터|드라이브|디스크|저장소|경로|에\s*(?:있|저장|둔)))[^\s,.·]*",
    re.IGNORECASE,
)


def _ascii_name_as_word(folded: str, name: str) -> bool:
    """영문·숫자 이름이 다른 영문·숫자에 **붙지 않은 채** 한 번이라도 나오는가(「16gb」 속의 `gb` 는 아니다)."""
    at = folded.find(name)
    while at >= 0:
        before = folded[at - 1] if at else " "
        after = folded[at + len(name)] if at + len(name) < len(folded) else " "
        if not (before.isascii() and before.isalnum()) and not (after.isascii() and after.isalnum()):
            return True
        at = folded.find(name, at + 1)
    return False


#: 이름 바로 뒤에 오면 그 이름은 **주제**다 — 「깃허브에 관한 규칙」 · 「쿠팡 관련 뉴스」 · 「네이버에 대한 평가」.
_SITE_AS_TOPIC_TAIL = re.compile(r"\s*(?:에\s*(?:관한|관해|관하여|대한|대해|대하여)|관련|얘기|이야기)")


def _only_as_topic(folded: str, name: str) -> bool:
    """이름이 나오는 자리가 **전부** 「…에 관한 · 에 대한 · 관련」 앞인가 — 그러면 갈 곳이 아니라 다룰 주제다.

    ★ 2026-10-18 r/LocalLLM 규칙 읽기 판(c9910364): 「자기 홍보 · 링크 · 프로젝트 소개 · 깃허브에 관한 규칙은 빠짐없이」의
      깃허브가 「깃허브에서 하기」가 됐다. 레딧 규칙을 다 읽고 옮겼는데 「❌ 깃허브에서 하기 — 거기서 아무것도 안 함」으로
      「일부만 끝냄」이 찍혔고, 되받기에 답하다 처음 답의 규칙 전문이 최종 답에서 사라졌다.
    """
    positions = [index for index in range(len(folded)) if folded.startswith(name, index)]
    return bool(positions) and all(_SITE_AS_TOPIC_TAIL.match(folded, index + len(name)) for index in positions)


def _only_in_example_lists(folded: str, name: str) -> bool:
    """이름이 나오는 자리가 **전부 괄호 속**인가 — 「채널(레딧·깃허브 등)」·「(깃허브 참고)」는 장소가 아니라 예시·참고다.

    「…등」으로 끝나는 목록까지 빼 봤다가 되돌렸다: 운영 요청 1,321개에 옛·새 판정을 돌려 보니 「사람인, 잡코리아 등등
    다양한 이력서 사이트를 조사」에서 **진짜 장소인 사람인**이 빠졌다(잡코리아는 남아 앞뒤도 안 맞았다). 괄호만 본다.
    """
    at = folded.find(name)
    if at < 0:
        return False
    while at >= 0:
        if folded.rfind("(", 0, at) <= folded.rfind(")", 0, at):
            return False
        at = folded.find(name, at + len(name))
    return True


#: 요청이 **웹의 어디서** 찾으라고 이름을 댔나 — 주소(`WEB_DOMAIN`)가 없어도 검색 엔진·「검색」·「인터넷」이면 웹이다.
_WEB_PLACE_NAMED = re.compile(r"구글|네이버|유튜브|깃허브|레딧|인터넷|검색(?:해|하|엔진|어|으로)|웹에서|google|naver|bing", re.IGNORECASE)
#: 「파일」이 **읽을 자료**로 쓰인 자리: 앞에 내 PC·로컬·첨부·폴더가 붙었거나, 바로 읽다·열다·분석하다가 따라온다.
#: 「파일을 다룬다」(제품 설명)·「그 파일은 머리줄만 있다」(만들 파일)·「파일로 정리」는 여기에 안 걸린다.
_FILE_AS_SOURCE = re.compile(
    r"(?:(?:내|이)\s*(?:PC|컴퓨터|노트북)\S{0,3}\s*|로컬\s*|첨부\S{0,2}\s*|업로드\S{0,2}\s*|올린\s*|폴더\S{0,6}\s*)파일"
    r"|파일(?:들)?(?:을|를|에서)?\s*(?:\S{1,4}개(?:를|을)?\s*)?(?:읽|열어|열고|분석|요약|비교|참고)"
)


def _last_file_as_source(text: str) -> int:
    """`text` 에서 읽을 자료로 쓰인 마지막 「파일」의 자리(없으면 -1)."""
    at = -1
    for match in _FILE_AS_SOURCE.finditer(text):
        at = max(at, match.start() + match.group(0).find("파일"))
    return at


def _scoped_count_targets(compact: str) -> list[tuple[str, int]]:
    """같은 요청 안의 로컬 자료 수와 웹 자료 수를 서로 다른 요구로 뽑는다.

    ``로컬 문서 3개 + 웹 자료 3개``를 전역 숫자 하나로 만들면 로컬 셋만 읽고도
    웹 셋을 다 읽은 것으로 끝난다. 각 숫자 앞에서 가장 가까운 출처 낱말을 보되,
    같은 범주의 뒤쪽 숫자(강점 3개 등)는 첫 요구 뒤에 버린다.
    """
    from .task_contract import WEB_DOMAIN

    # 주소(`lit.dev`)를 댄 문장의 「문서」는 그 사이트의 문서다 — 로컬 파일로 세면 모델이 정답 페이지를
    # 읽고도 「로컬 파일을 지정해 달라」고 답한다(2026-09-18 밤, task_contract.WEB_DOMAIN 주석).
    # ★ 웹의 어디서 찾으라고 이름을 댄 요청에서는 맨 낱말 「문서」·「파일」이 읽을 로컬 자료가 아니다 (2026-10-06).
    #   「구글과 네이버에서 … 한글 문서를 만들어 주는 AI 에이전트 … 5개 이상」이 「로컬 자료 5개 이상 읽기」가 됐고,
    #   이어서 시킨 판은 「(지금 그 파일은 머리줄만 있다)」의 「파일」 — 만들 파일 자신 — 을 읽을 자료로 셌다.
    #   개수 앞 180자 안에 「구글·네이버」가 없어도(문장이 길면 창 밖이다) 요청 전체가 웹을 가리키면 같은 규칙이다.
    #   그때 「파일」은 **읽을 자료로 쓰인 자리**(`_FILE_AS_SOURCE`)만 센다.
    named_web = bool(WEB_DOMAIN.search(compact)) or bool(_WEB_PLACE_NAMED.search(compact))
    local_terms = ("로컬", "폴더") if named_web else ("로컬", "파일", "문서", "폴더")
    web_terms = ("공개 웹", "웹", "사이트", "온라인", *(m.group(0) for m in WEB_DOMAIN.finditer(compact)))
    found: dict[str, int] = {}
    if _en.english_dominant(compact):
        local_en = _en.LOCAL_SCOPE_TERMS_NAMED_WEB if named_web else _en.LOCAL_SCOPE_TERMS
        web_en = (*_en.WEB_SCOPE_TERMS, *(m.group(0) for m in WEB_DOMAIN.finditer(compact)))
        folded = compact.casefold()
        for value, _q, role, start, end in _en.counts(compact):
            if not 1 < value <= 30 or role == _en.ROLE_OUTPUT:
                continue
            # ★ 영어는 세는 말이 숫자 **뒤**에 온다(「3 local files」) — 앞만 보면 「3 local files … and 3 web sources」에서 첫 3 은
            #   출처가 없고 둘째 3 이 「local files」를 주워 로컬이 된다(2026-10-04, 로컬·웹이 뒤바뀌어 읽혔다). 그 숫자의 **구절 안**을
            #   먼저 보고, 거기 없을 때만 앞을 본다.
            phrase = _en.LOCAL_NOT_SOURCE.sub(
                lambda m: " " * len(m.group(0)), re.split(r",|;|\.\s|\band\b|\bthen\b", folded[end : end + 60])[0]
            )
            local_in = min((at for at in (phrase.find(term.casefold()) for term in local_en) if at >= 0), default=-1)
            web_in = min((at for at in (phrase.find(term.casefold()) for term in web_en) if at >= 0), default=-1)
            if max(local_in, web_in) >= 0:
                scope = "local" if web_in < 0 or 0 <= local_in < web_in else "web"
                found.setdefault(scope, value)
                continue
            before = _en.LOCAL_NOT_SOURCE.sub(lambda m: " " * len(m.group(0)), folded[max(0, start - 180) : start])
            local_at = max((before.rfind(term.casefold()) for term in local_en), default=-1)
            web_at = max((before.rfind(term.casefold()) for term in web_en), default=-1)
            if max(local_at, web_at) < 0:
                continue
            found.setdefault("local" if local_at > web_at else "web", value)
        return list(found.items())
    for match in _COUNT.finditer(compact):
        value = _as_number(match.group(2))
        if not 1 < value <= 30:
            continue
        # ★ 여기도 같은 병이 있었다. 뒤 180자 안에 「파일」이 있으면 **무엇을 세든**
        #   로컬 읽기 개수가 됐다 — 「…파일을 읽고 핵심 3개를 정리해줘」가 문서
        #   3개 읽기로 둔갑한 자리가 바로 이 줄이다.
        if _counts_something_to_produce(compact, match) or _is_filter_condition(compact, match):
            continue
        before = _LOCAL_NOT_SOURCE.sub(lambda m: " " * len(m.group(0)), compact[max(0, match.start() - 180) : match.start()])
        before = _OUTPUT_FILE.sub(lambda m: " " * len(m.group(0)), before)
        local_at = max((_last_word_head(before, term) for term in local_terms), default=-1)
        if named_web:
            local_at = max(local_at, _last_file_as_source(before))
        web_at = max((before.rfind(term) for term in web_terms), default=-1)
        if max(local_at, web_at) < 0:
            continue
        scope = "local" if local_at > web_at else "web"
        # 「주제마다 2개」 — 묶음 수가 드러나면 곱한다(`_groups_for` · 2026-10-15). 주소를 댄 요청은 이 길로 와서
        #   `_count_target` 의 곱하기를 거치지 않았다(실제 판 b55c8113 이 고친 뒤에도 「2개 이상」이었다).
        found.setdefault(scope, min(30, value * _groups_for(compact, match)))
    return list(found.items())
# 한글 수사 → 숫자. 「세 군데」를 3 으로 읽는다.
#
# 「석·서·넉·너」까지 넣은 이유는 단위마다 꼴이 바뀌기 때문이다("석 잔", "넉 대").
# 국밥집을 세는 말로는 잘 안 쓰지만, 빠뜨려서 못 잡는 것보다 낫다 — 이 표에
# 잘못 걸릴 만한 다른 뜻이 없다.
_KOREAN_NUMERALS = {
    "한": 1, "하나": 1,
    "두": 2, "둘": 2,
    "세": 3, "셋": 3, "석": 3, "서": 3,
    "네": 4, "넷": 4, "넉": 4, "너": 4,
    "다섯": 5, "여섯": 6, "일곱": 7, "여덟": 8, "아홉": 9, "열": 10,
}


def _as_number(token: str) -> int:
    """정규식이 잡은 수 토큰을 숫자로. 못 읽으면 0(=버린다)."""
    token = token.strip()
    if token.isdigit():
        return int(token)
    return _KOREAN_NUMERALS.get(token, 0)


# ★ 2026-09-23 밤(영어판 짝 맞추기에서 드러남): 「구글 드라이브에 올려줘」·「구글 캘린더에 등록해줘」가 **「구글」의 「글」** 때문에
#   읽기 요구(`observe browser`)를 달았다. 메일의 「본문은 '…'」도 적을 내용을 불러준 것이지 읽으라는 말이 아니다.
_READ_INTENT = re.compile(r"(읽|열어|들어가|확인해|살펴|훑어|조사해|찾아|본문(?!\s*(?:은|:|：)\s*['\"“「‘])|원문|출처|공고|기사|(?<!구)글)")


def _read_intent(compact: str) -> bool:
    """읽으라는 말이 있는가 — 영어 요청은 영어 표로(`english.READ_INTENT`).

    ★ 따옴표 안은 **적을 내용**이지 시키는 말이 아니다(2026-09-23 밤) — 「'테스트 글'이라는 제목으로」의 「글」,
      「write 'scope check'」의 check 가 읽기 요구가 됐다. 두 언어 모두 따옴표 안을 지우고 본다.
    """
    if _en.english_dominant(compact):
        return bool(_en.READ_INTENT.search(_en.unquoted(compact)))
    return bool(_READ_INTENT.search(_en.unquoted(compact)))

# 개수가 **행위**에 붙었는가. 「5개 스크랩해줘」와 「5개 알려줘」는 다른 일이다.
#
# ★ 왜 필요한가 (2026-09-01 라이브). 「저 순위대로 5개 스크랩해줘」에 자비스는
#   2개만 스크랩하고 5개를 목록으로 적어 답했다. 그런데 개수 충족이
#   「읽은 쪽수 ≥ 5 **또는** 답에 적힌 항목 ≥ 5」였다. 답이 5줄이었으니
#   **충족으로 통과했다** — 실제로 한 것은 2개인데.
#
#   목록으로 답해서 되는 일(「3개 알려줘」)에는 그 규칙이 맞다. 하지만 행위에
#   붙은 개수는 **한 횟수로만** 세야 한다. 적어 놓는 것으로 대신할 수 없다.
_ACTION_COUNT = re.compile(
    # ★ 「지원자 수」의 지원은 행위가 아니다(2026-09-23 밤) — 「프로젝트 3개의 예상 기간과 지원자 수를 알려줘」가 「3개 실제로 처리하기」가 됐다.
    # ★ 「저장소」의 저장도 행위가 아니다(2026-10-08 실기 98fcd10c) — 「맨 위 5개 저장소의 이름과 별 수를 알려줘」가
    #   「5개 실제로 처리하기」가 됐고, 맞는 목록을 낸 판이 되받혀 목록이 사라진 채 「일부」로 끝났다.
    r"(스크랩|저장(?!소)|담아|담기|등록|신청|지원(?!자)|올려|올린|업로드|보내|발송|"
    r"삭제|지워|작성|만들어|제출|예약|추가)"
    # ★ 뒤에 **칸 이름을 만드는 말**이 붙으면 행위가 아니라 읽을 칸이다(2026-10-08) — 「올린 날짜」·「등록일」·「신청 기간」·
    #   「작성자」·「예약 현황」. 운영에서 「최근 게시물(최대 5개)마다 제목·게시판·올린 날짜·조회수…」가 「5개 실제로 처리하기」가
    #   됐다(읽기만 하는 요청). 위의 「지원자」도 같은 모양이다.
    #   한 글자짜리(일·자·수·순)는 뒤가 조사이거나 끝일 때만 본다 — 「등록일과」·「작성자를」은 칸, 「작성자료」·「등록일정」은 아니다.
    r"(?!\s*(?:날짜|일자|일시|시각|시간|기간|기한|마감|자격|여부|현황|상태|번호|방법"
    r"|[일자수순](?=[과와은는이가을를의도에로,·/\s)\]]|$)))"
)


#: 행위 낱말에 **바로 붙은** 금지 — 「예약은 하지 말고」·「보내지 마」·「등록은 빼고」. `_negated` 의 20자 창은 여기엔 넓다
#: (「5개 예약해줘 취소는 말고」의 「말고」가 예약까지 죽인다).
_NEGATED_RIGHT_AFTER = re.compile(
    r"\s*(?:은|는|을|를|도|까지)?\s*(?:하지\s*(?:는|도)?\s*(?:마(?![가-힣])|마세요|마라|말)|지\s*(?:는|도)?\s*(?:마(?![가-힣])|마세요|마라|말)"
    r"|빼고|제외|금지)"
)

#: 개수 둘레에서 동사를 찾는 창. 앞뒤로 이만큼만 본다.
_VERB_WINDOW_BEFORE = 10
_VERB_WINDOW_AFTER = 14


def _is_action_count(compact: str, target: int) -> bool:
    """개수 요구가 **행위에 붙었는가.**

    ⚠ 문장 어디든 행위 낱말이 있으면 참으로 두면 안 된다. 처음에 그렇게 짰다가
    검사가 잡았다:

        "위시켓과 크몽에서 공고를 4개 읽고 메일 보내줘"

    여기서 4개는 **읽기**에 붙어 있는데 뒤의 「보내」가 걸려 행위 개수가 됐다.
    그러면 4개를 다 읽고도 「4개 중 1개만 처리했습니다」가 답에 붙는다 — 맞는 판을
    틀렸다고 말하는 것이 제일 나쁘다.

    그래서 **개수에서 가장 가까운 동사**로 정한다. 앞뒤를 다 보는 이유는 사람이
    「스크랩 5개 해줘」라고도 「5개 스크랩해줘」라고도 말하기 때문이다.
    """
    if _en.english_dominant(compact):
        for value, _q, _role, start, end in _en.counts(compact):
            if value != target:
                continue
            before_en = compact[max(0, start - _en.VERB_WINDOW_BEFORE) : start]
            after_en = compact[end : end + _en.VERB_WINDOW_AFTER]

            def nearest(pattern: re.Pattern[str]) -> int | None:
                hits = [len(before_en) - m.end() for m in pattern.finditer(before_en)]
                hits += [m.start() for m in pattern.finditer(after_en)]
                return min(hits) if hits else None

            act_en = nearest(_en.ACTION_COUNT)
            read_en = nearest(_en.READ_INTENT)
            if act_en is not None and (read_en is None or act_en < read_en):
                return True
        return False
    for match in _COUNT.finditer(compact):
        if _as_number(match.group(2)) != target:
            continue
        # 조건값(「인원은 2명」·「후기 10개 이상인」)은 행위의 개수가 아니다 — `_count_target` 과 같은 눈으로 본다.
        if _counts_something_to_produce(compact, match) or _is_filter_condition(compact, match):
            continue
        before_at = max(0, match.start() - _VERB_WINDOW_BEFORE)
        before = compact[before_at : match.start()]
        # 개수로부터의 거리. 앞쪽은 뒤에서부터 세고, 뒤쪽은 앞에서부터 센다.
        # ★ **하지 말라는 행위**는 세지 않는다(2026-10-08) — 「2개만 알려줘. 예약은 하지 말고」의 예약은 시킨 일이 아니다.
        # 창 안에서 **시작해 끝나는** 낱말만 세되, 낱말 뒤는 원문으로 본다 — 창을 잘라 넘기면 「올린 날짜」의 「날짜」가
        # 창 밖이라 안 보인다(뒤를 보는 조건이 헛돈다).
        after_end = match.end() + _VERB_WINDOW_AFTER

        def distance(pattern: re.Pattern[str], skip_negated: bool = False) -> int | None:
            hits = [
                len(before) - m.end()
                for m in pattern.finditer(before)
                if not (skip_negated and _NEGATED_RIGHT_AFTER.match(compact, before_at + m.end()))
            ]
            hits += [
                m.start() - match.end()
                for m in pattern.finditer(compact, match.end())
                if m.end() <= after_end
                and not (skip_negated and _NEGATED_RIGHT_AFTER.match(compact, m.end()))
            ]
            return min(hits) if hits else None

        act = distance(_ACTION_COUNT, skip_negated=True)
        read = distance(_READ_INTENT)
        if act is not None and (read is None or act < read):
            return True
    return False


# 답이 **몇 개를 실제로 내놓았는가.** 결정론으로만 센다.
#
# ★ 왜 필요한가 (2026-08-23). 개수 요구의 충족을 「연 페이지 수」로 재고 있었다.
#   그런데 목록에 답이 다 있는 판에서는 상세를 열 이유가 없다:
#
#       무신사 "상위 3개의 제품명과 가격"
#           퀀3.8   시킨 대로 상세를 셋 열다 30걸음·401초·충족 0/1
#           젬마4   목록에서 뽑아 9걸음·124초·정답 3개
#
#   **시킨 대로 한 쪽이 졌다.** 재는 자가 「일한 방식」을 채점하고 있었기 때문이다.
#   시킨 것은 "3개를 알려 달라"였지 "3개를 열어 보라"가 아니었다.
#
#   그렇다고 개수 요구를 지우면 2026-08-13 의 병이 돌아온다(3개 시키고 1개 받기).
#   그래서 **두 길 중 하나면 충족**으로 바꾼다 — 충분히 읽었거나, 답이 실제로
#   N개를 담았거나.
#: 「**1. 제목**」(제목까지 굵게)은 세지 않는다 — 단락 제목(「**1. 확인된 내용**」 · 「**2. 실패한 작업**」)과 모양이 같아서, 세면
#: 「3개 중 2개」 답이 단락 제목 덕에 통과한다(2026-10-16 옛 · 새 대조: 운영 답 4,149개 중 93개가 바뀌었고 상당수가 단락 제목).
_LISTED_ITEM = re.compile(r"^\s*(?:[-*•]|\d{1,2}\s*[.)]|\*\*\d{1,2}\s*[.)]?\*\*)\s+\S")


def count_listed_items(answer: str) -> int:
    """답이 늘어놓은 항목 수. 번호 목록·불릿·표의 데이터 줄을 센다.

    세는 것만 센다 — "좋은 답인가"는 여전히 안 본다(그건 셀 수 없다).
    """
    if not answer:
        return 0
    lines = [line for line in answer.splitlines() if line.strip()]
    listed = sum(1 for line in lines if _LISTED_ITEM.match(line))
    # 표로 낸 경우 — 머리줄과 구분선을 빼고 센다.
    rows = [line for line in lines if line.strip().startswith("|")]
    body = [
        row for row in rows
        if not set(row.replace("|", "").strip()) <= set("-: ")
    ]
    table = max(0, len(body) - 1) if len(body) > 1 else 0
    # ★ **둘 다 세고 큰 쪽을 쓴다** (2026-09-12 역검증).
    #
    #   예전에는 불릿이 하나라도 있으면 거기서 멈추고 표를 아예 안 봤다. 그런데
    #   우리 답의 끝에는 거의 항상 이런 불릿이 붙는다:
    #
    #       **[산출물 정보]**
    #       *   **파일명:** wishket_ai_agent_projects.xlsx
    #       *   **경로:**   (문서 폴더 경로)
    #
    #   그래서 위에 3줄짜리 표를 완성해 놓고도 「답에 2개」로 찍혔다. 같은 판에
    #   채움 채점기는 12/12 만점을 줬다 — **우리가 붙인 꼬리가 본문을 가린 것이다.**
    #
    #   과거 판 445건에 대 보니 이 자리에서 3분의 1이 과소 계수였다. 「답에 N개」를
    #   근거로 무언가를 막는 규칙을 세우기 전에 이것부터 고쳐야 한다.
    #
    #   ★ `max` 가 아니라 **표가 있으면 표**다. 둘을 같은 표본에 대 봤다:
    #
    #       옛(불릿 우선)  18.9% 미달   ← 표를 못 봄
    #       max            2.7%  미달   ← 꼬리 불릿이 모자란 표를 덮음
    #       표 우선         5.3%  미달   ← 채택
    #
    #   max 가 놓친 10판을 전부 열어 봤더니 **열 건 다 진짜 미달**이었다. 답이
    #   스스로 「3개 중 2건」·「3권 중 1권」이라고 적어 놓았는데, 그 사정을 설명하는
    #   불릿이 4~6개라 max 가 통과시킨 것이다. 설명이 길수록 점수가 오르는 눈금은
    #   눈금이 아니다.
    return table if table else listed


#: 이것이 잡히면 저것은 안 잡는다. **더 구체적인 쪽이 이긴다.**
#
# 낱말표는 넓게 잡을수록 헛맞는다. 좁히면 못 잡는 것이 생기니, 넓게 두고
# **겹칠 때 누가 이기는지**만 여기 적는다.
_SUPERSEDED_BY: dict[str, tuple[str, ...]] = {
    # 「드라이브에 올려줘」는 블로그 발행이 아니다(발행 낱말표에 `올려줘`·`업로드`가 있다).
    "google_drive_write": ("naver_blog_publish",),
    # 열린 에디터 요청이면 레시피(배치 모드)가 아니라 에디터 도구가 요구다. 되돌리기는 바꾸기를 대신한다.
    "unity_editor_change": ("unity_run_recipe",),
    "unity_editor_look": ("unity_run_recipe",),
    "unity_editor_undo": ("unity_editor_change", "unity_run_recipe"),
}

#: 파일 이름의 **몸통**(확장자 앞) — 「커뮤니티_게시규칙.xlsx」의 「커뮤니티_게시규칙」.
_FILE_NAME_STEM = re.compile(
    r"[^\s「」『』\"'`()\[\],]+(?=\.(?:xlsx|xls|csv|docx|doc|hwpx|hwp|pptx|ppt|pdf|md|txt|json|html?|py|js|ts)\b)", re.IGNORECASE)


def _inside_a_file_name(text: str, start: int, end: int) -> bool:
    """`text[start:end]` 가 통째로 어떤 파일 이름의 몸통 안에 있는가."""
    return any(m.start() <= start and end <= m.end() for m in _FILE_NAME_STEM.finditer(text))


#: 사람이 **열거나 읽지 말라고** 한 말 — 「열어 읽지는 말고」·「열지 마」·「don't open it」.
_READ_FORBIDDEN = re.compile(
    r"(?:읽|열어\s*보|열)지[는도]?\s*(?:말|마)|(?i:\bdo(?:n['’]t|\s+not)\s+(?:open|read)\b|\bwithout\s+(?:opening|reading)\b)")


_ADDRESS_HOST = re.compile(r"https?://([^/\s\"'<>)]+)")


def _acts_on_another_site(compact: str, found: "list[Requirement]") -> bool:
    """이 요청이 **네이버가 아닌 곳**을 댔는가 — 적힌 주소의 호스트나, 이미 뽑힌 사이트 요구로 본다."""
    hosts = [host.lower().split(":")[0] for host in _ADDRESS_HOST.findall(compact)]
    if any(not host.endswith("naver.com") for host in hosts):
        return True
    for item in found:
        if item.kind != "site":
            continue
        domains = item.target if isinstance(item.target, (list, tuple, set, frozenset)) else (item.target,)
        if not any("naver" in str(domain).lower() for domain in domains):
            return True
    return False


# 행위 요구 -> 그것이 이뤄졌다고 말할 수 있는 도구.
#: 같은 문장 안의 아무 글자들 — 줄바꿈과 「마침표·물음표·느낌표 + 빈칸/끝」을 건너가지 않는다.
_SAME_SENTENCE = r"(?:(?![.!?](?:\s|$))[^\n])*"
_ACTIONS: tuple[tuple[re.Pattern[str], str, str], ...] = (
    (re.compile(r"임시\s*저장"), "naver_blog_save_draft", "임시저장"),
    (re.compile(r"(드라이브|구글\s?드라이브).{0,20}(올려|업로드|저장|넣어|만들어)"),
     "google_drive_write", "드라이브에 올리기"),
    # ★ 2026-09-18: 「MS 스토어에서 ChatGPT 앱의 **게시자**와 별점을 알려줘」가 발행 요구로 읽혀, 한 판은 네이버 블로그에
    #   글을 넣으러 갔고(로그인이 없어 막힘) 막고 나니 다음 판은 「발행 날짜」를 답했다. 명사 꼴(게시자·게시판·게시글·
    #   게시물·발행일·발행처…)은 행위가 아니다. 좁히기 전에 재 봤다: 30일 서로 다른 요청 564개 중 새로 놓치는 것은
    #   그 벤치 문장 1개, 새로 걸리는 것 0개, 실제 발행 요청 2개는 그대로 걸린다.
    (re.compile(r"발행(?!일|처|사|인|년|연도|부수|본|된|한\s*날)|게시(?!자|판|글|물|일|된|중|한\s*날)|포스팅해|올려\s*줘|업로드"),
     "naver_blog_publish", "발행"),
    # ★ 대상과 동사가 **한 문장 안에** 있어야 한다 (2026-10-08 전수조사 3-4 — 아래 일정 수정·삭제도 같다).
    #   `메일.*보내` · `일정.*수정` 의 `.*` 는 문장 끝을 건너갔다 — 「캘린더에서 다음 주 일정을 보여줘. 회의록은 따로
    #   수정할 거야.」가 「일정 수정」 요구가 됐다. 10-02 에 「일정 등록」에서 고친 병과 같은 꼴이 세 줄에 남아 있었다.
    #   문장 끝은 마침표·물음표·느낌표 **뒤에 빈칸이나 끝**이 올 때만 본다(메일 주소·파일 이름의 점에서 끊지 않는다).
    (re.compile(rf"메일{_SAME_SENTENCE}보내|메일{_SAME_SENTENCE}발송|이메일{_SAME_SENTENCE}보내"), "google_gmail_send", "메일 발송"),
    # ★ 2026-09-27 게임 팩 3모델 벤치: 4B 가 유니티 레시피를 한 번도 안 부르고 파일을 뒤지다 「텍스처가 없다」(115개 있음)·
    #   「버전 2022.3.17f1」(실제 6000.6.3f1)을 지어냈다. 35B 도 한 판은 「있을 가능성이 높다」로 끝냈다. 유니티 프로젝트를
    #   **점검·검사**하라는 요청이면 레시피 실행을 시킨 것으로 세서, 안 돌리고 답하면 되받게 한다.
    #   넣기 전에 재 봤다: 집 PC 서로 다른 요청 1,187개 중 걸리는 것 0개, 벤치 문장 둘은 걸림, 「유니티 설치 링크」·
    #   「유니티가 뭐야」·「텍스처 점검 방법 알려줘」는 안 걸림.
    # ★ 2026-09-28 데모: 「유니티 에디터에서 아까 만든 큐브 지워 줘」에 도구 0번으로 「이미 삭제된 상태입니다」(큐브는 그대로였다).
    #   앞 판이 도구 없이 「삭제했습니다」라고 거짓말한 것을 대화 기록에서 믿었다. 낱말표로 상태 단정을 다 막을 수는 없다 —
    #   열린 에디터 안을 바꾸거나 보라는 요청이면 **에디터 도구를 부른 기록**을 요구로 세운다(레시피 규칙과 같은 자리).
    #   넣기 전에 재 봤다: 운영 PC 서로 다른 요청 1,196개 중 셋 다 걸리는 것 0개, 데모 문장·「콘솔에 에러 있나」는 걸림,
    #   「유니티 설치 링크」·「텍스처 점검해줘」는 안 걸림. 되돌리기가 먼저다(되돌리기가 바꾸기를 대신한다).
    (re.compile(r"(?is)(?=.*(?:유니티|unity|에디터|editor)).*(?:되돌려|되돌리|\bundo\b|실행\s*취소|원래대로)"),
     "unity_editor_undo", "유니티 에디터에서 되돌리기"),
    (re.compile(
        r"(?is)(?=.*(?:유니티|unity|에디터|editor))"
        r"(?=.*(?:오브젝트|object|큐브|cube|스피어|sphere|프리팹|prefab|게임\s*오브젝트|위치|계층|씬\b|scene))"
        r".*(?:만들어|만들고|생성|지워|삭제|없애|옮겨|이동|바꿔|바꾸|수정|복제|추가|배치|놓아)"),
     "unity_editor_change", "유니티 에디터에서 실제로 바꾸기"),
    (re.compile(
        r"(?is)(?=.*(?:에디터|editor))"
        r"(?=.*(?:계층|hierarchy|콘솔|console|에러|오류|선택|selection|오브젝트|씬\b))"
        r".*(?:보여|보고|확인|알려|있어|있나|찾아|읽어)"),
     "unity_editor_look", "유니티 에디터 안 실제로 보기"),
    (re.compile(
        r"(?is)(?=.*(?:유니티|unity))"
        # ★ 2026-10-12 팩에 든 점검은 넷인데 이 규칙은 셋만 알았다 — 「안 쓰는 에셋 찾아줘」가 요구로 서지 않았다
        #   (점검 `unused_assets` 가 규칙보다 뒤에 들어왔다). 「에셋」 한 낱말로 넓히면 「유니티 에셋 스토어에서 …
        #   찾아줘」까지 점검을 요구하게 되므로(채울 길이 없는 요구) 「안 쓰는 에셋」 꼴만 받는다.
        #   팩의 점검마다 화면에 적는 예시 문장(`ask`)이 이 규칙에 걸리는지는 시험이 지킨다.
        r"(?=.*(?:텍스처|texture|빠진\s*참조|missing|프리팹|prefab|머티리얼|material|씬\b|scene"
        r"|(?:안\s*쓰는|쓰지\s*않는|안\s*쓰이는|미사용)\s*에셋|unused\s+assets?))"
        r".*(?:점검|찾아|확인|검사|살펴|check|find|audit|scan)"),
     "unity_run_recipe", "유니티 레시피로 실제 점검하기"),
    # ★ 2026-09-23 스토어 촬영(코드 모드): 「index.html 로 할 일 목록을 만들고 … 다 만들면 **미리보기로 보여 줘**」에서
    #   파일 쓰기·검사는 다 하고 `local_preview_page` 를 한 번도 안 불렀다. 「시킨 것 2/2」에 미리보기가 **아예 없어서**
    #   다 끝난 것처럼 보였다. 넓히기 전에 재 봤다: 최근 서로 다른 요청 1,181개 중 걸리는 것 4개 — 전부 웹페이지를
    #   만들고 띄워 보라는 요청, 잘못 걸리는 것 0개. 코워크에는 미리보기 도구가 없으니 「엑셀 만들고 미리보기로 보여줘」가
    #   영원히 못 채우는 요구가 되지 않게 **웹 페이지 맥락이 같이 있을 때만** 잡는다.
    (re.compile(
        r"(?is)(?=.*(?:웹\s*페이지|페이지|html|사이트|웹\s*앱|랜딩|프론트|\bui\b))"
        r".*(?:미리\s*보기\s*(?:로|를|로도)?\s*(?:보여|띄워|열어|확인)|미리\s*봐\s*줘|local_preview_page)"),
     "local_preview_page", "만든 화면 미리보기로 띄우기"),
    # ★ 2026-08-20. `만들`·`생성` 이 없었다. 그래서 **쓰기 벤치의 캘린더 과제 자체가**
    # ("일정을 30분짜리로 만들어줘") 요구 0개로 돌았다 — 계획 층이 통째로 눈을 감았고,
    # 네 걸음이 전부 실패했는데도 `model_answered`(대조할 요구를 못 뽑음)로 끝나
    # 위임 성공 쪽에 섰다. 재는 자가 자기가 시킨 일을 못 보고 있었던 것이다.
    #
    # 낱말을 넓히기 전에 **재 봤다**: 최근 30일 서로 다른 프롬프트 438개에 대고
    # 새 판과 옛 판을 돌려 새로 걸리는 것 1건(그 벤치 문장), 놓치는 것 0건.
    # 창을 좁히는 판(`일정…{0,12}만들`)도 같이 재 봤는데 수정 갈래에서 진짜 요청
    # 하나를 놓쳐서(18자짜리) 그쪽은 안 썼다 — 두 벌을 두면 갈린다.
    # ★ 2026-10-02 (공개 전 필수 ①): `일정.*만들` 이 문장 끝까지 건너가 「…본 서비스 **일정**·지금 할 수 있는 기능(…)을 정리한
    #   **문서**(.md)로 **만들**어줘」를 「일정 등록」으로 읽었고, 손자리가 시키지 않은 **예약 작업 만들기**를 준비했다(예약 기능이
    #   꺼져 있어서 막혔을 뿐). 「일정」 뒤 같은 마디(문장·쉼표·가운뎃점·괄호를 넘지 않고) 20자 안의 동사만, 사이에 문서·파일
    #   낱말이 끼면 아니다. 넣기 전에 쟀다: 운영 요청 1,420개 중 옛 판 3·새 판 2, 빠진 하나가 바로 그 오독, 새로 걸리는 것 0.
    (re.compile(
        r"일정(?:(?!문서|파일|엑셀|표로|보고서|정리한|docx|xlsx|한글)[^.!?\n,·、()（）]){0,20}?(등록|추가|잡아|넣어|만들|만드|생성)"),
     "google_calendar_create", "일정 등록"),
    # 2026-09-01: 드라이브 쓰기가 생겼다. 여기 없으면 「드라이브에 올려줘」가
    #   어느 요구에도 안 걸려 "올렸습니다"로 끝난다 — 캘린더에서 겪은 그대로다.
    # 2026-08-13: 수정·삭제가 없어서 "일정을 5시로 옮겨줘"가 어느 요구에도 안 걸렸다.
    # 라우터도 이 목록을 보고 갈래를 정하므로, 빠진 행위는 **모델 판단으로 떨어진다** —
    # 그날 실측에서 캘린더 수정 요청이 도구 0건으로 "준비를 마쳤습니다"가 됐다.
    (re.compile(rf"일정{_SAME_SENTENCE}(수정|변경|옮겨|바꿔|미뤄|당겨)"), "google_calendar_update", "일정 수정"),
    (re.compile(rf"일정{_SAME_SENTENCE}(삭제|지워|취소)"), "google_calendar_delete", "일정 삭제"),
    # 일반 폼은 전용 레시피가 없어도 최종 제출 도구가 명확하다. 이 요구를 계획에
    # 넣지 않았더니 공식 Selenium 폼에서 모든 칸을 채우고도 7걸음째에 답을 내며
    # Submit만 사람에게 남겼다. 금지형("제출하지 마")은 동작형 어미가 없어 잡히지
    # 않도록, 실제로 누르거나 실행하라는 표현만 좁게 받는다.
    (
        re.compile(
            r"(?:제출|submit)(?:\s*버튼)?(?:을|를)?[^.!?\n]{0,30}"
            r"(?:실행|클릭|눌러|누르|해(?:라|줘|주세요|))",
            re.IGNORECASE,
        ),
        "browser_click_confirmed",
        "폼 제출",
    ),
)


def mutating_google_tools(prompt: str) -> list[str]:
    """이 요청이 **구글 계정의 무언가를 바꾸라**는 것이면 그 도구 이름들.

    라우터가 갈래를 정할 때 쓴다. 목록을 새로 만들지 않고 위 `_ACTIONS` 를 그대로
    보는 이유: 두 벌을 두면 갈리고, 갈린 쪽은 조용히 틀린다. 여기 한 줄을 더하면
    계획 대조와 라우팅이 **같이** 좋아진다.
    """
    return [
        requirement.target
        for requirement in extract(prompt)
        if requirement.kind == "action" and str(requirement.target).startswith("google_")
    ]


@dataclass(frozen=True)
class Requirement:
    kind: str      # count | count_action | observe | site | filter | action | artifact | footnote | verify
    label: str     # 사람이 읽을 말
    target: Any    # 개수, 호스트 조각, 도구 이름
    # 같은 count라도 목록 자체가 답인 판과 상세 확인형은 최종 관문이 다르다.
    # 새 kind를 만들면 예산·이어받기·완료 판정의 모든 등록 자리가 갈린다.
    grounding_required: bool = False
    # 한 프롬프트 안에서 로컬 입력 수와 웹 근거 수를 따로 센다. 빈 값은 예전처럼
    # 모든 출처를 센다. 필드를 뒤에 두어 기존 위치 인자 계약은 그대로 유지한다.
    source_scope: str = ""
    # 산출물을 **어디에** 두라고 했나(사람이 자리와 이름을 둘 다 댄 경우만 · agent `requested_delivery`).
    # 비면 예전처럼 어디에 있든 연다. 있으면 그 폴더에 있어야 채운다 — 작업 폴더에 떨어진 것은 「못 한 것」이다.
    destination: str = ""
    # 사람이 「하나하나 열 필요 없다」고 한 개수 요구(`says_list_only`). 되받기가 「본문을 열어라」 대신 「목록으로 담아라」를 준다.
    list_only: bool = False


#: 앞 판 근거를 **이어받지 않을** 말 — 새것·더·다른 것을 달라는 요청 (2026-09-24).
#:   이어받기 자체는 문장이 아니라 **앞 판이 실제로 읽은 주소**로 정한다(agent `_carry_prior_reads`).
#:   여기는 그 반대쪽 문만 지킨다: 「사람인에서 3개 더 찾아줘」가 앞 판 3개로 채워져 조사를 건너뛰면 안 된다.
_NEW_RESULTS = re.compile(
    r"(?:더|추가로?|새로(?:운)?|다른|또)\s*(?:\d+\s*(?:개|건|곳|권)\s*)?(?:\S+\s+){0,2}?"
    r"(?:찾|조사|알아|뽑|가져|보여|검색|골라|추천)"
)
_NEW_RESULTS_EN = re.compile(r"\b(?:more|another|other|new|additional|different|else)\b", re.I)


def asks_for_new_results(prompt: str) -> bool:
    """앞 판에서 찾은 것 **말고** 새로 찾아 달라는 말인가."""
    compact = " ".join((prompt or "").split())
    if not compact:
        return False
    if _en.english_dominant(compact):
        return bool(_NEW_RESULTS_EN.search(compact))
    return bool(_NEW_RESULTS.search(compact))


_UNITY_ACTION_TOOLS = frozenset({"unity_run_recipe", "unity_editor_change", "unity_editor_look", "unity_editor_undo"})


def unity_work_asked(prompt: str) -> str:
    """계획 층이 유니티 도구 실행을 요구로 세울 문장이면 그 도구 이름. 아니면 빈 문자열.

    갈래 판정(`runtime._requires_unity_work`)이 이 표를 그대로 쓴다 — 갈래와 계획이 다른 낱말표를 쓰면
    「계획은 레시피를 요구하는데 갈래는 대화로 보낸」 판이 다시 생긴다(2026-10-01 테스터 시험).
    """
    compact = " ".join((prompt or "").split())
    for pattern, tool, _label in _ACTIONS:
        if tool not in _UNITY_ACTION_TOOLS:
            continue
        match = pattern.search(compact)
        if match and not _negated(compact, match.end()):
            return tool
    return ""


#: 옮겨 적은 말을 감싸는 따옴표 쌍 — 낫표 · 겹낫표 · 큰따옴표 · 작은따옴표(영문 소유격의 ' 는 낱말에 붙어 있어 안 걸린다).
_QUOTED_SPAN = re.compile(
    r"「[^「」]{1,240}」|『[^『』]{1,240}』|[\"“][^\"“”]{1,240}[\"”]|(?<![A-Za-z0-9])'[^']{1,240}'(?![A-Za-z0-9])|‘[^‘’]{1,240}’"
)


def _without_quoted_sentences(compact: str) -> str:
    """따옴표 안의 **문장·구**(빈칸이 든 것)를 같은 길이의 빈칸으로 바꾼 글. 따옴표 문자와 글 길이는 그대로 둔다.

    낱말 하나를 감싼 따옴표('사람인' · 'RAG')는 이름이나 낱말을 짚은 것이라 건드리지 않는다.
    """
    def blank(match: "re.Match[str]") -> str:
        text = match.group(0)
        inner = text[1:-1]
        if " " not in inner.strip():
            return text
        # 글의 대부분을 감싼 따옴표는 인용이 아니라 **요청 자체**다 — 요청을 통째로 낫표에 넣어 보내는 사람이 있다
        # (운영 요청에 대 보니 그런 요청 하나의 요구가 전부 사라졌다).
        if len(inner) * 2 > len(compact):
            return text
        return text[0] + " " * len(inner) + text[-1]
    return _QUOTED_SPAN.sub(blank, compact)


#: 읽을 곳이 **내 PC 의 프로그램 창**인 요청 — 「작업 관리자를 열어서」·「메모장 창에서」·「한글 프로그램 화면에서 읽어」.
#: 프로그램 이름만으로는 안 건다(「엑셀로 정리해 줘」는 만들 형식이다) — 열거나 · 창 · 화면이 같이 와야 한다.
_DESKTOP_SOURCE = re.compile(
    r"(?:작업\s*관리자|메모장|그림판|계산기|한글\s*프로그램|엑셀\s*프로그램|워드\s*프로그램|제어판|설정\s*앱"
    # ★ 2026-10-18 「모르는 앱」 측정: 「윈도우 설정의 「시스템 → 정보」 화면을 열어서」 · 「「시스템 정보」 프로그램(msinfo32)을
    #   열어서」가 이 표에 없어 「봐야 하는 일」이 안 됐고, 6판 중 4판이 도구 0걸음 · 기억으로 답했다(「이미 확인된 값」).
    r"|시스템\s*정보|msinfo32|(?:윈도우|windows)\s*설정|장치\s*관리자|디스크\s*관리|이벤트\s*뷰어)"
    r"[^.!?\n]{0,24}(?:열어|열고|켜서|켜고|창|화면)"
    r"|(?:프로그램|앱)\s*창(?:에서|을|의)"
    # 「…프로그램(msinfo32)을 열어서」 · 「…앱을 켜고」 — 프로그램 · 앱을 열라는 말(괄호 속 이름은 건너뛴다).
    r"|(?:프로그램|앱)\s*(?:\([^)]{1,40}\))?\s*(?:을|를)\s*(?:열어|열고|켜서|켜고)"
    r"|(?i:\b(?:task\s*manager|notepad)\b)"
)


#: 이름으로 여는 윈도우 화면 — (부른 말, `windows_open_path` 의 path 에 적을 이름). 앞의 것이 먼저다.
_NAMED_SCREENS = (
    (re.compile(r"시스템\s*정보|msinfo32", re.I), "시스템 정보"),
    (re.compile(r"(?:윈도우|windows)\s*설정[^.!?\n]{0,24}정보|시스템\s*(?:→|>|->)\s*정보", re.I), "시스템 → 정보"),
    (re.compile(r"(?:윈도우|windows)\s*설정|설정\s*앱", re.I), "설정"),
)


def named_screen_to_open(prompt: str) -> str | None:
    """「시스템 정보를 열어서」 · 「윈도우 설정의 시스템 → 정보 화면을 열어서」처럼 이름으로 여는 화면을 열라는 요청이면 그 이름(2026-10-18).

    「모르는 앱」 다시 잼: 모델이 「msinfo32 를 띄우는 도구가 없어」라고 쓰고 셸 명령으로만 읽었다 — 그래서 설정 화면에만 있는
    값(버전 이름)이 3판 중 2판에서 빠졌다. 열 수 있다는 것을 모르는 것이라 그 이름을 알려 준다.
    """
    compact = " ".join((prompt or "").split())
    if not names_a_desktop_source(compact):
        return None
    return next((name for pattern, name in _NAMED_SCREENS if pattern.search(compact)), None)


def names_a_desktop_source(prompt: str) -> bool:
    """내 PC 의 프로그램 · 설정 화면을 열어 읽으라는 요청인가(라우터가 「가서 봐야 하는 일」로 쓴다)."""
    return bool(_DESKTOP_SOURCE.search(" ".join((prompt or "").split())))


def extract(
    prompt: str,
    *,
    has_named_source: bool = False,
    # 라우터의 결정론 갈래가 「가서 봐야 아는 것」이라고 판정했는가.
    #
    # `has_named_source` 와 하는 일이 같다 — 읽기 동사가 없어도 개수를 읽기
    # 요구로 본다. 그런데 뜻이 다르므로 이름을 나눈다: 저쪽은 "출처를 댔다",
    # 이쪽은 "조사해야 하는 말이다".
    #
    # ★ 2026-08-24. 없던 동안 「여수 국밥집 세 곳 알려줘」처럼 **출처도 안 대고
    #   읽기 동사도 없는** 조사 요청이 계획 `{}` 으로 끝났다. 개수 요구가 없으니
    #   `_forced_read_call` 이 밀어 줄 것도 없어서, 검색 한 번으로 목록 요약만
    #   옮겨 적고 끝나는 판이 나왔다.
    must_observe: bool = False,
    known_sites: dict[str, tuple[str, ...]] | None = None,
    # 라우터가 must_observe 를 켠 **이유**(runtime `_chat_job`). 읽기 대상을 여기서 다시 추측하지 않는다.
    observe_from: "frozenset[str] | set[str] | tuple[str, ...]" = (),
    # 사람이 고른 Code 탭의 판인가. 코드 과제의 문장은 **만들 물건의 사양서**다(아래 끝 주석).
    code_workspace: bool = False,
) -> list[Requirement]:
    """이 요청에서 **대조 가능한** 요구만 뽑는다.

    `has_named_source` — 이 요청이 출처를 대고 물었는가. 부르는 쪽이 넣어 준다.
    그 판정(`agent.named_sources`)은 사이트 지식을 쥐고 있어 여기서 부르면 순환이
    되므로, 아는 쪽이 넘긴다.

    ★ 2026-08-13. 이 인자가 없어서 **계획 층이 개수 요구를 통째로 못 봤다.**

        "크몽에서 AI 에이전트 개발 서비스 3개의 제목과 가격을 알려줘"

    여기엔 읽기 동사가 하나도 없다("알려줘"는 목록에 없다). 그래서 개수 요구가 안
    뽑혔고, 그 결과 걸음 예산도 안 늘고 `remaining_note` 도 안 나가고 "다 채웠나"
    판정도 안 돌았다. 정작 정직성 검사(`_goal_shortfall_check`)는 같은 문장을 보고
    **부족하다고 답에 적고 있었다** — 두 층이 서로 다른 말을 하고 있었던 것이다.
    벤치 실측에서 크몽 2걸음·쿠팡 3걸음·교보문고 4걸음으로 끝난 원인이 여기다.

    출처를 대고 개수를 말한 것은 그 자체로 읽기 요구다("교보문고에서 책 3권").

    `known_sites` — 자비스가 **배운** 사이트까지 대조에 넣는다(`SiteAtlas.known`).

    ★ 2026-08-20. 여기 씨앗은 여덟 곳뿐인데, 자비스는 이미 다나와·쿠팡·교보문고·
      대학알리미·커리어넷을 배워 `site_atlas.json` 에 적립해 두고 있었다. 그걸
      계획 대조에 안 써서 **"다나와에서 …"·"쿠팡에서 …" 가 요구 0개로 돌았다**
      (실측: 최근 7일 276판 중 66.3%가 요구 0개). 배운 것을 안 쓰면 배운 값이 없다.

      애매한 이름(두 도메인에 걸린 것)은 `SiteAtlas.known` 이 이미 뺀다 — 그래서
      "블로그"·"상담" 같은 낱말이 요구로 둔갑하지 않는다.
    """
    compact = " ".join((prompt or "").split())
    found: list[Requirement] = []
    # ★ **따옴표로 옮겨 적은 문장은 시킨 말이 아니다** (2026-10-09 — 우리 숙제를 손자리에게 맡긴 판 88b47080).
    #   과제 글이 버그를 설명하며 예로 든 문장 — 「앞 판이 『사람인에서 공고 3개 열어 읽어줘』였으면 …」 — 에서 「3개 이상 직접
    #   열어 읽기」와 「사람인에서 하기」가 섰다. 코드 버그를 고치는 판이 「아직 읽은 본문 없음 · 3분 남음」으로 재촉받았고,
    #   끝에는 ❌ 둘이 붙을 판이었다. 바깥 AI 나 사람이 쓰는 긴 지시문에는 예시·제목·검색어가 따옴표로 늘 들어간다.
    #   개수와 사이트 이름은 따옴표 **밖**에서만 찾는다. 낱말 하나를 따옴표로 감싼 것('사람인'에서 · 'RAG' 로)은 이름을
    #   짚은 것이라 그대로 본다 — 빈칸이 든 것(문장·구)만 옮겨 적은 말로 친다.
    outside_quotes = _without_quoted_sentences(compact)

    target = _count_target(outside_quotes)
    # ⚠ 2026-08-23. 이 조건에서 `has_named_source` 를 빼려다 되돌렸다.
    #
    # 오늘 무신사 판에서 이런 걸 봤다 — "상위 3개의 제품명과 가격"은 목록에 이미
    # 있는데 계획 층이 "상세를 3개 열어 읽어라"로 읽어, 시킨 대로 한 퀀3.8 이
    # 30걸음·401초를 쓰고도 `충족 0/1` 이 됐다(젬마는 무시하고 9걸음에 끝냈다).
    #
    # 그래서 목록에 보이는 칸만 물었으면 요구를 안 만들게 고쳤더니 **검사 9개가
    # 걸렸다.** 그 인자는 2026-08-13 에 **일부러** 넣은 것이었다 — 없으면 크몽
    # 2걸음·쿠팡 3걸음·교보문고 4걸음으로 일찍 멈춰 3개 시키고 1개를 받는다.
    #
    #     요구 없음 → 너무 일찍 멈춘다        (2026-08-13 에 겪음)
    #     요구 있음 → 목록에 있는 걸 열어본다  (2026-08-23 에 겪음)
    #
    # 둘 다 진짜 병이므로 **한쪽을 지우는 것은 답이 아니다.** 고칠 자리는 여기가
    # 아니라 `check` 의 충족 판정이다 — 「연 페이지 수」가 아니라 「답에 담긴 항목
    # 수」로 세야 한다. 그건 별도 작업으로 남긴다.
    scoped_counts = _scoped_count_targets(outside_quotes)
    # Code/산출물 요청의 수량은 **만들 구조**다. ``목적지 카드 3개``를
    # ``자료 3개를 직접 열어 읽기``로 바꾸면 완성 뒤에도 조사 되받기가 돈다.
    # 명시적인 읽기나 지목한 출처가 있을 때만 그 수를 읽기 요구로 승격한다.
    code_or_artifact = wants_code(compact) or wants_an_artifact(compact)
    english = _en.english_dominant(compact)
    read_count_context = bool(
        _read_intent(compact)
        or has_named_source
        or (must_observe and not code_or_artifact)
    )
    # 긴 Code 요청에는 화면을 "열고 확인"하고 임시 "파일"을 지우라는 말이 흔하다.
    # 그 뒤에 있다고 해서 앞의 ``3개 카드``가 자료 읽기 수량은 아니다. Code/산출물
    # 요청에서는 수량 자체가 자료·문서·페이지 같은 입력 대상을 셀 때만 읽기 요구다.
    # ★ **출처를 댔으면 그 숫자는 밖에 있는 것을 센다** (2026-09-21).
    #
    #   위 규칙(산출물 요청이면 수량을 버린다)은 AURORA 의 「3개 관측 목적지 카드」를
    #   「자료 3개를 직접 열어 읽기」로 바꾸던 오염을 막으려고 생겼다. 맞는 방향인데
    #   **너무 넓었다** — 이런 것까지 같이 버렸다:
    #
    #       "YES24에서 파이썬 책 3권의 쪽수와 ISBN 을 조사해서 마크다운 문서로 만들어줘"
    #
    #   기존 시험 다섯 개가 그래서 깨졌다. 둘 다 「만들어줘」로 끝나고 둘 다 숫자가
    #   있지만, 앞엣것의 3 은 **우리가 만들 모양**이고 뒤엣것의 3 은 **읽어야 할 개수**다.
    #
    #   가르는 선은 낱말(권·개·건)이 아니라 **출처를 댔는가**이다. 낱말로 가르면 또
    #   구멍이 난다 — 이 저장소가 여러 번 겪은 자리다. AURORA 에는 댈 출처가 없다.
    # ※ 2026-09-30 코드 벤치의 「모바일에서도」(출처로 오독)는 여기서 막지 않는다. 「만들 것이 있으면 아는 사이트만 출처」로
    #   좁혔더니 코워크의 「마이크로소프트 스토어에서 메모 앱 3개 → 엑셀」·「알라딘에서 책 3권 → 워드」(아는 사이트 목록 밖의
    #   진짜 출처) 5건이 읽기 개수를 잃었다(역검증). 사양서 문장이 오독되는 곳은 **Code 탭**이라 거기서만 거른다(맨 아래).
    if code_or_artifact and not has_named_source and not _has_explicit_input_count(compact):
        read_count_context = False
    # ★ 2026-09-30: 「폴더의 파일 이름을 세 개만 알려 줘」 — 이름은 목록 보기로 끝난다. 본문을 보라는 말이 없으면
    #   읽기 개수 요구를 만들지 않는다(이름만 볼 일에 파일을 열었고, 파일이 하나뿐이라 부분 완료로 끝났다).
    if _FILE_NAMES_ASK.search(compact) and not _wants_detail(compact) and not _read_intent(compact):
        read_count_context = False
    # 긴 글의 진술 문장에 든 개수(붙여넣은 글의 내용)는 읽을 개수가 아니다 — `_count_only_in_a_statement`.
    if read_count_context and _count_only_in_a_statement(compact):
        read_count_context = False
    # ★ **읽을 곳이 내 PC 의 프로그램 창이면 그 개수는 웹 자료 개수가 아니다** (2026-10-17 화면 자동화 측정 5f7d4d63).
    #   「작업 관리자를 열어서 메모리를 가장 많이 쓰는 프로그램 다섯 개를 … 화면에서 읽어서 알려 줘」가 「5개 이상 직접 열어 읽기」가
    #   되어 판 전체가 웹 조사 갈래로 묶였다 — 작업 관리자를 여는 손이 「이 갈래가 아니다」로 막혔고, 끝에는 「5개를 열어 보라고
    #   하셨는데 하나도 안 열었다」 헛경고가 붙고 모델은 엉뚱하게 웹 결과 열기를 시도했다. 사이트나 주소를 댄 요청은 그대로 센다.
    if read_count_context and not has_named_source and _DESKTOP_SOURCE.search(compact) and not _EMBEDDED_ADDRESS.search(compact):
        read_count_context = False
    if target and _is_action_count(outside_quotes, target):
        # 행위에 붙은 개수. **한 횟수로만** 센다 — 적어 놓는 것으로는 못 채운다.
        found.append(Requirement("count_action", f"{target}개 실제로 처리하기", target))
    elif scoped_counts and read_count_context:
        for scope, scoped_target in scoped_counts:
            found.append(
                Requirement(
                    "count",
                    f"{'로컬 자료' if scope == 'local' else '웹 자료'} "
                    f"{scoped_target}개 이상 {_count_label_tail(compact)}",
                    scoped_target,
                    grounding_required=not wants_a_count(compact) and not says_list_only(compact),
                    source_scope=scope,
                    list_only=says_list_only(compact),
                )
            )
    elif target and read_count_context:
        # 목록 자체가 답인 요청(제목·가격·건수)은 답에 N개가 있으면 충분하다.
        # 그 밖의 조사 요청은 **읽은 N개와 답의 N개가 같은 것**이어야 한다.
        # 두 갈래를 같은 `count`로 두면 상세 셋을 읽고 목록의 다른 회사 셋을
        # 베껴도 숫자만 맞아 통과한다(2026-09-02 사람인 라이브).
        found.append(
            Requirement(
                "count",
                f"{target}개 이상 {_count_label_tail(compact)}",
                target,
                # 사람이 **직접** 「열지 마라 · 열 필요 없다」고 한 요구는 읽기 근거를 따로 따지지 않는다 — 추측보다 지시가 먼저다.
                grounding_required=not wants_a_count(compact) and not says_list_only(compact),
                list_only=says_list_only(compact),
            )
        )

    aliases = dict(SITE_ALIASES)
    aliases.update(known_sites or {})
    if english:
        # 영문 이름(「Saramin」)도 같은 호스트다. 영어 요청에만 — 한국어 요청의 판정은 그대로 둔다.
        for name, host in _en.SITE_NAMES.items():
            aliases.setdefault(name, (host,))
    # 긴 이름부터 본다. "네이버 블로그" 가 걸리면 "네이버" 는 안 건다 —
    # 둘 다 걸면 같은 요구가 두 줄이 되고, 화면의 「시킨 것」이 1/4 처럼 부풀어
    # 실제보다 못한 것처럼 보인다(2026-08-20 쓰기 벤치 실측).
    matched: list[str] = []
    # ★ 영문 이름은 대소문자를 가리지 않는다 (2026-09-14). 사람은 「YES24」·「Yes24」·
    #   「yes24」를 다 쓴다. 한글에는 대소문자가 없으니 casefold 가 아무것도 안 바꾼다.
    folded = outside_quotes.casefold()
    for name in sorted(aliases, key=len, reverse=True):
        if name.casefold() not in folded:
            continue
        # ★ 영문·숫자로 된 이름은 **낱말로** 있을 때만 건다 (2026-10-04 레딧 조사 판 aa741c81). 배운 사이트 이름 「GB」가
        #   「16GB 이하에서」 속에서 걸려 대조표에 「GB에서 하기」가 찍혔다. 한글 이름은 조사가 바로 붙으므로 그대로 둔다.
        if name.isascii() and not _ascii_name_as_word(folded, name.casefold()):
            continue
        if any(name.casefold() in longer.casefold() for longer in matched):
            continue
        # ★ 예시 목록 속 이름은 장소가 아니다 (2026-09-30 손자리 조사 판 e7dd3831). 「처음 쓴 채널(레딧·해커뉴스·깃허브·트위터 등)」의
        #   깃허브가 「깃허브에서 하기」가 됐고, 웹 글을 10개 읽고도 「시킨 곳에서 안 읽었다」로 41걸음을 헤맸다.
        if _only_in_example_lists(folded, name.casefold()):
            continue
        if _only_as_topic(folded, name.casefold()):
            continue
        matched.append(name)
        found.append(Requirement("site", f"{name}에서 하기", aliases[name]))

    # ★ 2026-10-03 Online-Mind2Web TED 판: 「가장 많이 본 · 12~18분」을 검색어 결과에서 골라 틀렸다(정답은 제목에 robot 이 없는
    #   쿼드콥터 강연 — 사이트 필터 목록의 1위). 지름길(검색)이 「되는 것처럼」 보이면 모델은 필터 방에 안 들어간다.
    #   최상급·범위 조건이면 **사이트의 정렬·필터를 실제로 걸었는가**를 대조표에 둔다(말이 아니라 연 주소로 채운다).
    if asks_for_site_filter(compact):
        found.append(Requirement("filter", "사이트 정렬·필터로 보기", site_filter_needs(compact)))

    picked: set[str] = set()
    for pattern, tool, label in (_en.ACTIONS if english else _ACTIONS):
        # 영어는 부정이 동사 **앞**에 온다(「Don't publish」) — 매치 앞을 본다. 부정된 매치 말고 긍정 매치가
        # 따로 있을 수 있으니(「Don't publish yet — publish it tomorrow」) 영어는 매치를 전부 훑는다.
        if english:
            match = next((m for m in pattern.finditer(compact) if not _en.negated_before(compact, m.start())), None)
        else:
            # ★ **파일 이름 속 낱말은 이름이지 시킨 일이 아니다** (2026-10-04 밤 판 98ae0a6b). 「…엑셀 파일 「커뮤니티_게시규칙.xlsx」로
            #   만들어 줘」의 「게시」가 발행 요구로 잡혀, 규칙을 다 읽고 파일까지 만든 판이 「발행 — 실행 안 됨 · 일부만 끝냄」으로 나갔다.
            #   이름의 **몸통 안**에서만 걸린 것을 건너뛴다 — 확장자까지 걸친 것(「index.html 을 만들어 띄워」)은 그대로 요구다.
            match = next((m for m in pattern.finditer(compact) if not _inside_a_file_name(compact, m.start(), m.end())), None)
        # "발행은 하지 마" 의 '발행' 은 요구가 아니다. 뽑으면 시키지도 않은 것을
        # 못 했다고 답에 적게 된다(2026-08-20 쓰기 벤치 실측).
        if not match or (not english and _negated(compact, match.end())):
            continue
        # ★ 더 구체적인 것이 이겼으면 넘어간다 (2026-09-01).
        #
        #   「드라이브에 올려줘」가 블로그 **발행** 요구로도 잡혔다 — 발행 낱말표에
        #   `올려\s*줘|업로드` 가 들어 있기 때문이다. 같은 뜻을 두 줄로 세우면
        #   화면의 「시킨 것」이 1/2 처럼 부풀어 **된 판이 못한 것처럼 보인다**
        #   (사이트 요구에서 같은 이유로 이미 한 번 겪었다).
        if any(tool in _SUPERSEDED_BY.get(chosen, ()) for chosen in picked):
            continue
        # ★ **다른 사이트의 「임시저장」은 그 사이트의 단추다** (2026-10-07 폼 벤치). 주소를 준 이력서 수정 화면에서
        #   「…넣고 임시저장해줘」가 네이버 블로그 임시저장 요구가 됐다 — 그 도구로만 채워지는 칸이라 화면의 단추를 눌러도
        #   「임시저장 — 실행 안 됨」이었고, 판은 「남은 요구에 진전 없음」으로 접혔다. 주소나 지목한 사이트가 네이버가
        #   아니면 블로그 행위로 세우지 않는다. 운영 요청 1,644개에서 블로그 행위 요구 18개 중 이 꼴은 벤치 문장 둘뿐이다.
        if str(tool).startswith("naver_blog_") and _acts_on_another_site(compact, found):
            continue
        picked.add(tool)
        found.append(Requirement("action", label, tool))
    # ★ 셈을 시켰으면 「직접 계산해 보기」도 요구다 (2026-08-23).
    #
    # 왜 여기 넣나. 라우터에 계산 갈래를 열었더니 도구를 **쥐어주기만** 하고
    # 부르라고는 안 해서, 모델이 상자를 손에 들고 암산했다. 실측: P2·P3·P4 를
    # 다 맞혔는데 `audit_log` 에 `sandbox_run` 이 0건 — 맞았지만 근거가 없었다
    # (자비스 자신이 "이 숫자는 열어 본 자료 어디에도 없다"고 경고까지 달았다).
    #
    # 요구로 넣으면 기존 기계가 그대로 일한다: `remaining_note` 가 "아직 계산을
    # 안 했다"고 흘려 넣고, `all_met` 이 거짓이라 답하고 끝내지 않으며, 끝내
    # 안 하면 대조표에 ❌ 로 남는다. **새 장치를 만들지 않는다.**
    if wants_a_computation(compact):
        found.append(Requirement("action", "직접 계산해 보기", "sandbox_run"))
    # 파일·코드는 말로 썼다고 끝난 것이 아니다. 성공 걸음이 남긴 **실제 경로**를
    # 아래 `check`가 다시 열 수 있어야 산출물 요구를 채운다. `wants_an_artifact`를
    # 라우터와 같이 쓰므로 갈래와 계획의 낱말표가 갈라지지 않는다.
    if wants_an_artifact(compact):
        found.append(
            Requirement(
                "artifact",
                "실제 산출물 파일 만들기",
                _artifact_extensions(compact),
            )
        )
    # ★ 각주는 **흉내 낼 수 있다** (2026-09-01 라이브).
    #
    #   「고칠 만한 곳마다 각주로 수정 제안을 달아줘」에 자비스는 본문 안에
    #   `[수정 제안: …]` 라고 써 넣었다. 그건 각주가 아니다 — 그때는 도구에
    #   각주 블록이 아예 없어서 그것밖에 할 수 없었다. 도구를 만든 지금, 진짜로
    #   달렸는지는 **파일을 열어 세면** 안다. 세지 않으면 또 흉내 낸다.
    footnote_pattern, everywhere_pattern = (
        (_en.FOOTNOTE_REQUEST, _en.EVERYWHERE) if english else (_FOOTNOTE_REQUEST, _EVERYWHERE)
    )
    # ★ 코드 이야기의 「주석」은 각주가 아니다 (2026-10-08 — 손자리에게 맡긴 숙제 abed2faf).
    #   「SEED 에 여섯을 넣어라. 왜 넣는지 … 짧은 **주석을 달아라**」에 「각주 3개 이상 달기 — 워드는 각주가 안 되니
    #   한글(hwpx)로」가 섰고, 시킨 것을 다 하고 시험까지 통과시킨 판이 「❌ 각주 없음」·덜 끝남으로 찍혔다.
    #   「각주」는 한 뜻이지만 「주석」은 두 뜻이다 — 「각주」라는 말이 없고 코드 대상(파일 이름 · 함수 · 코드 …)이 같이 있으면 코드 주석이다.
    comment_in_code = (
        not english and "각주" not in compact and bool(_CODE_TARGET.search(compact))
    )
    if footnote_pattern.search(compact) and wants_an_artifact(compact) and not comment_in_code:
        # 「곳마다」·「전부」는 **하나로 안 끝난다.**
        #
        # 라이브에서 「고칠 만한 곳마다 각주로 수정 제안을 달아줘」에 각주 1개가
        # 왔다. 이력서 한 장에서 한 줄만 손보고 「완벽히 반영했다」고 답했다.
        # 목표를 1 로 두면 그게 충족으로 통과한다.
        #
        # 3 은 자의적인 수다. 다만 **한 번은 시늉이고 셋은 패턴**이라는 선은
        # 그을 수 있다. 더 정확한 수를 아는 척하는 것보다 낫고, 미달이면 도는
        # 중에 「아직 남았다」가 들어가 자비스가 더 단다.
        # ★ 개수를 말했으면 **그 수만큼** 이다 (2026-09-02).
        #
        #   「1순위부터 5순위까지 찾아서 각주로 수정 제안 달아줘」는 각주 5개다.
        #   전에는 「곳마다」면 3, 아니면 1 이라 5순위를 시켜도 하나로 통과했다.
        #   시킨 수가 프롬프트에 **적혀 있는데** 다른 수를 목표로 삼을 이유가 없다.
        counted = _count_target(compact)
        target = counted if counted > 1 else (3 if everywhere_pattern.search(compact) else 1)
        label = f"각주 {target}개 이상 달기" if target > 1 else "각주로 달기"
        # ★ **워드 + 각주는 불가능한 조합이다** (2026-09-01 실측).
        #
        #   python-docx 에는 각주 API 가 없다. 「워드 파일로 옮기고 각주 달아줘」를
        #   그대로 두면 산출물은 `.docx` 를 요구하고 각주는 3개를 요구해서 **어떻게
        #   해도 못 채우는 체크리스트**가 된다. 못 채울 요구를 세워 두는 것은
        #   재는 자가 하는 일이 아니다.
        #
        #   그래서 hwpx 를 **대안으로 열어 주고** 이름에 이유를 적는다. 자비스는
        #   `remaining_note` 로 이 말을 받아 한글로 만든다 — 막는 대신 알려 준다.
        if any(item.kind == "artifact" for item in found):
            index = next(i for i, item in enumerate(found) if item.kind == "artifact")
            exts = tuple(found[index].target or ())
            if exts and ".hwpx" not in exts:
                found[index] = Requirement(
                    "artifact", "실제 산출물 파일 만들기", (".hwpx", *exts)
                )
                label += " — 워드는 각주가 안 되니 한글(hwpx)로"
        found.append(Requirement("footnote", label, target))
    if wants_verification(compact):
        found.append(Requirement("verify", "실행·테스트 통과 확인", "execution"))
    # 라우터가 이미 "가서 봐야 아는 일"이라고 결정했는데 위의 결정론 요구가
    # 하나도 없으면 plan={}가 된다. 그러면 미달 판정과 되먹임이 모두 꺼진다.
    # 라우터의 판정을 새로 추측하지 않고 그대로 요구 한 칸으로 옮긴다.
    #
    # 목록·건수 자체가 답인 요청은 상세 본문을 강제하지 않는다. 그 구분은
    # `wants_a_count` 한 곳이 SSOT다(무신사 목록 사건의 역회귀 방지).
    # ★ 2026-09-23 스토어 촬영: 「바탕화면에 회의 메모.txt 만들고 …라고 적어 줘」에 「근거가 되는 대상 실제로 읽기」가
    #   붙었다(라우터가 파일 만들기를 전부 「봐야 아는 일」로 본다). 읽을 게 없으니 영원히 미달이고, 되받기가 조사용
    #   재촉을 보내자 모델이 **시킨 일 대신 그 재촉에** 「지침을 완벽히 이해했습니다」로 답했다.
    #   사람이 **내용을 불러준** 파일 만들기는 읽을 대상이 없다. 재 봤다: 요청 1,181개 중 이 규칙으로 읽기 요구가
    #   빠지는 것 2개 — 둘 다 「memo.txt 를 만들고 '범위 확인'이라고 적어줘」.
    embedded_address = not english and bool(_EMBEDDED_ADDRESS.search(compact))
    dictated_artifact = (
        wants_an_artifact(compact)
        and bool((_en.DICTATED_CONTENT if english else _DICTATED_CONTENT).search(compact))
        and not _read_intent(compact)
        and (not has_named_source or embedded_address)
    )
    if (
        must_observe
        and not dictated_artifact
        # ★ 사람이 **읽지 말라고** 했으면 「실제로 읽기」를 요구로 세우지 않는다 (2026-10-05 판 b4b71589).
        #   「E 드라이브에서 이름에 JARVIS_전략안이 든 파일을 찾아 경로만 알려줘. 파일을 열어 읽지는 말고」에 첫 걸음에서 파일을
        #   찾고도 「근거가 되는 대상 실제로 읽기 — 0개」가 남아 「일부만 끝냄」이 됐고, 하네스가 그 요구를 채우려고 브라우저
        #   결과 열기까지 시켰다. 요청과 어긋나는 요구는 하네스의 버그다. 도구를 한 번은 써야 한다는 것(`must_observe`)은 그대로다.
        #   **이 PC 의 파일 일**(이유가 local·workspace 뿐)일 때만 — 웹 조사의 「이 두 건은 다시 열지 말고」·「다른 페이지는 열지 마」는
        #   일부만 막는 말이라 나머지는 읽어야 한다(운영 요청 1,538개 중 「읽지 말라」 3개, 그 둘이 이 모양이었다).
        and not (
            observe_from and frozenset(observe_from) <= {"local", "workspace"} and _READ_FORBIDDEN.search(compact)
        )
        and not wants_a_count(compact)
        and not any(item.kind == "count" and item.grounding_required for item in found)
    ):
        reasons = frozenset(observe_from or ())
        target = (
            "browser"
            if (_read_intent(compact) or has_named_source)
            and not wants_code(compact)
            and not wants_an_artifact(compact)
            # ★ 2026-09-30: 라우터가 웹 때문에 켠 것이 아니면 브라우저 읽기를 요구하지 않는다. 「E:\…\CLAUDE.md 파일을
            #   열어서 첫 줄만 알려 줘」가 「열어서」(읽기 동사) 때문에 browser 대상이 돼, 브라우저로 억지로 읽히는
            #   쪽(`_forced_read_call`·본문 열기 재촉)으로 밀렸다. 이유가 없으면(옛 호출) 옛 추측 그대로다.
            and (not reasons or bool(reasons & {"web", "named"}))
            else "local"
            if reasons and reasons <= {"local", "document", "workspace"}
            else "any"
        )
        found.append(Requirement("observe", "근거가 되는 대상 실제로 읽기", target))
    if code_workspace:
        # ★ Code 탭에서는 **읽기 개수·셈 요구를 낱말로 뽑지 않는다** (2026-09-30, 대표 「코드 모드에선 두더지가 특히 나쁘다」).
        #   코드 과제의 문장은 사양서다 — 「참가자 3명」은 넣을 예시, 「모바일에서도」는 반응형 조건, 「카드 3개」는 화면 구성.
        #   재 봤다(운영 코드 판 196개): 헛걸린 것은 읽기 개수(6판)와 「직접 계산해 보기」(2판, 둘 다 못 채움)뿐이고, 못 채운
        #   5판이 되받기 중앙 2·525초, 최악 AURORA 4번·66걸음·913초. 코드 판의 채점자는 이미 있다 — 「실행·테스트 통과 확인」
        #   (163판 중 헛걸림 1)과 「산출물 만들기」(118판 0). 이 둘과 저장소 읽기(근거 읽기, 이유로 붙는다)는 그대로 둔다.
        #   읽기 개수는 **실제로 읽을 자리를 댔을 때만** 남긴다(「깃허브에서 라이브러리 3개 비교」 같은 코드 모드 조사).
        readable = _names_a_readable_place(compact, known_sites)
        found = [
            item for item in found
            if not (item.kind == "count" and not readable)
            and not (item.kind == "action" and item.target == "sandbox_run")
        ]
    return found


# 사람이 조건으로 말하는 것들과, **그 조건과 같이 설 수 없는 말.**
#
# ★ 2026-08-16 실측. "정규직 공고 3개를 골라줘"라고 시켰더니 계약직 둘과 파견직
# 하나를 가져왔다. 개수(3개)는 맞았고 계획 대조도 통과했다 — **개수는 세는데
# 조건은 아무도 안 봤기 때문이다.** 그리고 못 지켰다는 말도 안 했다.
#
# 여기서 하는 것은 "그 공고가 정말 정규직인가"를 확인하는 것이 **아니다.** 그건
# 본문을 읽어야 알고, 읽었는지는 이미 다른 검사가 본다. 여기서 잡는 것은 더 싸고
# 더 확실한 것 — **답이 스스로 어긴 경우**다. 정규직을 시켰는데 답에 "계약직"이라고
# 적혀 있으면, 그건 자비스 자신이 아니라고 적어 놓은 것이다.
#
# 이 방향으로만 판정하면 오탐이 없다. 못 잡는 것(답에 형태를 안 적은 경우)은
# 있지만, 없는 위반을 만들어 내지는 않는다.
_CONDITION_CONFLICTS: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("정규직", ("계약직", "파견직", "인턴", "아르바이트", "프리랜서", "기간제")),
    ("계약직", ("정규직",)),
    ("신입", ("경력자만", "경력 필수", "경력만")),
    ("재택", ("출근 필수", "상주",)),
    ("무료", ("유료",)),
)


#: 결과 전체를 설명하거나, 못 찾았다고 하거나, 사과하는 문장 — 항목을 내놓는 말이 아니다.
_EXPLAINS_ABSENCE = re.compile(r"모두|전부|대부분|찾지\s*못|없었|없습니다|죄송|사과|바로잡|혼란을")
#: 답이 항목을 내놓는 줄 — 표 줄 · 글머리표 · 번호 목록.
_ITEM_LINE = re.compile(r"^\s*(?:\||[-*•·▪◦]\s|\d{1,2}[.)]\s|[①-⑳])")
#: 「뺐다」는 말 — 이 말이 든 줄과, 이 말이 든 제목 아래의 줄은 **빼 놓은 것의 설명**이다.
_EXCLUDED_WORDS = re.compile(r"제외|빼고|뺐|뺀\s|빼\s*놓|넣지\s*않|아닌\s*항목|해당\s*없|조건\s*불충|미충족|맞지\s*않아")
_HEADING_LINE = re.compile(r"^\s*(?:(?P<hashes>#{1,6})\s|\*\*[^*\n]{2,60}\*\*\s*$)")


def without_excluded_parts(answer: str) -> str:
    """답에서 **빼 놓은 것을 설명하는 자리**를 뗀 글.

    ★ 2026-10-04 채용 공고 조사 판(788ef850): 답이 파트타임·기간제 공고 셋을 추천 표에서 **빼고** 「표에서 제외한 항목 (정규직 조건
      불충분)」 아래에 이유와 함께 적었다 — 조건을 지킨 것이다. 그런데 그 「기간제」 글자를 보고 「'정규직'을 시키셨는데 답에 기간제가
      들어 있습니다」가 붙었다. 잘한 일에 경고가 붙으면 다음부터 빼 놓은 이유를 안 적게 된다.
    제목에 「뺐다」는 말이 있으면 다음 제목까지, 아니면 그 말이 든 줄만 뗀다. 못 잡는 것(같은 줄에 진짜 위반과 「제외」가 같이 있는 경우)은
    있어도 없는 위반을 만들지 않는다 — 이 검사의 원래 약속 그대로다.
    """
    # ★ **제목의 층을 본다** (같은 날 실전 판 574826ac). 「### 표에 넣지 않고 뺀 공고 (15건 전체)」 아래에 굵은 글씨 소제목
    #   「**계약직이고 지역도 다른 곳**」이 나오자 「빼 놓은 자리가 끝났다」고 보고, 빼 놓은 계약직 공고 열 건에 경고를 붙였다.
    #   빼 놓은 자리는 **같거나 더 높은 층의 제목**이 와야 끝난다 — 굵은 글씨 줄은 `#` 제목보다 낮은 층이다.
    kept: list[str] = []
    skip_level = 0  # 0 = 안 건너뜀 · 1~6 = 그 층의 `#` 제목 아래 · 7 = 굵은 글씨 소제목 아래
    for line in str(answer or "").split("\n"):
        heading = _HEADING_LINE.match(line)
        if heading:
            level = len(heading.group("hashes")) if heading.group("hashes") else 7
            if _EXCLUDED_WORDS.search(line):
                skip_level = level if not skip_level else min(skip_level, level)
                continue
            if skip_level and level <= skip_level:
                skip_level = 0
        if skip_level or _EXCLUDED_WORDS.search(line):
            continue
        kept.append(line)
    return "\n".join(kept)


def broken_conditions(prompt: str, answer: str) -> list[tuple[str, str]]:
    """시킨 조건과 **답이 스스로 어긋난** 자리. `(시킨 것, 답에 적힌 것)`.

    개수 대조(`check`)의 빈 자리를 메운다 — 그쪽은 몇 개인지만 보고 **무엇인지**는
    안 본다. 실측에서 "정규직 3개"에 계약직·파견직이 섞여 왔는데 개수가 맞아서
    전부 통과했다.

    ★ 답이 그 조건을 **지켰는지**는 여기서 판정하지 않는다(본문을 읽어야 안다).
    **어겼다고 답에 적혀 있는 것**만 잡는다. 그래서 없는 위반을 만들지 않는다.
    """
    asked = " ".join((prompt or "").split())
    # ★ **왜 없는지 설명한 문장**은 어긴 것이 아니다 (2026-10-04 실전 판 574826ac). 조건에 맞는 것이 없다고 답하며
    #   「검색 결과 15건은 모두 서울·경기이거나 계약직이었습니다」라고 적은 문장에 「답에 계약직이 들어 있습니다」가 붙었다.
    #   항목 줄(표·목록)은 그대로 다 본다. 글 문장은 결과 전체를 설명하거나(모두·전부·대부분) 못 찾았다·사과하는 문장만 뺀다 —
    #   「계약직 공고를 포함해 3개를 정리했습니다」처럼 **스스로 어겼다고 말한 문장**은 그대로 잡는다.
    kept: list[str] = []
    for line in without_excluded_parts(answer or "").split("\n"):
        if _ITEM_LINE.match(line):
            kept.append(line.strip())
            continue
        kept.extend(
            sentence for sentence in re.split(r"(?<=[.!?。])\s+", line.strip())
            if sentence and not _EXPLAINS_ABSENCE.search(sentence)
        )
    said = " ".join(kept)
    if not asked or not said:
        return []
    broken: list[tuple[str, str]] = []
    for wanted, conflicts in _CONDITION_CONFLICTS:
        if wanted not in asked:
            continue
        for bad in conflicts:
            if bad in said and (wanted, bad) not in broken:
                broken.append((wanted, bad))
    return broken


def broken_condition_note(prompt: str, answer: str) -> str:
    """어긴 것이 있으면 답 끝에 붙일 말. 없으면 빈 문자열."""
    broken = broken_conditions(prompt, answer)
    if not broken:
        return ""
    lines = "\n".join(f"- '{wanted}' 를 시키셨는데 답에 **{bad}** 이 들어 있습니다" for wanted, bad in broken)
    return (
        "\n\n⚠ **시킨 조건과 다른 것이 섞여 있습니다.**\n"
        + lines
        + "\n조건에 맞는 것만 다시 찾아드릴까요?"
    )



# ── 분야 조건의 근거 (2026-10-03 K-Startup 5차 9ac405c8) ────────────────────────────────────────────
# 「AI나 소프트웨어 분야에 신청할 수 있는 것 5개」를 시켰더니 사업명·본문 어디에도 AI·SW 가 없는 일반 사업 5개(팁스·SVC…)를
# 「AI/소프트웨어 창업지원사업」이라는 제목 아래 적었다. 개수 압박이 조건을 이겼다 — 안내문으로는 안 됐다(사고 R).
# `broken_conditions` 는 답이 **스스로 어긴** 것(「정규직」인데 「계약직」)만 잡는다. 여기서는 반대로 **조건에 맞는다는 근거가
# 없는 줄**을 찾는다: 줄의 이름과 그 줄이 가리키는 읽은 본문 어디에도 조건 낱말이 없으면 「확인 안 됨」이다.
_TOPIC_TERM = r"(?:[A-Za-z][A-Za-z0-9+#.\-]{0,15}|[가-힣]{2,8})"
_TOPIC_JOIN = r"\s*(?:이나|나|또는|혹은|및|와|과|,|·|/)\s*"
_TOPIC_UNIT = _TOPIC_TERM + r"(?:\s+" + _TOPIC_TERM + r")?"
_TOPIC_BEFORE_KO = re.compile(r"(?:^|\s)(" + _TOPIC_UNIT + r"(?:" + _TOPIC_JOIN + _TOPIC_UNIT + r")*)\s*$")
#: 「분야별로」(분야마다)는 조건이 아니다.
_TOPIC_KEYWORD_KO = re.compile(r"\s*(?:분야|관련|부문|업종|계열)(?!별)")
_TOPIC_EN = re.compile(
    r"(?i)\b(?:related\s+to|in\s+the\s+(?:field|area)\s+of)\s+"
    r"([A-Za-z][A-Za-z0-9+#.\-]{0,15}(?:\s*(?:or|and|,|/)\s*[A-Za-z][A-Za-z0-9+#.\-]{0,15})*)"
    r"|\b([A-Za-z][A-Za-z0-9+#.]{0,15}(?:\s*(?:or|/)\s*[A-Za-z][A-Za-z0-9+#.]{0,15})*)-related\b")
#: 조건이 아닌 낱말(요청의 뼈대) — 「창업 관련 지원사업」의 창업처럼 거의 모든 본문에 있는 말은 조건으로 잴 수 없다.
_TOPIC_STOP = frozenset({"관련", "분야", "사업", "공고", "지원", "창업", "지원사업", "중", "중에서", "특정", "정도", "여러", "해당",
                         "다른", "모든", "기타", "직접", "최신", "오늘", "국내", "해외", "이런", "그런", "어떤", "지역", "사이트",
                         "the", "a", "an", "of", "and", "or"})
#: 같은 조건을 가리키는 말. 없으면 낱말 그대로.
_TOPIC_SYNONYMS: dict[str, tuple[str, ...]] = {
    "ai": ("ai", "인공지능", "에이아이", "머신러닝", "딥러닝", "llm", "생성형"),
    "소프트웨어": ("소프트웨어", "sw", "s/w", "software"),
    "데이터": ("데이터", "data", "빅데이터"),
    "디지털": ("디지털", "digital"),
    "바이오": ("바이오", "bio", "생명공학"),
    "it": ("it", "ict", "정보통신"),
}
for _key, _same in (("인공지능", "ai"), ("sw", "소프트웨어"), ("software", "소프트웨어"), ("data", "데이터"),
                    ("digital", "디지털"), ("bio", "바이오"), ("ict", "it"), ("정보통신", "it")):
    _TOPIC_SYNONYMS[_key] = _TOPIC_SYNONYMS[_same]
_ASKS_A_NUMBER = re.compile(r"(?i)\d+\s*(?:개|곳|건|권|명|가지)|\btop\s+\d+\b|\b\d+\s+[a-z][a-z\-]*")


def topic_conditions(prompt: str) -> list[tuple[str, tuple[str, ...]]]:
    """개수와 같이 시킨 **분야 조건** — `[(이름, 같은 말들)]`. 「AI나 소프트웨어 분야 5개」 → AI · 소프트웨어."""
    text = " ".join(str(prompt or "").split())
    if not text or not _ASKS_A_NUMBER.search(text):
        return []
    phrases: list[str] = []
    for keyword in _TOPIC_KEYWORD_KO.finditer(text):
        before = text[: keyword.start()]
        found = _TOPIC_BEFORE_KO.search(" ".join(before.split()[-5:]))
        if found:
            phrases.append(found.group(1))
    for match in _TOPIC_EN.finditer(text):
        phrases.append(match.group(1) or match.group(2) or "")
    out: list[tuple[str, tuple[str, ...]]] = []
    seen: set[tuple[str, ...]] = set()
    for phrase in phrases:
        for raw in re.split(_TOPIC_JOIN + r"|\s+(?:or|and)\s+", phrase):
            # 한 조건이 두 낱말일 수 있다(「AI 에이전트」) — **어느 낱말이든** 있으면 근거로 친다(경고를 줄이는 쪽).
            words: list[str] = []
            for piece in raw.split():
                # 한국어 낱말 끝의 이음 조사(「소프트웨어나」)는 위 split 이 못 뗀다 — 남은 「나」·「이나」를 뗀다.
                if re.fullmatch(r"[가-힣]{3,}", piece):
                    piece = re.sub(r"(?:이나|나|와|과)$", "", piece)
                # 「모집 중인 AI」의 「중인」·「있는」·「가능한」 같은 꾸밈말은 조건이 아니다(짧고 어미로 끝남 — 「블록체인」은 남는다).
                if re.fullmatch(r"[가-힣]{1,2}(?:는|인|한|된|할|을|를|의|에|로|은)", piece) or re.search(
                        r"[가-힣](?:에서|에게|으로|이나|야|요|다)$", piece):
                    continue
                if len(piece) >= 2 and piece.lower() not in _TOPIC_STOP:
                    words.append(piece)
            if not words:
                continue
            same = tuple(dict.fromkeys(word for piece in words for word in _TOPIC_SYNONYMS.get(piece.lower(), (piece.lower(),))))
            if same in seen:
                continue
            seen.add(same)
            out.append((" ".join(words), same))
    return out


def mentions_topic(text: str, same: tuple[str, ...]) -> bool:
    """글에 그 조건 말이 있나. 영문 약어(AI·SW)는 낱말 경계로만 — 「main」 속 ai 를 세지 않는다."""
    low = str(text or "").lower()
    for word in same:
        if re.search(r"[a-z]", word):
            if re.search(r"(?<![a-z0-9])" + re.escape(word) + r"(?![a-z0-9])", low):
                return True
        elif word in low:
            return True
    return False


def _hosts_read(sources: list[str]) -> set[str]:
    return {(urlparse(url).hostname or "").lower() for url in sources}


def _hosts_touched(steps: list[dict[str, Any]]) -> set[str]:
    """성공한 브라우저 걸음이 **실제로 서 있던** 호스트(도구가 돌려준 주소 기준 — 모델 인자가 아니다)."""
    out: set[str] = set()
    for step in steps:
        if not step.get("ok") or not str(step.get("tool") or "").startswith("browser_"):
            continue
        evidence = step.get("evidence")
        url = str(evidence.get("url") or "") if isinstance(evidence, dict) else ""
        host = (urlparse(url).hostname or "").lower()
        if host:
            out.add(host)
    return out


_ARTIFACT_TOOLS = frozenset(
    {
        "local_write_file",
        "local_apply_edit",
        # ★ **여기 없어서 만든 문서가 안 보였다** (2026-09-01).
        #
        #   어제 문서 도구 셋을 만들면서 이 표를 안 고쳤다. 그래서 hwpx·docx 를
        #   제대로 만들어도 계획 층에는 **아무것도 안 만든 것**이었고, 「실제 산출물
        #   파일 만들기 — 열 수 있는 산출물 없음」이 찍혔다. 등록 자리가 열 곳이라는
        #   말은 이런 뜻이다: 한 곳만 빠져도 조용히, 그리고 엉뚱한 모양으로 틀린다.
        "local_write_document",
        "local_edit_document",
        "local_edit_spreadsheet",
        "artifact_from_observations",
        "sandbox_run",
        "claude_code_delegate",
        "hwp_apply_edit",
        "hwp_apply_format",
        "hwp_export_pdf",
        "browser_screenshot",
    }
)
_VERIFY_TOOLS = frozenset(
    {"sandbox_run", "local_run_project_command", "claude_code_delegate"}
)


def _successful_step(step: dict[str, Any]) -> bool:
    evidence = step.get("evidence")
    evidence = evidence if isinstance(evidence, dict) else {}
    return bool(
        step.get("ok")
        and step.get("status") != "executed_unverified"
        and evidence.get("verified") is not False
    )


def _step_output_paths(step: dict[str, Any], *, even_unverified: bool = False) -> list[Path]:
    """파일을 **만드는 도구가 선언한 경로만** 꺼낸다.

    읽기 도구에도 `path`가 있으므로 모든 걸음에서 경로를 모으면 읽은 파일이 산출물로
    둔갑한다. R4 출력 카드와 같은 규율로 생산 도구를 먼저 제한한다.
    """
    if str(step.get("tool") or "") not in _ARTIFACT_TOOLS:
        return []
    if not (_successful_step(step) or (even_unverified and step.get("ok"))):
        return []
    evidence = step.get("evidence")
    evidence = evidence if isinstance(evidence, dict) else {}
    raw: list[str] = []
    for key in ("output_path", "path"):
        value = str(evidence.get(key) or "").strip()
        if value:
            raw.append(value)
    for key in ("output_paths", "changed_paths"):
        values = evidence.get(key)
        if isinstance(values, (list, tuple)):
            raw.extend(str(value or "").strip() for value in values)
    # 원문 ToolResult가 들어오는 단위 검사도 받는다. 장부에서는 `_step_evidence`가
    # 이미 saved_to를 output_paths로 평평하게 만든다.
    saved = evidence.get("saved_to")
    if isinstance(saved, list):
        raw.extend(
            str(item.get("path") or "").strip()
            for item in saved
            if isinstance(item, dict)
        )
    return list(dict.fromkeys(Path(value) for value in raw if value))


#: 각주를 달라고 했는가. 「주석」도 같은 뜻으로 쓴다.
_FOOTNOTE_REQUEST = re.compile(r"각주|주석")

#: 「하나만」이 아니라는 말. 「곳마다」·「전부」면 각주도 여러 개여야 한다.
_EVERYWHERE = re.compile(r"곳마다|군데마다|부분마다|줄마다|항목마다|모든|전부|여러|각각|하나하나")

#: 각주 **태그**. `<hp:footNotePr>`(각주 설정)에 헛맞지 않도록 경계를 요구한다.
_FOOTNOTE_TAG = re.compile(r"<hp:footNote[\s/>]")


def _footnotes_in(path: Path) -> int:
    """그 문서에 **진짜 각주**가 몇 개 있는가.

    HWPX 는 ZIP+XML(OWPML) 이라 `<hp:footNote` 태그를 세면 된다. 본문에 써 넣은
    `[수정 제안: …]` 는 여기 안 걸린다 — 그게 이 눈금의 존재 이유다.

    ⚠ **태그 경계를 요구한다.** 그냥 `"<hp:footNote"` 로 세면 빈 문서에도 있는
    각주 **설정**(`<hp:footNotePr>`)이 접두사로 걸려 바닥값이 1 이 된다. 실측에서
    각주를 하나도 안 단 문서가 「각주 1개」로 통과했다 — 흉내 낸 것을 진짜로 세는,
    이 눈금이 잡으려던 바로 그 고장이다.
    """
    if path.suffix.lower() != ".hwpx":
        return 0
    try:
        import zipfile

        with zipfile.ZipFile(path) as bundle:
            return sum(
                len(_FOOTNOTE_TAG.findall(bundle.read(name).decode("utf-8", "ignore")))
                for name in bundle.namelist()
                if name.startswith("Contents/section") and name.endswith(".xml")
            )
    except Exception:
        return 0


#: 프로그램 창에서 실행한 동작 — 그 **뒤에** 생긴 파일만 창으로 만든 산출물 후보다.
_GUI_ACTION_TOOLS = frozenset({"windows_execute_preview"})
#: 그 파일을 실제로 **열어 본** 걸음 — 이름만 짐작한 파일은 세지 않는다.
_GUI_MADE_READERS = frozenset({"local_read_document", "local_read_text_file"})


def _gui_made_paths(steps: list[dict[str, Any]]) -> list[Path]:
    """**프로그램 창으로 만든 파일** — 창 동작을 실행한 뒤 손자리가 직접 열어 읽었고, 그 동작 뒤에 쓰인 파일 (2026-10-17).

    한글을 메뉴로 조작해 PDF 를 저장하고 그 PDF 를 열어 5쪽까지 확인한 판(화면 자동화 측정)에 「❌ 실제 산출물 파일 만들기 —
    열 수 있는 산출물 없음」이 붙었다 — 산출물은 파일 쓰기 도구가 선언한 경로만 셌다. 지어낸 파일이 끼지 않게 셋을 다 본다:
    창 동작이 먼저 있었고 · 그 파일을 실제로 읽었고(성공) · 파일이 첫 창 동작보다 나중에 쓰였다.
    """
    started = [
        float((step.get("evidence") or {}).get("executed_at") or 0)
        for step in steps
        if step.get("ok") and str(step.get("tool") or "") in _GUI_ACTION_TOOLS
    ]
    started = [value for value in started if value > 0]
    if not started:
        return []
    since = min(started) - 2.0
    made: list[Path] = []
    # 저장 창이 닫힌 뒤 창 도구가 **그 자리에 파일이 실제로 생긴 것을 확인해** 적은 경로(`saved_file` · 5046e2e7)도 센다
    # (2026-10-17 앱 묶음 · 그림판 cf93ca97 — 저장하고 폴더 목록만 본 판에 「산출물 없음」). 도구가 적은 값이라 모델 말이 아니다.
    for step in steps:
        if step.get("ok") and str(step.get("tool") or "") in _GUI_ACTION_TOOLS:
            saved = str((step.get("evidence") or {}).get("saved_file") or "").strip()
            if saved:
                try:
                    if Path(saved).is_file():
                        made.append(Path(saved))
                except OSError:
                    continue
    for step in steps:
        if not step.get("ok") or str(step.get("tool") or "") not in _GUI_MADE_READERS:
            continue
        raw = str((step.get("arguments") or {}).get("path") or "").strip()
        if not raw:
            continue
        path = Path(raw)
        try:
            if path.is_file() and path.stat().st_mtime >= since:
                made.append(path)
        except OSError:
            continue
    return list(dict.fromkeys(made))


def _open_artifacts(
    steps: list[dict[str, Any]], extensions: tuple[str, ...]
) -> list[Path]:
    opened: list[Path] = []
    wanted = {suffix.lower() for suffix in extensions}
    candidates = [path for step in steps for path in _step_output_paths(step)] + _gui_made_paths(steps)
    for path in candidates:
        try:
            if wanted and path.suffix.lower() not in wanted:
                continue
            if not path.is_file():
                continue
            # 존재 확인과 열기 확인은 다르다. 잠긴 파일·깨진 경로를 성공으로 세지
            # 않도록 실제로 파일 핸들을 연다. 내용 품질은 이 눈금의 범위가 아니다.
            with path.open("rb") as stream:
                stream.read(1)
        except OSError:
            continue
        opened.append(path)
    return list(dict.fromkeys(opened))


def _made_but_unverified(steps: list[dict[str, Any]], extensions: tuple[str, ...]) -> list[Path]:
    """도구가 **만들기는 했지만 「확인 안 됨」으로 돌려준** 파일 가운데 실제로 있는 것.

    ★ 2026-10-05 AWS 영어 윈도우(9B): 문서 내용을 글자열로 보내 블록 일부가 깨졌다. 도구는 나머지로 37KB 문서를 만들고
      `verified=false`(빠진 것이 있다)를 돌려줬는데, 대조표는 「열 수 있는 산출물 없음」이라고 했다 — 파일은 있고 열린다.
      요구는 여전히 **미달**이다(시킨 것이 빠졌다). 다만 본 것은 사실대로 적는다.
    """
    wanted = {suffix.lower() for suffix in extensions}
    found: list[Path] = []
    for step in steps:
        if _successful_step(step):
            continue
        for path in _step_output_paths(step, even_unverified=True):
            try:
                if (not wanted or path.suffix.lower() in wanted) and path.is_file() and path.stat().st_size > 0:
                    found.append(path)
            except OSError:
                continue
    return list(dict.fromkeys(found))


def _execution_verified(steps: list[dict[str, Any]]) -> tuple[bool, str]:
    # 이 PC 안 페이지를 만지고 읽은 확인도 실행 검증이다(2026-09-30 대표 결정, `code_runtime` 과 같은 판정).
    from .code_runtime import browser_verification_indices

    if browser_verification_indices(steps):
        return True, "만든 페이지를 눌러 보고 결과를 읽음"
    for step in reversed(steps):
        tool = str(step.get("tool") or "")
        if tool not in _VERIFY_TOOLS or not _successful_step(step):
            continue
        evidence = step.get("evidence")
        evidence = evidence if isinstance(evidence, dict) else {}
        if tool in {"sandbox_run", "local_run_project_command"}:
            try:
                exit_code = int(evidence.get("exit_code") or 0)
            except (TypeError, ValueError):
                continue
            if exit_code == 0:
                return True, (
                    "프로젝트 검증 종료코드 0"
                    if tool == "local_run_project_command"
                    else "sandbox 종료코드 0"
                )
        elif evidence.get("status") == "completed" and evidence.get("verified") is True:
            return True, "Claude Code 실행·검증 완료"
    return False, "성공한 실행 결과 없음"


#: 「했다」로 셀 수 있는 도구. 화면 조작(누르기·치기)은 여기 없다 —
#  누른 것이 곧 처리된 것은 아니기 때문이다. **끝맺는 도구**만 센다.
_ACTION_TOOLS = frozenset({
    "browser_submit_form",
    "site_promote_draft",
    "naver_blog_save_draft",
    "naver_blog_publish",
    "google_gmail_send",
    "google_calendar_create",
    "google_calendar_update",
    "google_calendar_delete",
    "google_drive_write",
    "local_write_document",
    "local_edit_document",
    "local_edit_spreadsheet",
    "local_write_file",
    "local_apply_edit",
    "local_delete_file",
    "hwp_apply_edit",
    "hwp_apply_format",
    "hwp_export_pdf",
    "windows_execute_preview",
    "artifact_from_observations",
})


def actions_done(steps: list[dict[str, Any]]) -> int:
    """**끝맺은** 행위가 몇 번인가. 성공한 것만, 도구 한 번을 한 개로 센다.

    `executed_unverified` 와 `verified: false` 는 빼는데, 그건 「했는지 모른다」이지
    「했다」가 아니기 때문이다 — 이 저장소가 그 둘을 섞어서 여러 번 속았다.
    """
    done = 0
    for step in steps:
        if str(step.get("tool") or "") not in _ACTION_TOOLS:
            continue
        if not step.get("ok") or step.get("status") == "executed_unverified":
            continue
        evidence = step.get("evidence")
        if isinstance(evidence, dict) and evidence.get("verified") is False:
            continue
        done += 1
    return done


#: 표를 담아 파일을 만드는 손 — 인자에 표가 그대로 들어 있다.
_TABLE_WRITE_TOOLS = frozenset({"local_write_document"})


def _rows_written(steps: list[dict[str, Any]]) -> int:
    """성공한 파일 쓰기에 담긴 표 중 가장 긴 표의 **내용 줄 수**(머리줄 뺌). 표가 없으면 0."""
    best = 0
    for step in steps:
        if not step.get("ok") or str(step.get("tool") or "") not in _TABLE_WRITE_TOOLS:
            continue
        args = step.get("arguments") or {}
        if not isinstance(args, dict):
            continue
        tables: list[Any] = [args.get("rows")]
        sheets = args.get("sheets")
        if isinstance(sheets, dict):
            tables.extend(sheets.values())
        for block in args.get("blocks") or []:
            if isinstance(block, dict) and block.get("type") == "table":
                tables.append(block.get("rows"))
        for rows in tables:
            if isinstance(rows, list) and len(rows) >= 2:
                best = max(best, len(rows) - 1)
    # ★ 문서 쓰기 손으로 안 쓴 표도 센다 (2026-10-04 레딧 기준선 판 92b3776a). 엑셀을 상자(`sandbox_run`)에서 만들자
    #   인자에 표가 없어 「답에 5개」로 세었고, 25줄을 제대로 낸 판이 미달로 찍혔다. 만든 파일을 **열어서** 센다.
    for step in steps:
        if str(step.get("tool") or "") in _TABLE_WRITE_TOOLS:
            continue
        for path in _step_output_paths(step):
            best = max(best, _rows_in_file(path))
    return best


#: (경로, 고친 시각, 크기) → 내용 줄 수. 대조는 걸음마다 돈다 — 같은 파일을 매번 다시 열지 않는다.
_ROWS_IN_FILE_CACHE: dict[tuple[str, int, int], int] = {}


def _rows_in_file(path: Path) -> int:
    """표 파일(.xlsx·.csv·.tsv)의 **내용 줄 수**(머리줄 뺌, 시트가 여럿이면 가장 긴 것). 못 열면 0."""
    suffix = path.suffix.lower()
    if suffix not in {".xlsx", ".csv", ".tsv"}:
        return 0
    try:
        stat = path.stat()
    except OSError:
        return 0
    key = (str(path), stat.st_mtime_ns, stat.st_size)
    if key in _ROWS_IN_FILE_CACHE:
        return _ROWS_IN_FILE_CACHE[key]
    rows = 0
    try:
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            book = load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in book.worksheets:
                    filled = sum(
                        1 for row in sheet.iter_rows(values_only=True)
                        if any(cell is not None and str(cell).strip() for cell in row)
                    )
                    rows = max(rows, filled - 1)
            finally:
                book.close()
        else:
            with path.open("r", encoding="utf-8-sig", errors="replace") as stream:
                rows = sum(1 for line in stream if line.strip()) - 1
    except Exception:  # noqa: BLE001 — 못 여는 파일은 0줄로 본다(없는 줄을 세지 않는다)
        rows = 0
    rows = max(0, rows)
    if len(_ROWS_IN_FILE_CACHE) > 256:
        _ROWS_IN_FILE_CACHE.clear()
    _ROWS_IN_FILE_CACHE[key] = rows
    return rows


#: 화면의 칸에 **값을 넣는** 도구 — 성공했으면 그 화면을 보고 한 일이다.
#: 열린 프로그램 창을 **읽는** 도구 — 내 PC 를 보는 판의 「읽기」로 센다(`answer_claims._DESKTOP_READ_TOOLS` 와 같은 둘).
_WINDOW_READ_TOOLS = frozenset({"windows_uia_snapshot", "windows_copy_from_window"})
#: 이 PC 의 상태를 **읽기 전용 명령**으로 읽는 도구 — 허용목록 실행 파일 · 읽기 하위명령만 돈다(`shell_probe.py` 서두).
_SYSTEM_READ_TOOLS = frozenset({"shell_probe"})
#: 원본 파일을 **열어서** 다른 형식으로 굽는 도구 — 성공했으면 그 원본을 본 것이다.
_SOURCE_OPENING_TOOLS = frozenset({"hwp_export_pdf"})
_DATA_ENTRY_TOOLS = frozenset({
    "browser_fill", "browser_fill_form", "browser_select_option", "browser_pick_suggestion", "browser_type_at",
    "browser_file_upload",
})


def _table_without_rows(path: Path) -> bool:
    """표 파일인데 값 줄이 하나도 없다(머리줄뿐이거나 비었다)."""
    return path.suffix.lower() in {".xlsx", ".csv", ".tsv"} and _rows_in_file(path) == 0


#: 결과 **목록을 만들어 내는** 도구 — 검색 · 사이트 안 검색 · 더 불러오기. (레시피 실행은 이름에 search 가 든 것만.)
_LISTING_MAKERS = frozenset({"browser_search", "browser_site_search", "browser_load_more"})
#: 화면을 옮기지 않고 그 자리에서 읽거나 보는 도구.
_STAYS_ON_THE_PAGE = frozenset({"browser_read", "browser_extract", "browser_snapshot", "browser_links", "browser_find",
                                "browser_scroll", "browser_page_data", "browser_annotate"})


def pages_results_were_opened_from(steps: list[dict[str, Any]]) -> set[str]:
    """**목록 화면**의 주소들 — 검색 도구(검색 · 사이트 안 검색 · 검색 레시피 · 더 불러오기)가 **도착한** 화면.

    목록인지 상세인지는 주소나 낱말로 알 수 없다(사이트마다 다르다). 도구가 말해 준다: 검색 도구가 돌려준 화면이 결과 목록이다.
    검색 레시피처럼 도착한 주소를 결과에 안 싣는 도구는, 바로 다음 성공한 걸음이 서 있는 주소가 그 화면이다.

    ★ 처음 두 번은 「결과를 열기 직전에 서 있던 화면」으로 잡았다가 틀렸다(2026-10-08 — 운영 판에 대 보고 알았다).
      상세 1번을 읽고 곧바로 결과 2번을 열면 그때 서 있던 것은 상세다: 첫 규칙은 충족 739판 중 689판을 뒤집었고,
      「결과 열기로 도착한 화면은 뺀다」로 조인 둘째 규칙도 직접 연 상세 화면을 목록으로 찍었다. 서 있던 곳은 증거가 아니다.
    """
    listings: set[str] = set()
    awaiting_landing = False
    for step in steps or []:
        if not step.get("ok"):
            continue
        tool = str(step.get("tool") or "")
        evidence = step.get("evidence") or {}
        url = str(evidence.get("url") or "")
        makes_listing = tool in _LISTING_MAKERS or (tool == "site_run" and "search" in str(evidence.get("action") or "").lower())
        # ★ 읽는 도구가 **글 목록을 돌려준** 화면도 목록이다 (2026-10-15 — 운영 판 ba5d9277). 검색 주소를 browser_open 으로
        #   직접 열면 검색 도구를 안 거쳐 목록 4곳이 「읽은 글」로 세어졌고(12개 읽음 = 목록 4 + 글 8), 다섯째 주제는
        #   열지도 않은 채 「시킨 것 10」이 찼다. 주소 모양이 아니라 도구가 붙인 `listing` 칸으로 가른다.
        if url and tool in _STAYS_ON_THE_PAGE and evidence.get("listing"):
            listings.add(url)
            awaiting_landing = False
            continue
        if makes_listing:
            if url:
                listings.add(url)
                awaiting_landing = False
            else:
                awaiting_landing = True
            continue
        if awaiting_landing and tool.startswith("browser_"):
            # 그 화면을 **읽거나 본** 걸음일 때만 목록으로 잡는다 — 레시피 바로 다음이 결과 열기면 그 주소는 상세다
            # (운영 판에 대 보니 사람인 두 판의 상세 주소가 목록으로 찍혔다).
            if url and tool in _STAYS_ON_THE_PAGE:
                listings.add(url)
            awaiting_landing = False
    return listings


#: 검색 엔진의 바탕 주소. 이것을 「○○에서 하기」로 지목한 요청은 「거기서 찾아라」이지 「그 주소의 글만」이 아니다(`check`).
_SEARCH_PORTAL_DOMAINS = frozenset({"naver.com", "google.com", "daum.net", "bing.com", "duckduckgo.com"})


def check(
    requirements: list[Requirement],
    *,
    sources: list[str],
    steps: list[dict[str, Any]],
    answer: str = "",
    listed_entities_grounded: bool | None = None,
) -> list[tuple[Requirement, bool, str]]:
    """각 요구가 이뤄졌는가. `(요구, 됐나, 실제로 어땠나)`."""
    done_tools = {
        str(step.get("tool") or "")
        for step in steps
        if step.get("ok")
        and step.get("status") != "executed_unverified"
        and not (
            isinstance(step.get("evidence"), dict)
            and step["evidence"].get("verified") is False
        )
    }
    hosts = _hosts_read(sources)
    # 복합 작업의 입력 자료를 결과 개수로 섞지 않는다. 예: Drive 이력서 1건을 읽고
    # 사람인 공고 2건을 읽은 판이 "공고 3개"로 통과했다. 사용자가 사이트를 지목한
    # 개수 요구는 그 사이트들에서 실제로 읽은 주소만 센다.
    # ★ **검색 엔진 자체를 지목한 것은 장소가 아니다** (2026-10-16 튜토리얼 후보 측정 · 185f5208).
    #   「네이버에서 오늘 국내 뉴스 세 개」 — 검색 결과는 언론사 누리집으로 나간다. 기사 12곳을 열어 읽었는데 네이버 주소만
    #   세서 「1개 읽음」이 됐고, 되받기에 밀려 기사를 더 열다 걸음 한도(48)에 닿아 「덜 끝남」으로 닫혔다(답은 멀쩡했다).
    #   같은 문장 둘째 판은 기사 둘이 `n.news.naver.com` 에 걸려 통과했다 — 운이 판정을 갈랐다.
    #   `agent.site_first_target` 은 이미 「"네이버에서 찾아줘"는 검색엔진으로 가라는 뜻」으로 다룬다. 바탕 주소(`naver.com`)
    #   그대로일 때만 푼다 — 「네이버 블로그」(`blog.naver.com`)처럼 서비스를 지목한 것은 그대로 그 주소만 센다.
    named_hosts = tuple(
        str(host).lower()
        for req in requirements
        if req.kind == "site"
        for host in (req.target if isinstance(req.target, (list, tuple)) else ())
    )
    # 검색 엔진을 하나라도 지목했으면 주소로 가르지 않는다 — 「크몽 후기를 네이버에서 찾아서」의 크몽은 주제다(옛 · 새 대조 edcaf090:
    # 네이버만 빼면 크몽 주소만 세서 더 엉뚱해졌다).
    if any(host in _SEARCH_PORTAL_DOMAINS for host in named_hosts):
        named_hosts = ()
    counted_sources = (
        [
            source
            for source in sources
            if any(
                (host := (urlparse(source).hostname or "").lower()) == wanted
                or host.endswith("." + wanted)
                for wanted in named_hosts
            )
        ]
        if named_hosts
        else sources
    )
    out: list[tuple[Requirement, bool, str]] = []
    for req in requirements:
        if req.kind == "count_action":
            target = int(req.target)
            acted = actions_done(steps)
            # 답에 몇 줄을 적었는지는 **보지 않는다.** 그게 이 갈래의 존재 이유다.
            out.append((
                req,
                acted >= target,
                f"{acted}개 처리함" + (f" (시킨 것 {target}개)" if acted < target else ""),
            ))
        elif req.kind == "count":
            target = int(req.target)
            scoped_sources = counted_sources
            if req.source_scope == "local":
                scoped_sources = [
                    source for source in counted_sources
                    if urlparse(source).scheme.lower() not in {"http", "https"}
                ]
            elif req.source_scope == "web":
                scoped_sources = [
                    source for source in counted_sources
                    if urlparse(source).scheme.lower() in {"http", "https"}
                ]
            # ★ **출처 이름표가 틀려도 판이 미달로 끝나지 않게** (2026-10-04 허깅페이스 판 5749ff16). 「로컬 자료 5개」로 잘못 읽힌 요구는
            #   웹 페이지를 열일곱 장 읽어도 「0개 읽음」이었다. 출처를 **둘로 나눠** 시킨 요청(로컬 3 + 웹 3)에서만 이름표가 뜻을 가진다 —
            #   한쪽만 시켰는데 그쪽으로는 하나도 안 읽고 다른 쪽으로만 읽었으면, 틀린 것은 읽은 쪽이 아니라 이름표다. 읽은 것을 센다.
            scope_fell_back = False
            if (
                req.source_scope in {"local", "web"}
                and not scoped_sources
                and counted_sources
                and not any(
                    other is not req and other.kind == "count"
                    and getattr(other, "source_scope", "") in {"local", "web"}
                    and other.source_scope != req.source_scope
                    for other in requirements
                )
            ):
                scoped_sources = counted_sources
                scope_fell_back = True
            # ★ **결과를 열어 들어간 출발 화면은 「직접 열어 읽은 것」이 아니다** (2026-10-08 실기 — 사람인 판 202ee52f).
            #   「공고 3개의 고용형태와 마감일」에 계획이 「3개 이상 직접 열어 읽기 — 3개 읽음」으로 충족을 찍었는데, 연 상세는
            #   하나뿐이었다(그것도 다른 지역 공고). 검색 목록 화면을 읽은 것 둘이 「읽은 출처」로 세어져 3이 됐고, 「목표 달성」이
            #   켜져 모델이 맞는 공고의 상세를 막 열기 시작한 걸음에서 판이 닫혔다. 모델은 답에 「상세 화면을 각각 열어
            #   확인하지는 못했다」고 스스로 적었다.
            #   낱말이 아니라 걸음 순서로 가른다 — 결과를 열어 들어간 출발 화면과 「더 불러오기」를 누른 화면은 목록이다.
            #   읽기 근거를 따로 따지는 요구에만 건다. 목록이 답인 요구(답에 N개면 충족)는 그대로다.
            skipped_listings = 0
            if req.grounding_required:
                listings = pages_results_were_opened_from(steps)
                detail_sources = [source for source in scoped_sources if source not in listings]
                skipped_listings = len(scoped_sources) - len(detail_sources)
                scoped_sources = detail_sources
            read = len(scoped_sources)
            listed = count_listed_items(answer)
            # ★ 2026-10-03: 목록이 답인 요청은 **파일에 쓴 표**도 답이다. 영어 GitHub 트렌딩 판이 세 줄을 파일 표에만
            #   쓰고 답에는 「파일을 저장했다」만 적어 「답에 0개」→ 미달(partial)로 찍혔다(원문 대조 3/3 맞음).
            #   읽기 근거가 필요한 요구(grounding_required)에는 쓰지 않는다 — 그쪽은 읽은 수가 조건이다.
            in_file = 0 if req.grounding_required else _rows_written(steps)
            listed_in_file = in_file > listed
            listed = max(listed, in_file)
            # ★ 두 길 중 하나면 충족이다. 충분히 읽었거나, 답이 N개를 담았거나.
            #   "일한 방식"이 아니라 "시킨 것을 냈는가"를 본다.
            ok = read >= target or listed >= target
            if req.grounding_required:
                # 상세 확인형은 답에 세 줄을 쓰는 것으로 읽기를 대신할 수 없다.
                # 서로 다른 웹 본문 하나를 세 번 읽고 표에 세 줄을 쓰는 판이
                # complete로 통과한 라이브 회귀(2026-09-03)를 막는다.
                ok = read >= target
                # ★ **그 반대쪽 구멍도 막는다** (2026-09-12 실측).
                #
                #   yes24 판이 상세를 여섯 곳 읽고 **표에 한 권만 써서** 냈는데
                #   이 요구가 충족으로 통과했다. 자기 입으로 「6개 읽음 · 답에 1개」
                #   라고 적어 놓고도 통과시킨 것이다 — 읽은 수로만 판정했으니까.
                #
                #   읽기는 조건이고 목적이 아니다. 시킨 것은 「세 권을 내라」였다.
                #
                #   답이 아직 없으면(판 중간 대조) 걸지 않는다. `listed_entities_grounded`
                #   가 None 인 자리와 같은 이유다 — 쓰기 전에 모자라다고 말할 수 없다.
                #
                #   역검증(과거 445판): 이 규칙만 켜면 33.6% 가 뒤집혔는데, 그 중
                #   대부분이 `count_listed_items` 의 과소 계수였다(불릿이 표를 가렸다).
                #   그것을 먼저 고치니 남은 20건이 되었고, 전부 열어 보니 진짜였다
                #   (답이 없거나, 스스로 「3개 중 2건」이라 적어 놓았거나).
                #
                #   ★ **세지 못하는 답에는 걸지 않는다** (`0 < listed`).
                #
                #     "정규직 공고 1, 정규직 공고 2, 정규직 공고 3을 모두 확인했습니다"
                #     처럼 한 문장에 늘어놓은 답은 표도 불릿도 없어 0개로 세어진다.
                #     그것을 미달로 찍으면 **제대로 낸 판을 막는다** — 실제로 기존
                #     시험 하나가 그렇게 깨져서 잡았다.
                #
                #     못 세는 것과 모자란 것은 다르다. 세어지는데 모자랄 때만 막는다.
                if answer.strip() and 0 < listed < target:
                    ok = False
            # 상세 확인형은 최종 답의 이름이 읽은 원문과 다르면 숫자가 맞아도
            # 통과하지 않는다. None은 아직 답을 쓰기 전의 중간 대조다.
            if req.grounding_required and listed_entities_grounded is False:
                ok = False
            saw = f"{read}개 읽음" + (
                (f" · 파일 표에 {listed}줄" if listed_in_file else f" · 답에 {listed}개") if listed else "")
            if skipped_listings and read < target:
                saw += f" (목록 화면 {skipped_listings}곳은 세지 않음)"
            if scope_fell_back:
                saw += " (요청을 「로컬」로 읽었지만 웹에서 읽은 것을 셈)" if req.source_scope == "local" else " (요청을 「웹」으로 읽었지만 이 PC 에서 읽은 것을 셈)"
            # ★ **0이면 왜 0인지 말한다** (2026-09-21 실전).
            #
            #   「깃허브 오픈소스 10개」 판에서 모델이 7쪽을 열어 읽었는데 체크리스트는
            #   「0개 읽음」 한 마디였다. 숫자는 맞다 — 시킨 곳(github.com) 것만 세니까.
            #   하지만 모델도 사람도 그것만 보면 **아무것도 안 읽은 줄** 안다. 모델은
            #   실제로 일곱 번 읽었으므로 「더 읽어야 하나」로 읽고 같은 짓을 반복했다.
            #   필요한 말은 「깃허브로 가라」였다.
            #
            #   빈 값 하나가 두 사실을 뜻하면 조용히 고장 난다 — 「봤는데 0건」과
            #   「엉뚱한 데서 봤다」는 다른 말이다.
            if not read and named_hosts and sources:
                outside = len(sources) - len(counted_sources)
                if outside > 0:
                    saw += (
                        f" · {', '.join(sorted(named_hosts))} 밖에서 {outside}개 읽음"
                        " — 시킨 곳에서 읽어야 센다"
                    )
            if req.grounding_required and listed_entities_grounded is False:
                saw += " · 답 항목 중 읽은 원문에 없는 이름 있음"
            out.append((req, ok, saw))
        elif req.kind == "observe":
            read = len(sources)
            # ★ **화면에 값을 넣은 판은 본 것이다** (2026-10-07 폼 벤치). 「경력기술서 칸에 이 글을 넣고 임시저장해줘」에
            #   칸을 찾아 채운 판마다 「근거가 되는 대상 실제로 읽기 — 0개 실제로 읽음」이 붙었다. 이 요구는 **안 보고 답하는 것**을
            #   막으려는 것이다 — 그 화면의 칸에 값을 넣는 데 성공했다면 그 화면을 본 것이다. 누르기만 한 것은 세지 않는다
            #   (링크 하나 누르고 기억으로 답하는 조사 판이 통과하면 안 된다).
            entered = sum(1 for step in steps if step.get("ok") and str(step.get("tool") or "") in _DATA_ENTRY_TOOLS)
            if read < 1 and entered:
                out.append((req, True, f"{_ENTERED_ON_SCREEN} {entered}번"))
                continue
            # ★ **내 PC 를 보는 판에서는 프로그램 창을 읽은 것도 본 것이다** (2026-10-17 앱 묶음 측정 · 그림판 cf93ca97). 파일을 그림판으로
            #   열어 창을 읽고 다뤘는데 「근거가 되는 대상 실제로 읽기 — 0개」가 붙었다. 웹을 보라는 판(대상이 웹)은 그대로 둔다 —
            #   창 하나 읽고 기억으로 웹 조사 답을 쓰는 판이 통과하면 안 된다.
            window_reads = sum(1 for step in steps if step.get("ok") and str(step.get("tool") or "") in _WINDOW_READ_TOOLS)
            if read < 1 and window_reads and str(req.target or "") == "local":
                out.append((req, True, f"프로그램 창 {window_reads}번 읽음"))
                continue
            # ★ **이 PC 의 상태를 읽기 전용 명령으로 읽은 것도 본 것이다** (2026-10-18 「모르는 앱」 다시 잼). 「시스템 정보를 열어서 BIOS
            #   버전을 읽어 줘」에 `systeminfo` 로 맞는 값을 읽었는데 「0개 실제로 읽음」이 붙어 ❌ 로 나갔고, 모델은 그 칸을 채우려
            #   「열 대상 파일이 없습니다」를 답 첫 줄에 썼다. 내 PC 대상일 때만 — 웹을 보라는 판은 그대로다.
            system_reads = sum(1 for step in steps if step.get("ok") and str(step.get("tool") or "") in _SYSTEM_READ_TOOLS)
            if read < 1 and system_reads and str(req.target or "") == "local":
                out.append((req, True, f"이 PC 상태를 읽기 전용 명령으로 {system_reads}번 읽음"))
                continue
            # ★ **원본을 열어 바꾼 것도 본 것이다** (2026-10-17 한글→PDF). 「한글 파일을 PDF 로 저장해 줘」에 읽기 요구가 붙고,
            #   변환 도구는 원본을 열어 굽는데도 「0개 실제로 읽음」이 남아 모델이 원본을 따로 뜯어 읽으러 갔다. 웹 대상은 그대로다.
            converted = sum(
                1 for step in steps
                if step.get("ok") and str(step.get("tool") or "") in _SOURCE_OPENING_TOOLS
                and not (isinstance(step.get("evidence"), dict) and step["evidence"].get("verified") is False)
            )
            if read < 1 and converted and str(req.target or "") in {"local", "any"}:
                out.append((req, True, f"원본을 열어 {converted}번 변환함"))
                continue
            out.append((req, read >= 1, f"{read}개 실제로 읽음"))
        elif req.kind == "filter":
            browsed = [step for step in steps if str(step.get("tool") or "").startswith("browser_")]
            if not browsed:
                # 브라우저를 안 쓴 판(로컬 파일·계산) — 사이트 필터를 물을 자리가 아니다.
                out.append((req, True, "브라우저 안 씀"))
                continue
            urls = {str(u) for u in sources}
            for step in browsed:
                evidence = step.get("evidence") if isinstance(step.get("evidence"), dict) else {}
                args = step.get("arguments") if isinstance(step.get("arguments"), dict) else {}
                urls.update(str(v) for v in (evidence.get("url"), args.get("url")) if v)
            # ★ 2026-10-03 TED 2차: 「주제=로봇」 하나만 걸고도 ✅ 였다 — 요청은 「많이 본 순 정렬」과 「12~18분 범위」 둘 다였다.
            #   **요청이 필요로 하는 것(정렬·범위)을 다 걸었을 때만** 채운다. 필요한 것을 모르는 옛 요구(빈 target)는 하나로 충분.
            needs = set(req.target or ()) or {"sort:any"}
            applied: set[str] = set()
            for u in urls:
                applied |= url_filter_kinds(u)
            for s in browsed:
                if not s.get("ok"):
                    continue
                tool = str(s.get("tool") or "")
                text = json.dumps(
                    [s.get("arguments"), (s.get("evidence") or {}).get("label") if isinstance(s.get("evidence"), dict) else ""],
                    ensure_ascii=False)
                if tool == "browser_select_option" or (tool.startswith("browser_click") and _FILTER_CLICK_WORDS.search(text)):
                    # 「Rent (low to high)」·「Lowest」처럼 정렬 낱말 없이 **방향만** 있는 칸도 정렬이다(2026-10-03 A 판 apartments.com).
                    word = _FILTER_CLICK_WORDS.search(text)
                    direction = _sort_direction(word.group(0)) if word else "any"
                    if direction == "any":
                        direction = _sort_direction(text)  # 「낮은가격순」은 걸린 낱말이 「순」뿐이다 — 칸 글 전체에서 읽는다
                    # 「Experience (High to Low)」처럼 방향 표에 없는 정렬도 정렬이다(방향은 「any」 — 「most experienced」 요구를 채운다).
                    if (re.search(r"순|정렬|(?i:sort)", text) or tool == "browser_select_option"
                            or (word and not re.match(r"필터|(?i:filter)", word.group(0)))):
                        applied.add("sort:" + direction)
                    if re.search(r"필터|(?i:filter)", text) or tool == "browser_select_option":
                        applied.add("category")  # 무엇을 거른 칸인지 모른다 — 「filter」 요구만 채운다
            missing = filter_needs_met(needs, applied)
            names = {"range": "범위 필터", "filter": "필터", "sort:any": "정렬", "sort:popular": "많이 본 순 정렬",
                     "sort:price_low": "낮은 가격순 정렬", "sort:rating": "평점순 정렬", "sort:newest": "최신순 정렬"}
            if not applied:
                out.append((req, False, "주소·누른 칸에 정렬·필터 흔적 없음"))
            elif missing:
                out.append((req, False, " · ".join(f"{names.get(m, m)} 안 걸림" for m in missing)))
            else:
                out.append((req, True, " · ".join(f"{names.get(n, n)} 걸림" for n in sorted(needs))))
        elif req.kind == "site":
            hit = [h for h in hosts if any(h == part or h.endswith("." + part) for part in req.target)]
            if hit:
                out.append((req, True, f"{len(hit)}건"))
                continue
            # ★ **거기서 움직인 것도 「거기서 일했다」이다** (2026-09-26 노트북). 네이버 지도에서 검색·길찾기까지 13걸음을
            #   했는데 본문 읽기(browser_extract)가 없다고 「네이버에서 하기 — 거기서 아무것도 안 함」이 찍혔다.
            #   읽었는지는 「근거가 되는 대상 실제로 읽기」가 따로 본다 — 여기는 **어디서** 일했나만 본다.
            touched = sorted({
                host for host in _hosts_touched(steps)
                if any(host == part or host.endswith("." + part) for part in req.target)
            })
            if touched:
                out.append((req, True, f"{touched[0]} 에서 작업"))
                continue
            # ★ **그 사이트에서 쓴 것도 「거기서 일했다」이다** (2026-08-20).
            #
            #   "네이버 블로그에 글 쓰고 임시저장해줘" 는 읽으라는 말이 아닌데,
            #   사이트 요구가 읽은 출처만 세어 **성공한 판이 스스로 미이행이라고
            #   답했다**(쓰기 벤치 실측: 임시저장까지 마치고 "0건 읽음"으로 X).
            #
            #   판정은 도구 이름으로 한다. 사이트 이름의 몸통(blog.naver.com → naver)이
            #   성공한 도구 이름에 들어 있으면 거기서 일한 것이다. 브라우저 범용 도구
            #   (browser_*)에는 사이트 이름이 없으므로 헛맞지 않는다.
            stems = {part.split(".")[-2] for part in req.target if part.count(".") >= 1}
            worked = sorted(
                tool for tool in done_tools if any(stem and stem in tool for stem in stems)
            )
            out.append((
                req,
                bool(worked),
                f"{worked[0]} 로 작업" if worked else "거기서 아무것도 안 함",
            ))
        elif req.kind == "footnote":
            made = _open_artifacts(steps, ())
            total = sum(_footnotes_in(path) for path in made)
            target = int(req.target)
            out.append((
                req,
                total >= target,
                (
                    f"각주 {total}개" + (f" (시킨 것 {target}개 이상)" if total < target else "")
                    if total
                    else "각주 없음(본문에 써 넣은 것은 각주가 아니다)"
                ),
            ))
        elif req.kind == "artifact":
            extensions = tuple(req.target) if isinstance(req.target, (list, tuple)) else ()
            opened = _open_artifacts(steps, extensions)
            # ★ **몇 개를 찾아 채우라고 한 판**에서 머리줄만 있는 표는 산출물이 아니다 (2026-10-06 국내 조사 첫 판).
            #   8쪽을 읽고 머리줄만 든 엑셀을 써 둔 순간 이 칸이 ✅ 가 됐다. 세 칸이 다 차 「시킨 것을 다 채웠다, 그만 찾고
            #   답해라」가 나갔고, 판은 글로만 쓰는 마무리로 접혀 「다시 읽겠습니다」로 끝났다 — 빈 표를 남긴 채.
            #   「있고 열린다」만 보던 눈금이 빈 그릇을 다 된 일로 셌다.
            #   개수 요구가 없는 판(빈 양식을 만들어 달라는 요청)은 그대로 둔다 — 거기서는 머리줄만이 곧 시킨 것이다.
            if opened and any(other.kind in {"count", "count_action"} for other in requirements):
                hollow = [path for path in opened if _table_without_rows(path)]
                if len(hollow) == len(opened):
                    out.append((req, False, f"{len(hollow)}개 만들었지만 표에 값 줄이 없음(머리줄만)"))
                    continue
            if opened and req.destination:
                # ★ 이름은 너그럽게, **폴더는 엄격하게** (2026-09-30). 워드클라우드가 잡 작업 폴더에 떨어졌는데 ✅ 였다.
                #   이름까지 맞추라 하면 「회의 메모.txt」처럼 띄어 쓴 이름이 반만 읽혀 된 판을 못 했다고 한다.
                folder = str(Path(req.destination).parent).casefold()
                there = [path for path in opened if str(path.parent).casefold() == folder]
                out.append((
                    req,
                    bool(there),
                    f"{len(there)}개 요청한 자리에 있음" if there
                    else f"요청한 자리({Path(req.destination).parent})에 없음 — 만든 곳: {opened[-1]}",
                ))
                continue
            out.append(
                (
                    req,
                    bool(opened),
                    f"{len(opened)}개 존재·열기 확인" if opened
                    else f"{len(partial)}개 만들었지만 내용 일부가 빠짐" if (partial := _made_but_unverified(steps, extensions))
                    else "열 수 있는 산출물 없음",
                )
            )
        elif req.kind == "verify":
            ok, saw = _execution_verified(steps)
            out.append((req, ok, saw))
        else:
            ok = req.target in done_tools
            out.append((req, ok, "완료" if ok else "실행 안 됨"))
    return out


def remaining_note(results: list[tuple[Requirement, bool, str]]) -> str:
    """루프에 되먹일 말. 아직 남은 것만 짚는다. 다 됐으면 빈 문자열.

    ★ 이걸 **작업 중에** 흘려 넣는 것이 계획 층의 핵심이다. 끝나고 나서 "못 했네"
    하는 것은 정직할 뿐 도움이 안 된다. 아직 할 수 있을 때 "크몽에서 아직 하나도
    못 읽었다"고 말해 줘야 검색 루프에서 빠져나온다.
    """
    missing = [(req, saw) for req, ok, saw in results if not ok]
    if not missing:
        return ""
    lines = "\n".join(f"  · {req.label} — 지금까지 {saw}" for req, saw in missing)
    return "★ 아직 남은 것이 있다. 검색을 되풀이하지 말고 이것부터 채워라:\n" + lines


def all_met(results: list[tuple[Requirement, bool, str]]) -> bool:
    """뽑아낸 요구가 있고, 그게 전부 이뤄졌는가.

    요구가 하나도 없으면 **False** 다. 아무것도 못 뽑았다는 것은 "다 했다"가 아니라
    "모른다"이기 때문이다 — 그걸 다 됐다고 읽으면 문체·구성처럼 애초에 안 뽑는 요구가
    걸린 요청이 첫 걸음에 끝나 버린다.
    """
    return bool(results) and all(ok for _, ok, _ in results)


#: 화면에 값을 넣어 채운 「대상 읽기」의 봄 글 머리 — `check` 가 쓰고 `enough_to_stop` 이 알아본다.
_ENTERED_ON_SCREEN = "화면의 칸에 직접 넣음"


def enough_to_stop(results: list[tuple[Requirement, bool, str]]) -> bool:
    """요구를 다 채웠고, 그 사실이 **그만해도 된다**는 뜻인가.

    ★ 2026-10-07 폼 벤치 넷째 판. 「대상 실제로 읽기」는 **안 보고 답하는 것**을 막는 바닥이다. 그 하나뿐인 계획을 화면의
      칸에 값을 넣어 채운 판은 「할 일을 다 했다」가 아니다 — 폼의 첫 칸을 채운 순간 충족이 되기 때문이다. 그런데 루프는
      그 순간 모델에게 「★ 시킨 것을 다 채웠다. 더 찾지 말고 답해라」를 끼워 넣었고, 걸음 상한에서는 `goal_met` 으로 접었다:
      칸 여섯 개짜리 지원서가 넷을 채운 13걸음에서 끝났다(세 판 중 세 판, 그중 둘은 마무리 답이 비어 작업 실패).
      「완료」 표시(`all_met`)는 그대로 둔다 — 여기는 **일하는 도중에 멈추라고 말할 근거**만 가른다.
    """
    if not all_met(results):
        return False
    return not all(
        getattr(req, "kind", "") == "observe" and str(saw).startswith(_ENTERED_ON_SCREEN) for req, _ok, saw in results
    )


def satisfied_note(results: list[tuple[Requirement, bool, str]]) -> str:
    """다 채웠다고 루프에 알리는 말. `remaining_note` 의 짝이다.

    남은 것을 짚어 주는 말은 있었는데 **다 됐다고 말해 주는 말은 없었다.** 그래서
    시킨 것을 다 하고도 예산이 남아 있으면 계속 찾았다 — 실측 11.7분의 꼬리가
    거기였다. 더 찾아서 나아지는 것이 없는데 계속 도는 것은 자율이 아니라 낭비다.
    """
    if not enough_to_stop(results):
        return ""
    lines = "\n".join(f"  · {req.label} — {saw}" for req, _, saw in results)
    return (
        "★ 시킨 것을 다 채웠다. **더 찾지 말고 지금 가진 것으로 답해라.**\n"
        + lines
    )


#: 행동 요구 이름의 영어판 — **사람에게 보여 줄 때만** 쓴다(2026-10-01 영어판 ③-1).
#: 칸 이름 자체(`Requirement.label`)는 한국어 그대로 둔다: 모델에게 「아직 남은 것」을 알리는 말·장부·시험이 그 이름을 쓴다.
_LABEL_EN = {
    "임시저장": "Save as a draft",
    "드라이브에 올리기": "Upload to Google Drive",
    "발행": "Publish",
    "메일 발송": "Send the email",
    "만든 화면 미리보기로 띄우기": "Show the page in a preview",
    "일정 등록": "Add the calendar event",
    "일정 수정": "Update the calendar event",
    "일정 삭제": "Delete the calendar event",
    "유니티 에디터에서 되돌리기": "Undo in the Unity editor",
    "유니티 에디터에서 실제로 바꾸기": "Actually change it in the Unity editor",
    "유니티 에디터 안 실제로 보기": "Actually look inside the Unity editor",
    "유니티 레시피로 실제 점검하기": "Actually check it with a Unity recipe",
    "직접 계산해 보기": "Calculate it yourself",
    "실제 산출물 파일 만들기": "Make the actual file",
    "실행·테스트 통과 확인": "Confirm it runs / tests pass",
    "근거가 되는 대상 실제로 읽기": "Actually read the sources",
    "사이트 정렬·필터로 보기": "Use the site's own sort / filters",
    "각주로 달기": "Add footnotes",
}


def display_label(req: Requirement, lang: str = "ko") -> str:
    """대조표 칸 이름을 사람의 언어로. 모르는 이름은 **그대로** 둔다(지어내지 않는다)."""
    label = req.label
    if lang != "en":
        return label
    if label in _LABEL_EN:
        return _LABEL_EN[label]
    if req.kind == "count_action":
        return f"Actually do {req.target}"
    if req.kind == "count":
        where = {"local": " local files", "web": " web sources"}.get(req.source_scope, "")
        return f"Open and read at least {req.target}{where} yourself"
    if req.kind == "site" and label.endswith("에서 하기"):
        return f"Do it on {label[: -len('에서 하기')]}"
    if req.kind == "footnote" and label.startswith("각주 "):
        tail = " — Word can't hold footnotes, so as HWPX" if "hwpx" in label else ""
        return f"Add at least {req.target} footnotes{tail}"
    return label


_SAW_EN = (
    (re.compile(r"^완료$"), "done"),
    (re.compile(r"^실행 안 됨$"), "not run"),
    (re.compile(r"^(\d+)건$"), r"\1"),
    (re.compile(r"^(\d+)개 실제로 읽음$"), r"\1 actually read"),
    (re.compile(r"^화면의 칸에 직접 넣음 (\d+)번$"), r"entered into the page's fields \1 time(s)"),
    (re.compile(r"^프로그램 창 (\d+)번 읽음$"), r"read the program window \1 time(s)"),
    (re.compile(r"^원본을 열어 (\d+)번 변환함$"), r"opened and converted the original \1 time(s)"),
    (re.compile(r"^이 PC 상태를 읽기 전용 명령으로 (\d+)번 읽음$"), r"read this PC's state with read-only commands \1 time(s)"),
    (re.compile(r"^(\d+)개 만들었지만 내용 일부가 빠짐$"), r"\1 made, but part of the content was left out"),
    (re.compile(r"^(\d+)개 읽음 · 답에 (\d+)개$"), r"\1 read · \2 in the answer"),
    (re.compile(r"^(\d+)개 읽음 · 파일 표에 (\d+)줄$"), r"\1 read · \2 rows in the file"),
    (re.compile(r"^(\d+)개 읽음$"), r"\1 read"),
    (re.compile(r"^정렬 걸림 · 범위 필터 걸림$"), "sorted and filtered"),
    (re.compile(r"^정렬 걸림$"), "sorted"),
    (re.compile(r"^범위 필터 걸림$"), "filtered"),
    (re.compile(r"^정렬 안 걸림$"), "sort not applied"),
    (re.compile(r"^범위 필터 안 걸림$"), "range filter not applied"),
    (re.compile(r"^정렬 안 걸림 · 범위 필터 안 걸림$"), "neither sort nor range filter applied"),
    (re.compile(r"^주소·누른 칸에 정렬·필터 흔적 없음$"), "no sort/filter found in the pages opened or the buttons pressed"),
    (re.compile(r"^브라우저 안 씀$"), "browser not used"),
    (re.compile(r"^(.+) 에서 작업$"), r"worked on \1"),
)


def _saw_en(saw: str) -> str | None:
    """「본 것」 칸의 영어판 — 아는 꼴만. 모르면 None(지어내지 않는다)."""
    for pattern, replacement in _SAW_EN:
        if pattern.match(saw):
            return pattern.sub(replacement, saw)
    if re.fullmatch(r"(?:[^·]+? (?:안 )?걸림)(?: · [^·]+? (?:안 )?걸림)*", saw):
        words = {"범위 필터": "range filter", "필터": "filter", "많이 본 순 정렬": "most-viewed sort",
                 "낮은 가격순 정렬": "lowest-price sort", "평점순 정렬": "rating sort", "최신순 정렬": "newest sort", "정렬": "sort"}
        parts = []
        for part in saw.split(" · "):
            missing = part.endswith(" 안 걸림")
            name = part[: -len(" 안 걸림")] if missing else part[: -len(" 걸림")]
            label = words.get(name.strip(), name.strip())
            parts.append(f"{label} not applied" if missing else f"{label} applied")
        return " · ".join(parts)
    return None


def english_for_screen(text: str) -> str | None:
    """화면 사전에 없는 **대조표 글**(칸 이름·본 것)의 영어판. 모르면 None.

    ★ 2026-10-05 AWS 영어 윈도우 첫 실행: 대조표의 영어판은 여기(`_LABEL_EN`·`_SAW_EN`)에 있는데 **바깥 AI 로 가는 길만**
      이것을 썼다. 화면으로 가는 길은 사전(`ui_i18n_en.json`)만 봐서, 영어 화면의 「Requested」 칸에
      「사이트 정렬·필터로 보기 — 최신순 정렬 안 걸림」·「1건」이 한국어로 남았다. 같은 말의 영어판을 두 군데서 따로
      챙기면 갈린다 — 화면 번역(`ui_i18n.translate_text`)이 사전에 없을 때 이 함수를 본다.
    """
    return _LABEL_EN.get(text) or _saw_en(text)


def display_saw(saw: str, lang: str = "ko") -> str:
    """대조표의 「본 것」 칸을 사람의 언어로. 모르는 꼴은 그대로 둔다."""
    if lang != "en":
        return saw
    known = _saw_en(saw)
    if known is not None:
        return known
    # ★ 2026-10-03: 나머지 꼴(「1개 존재·열기 확인」 등)은 화면 사전으로. MCP 결과는 화면 번역을 안 거쳐서
    #   영어로 시킨 GitHub 판의 체크리스트가 「Make the actual file — 1개 존재·열기 확인」으로 바깥 AI 에게 갔다.
    from . import ui_i18n  # 늦게 가져온다 — 계획 층이 화면 사전을 미리 읽을 까닭은 없다

    return ui_i18n.translate_text(saw)


def summary_note(results: list[tuple[Requirement, bool, str]], lang: str = "ko") -> str:
    """답변 끝에 붙일 체크리스트. 다 됐으면 붙이지 않는다."""
    if not results or all(ok for _, ok, _ in results):
        return ""
    lines = "\n".join(
        f"- {'✅' if ok else '❌'} {display_label(req, lang)} — {display_saw(saw, lang)}" for req, ok, saw in results
    )
    return ("\n\n**Checklist**\n" if lang == "en" else "\n\n**시킨 것 대조**\n") + lines


# ═══════════════════════════════════════════════════════════════════════════
# 「요구를 못 뽑았다」가 두 가지를 뜻하고 있었다
# ═══════════════════════════════════════════════════════════════════════════
#
# `extract` 가 빈 목록을 돌려주면 `completion_state` 는 `answered` 가 된다. 그런데
# 그 한 칸에 **완전히 다른 둘**이 들어 있었다:
#
#     "지금 몇 시야"              → 대조할 요구가 **원래 없다.** 정상이다
#     "그래서 공고를 좀 찾아줘"    → 일을 시켰는데 **못 뽑았다.** 눈금의 구멍이다
#
# 빈 값 하나가 두 사실을 뜻하면 조용히 고장난다. 그리고 이 경우는 더 나빴다 —
# `measure_autonomy._goal_reached` 가 `answered` 를 **성공으로 셌다.** 즉 못 잰
# 판이 L4 분자에 들어가 있었다(2026-08-21 실측: 위임 54판 중 17판 = 31.5%).
#
# 같은 파일이 바로 위에서 "모르는 것을 성공으로도 실패로도 세지 않고 분모에서
# 뺀다"고 적고 있었다. 그 원칙이 이 칸에만 안 지켜지고 있었던 것이다.

# 일을 시키는 말. **묻는 것과 시키는 것을 가른다.**
#
# 이 목록이 넓어도 되는 이유: 여기 걸린 판은 이미 `extract` 가 요구를 0개로 돌려준
# 판이다. 즉 "일을 시켰는데 대조할 것을 못 뽑았다"를 찾는 자리이지, 갈래를 여는
# 자리가 아니다. 오탐의 값이 싸다 — 눈금에 「못 쟀다」가 하나 더 찍힐 뿐이다.
_TASK_REQUEST = re.compile(
    r"(찾아|알아봐|알아보|조사|검색|읽어|열어|확인해|살펴|훑어"
    r"|정리해|만들어|만들자|생성|써\s*줘|써줘|작성|뽑아|계산"
    r"|보내|발송|올려|업로드|등록|예약|잡아\s*줘|잡아줘"
    r"|해\s*줘|해줘|해\s*주라|해주라|부탁|시작하자|진행해"
    r"|비교해|추천해|골라|분석해|요약해)"
)


def asked_for_work(prompt: str) -> bool:
    """이 말이 **일을 시킨 것**인가, 그냥 물어본 것인가.

    `extract` 가 빈 목록을 준 판에서만 쓴다. 참이면 「일을 시켰는데 못 뽑았다」
    (눈금의 구멍), 거짓이면 「대조할 요구가 원래 없다」(정상).
    """
    compact = " ".join(prompt.split())
    if _en.english_dominant(compact):
        return bool(_en.TASK_REQUEST.search(compact))
    return bool(_TASK_REQUEST.search(compact))


# ═══════════════════════════════════════════════════════════════════════════
# 「세는 일」 — 목록 자체가 답인 요청
# ═══════════════════════════════════════════════════════════════════════════
#
# 2026-08-22 실측. 검색어 다섯으로 크몽 건수를 표로 적으라고 시켰더니:
#
#     검색 5번 정확히 수행 ✓  →  그런데 상세 3건을 열어 읽고 그 요약을 답으로 냄
#
# 답 끝에 흔적이 그대로 남았다: "이제 목표한 3개 읽기를 완료했습니다."
#
# `_forced_read_call` 은 「목록만 보고 답」을 막으려고 만든 장치인데, **목록 자체가
# 답인 요청**에서는 반대로 일을 가로챈다. 세라고 했는데 읽고 앉아 있는 것이다.
#
# 그리고 더 나쁜 일이 같이 났다 — `browser_site_search` 의 `result_count` 가 도구
# 결과에 멀쩡히 있었는데 답에는 한 줄도 안 들어갔다. 사람이 DB 를 열어야 알 수
# 있었다. **자료가 왔는데 답이 못 담으면 그건 답이 아니다.**

# 무엇을 세라는 말인가.
_COUNTING_ASK = re.compile(
    r"(몇\s*(?:개|건|곳|종|번|차례)|건수|개수|갯수|몇\s*개나"
    r"|목록(?:으로|을|만)?\s*(?:줘|보여|뽑|적|정리)"
    r"|표로\s*(?:줘|적|정리|만들)|리스트로|나열해"
    r"|분포|얼마나\s*(?:있|되)|세\s*(?:줘|어)|세어"
    # ★ 2026-08-22 2차. "제목만 적어줘"가 안 걸려서 또 상세 5개를 읽었다.
    #   「무엇**만**」은 그것 말고는 보지 말라는 뜻이다 — 세는 일과 같은 성질이다.
    r"|(?:제목|이름|가격|링크|주소)만\s*(?:줘|적|알려|보여|뽑)"
    # ★ 2026-09-30: 「strategy 폴더에 있는 파일 이름을 세 개만 알려 줘」가 「로컬 자료 3개 이상 직접 열어 읽기」가 돼
    #   이름만 보면 될 일에 파일을 열었고, 폴더에 파일이 하나뿐이라 부분 완료로 끝났다. **파일·폴더의 이름**은
    #   목록이 답이다. 넓게 「이름」 전부로 잡지 않는다 — 「책 이름 3개」 같은 웹 조사는 그대로 둔다.
    r"|(?:파일|폴더|문서)\s*(?:이름|명)(?:을|를|만)?\s*[^.!?\n]{0,12}(?:줘|적|알려|보여|뽑|나열)"
    r"|상위\s*\d+\s*개?\s*(?:제목|이름|만)"
    # ★ 2026-09-03 라이브. 「docs 폴더에 어떤 문서들이 있는지 알려줘」에
    #   「근거가 되는 대상 실제로 읽기 ❌」가 붙었다. **목록이 곧 답인 요청**인데
    #   본문을 안 읽었다고 미달로 찍은 것이다.
    #   「무엇이 있는가」는 열거가 답이다. 속성(마감일·가격)을 물으면 그때 본문이
    #   필요한데, 그건 아래 `_WANTS_DETAIL` 이 브레이크로 잡는다.
    #   실측: 과거 89문장 중 뒤집힌 것 1개(바로 그 문장), 만든 문장 5/5.
    r"|어떤\s*[^.!?\n]{0,20}(?:있는지|있나|있어|있니|들이)\s*(?:알려|보여|말해|줘)"
    # ★ 2026-10-04 레딧 기준선 판(8faf3a05). 「글을 하나하나 **열 필요는 없고** 목록에 보이는 숫자만 적어」라고 직접 말했는데
    #   「25개 이상 직접 열어 읽기」가 붙었다. 사람이 **열지 말라고 한 말**은 낱말 추측이 아니라 지시다 — 목록이 답이다.
    #   좁게 잡는다: 처음엔 「열지 마」·「목록만 보고」까지 넣었는데 운영 요청 1,512개에 돌려 보니 「검색 결과 **목록만 보고
    #   답하지 마**」(읽으라는 말)를 목록 과제로 뒤집었다. 뜻이 한쪽뿐인 두 꼴만 남긴다.
    r"|열\s*필요(?:는|가)?\s*없|목록에\s*(?:보이는|나온)\s*[^.!?\n]{0,8}만\s*(?:적|써|옮|정리))"
)

#: 파일·폴더의 **이름만** 묻는 말(2026-09-30). 목록 보기로 끝나는 일이라 「N개 열어 읽기」를 만들지 않는다.
_FILE_NAMES_ASK = re.compile(r"(?:파일|폴더|문서)\s*(?:이름|명)(?:을|를|만)?\s*[^.!?\n]{0,12}(?:줘|적|알려|보여|뽑|나열)")

# 반대로 **본문을 보라**고 한 말. 이게 있으면 세는 일이 아니다.
_WANTS_DETAIL = re.compile(
    r"(본문|원문|상세(?:\s*페이지|\s*내용)?|내용을?\s*(?:읽|확인|봐|보고)"
    r"|자세히\s*(?:읽|봐|확인)|들어가\s*(?:서|봐)|열어\s*(?:서|봐|읽))"
)


def wants_a_count(prompt: str) -> bool:
    """목록·건수가 **그 자체로 답**인 요청인가.

    참이면 상세를 강제로 열지 않는다 — 목록이 답인데 상세를 열면, 열린 몇 건이
    답을 통째로 차지하고 정작 센 숫자는 사라진다(2026-08-22 실측).

    ★ 「본문을 읽어라」가 같이 있으면 거짓이다. "15건 중 3개는 내용도 봐줘" 같은
      요청은 세는 일이 **아니라** 세고 읽는 일이고, 그때는 강제 읽기가 맞다.
    """
    compact = " ".join(prompt.split())
    if _en.english_dominant(compact):
        if _en.WANTS_DETAIL.search(compact):
            return False
        return bool(_en.COUNTING_ASK.search(compact))
    if _wants_detail(compact):
        return False
    return bool(_COUNTING_ASK.search(compact))

# ═══════════════════════════════════════════════════════════════════════════
# 셈을 시킨 요청인가 — **이 표가 유일한 원본이다**
# ═══════════════════════════════════════════════════════════════════════════
#
# `runtime._requires_sandbox` 가 이걸 가져다 쓴다. 예전에는 runtime 에만 있었는데,
# 계획 층에도 같은 판정이 필요해지면서 두 벌이 될 뻔했다. 이 저장소가 여러 번
# 당한 자리라(`mutating_google_tools` 주석 참고) 한 곳에 둔다.
#
# ★ 2026-08-23. 넓히기 전에 **실제 문장으로** 양방향 역검증했다. 처음엔 "…계산해줘"로
#   고쳐 쓴 표본을 썼는데 게이트는 통과인데 라이브는 그대로 0걸음이었다 —
#   표본이 진짜가 아니면 역검증도 가짜다.
# ── 둘은 다른 것이다 (2026-08-23) ──────────────────────────────────────
#
#   산출물   "엑셀로 만들어줘" · "csv 로 정리해줘"     → 상자가 **파일을 만든다**
#   셈       "평균 얼마야" · "몇 개인지 구해라"        → 상자가 **계산한다**
#
# 처음엔 한 표에 섞어 뒀는데, 그러면 "표로 정리해줘"가 셈으로 잡힌다. 그건 웹
# 조사 결과를 표 **모양**으로 써 달라는 말이지 계산이 아니다(기존 검사
# `test_plan.py::_WISHKET` 이 바로 그 문장이라 그 자리에서 걸렸다).
#
# 갈라 두면 각자 제 일만 한다:
#   라우터    둘 중 하나면 상자를 쥐어준다(예전 동작 그대로)
#   계획 요구 **셈일 때만** `sandbox_run` 을 요구한다 — 산출물은 안 그렇다.
#             표로 정리하는 판까지 상자를 돌려야 끝나게 만들면 오탐의 값이 크다.
_ARTIFACT_TARGET = re.compile(
    r"(엑셀|xlsx|스프레드시트|시트로|csv|워드\s*파일|docx|문서\s*파일"
    r"|그래프|차트|도표|시뮬레이션"
    r"|파일\s*(?:로|을|를|하나)|표\s*(?:로|를|을)\s*(?:만들|정리|뽑))",
    re.I,
)
_ARITHMETIC_TARGET = re.compile(
    r"(소수|약수|배수|최대공약수|최소공배수|경우의\s*수|확률|통계"
    r"|평균|가중평균|중앙값|표준편차|합계|총액|원가|손익"
    r"|몇\s*(?:개|명|시간|분|번|가지|퍼센트|%)|개수|비율|퍼센트"
    r"|등급|점수|할인|인상|이자|복리|계산기"
    # ★ 2026-10-18: 「1부터 100까지 합을 계산해 줘」가 셈으로 안 잡혔다. 「합」 한 글자는 합격 · 합의 · 종합에 걸리므로 꼴로만 받는다.
    r"|총합|합산|까지\s*(?:의\s*)?합(?![가-힣])|까지\s*(?:의\s*)?합(?:을|이|은)|(?<![가-힣])합(?:을|이)?\s*(?:계산|구해|구하))",
    re.I,
)
_COMPUTE_VERB = re.compile(
    r"(만들어|만들자|생성|생성해|뽑아|뽑아줘|계산해|계산\s*좀|돌려\s*봐|돌려봐"
    r"|정리해\s*(?:서|줘)|저장해|출력해"
    r"|구해|구하|셈해|산출|답해|답하|얼마(?:인가|야|예요|입니까))"
)
#: `_COMPUTE_VERB` 중 **만드는·옮기는** 동사. 셈 낱말과 짝지어도 「그 값을 파일·표에 담아라」일 수 있다 — 아래 `_a_listed_column`.
_ARRANGE_VERB = re.compile(r"만들어|만들자|생성해?|뽑아(?:줘)?|정리해\s*(?:서|줘)|저장해|출력해")

# ★ 늘어놓은 칸 이름 (2026-10-16 r/LLMDevs 판 · job 759ba894). 「25개의 제목 · 점수(업보트) · 댓글 수 · 올린 지 얼마나
#   됐는지를 엑셀 파일 하나로 정리해 줘」에 「직접 계산해 보기」가 요구로 붙어 ❌ 로 남고 판이 partial 이 됐다 — 셈 대상
#   낱말표의 「점수」와 「정리해 줘」가 짝지어졌다(「얼마나」는 낱말표에 없다). 그 「점수」는 화면에 적힌 값을 옮길 **칸 이름**이다.
#   낱말을 빼지 않는다(「평균 점수를 계산해줘」는 셈이다). 모양으로 가른다: 만드는 동사와만 짝지었고, 그 낱말이 구분자
#   (· , 、 / 와·과·랑)로 **늘어놓은 항목 중 하나**면 칸 이름이다. 셈 동사(계산해 · 구해 · 얼마야)와 짝지으면 그대로 셈이다.
_LIST_SEP = r"[·・,、/]"
_LIST_ITEM = r"[^\s·・,、/]+"
#: 쉼표 앞 낱말이 이것으로 끝나면 목록이 아니라 절을 이은 쉼표다(「페이지를 열고, 평균을 …」).
_CLAUSE_END = re.compile(r"(?:고|서|면|며|데|요|다)$")


def _a_listed_column(compact: str, start: int, end: int, *, english: bool) -> bool:
    """`compact[start:end]`(셈 낱말)가 「A · B · C」처럼 늘어놓은 이름 중 하나인가 — 이웃 항목이 짧은 이름이어야 한다."""
    sep = rf"(?:{_LIST_SEP}|\s(?:and|or)\s)" if english else rf"(?:{_LIST_SEP}|(?<=[가-힣])(?:와|과|랑|이랑|하고)\s)"
    before = re.search(rf"({_LIST_ITEM})\s*{sep}\s*(?:{_LIST_ITEM}\s+){{0,2}}$", compact[max(0, start - 40):start], re.I)
    if before and (english or not _CLAUSE_END.search(before.group(1))):
        return True
    # 뒤: 붙은 풀이(「점수(업보트)」) · 이어지는 낱말 둘까지 지나 구분자가 오고, 그 뒤에 이름이 있다.
    return bool(re.match(rf"\s*(?:\([^)]{{0,20}}\))?\s*(?:{_LIST_ITEM}\s+){{0,2}}?{sep}\s*{_LIST_ITEM}",
                         compact[end:end + 40], re.I))


# 산출물은 셈과 다른 갈래다. 대상 하나만으로 열면 "보고서 작성법을 조사해줘"까지
# 파일 만들기로 새고, 동사 하나만으로 열면 평범한 글쓰기 대화가 전부 도구 판이 된다.
# 실제 명령형 어미만 대상과 **논리곱**한다.
# ★ 2026-09-23 밤: 「워드로 만들어줘」·「이력서 초안을 워드로 써줘」가 **산출물 요구 0개**였다 — 대상에 워드·PPT 가 없었고
#   동사에 「써줘」가 없었다(엑셀은 있는데 워드는 없는 비대칭). 영어판 짝을 맞추다 드러났다.
_ARTIFACT_REQUEST_TARGET = re.compile(
    r"(파일|엑셀|xlsx|스프레드시트|문서|(?<![가-힣])표(?:로|를)?(?![가-힣])"
    r"|스크립트|코드|보고서|초안|이미지|csv|한글|hwp|워드|docx|pptx?|피피티|파워포인트)",
    re.I,
)
_ARTIFACT_REQUEST_VERB = re.compile(
    r"(만들(?:어|고|어서|자|어줘)|생성해|작성해|저장해|뽑아|짜\s*줘|써\s*줘|고쳐"
    r"|수정(?:해|하고|해서|하라|해주세요)|정리해서\s*(?:줘|달라)"
    # ★ 2026-09-01. 「한글 파일 하나에 본문을 전부 **옮기고** … 각주로 **달아줘**」가
    #   요구 0개로 돌았다(`plan: {}`). 산출물 요구가 없으니 워드로 만들어도, 각주를
    #   본문에 흉내 내도 **아무도 대조하지 않았다.**
    #
    #   넓히기 전에 재 봤다: 최근 서로 다른 프롬프트 320개에 새로 걸리는 것 3건,
    #   전부 진짜 요청(이력서를 한글/HWPX 로 만들어 달라는 것)이고 오탐 0.
    r"|옮기(?:고|어|자|어줘|세요)|옮겨|담아|달아(?:줘|주세요|라)?|넣어(?:줘|주세요)?)",
    re.I,
)
_CODE_TARGET = re.compile(
    r"(파이썬|python|자바스크립트|javascript|typescript|타입스크립트"
    r"|(?<![가-힣])html(?![가-힣])|(?<![가-힣])css(?![가-힣])|함수|코드|스크립트|프로그램"
    # 실제 코딩 지시는 언어 이름 대신 파일명과 테스트 러너를 말하는 경우가 더 많다.
    # 이 표가 그걸 못 보면 다중 파일 판정이 맞아도 `_requires_coding` 논리곱에서
    # 지워져 Claude 위임 전에 로컬 편집 승인을 요구한다(2026-08-25 라이브 재현).
    # 문장 끝 마침표는 이름의 일부가 아니다(「implement beer_song.py. Don't …」) — 확장자 규칙과 같은 고침(2026-09-23 밤).
    r"|(?<![A-Za-z0-9_.])[^\s/\\]+\.(?:py|js|jsx|ts|tsx|java|kt|go|rs|cs|cpp|c|h|html|css|vue|svelte)(?![A-Za-z0-9_]|\.[A-Za-z0-9])"
    # 한국어 조사는 붙여 쓰므로 `pytest를`, `alpha.py에` 뒤의 한글은 경계로 허용한다.
    r"|(?<![A-Za-z0-9_])pytest(?![A-Za-z0-9_])|(?<![A-Za-z0-9_])unittest(?![A-Za-z0-9_]))",
    re.I,
)
_CODE_REQUEST_VERB = re.compile(
    # ★ 「원인을 찾아 calc.py 를 고치고, … 통과를 확인해」가 코드 요청이 아니었다(「고쳐」만 알았다) — 2026-09-23 밤.
    r"(짜\s*줘|작성해|구현해|만들(?:어|고|어서)|생성해|고쳐|고치고"
    r"|수정(?:해|하고|해서|하라|해주세요)|리팩터|실행해|테스트해)",
    re.I,
)
_NAMED_FILE_SUFFIX = re.compile(
    r"\.((?:py|js|jsx|ts|tsx|java|kt|go|rs|cs|cpp|c|h|html|css|vue|svelte|"
    r"json|ya?ml|toml|csv|xlsx|docx|hwp|md|png|jpe?g))"
    # ★ 문장 끝 마침표는 확장자의 일부가 아니다(2026-09-23 밤) — 「beer_song.py 를 구현해줘.」는 잡고
    #   「implement beer_song.py. Don't …」는 못 잡았다. 마침표 **뒤에 글자가 이어질 때만**(`a.py.bak`) 막는다.
    r"(?![A-Za-z0-9_]|\.[A-Za-z0-9])",
    re.I,
)
_VERIFY_REQUEST = re.compile(
    r"((?:테스트|pytest|unittest).{0,24}(?:통과|성공|붙|돌려|실행|확인)"
    r"|(?:통과|성공).{0,16}(?:테스트|pytest|unittest)"
    r"|돌아가게|동작하게|실행되게|실행\s*가능하게|에러\s*없이\s*(?:돌|실행)"
    r"|실제로\s*(?:돌려|실행해|실행해서)|실행\s*결과.{0,12}(?:확인|검증))",
    re.I,
)


def _language_is_topic(compact: str, match: "re.Match[str]") -> bool:
    """이 언어 이름이 **만들 것**이 아니라 **소재**로 쓰였는가.

    `wants_code` 의 `_CODE_TARGET_IS_TOPIC` 과 같은 규율이다. 여기서만 쓰는
    이유는 확장자 규칙에 엑셀·워드처럼 언어가 아닌 것도 섞여 있어서다 —
    「엑셀 책」은 여전히 `.xlsx` 로 두는 게 맞다(그 요청의 산출물은 엑셀이다).
    """
    if match.group(0).lower() not in _TOPIC_PRONE_LANGUAGES:
        return False
    if _en.english_dominant(compact):
        return bool(_en.CODE_TARGET_IS_TOPIC.match(compact, match.end()))
    return bool(_CODE_TARGET_IS_TOPIC.match(compact, match.end()))


#: 소재로도 자주 쓰이는 이름들. 「파이썬 책」·「자바스크립트 강의」.
_TOPIC_PRONE_LANGUAGES = frozenset(
    {"파이썬", "python", "타입스크립트", "typescript", "자바스크립트", "javascript"}
)


#: 「한글이나 워드 파일로」처럼 **둘 중 하나**를 달라는 말 — 어느 쪽으로 만들어도 된다.
_FORMAT_WORD_EXTENSIONS: dict[str, tuple[str, ...]] = {
    "한글": (".hwpx", ".hwp"), "hwp": (".hwpx", ".hwp"), "hwpx": (".hwpx", ".hwp"),
    "워드": (".docx",), "docx": (".docx",), "엑셀": (".xlsx",), "xlsx": (".xlsx",),
    "피피티": (".pptx",), "ppt": (".pptx",), "pptx": (".pptx",), "마크다운": (".md",),
}
_FORMAT_WORD = r"(한글|워드|엑셀|피피티|마크다운|hwpx?|docx|xlsx|pptx?)"
_EITHER_FORMAT = re.compile(
    _FORMAT_WORD + r"\s*(?:파일|문서)?\s*(?:이나|나|또는|혹은|아니면|or|/)\s*" + _FORMAT_WORD + r"\s*(?:파일|문서)?\s*(?:형식|형태)?\s*(?:으로|로)(?![가-힣])",
    re.I,
)


#: 형식 낱말 **바로 뒤**가 이것이면 그 파일은 만들 것이 아니라 **읽을 재료**다 — 「엑셀을 읽어서」 · 「엑셀 파일에서」 · 「한글 파일의 내용을」.
#:
#: ★ 2026-10-13 촬영 판(코드 모드): 「이 폴더의 가격표 엑셀을 읽어서 한눈에 보는 웹페이지로 만들어 줘」가 .xlsx 를 요구했다.
#:   html 을 만들어 화면에 띄우고도 「열 수 있는 산출물 없음 · partial」로 끝났다. 「파이썬 책」의 파이썬(소재)과 같은 자리다.
#:   만들 형식은 「로」가 붙거나(「엑셀로」) 만드는 말이 따른다(「엑셀 파일 만들어줘」) — 그쪽은 여기에 안 걸린다.
_FORMAT_IS_SOURCE = re.compile(
    r"\s*(?:파일|문서|시트|자료)?\s*"      # `match(글, 자리)` 로 부른다 — `^` 를 쓰면 글의 맨 앞만 가리킨다
    r"(?:(?:을|를)\s*(?:읽|열어|보고|보면|보니|참고|분석|확인하|불러|가져)"
    r"|에서(?![가-힣])|에\s*(?:있는|든|적힌|담긴)|의\s*(?:내용|값|데이터|표|숫자|줄|행|열)"
    r"|(?:을|를)?\s*(?:기반|바탕|참고)\s*(?:으로|해|하))"
)


#: 이름이 붙은 파일(「매출표.xlsx」 · 「INSTRUCTIONS.md」)은 **읽는 말이 분명할 때만** 재료다.
#: 「tidy.py 에 있는 함수를 완성해줘」의 `tidy.py` 는 고칠 파일이다 — 「에 있는」 · 「에서」로는 재료라고 못 한다
#: (운영 요청에 대 보니 그 모양이 코드 과제에 있었다).
_NAMED_FILE_IS_READ = re.compile(
    # 「…png」 **그림을** 그림판으로 열어서 (2026-10-17 앱 묶음 측정 cf93ca97) — 파일 · 문서 말고도 그림 · 사진 · 이미지 · 표가 재료 이름 자리에 온다.
    # ★ 「…welcome.hwp」 **한글 파일을** 같은 폴더에 … PDF 로 저장해 줘」 (2026-10-17 한글→PDF 3판) — 따옴표와 「파일」 사이의 형식 이름
    #   (한글 · 워드 · 엑셀 · 파워포인트)을 못 넘어 원본 .hwp 가 만들 형식이 됐고, PDF 를 제대로 만든 판마다 「열 수 있는 산출물 없음」 ·
    #   「일부만 끝냄」이 찍혀 모델이 한 번 더 재촉받았다.
    r"[」』\"'”’)\]]?\s*(?:(?:한글|워드|엑셀|파워\s*포인트|텍스트)\s*)?(?:파일|문서|그림|사진|이미지|표)?\s*(?:을|를)\s*"
    # 「FAQ 문서(docs/faq.md)를 **찾아 읽고**」 (2026-10-19 판 029fae29) — 찾는 말이 끼어도 읽을 재료다.
    r"(?:찾아\s*(?:서\s*)?)?"
    r"(?:읽|열어\s*보|보고|참고|분석"
    # 그 파일을 **다른 형식으로** 저장 · 변환하라는 말이면 그 파일은 재료다(만들 형식은 뒤의 형식 규칙이 잡는다).
    r"|(?:(?!만들|작성|써)[^,.。!?]){0,40}?(?<![A-Za-z])(?i:pdf|jpe?g|png|hwpx|docx|xlsx|pptx|txt|csv)(?![A-Za-z])"
    r"[^,.。!?]{0,12}?(?:로|으로)\s*(?:저장|변환|바꿔|바꾸|내보내)"
    # ★ 「「…welcome.hwp」 파일을 한글 프로그램으로 **열어서** PDF로 저장해 줘」 (2026-10-17 화면 자동화 측정) — 원본 이름을 만들 형식으로
    #   읽어 「.hwp 산출물」을 찾았고, PDF 를 만들어 열어 읽은 세 판 모두 「❌ 열 수 있는 산출물 없음」이 붙었다. 닫는 따옴표와
    #   「(프로그램으로) 열어」를 재료로 본다. 사이에 만들기 · 저장 말이 끼면 그것은 만들 파일이다(「결과.md 를 만들어서 열어 줘」).
    r"|(?:(?!만들|저장|작성|써)[^,.。!?]){0,24}?(?:열어|열고|띄워))"
)


def _format_is_source(compact: str, end: int, *, named_file: bool = False) -> bool:
    """`end` 에서 끝나는 형식 낱말(「엑셀」)이나 파일 이름(「매출표.xlsx」)이 읽을 재료로 쓰였는가."""
    pattern = _NAMED_FILE_IS_READ if named_file else _FORMAT_IS_SOURCE
    if pattern.match(compact, end):
        return True
    if not named_file:
        return False
    # 영어는 읽는 말이 파일 이름 **앞**에 온다 — 「Read INSTRUCTIONS.md and implement beer_song.py」.
    start = max(compact.rfind(" ", 0, end), compact.rfind("\t", 0, end)) + 1
    return bool(_en.NAMED_FILE_IS_READ_BEFORE.search(compact[:start]))


def _artifact_extensions(prompt: str) -> tuple[str, ...]:
    """요청이 형식을 못 박았을 때만 확장자를 제한한다."""
    # 「엑셀은 안 만들어도 되고 … 워드 파일에」 — 거둔 형식은 시킨 형식이 아니다(`_FILE_DECLINED`, 2026-10-19).
    compact = " ".join(_without_declined_files(" ".join(prompt.split())).split())
    # ★ 둘 중 하나를 달라고 했으면 둘 다 받는다 (2026-10-13). 예전에는 규칙 순서가 앞선 쪽 하나만 요구가 돼서, 「한글이나 워드
    #   파일로」에 한글 파일을 만들면 「열 수 있는 산출물 없음」이 됐다.
    either = _EITHER_FORMAT.search(compact)
    if either:
        both = tuple(dict.fromkeys(
            extension for word in either.groups() for extension in _FORMAT_WORD_EXTENSIONS.get(word.lower(), ())
        ))
        if both:
            return both
    # ★ 「llama.cpp 저장소 문서에서 … 워드 파일로」의 `.cpp` 는 형식이 아니라 **읽을 곳의 이름**이다 (2026-09-28, MCP 위임).
    #   워드를 제대로 만들었는데 `.cpp` 파일을 찾느라 「열 수 있는 산출물 없음」이 찍히고 산출물 되받기가 돌았다.
    #   `wants_code` 와 같은 소재 규율로 거른다 — 이름 뒤에 저장소·README·책 같은 말이 오면 소재다.
    topic = _en.CODE_TARGET_IS_TOPIC if _en.english_dominant(compact) else _CODE_TARGET_IS_TOPIC
    named = [
        f".{match.group(1).lower()}"
        for match in _NAMED_FILE_SUFFIX.finditer(compact)
        if not topic.match(compact, match.end()) and not _format_is_source(compact, match.end(), named_file=True)
    ]
    if named:
        return tuple(dict.fromkeys(named))
    rules: tuple[tuple[re.Pattern[str], tuple[str, ...]], ...] = _en.EXTENSION_RULES if _en.english_dominant(compact) else (
        # ★ PDF 규칙이 **아예 없었다** (2026-10-17 화면 자동화 측정). 「한글 창에서 … PDF로 저장해 줘」가 앞의 「한글」로 흘러갔다.
        #   PDF 는 읽을 재료로 훨씬 자주 나온다(「이 PDF 요약해 줘」) — 그래서 **만들 형식 자리**(로 저장 · 변환 · 내보내기)일 때만 잡는다.
        (re.compile(r"(?<![A-Za-z])pdf(?![A-Za-z])\s*(?:파일|문서)?\s*(?:로|으로)\s*(?:저장|만들|변환|바꿔|바꾸|내보내|뽑|출력|인쇄)", re.I),
         (".pdf",)),
        (re.compile(r"엑셀|스프레드시트", re.I), (".xlsx",)),
        (re.compile(r"(?<![A-Za-z])csv(?![A-Za-z])", re.I), (".csv",)),
        (re.compile(r"워드", re.I), (".docx",)),
        # ★ PPT 규칙이 아예 없었다 (2026-09-01). 확장자가 비면 **아무 파일이나**
        #   산출물 요구를 채운다 — 「PPT로 만들어줘」에 워드를 내놔도 통과한다.
        (
            re.compile(r"(?<![A-Za-z])pptx?(?![A-Za-z])|피피티|파워포인트|발표\s*자료|슬라이드", re.I),
            (".pptx",),
        ),
        # ★ 자비스가 만드는 한글 문서는 **hwpx** 다(`documents._write_hwpx`).
        #   `.hwp` 만 요구로 잡으면 제대로 만든 파일이 「산출물 없음」으로 찍힌다 —
        #   재는 자와 만드는 자가 갈리면 반드시 한쪽이 조용히 틀린다.
        (re.compile(r"한글|(?<![A-Za-z])hwpx?(?![A-Za-z])", re.I), (".hwpx", ".hwp")),
        # ★ 마크다운 규칙이 **아예 없었다** (2026-09-07). 그래서 「마크다운 문서로
        #   만들어줘」가 아래 언어 규칙으로 흘러갔다 — PPT 가 없어서 워드로도
        #   통과하던 것과 같은 구멍이고, 이번엔 반대로 **맞는 파일이 거부됐다.**
        (re.compile(r"마크다운|markdown|(?<![A-Za-z])md(?![A-Za-z])", re.I), (".md",)),
        # 아래 언어 규칙은 **소재로 쓰인 이름을 건너뛴다** — `_language_is_topic`.
        (re.compile(r"파이썬|python", re.I), (".py",)),
        (re.compile(r"타입스크립트|typescript", re.I), (".ts", ".tsx")),
        (re.compile(r"자바스크립트|javascript", re.I), (".js", ".jsx")),
        (re.compile(r"이미지", re.I), (".png", ".jpg", ".jpeg", ".webp")),
    )
    for pattern, extensions in rules:
        # 읽을 재료로 쓰인 것은 건너뛰고, 같은 낱말이 뒤에서 만들 형식으로 다시 나오면 그것을 본다
        # (「이 엑셀을 읽고 … 새 엑셀 파일로 저장해줘」). 영어 규칙에는 이 판정이 없다 — 한국어 조사를 본다.
        match = next((found for found in pattern.finditer(compact) if not _format_is_source(compact, found.end())), None)
        if not match:
            continue
        # ★ 「파이썬 **책**」의 파이썬은 형식이 아니라 소재다 (2026-09-07 라이브).
        #
        #   「YES24에서 파이썬 책 2권 … 마크다운 문서로 만들어줘」가 `.py` 를
        #   요구했다. 자비스는 `python_books.md` 를 제대로 만들었는데 계획 층은
        #   「열 수 있는 산출물 없음」이라 했고, 되받기 루프에 빠져 **28분** 을
        #   답만 다시 쓰며 보냈다. 같은 낱말이 같은 날 `wants_code` 도 속였다.
        if _language_is_topic(compact, match):
            continue
        return extensions
    return ()


def artifact_extensions(prompt: str) -> tuple[str, ...]:
    """시킨 형식. 없으면 빈 튜플(아무 형식이나 된다는 뜻).

    `_artifact_extensions` 의 공개 이름이다. 밖에서 사적 이름을 부르면
    이름이 바뀌는 날 조용히 깨진다.
    """
    return _artifact_extensions(prompt)


def wants_verification(prompt: str) -> bool:
    """실행 성공을 결과로 요구했는가. 단순히 '테스트가 뭐야'는 잡지 않는다."""
    compact = " ".join(prompt.split())
    if _en.english_dominant(compact):
        return bool(_en.VERIFY_REQUEST.search(compact))
    return bool(_VERIFY_REQUEST.search(compact))


#: 「정리해줘」류. 이 동사는 **파일 낱말이 있을 때만** 산출물 요구다.
#
# ★ 2026-09-01. 「…워드 파일로 정리해줘」가 요구 0개로 돌아 낱말표에 넣었더니
#   「위시켓과 크몽에서 … 표로 정리해줘」까지 파일 요구가 됐다(검사가 잡았다).
#   그건 **답에 표를 그리면 되는 일**이지 파일을 만들라는 말이 아니다.
#
#   그래서 이 동사만 대상을 좁힌다: 「표」·「문서」처럼 답으로도 되는 말 말고
#   **파일·확장자·프로그램 이름**이 있을 때만.
_TIDY_VERB = re.compile(r"정리해\s*(?:줘|주세요|주라|줄래)")
#: 「**~문서로** 정리해줘」 — 조사 「로」가 **산출물의 형식**을 가리키는 자리.
#
# ★ 2026-09-10 사슬 벤치에서 찾았다. `competitor_skill`("경쟁 분석 문서로 정리해줘")이
#   2판 중 2판 실패했는데, 조사는 다 하고 표까지 짜 놓고 **파일을 안 만들었다.**
#   원인은 모델이 아니라 여기였다 — `wants_an_artifact` 가 False 를 내서 산출물
#   요구가 아예 안 생겼고, 그래서 채팅으로 답해도 **아무도 되묻지 않았다.**
#
#   같은 벤치의 `saramin_doc`("문서로 정리해줘")은 통과했다. 하지만 그건 규칙이
#   맞아서가 아니라 뒷문장의 「표로 넣어줘」가 대신 걸려 준 **운**이었다. 한 판이
#   통과한 것은 한 판만 증명한다.
#
#   「문서」를 `_REAL_FILE_TARGET` 에 넣지 않는 이유는 그대로 둔다 — 「이 **문서를**
#   정리해줘」는 파일 요청이 아니라 읽고 요약해 달라는 말이다. 그래서 **조사 「로」를
#   필수로 잡는다.** 「문서를」은 안 걸리고 「문서로」만 걸린다.
_DOC_FORMAT_REQUEST = re.compile(
    r"(문서|보고서|리포트|장표|기획서|제안서)\s*(?:형식|형태|양식)?\s*으?로"
)
#: 문서 파일 이름 + 「에·에다·로」 — 결과를 받는 곳. 「에서」(읽을 곳)·「을/를」(읽을 것)은 아니다.
_NAMED_DOC_DESTINATION = re.compile(
    r"[\w가-힣.\-\/:]*[\w가-힣]\.(?:xlsx|xls|csv|docx|doc|hwpx?|pptx?|md|txt)\s*(?:에다가|에다|에(?!서)|으로|로)(?![가-힣])",
    re.I,
)
_REAL_FILE_TARGET = re.compile(
    r"(파일|엑셀|xlsx|스프레드시트|워드|docx|한글|hwpx?|csv|ppt|피피티|파워포인트|이미지)",
    re.I,
)
#: 파일 형식이 **받는 꼴**(「워드 파일로」·「엑셀로」·「파일로」)인 자리 — 뒤의 동사가 무엇이든 산출물이다.
#
# ★ 2026-10-13 네 번 이어 묻기(4B · 창 32K)의 셋째 판: 「방금 내용을 한글이나 워드 파일로 **전달해 줘**」가 요구 0개로 돌았다.
#   요구가 없으니 도구 없이 바로 답하는 길로 갔고, 모델은 「…txt 파일로 저장되었습니다」라며 **있지도 않은 경로**를 댔다
#   (작은 카드로 쓰는 사람이 같은 날 말한 「존재하지 않는 위치에 파일을 저장하려 한다」). 산출물 판정은 동사를 본다 —
#   「정리해 줘」·「만들어 줘」는 걸리고 「전달해 줘」·「줘」·「보내줘」·「달라」·「받고 싶어」는 안 걸렸다.
#   동사를 하나씩 더하지 않는다(10-02 `_NAMED_DOC_DESTINATION` 과 같은 길) — **형식 이름에 「로」가 붙은 모양**으로 본다.
#
#   좁게 둔다(넣기 전에 운영의 서로 다른 요청 1,931개에 댔다 — 새로 걸리는 것 2개, 둘 다 진짜 파일 요청):
#   - 「키워드로」·「프로파일로」는 형식 이름이 아니다(앞이 한글 · 영문이면 아니다).
#   - 「한글로 답해줘」의 한글은 말이다 — 「한글 파일 / 문서」일 때만.
#   - 이미 있는 것을 가리키는 말(「워드로 된」·「엑셀로 저장된」)과 묻는 말(「엑셀로 뭘 할 수」)은 아니다.
#   - 「파일로 만들 필요는 없고」처럼 바로 뒤에서 거두는 말은 아니다.
#   - PDF 는 넣지 않았다 — 이 도구가 만들 수 없는 형식을 요구로 세우지 않는다.
_FORMAT_AS_DESTINATION = re.compile(
    r"(?:(?<![가-힣A-Za-z])(?:워드|엑셀|피피티|파워\s*포인트|마크다운|docx?|xlsx?|hwpx?|pptx?|csv|md)\s*(?:파일|문서)?"
    r"|(?<![가-힣])한글\s*(?:파일|문서)"
    r"|(?<![가-힣A-Za-z])파일)"
    r"\s*(?:형식|형태)?\s*(?:으로|로)(?![가-힣])"
    r"(?!\s*(?:된|되어|돼|저장된|만든|만들어진|적힌|쓰인|쓰여|열|읽|뭘|무엇|무슨|어떻게|할\s*수|하는|가능))"
    r"(?![^.?!]{0,14}(?:필요\s*(?:는|가)?\s*없|말고|하지\s*말|안\s*해도))",
    re.I,
)

# 명시적인 읽기 전용 계약. 산출물 단어와 뒤쪽의 다른 동사를 문장 전체에서
# 아무렇게나 짝지으면 이런 실제 문장이 파일 생성으로 뒤집힌다:
#
#   "파일을 수정하거나 외부 서비스에 쓰지 말고 조사·분석만 해라. 근거표를 만들어라."
#
# 앞의 ``파일`` 과 뒤 문장의 ``만들어라`` 가 결합되어 xlsx/docx 두 개가 생겼다
# (2026-09-02 라이브). 금지는 단순 선호가 아니라 도구 표면과 실행 관문이 함께
# 지켜야 하는 계약이므로 한 곳에서 판정해 라우터·계획·에이전트가 같이 쓴다.
_WRITE_PROHIBITION = re.compile(
    r"(?:"
    r"(?:파일|문서|산출물)(?:을|를)?[^.!?\n]{0,18}"
    r"(?:수정|편집|작성|생성|만들|저장|쓰|건드리|손대|고치|바꾸|변경|지우|삭제)[^.!?\n]{0,12}"
    r"(?:하지\s*(?:마|말)|지는\s*(?:마|말)|지\s*(?:마|말)|말고|말아|금지|않(?:게|고|도록))"
    r"|(?:파일|문서|산출물)(?:을|를)?[^.!?\n]{0,18}"
    r"(?:수정|편집|작성|생성|만들|저장|쓰|건드리|손대|고치|바꾸|변경|지우|삭제)(?:하거나)[^.!?\n]{0,40}"
    r"(?:하지\s*(?:마|말)|말고|말아|금지|않(?:게|고|도록))"
    # ★ 대상 없는 금지(2026-09-27 게임 팩 벤치): 「고치지는 마」·「아무것도 고치지 마」가 **금지로 안 읽혀**,
    #   「파일은 고치지 마」라고 한 판에서 모델이 프로젝트에 파이썬 파일을 쓰려 했다(다른 검사에 우연히 막혔다).
    #   문장 첫머리이거나 「아무것도·절대·하나도·전혀」 뒤일 때만 잡는다 — 「테스트 코드는 고치지 마」처럼
    #   대상을 한정한 말은 여기 안 걸려 판을 통째로 묶지 않는다(2026-09-03 SWE 벤치 사고).
    r"|(?:(?<![^\s.!?,])(?=(?:아무것도|절대로?|하나도|전혀)\s*(?:고치|수정하|바꾸|건드리|손대|변경하)))"
    r"(?:아무것도|절대로?|하나도|전혀)\s*(?:고치|수정하|바꾸|건드리|손대|변경하)(?:지는|지도|지)\s*(?:마|말)"
    r"|(?:^|(?<=[.!?,\n])\s*|(?<=[.!?,\n]\s))(?:고치|수정하|바꾸|건드리|손대|변경하)(?:지는|지도|지)\s*(?:마|말)"
    r"|(?:조사|조회|분석|검토)(?:[·/&\s]*(?:와|및)?[·/&\s]*(?:조사|조회|분석|검토))*"
    r"\s*만\s*(?:해|하라|해주세요)"
    r"|읽기\s*전용|read[ -]?only"
    r"|do\s+not\s+(?:write|modify|edit|create)\s+(?:any\s+)?files?"
    r")",
    re.I,
)


#: 금지 대상이 **좁혀졌는지** 보는 자리. 「테스트 파일」·「기존 공개 테스트 파일」
#: 처럼 앞에 맨 명사가 붙으면 그것은 그 대상에 대한 금지이지 판 전체의 정책이 아니다.
#: 반대로 「이 판에서는 파일을」처럼 조사로 끝난 말이 앞에 오면 수식이 아니라 배경이라
#: 금지는 여전히 판 전체를 묶는다.
_OBJECT_QUALIFIER = re.compile(r"([가-힣]{1,8})\s*$")
#: 조사로 끝나는 말은 수식어가 아니라 배경이다. 「테스트 파일」은 수식이지만
#: 「이 판에서는 파일을」의 「판에서는」은 아니다 — 이 구분을 못 하면 판 전역
#: 금지가 대상 금지로 새어 나간다.
_TRAILING_PARTICLE = re.compile(
    r"(?:은|는|이|가|을|를|의|도|만|와|과|에|에서|에게|한테|으로|로|까지|부터|보다|처럼|라도)$"
)

#: 「○○ 밖의 · 외의 · 이외의 (파일)」 — 금지 대상 앞에서, 건드려도 되는 것(○○)을 남겨 두는 한정.
_COMPLEMENT_QUALIFIER = re.compile(r"(?:밖|외|이외)의\s*$")

#: 그 자체로 판 전체를 묶는 말. 대상이 아니라 **작업 성격**을 못 박는다.
_JOB_WIDE_PROHIBITION = re.compile(
    r"(?:조사|조회|분석|검토)(?:[·/&\s]*(?:와|및)?[·/&\s]*(?:조사|조회|분석|검토))*"
    r"\s*만\s*(?:해|하라|해주세요)"
    r"|읽기\s*전용|read[ -]?only"
    r"|do\s+not\s+(?:write|modify|edit|create)\s+(?:any\s+)?files?",
    re.I,
)


#: 금지 동사 — `_WRITE_PROHIBITION` 의 동사 묶음과 같은 말.
_PROHIBITED_VERB = re.compile(r"수정|편집|작성|생성|만들|저장|쓰|건드리|손대|고치|바꾸|변경|지우|삭제")
_PROHIBITION_HEAD = re.compile(r"(?:파일|문서|산출물)(?:을|를|이|가|은|는)?")
#: 「파일」과 금지 동사 사이에 끼어든 **다른 대상**(「기존 행은」·「내용을」).
_INNER_OBJECT = re.compile(r"[가-힣]{1,8}(?:은|는|을|를)(?=\s)")


def _prohibition_object_switched(matched: str) -> bool:
    """「파일이 이미 있으면 **기존 행은** 지우지 말고」 — 금지 대상이 파일이 아니라 사이에 낀 말인가.

    ★ 2026-10-02 레딧 기록 예약 시험: 「…xlsx 에 기록해줘. 파일이 이미 있으면 기존 행은 지우지 말고 아래에 덧붙여」를
      「파일 … 지우지 말고」로 읽어 판 전체를 읽기 전용으로 묶었다. 손자리는 게시물 5개를 다 읽고도 「파일 쓰기가 금지됐다」며
      엑셀을 안 썼다. 지우지 말라는 것은 행이고, 파일은 오히려 써 달라는 것이었다.
    """
    head = _PROHIBITION_HEAD.match(matched)
    if not head:
        return False
    # ★ 첫 동사가 아니라 **금지 꼬리(「~지 말고」) 바로 앞의 동사**까지 본다 (2026-10-17 화면 자동화 첫 시험 d079f8b8).
    #   「메모장으로 저장해 줘. **파일을** 바로 **쓰는** 도구는 쓰지 말고 메모장 창을 조작해서」 — 첫 동사(「쓰는」 · 도구를 꾸미는 말)
    #   앞에는 다른 대상이 없어서 판 전체 쓰기 금지가 됐고, 손자리는 메모장을 열지도 않았다. 금지된 것은 끝의 「쓰지 말고」이고 그 앞에
    #   「도구는」이 끼어 있다 — 「쓰다」가 「사용하다」인 문장이다.
    tail = _NEGATION_TAIL.search(matched, head.end())
    stop = tail.start() if tail else len(matched)
    verbs = [found for found in _PROHIBITED_VERB.finditer(matched, head.end()) if found.start() < stop]
    if not verbs:
        return False
    return bool(_INNER_OBJECT.search(matched[head.end():verbs[-1].start()]))


#: 금지 꼬리 — 「하지 마」·「지는 말고」·「지 말고」·「말고」·「금지」·「않게」.
_NEGATION_TAIL = re.compile(r"(?:하지|지는|지도|지)\s*(?:마|말)|말고|말아|금지|않(?:게|고|도록)")


def write_prohibition_scope(prompt: str) -> str:
    """쓰기 금지가 **판 전체**인가, 지목한 대상 하나인가.

    ★ 2026-09-03. 이 구분이 없어서 SWE 벤치가 통째로 막혔다. 벤치는 모든 판에

        테스트 파일은 절대 수정하지 마 — 구현을 고쳐야 한다.

    를 붙이는데, 「파일 … 수정하지 마」 하나만 보고 판 전체를 읽기 전용으로 묶어
    `local_apply_edit`·`local_write_file`·`local_run_project_command`·`sandbox_run`
    을 전부 막았다. 고칠 손도 실행할 손도 없이 「고쳐서 통과시켜줘」를 시킨 셈이다.

    돌려주는 값은 셋이다.

        ""        금지 없음
        "clause"  지목한 대상만 — 그 대상은 지키되 판은 계속 일한다
        "job"     판 전체가 조사·분석 전용
    """
    # ★ **따옴표로 옮겨 적은 문장은 시킨 말이 아니다** — 쓰기 금지도 따옴표 밖에서만 읽는다 (2026-10-09 새벽 · 판 hw3).
    #   「쓰기 금지를 잘못 읽는 버그를 고쳐라」는 과제 글이 그 버그의 예시(「…파일은 만들지 마」)를 따옴표로 옮겨 적자,
    #   그 예시가 시킨 말로 읽혀 **판 전체가 읽기 전용**이 됐다. 손자리는 원인을 다 읽고도 명령 하나 못 돌리고 22분을 썼다
    #   (답에 스스로 적었다: 「고쳐야 할 판정 규칙이 이 판의 프롬프트를 읽기 전용으로 묶었다」). 개수·사이트 요구는
    #   2026-10-09 에 이미 따옴표 밖에서만 뽑게 했는데(`extract`) 금지 쪽이 빠져 있었다. 글의 대부분을 감싼 따옴표와
    #   낱말 하나짜리 따옴표는 그대로 본다(`_without_quoted_sentences`).
    compact = _without_quoted_sentences(" ".join(prompt.split()))
    if _en.english_dominant(compact):
        return _en.write_prohibition_scope(compact)
    scope, match, wide = _write_prohibition_scope_by_words(compact)
    # ★ **금지 문장 밖에서 쓰라고 시킨 요청은 판 전체 쓰기 금지일 수 없다** (2026-10-08 전수조사 3-4).
    #   아래 낱말 규칙은 금지 앞 낱말의 끝 글자로 갈린다 — 「설정 파일은 건드리지 마」는 지목(clause)인데 「지도 파일」·
    #   「회의 파일」·「원가 파일」은 도·의·가(조사와 같은 글자)로 끝나 판 전체(job)가 됐다. 「지도 파일은 건드리지 마.
    #   요약 문서를 새로 만들어줘」가 읽기 전용이 되어, 만들어 달라는 문서를 못 만드는 판이 된다.
    #   끝 글자 표를 늘리지 않고 요청 자체를 본다: 금지 구절을 걷어 낸 나머지가 **새 파일·문서를 만들라고 분명히 시키면**
    #   그 금지는 지목한 대상의 것이다. 「조사만 해줘. 파일은 만들지 마」처럼 나머지에 그런 말이 없으면 그대로 판 전체다.
    #   ★ 두 가지로 조였다(운영 요청 1,810개에 대 보니 처음 규칙은 넷을 바꿨고 셋이 틀린 방향이었다):
    #     · 「조사·분석만 해라」·「읽기 전용」처럼 **판 전체를 직접 말한** 금지는 낮추지 않는다.
    #     · 「채팅으로 정리해줘 … 파일 수정하지 마」의 「정리해줘」는 파일을 만들라는 말이 아니다 — 「만들어·작성해·
    #       파일로 저장·새 문서」처럼 새것을 분명히 시킬 때만 낮춘다.
    if scope == "job" and match is not None and not wide:
        rest = (compact[: match.start()] + " " + compact[match.end():]).strip()
        if rest and (_MAKES_A_FILE.search(rest) or _SOMETHING_NEW.search(rest)):
            return "clause"
    return scope


#: 새 파일·문서를 가리키는 말 — 「새 문서로」·「새 엑셀」·「새 파일로 저장」.
_SOMETHING_NEW = re.compile(r"새(?:로운)?\s*(?:문서|파일|엑셀|시트|워드|한글|폴더|보고서)")


def _write_prohibition_scope_by_words(compact: str) -> tuple[str, "re.Match[str] | None", bool]:
    """낱말 규칙만으로 본 범위 · 그 판정을 낸 금지 구절 · 판 전체를 **직접 말한** 금지인가(「조사만」·「읽기 전용」)."""
    wide = _JOB_WIDE_PROHIBITION.search(compact)
    if wide:
        return "job", wide, True
    match = _WRITE_PROHIBITION.search(compact)
    if not match:
        return "", None, False
    if _prohibition_object_switched(match.group(0)):
        return "clause", match, False
    # ★ 「○○ **밖의** 파일은 건드리지 마라」는 ○○ 을 **고치라는 판**이다 (2026-10-08 — 숙제를 손자리에게 맡긴 첫 판).
    #   「SEED 에 여섯을 넣어라 … 시험을 고쳐라 … 위 세 가지 밖의 파일은 건드리지 마라」를 판 전체 쓰기 금지로 읽어,
    #   손자리가 「이 판은 읽기 전용이라 파일은 수정하지 않고 조사만 하겠습니다」로 갔다. 「밖의 · 외의 · 이외의」는 조사(의)로
    #   끝나 아래에서 「배경」으로 떨어졌는데, 이 말들은 배경이 아니라 **건드려도 되는 것을 남겨 두는 한정**이다.
    if _COMPLEMENT_QUALIFIER.search(compact[: match.start()]):
        return "clause", match, False
    qualifier = _OBJECT_QUALIFIER.search(compact[: match.start()])
    if not qualifier:
        return "job", match, False
    # 조사가 붙은 말은 수식어가 아니다. 한 글자짜리 지시어(「이 파일」의 「이」)는
    # 조사와 글자가 겹치므로 수식어로 본다 — 그쪽이 실제 쓰임이다.
    word = qualifier.group(1)
    if len(word) > 1 and _TRAILING_PARTICLE.search(word):
        return "job", match, False
    return "clause", match, False


#: 지목 금지의 대상이 **시험 코드**인가.
_TEST_TARGET_WORD = re.compile(r"테스트|시험\s*(?:파일|코드)|\btests?\b|\bspecs?\b|conftest", re.I)
#: 지목 금지 문장에 적힌 파일 이름(`config.py`·`src/app.ts`).
_TARGET_FILE_NAME = re.compile(r"(?<![\w.])[\w./\\-]*[\w-]\.[A-Za-z][A-Za-z0-9]{0,5}(?![\w.])")
#: 문장 끝 — 마침표 **뒤에 공백**이 올 때만(`test_report.py` 의 마침표에서 자르지 않는다).
_SENTENCE_START = re.compile(r"[.!?](?:\s+|$)|\n")
_EN_CLAUSE_SENTENCE = re.compile(
    r"[^.!?\n]*\b(?:do\s+not|don'?t|never|must\s+not)\s+(?:modify|edit|change|touch|alter|rewrite)\b[^.!?\n]*", re.I
)
#: 「a.py 와 b.csv 는 고치지 마」 — 파일 이름을 늘어놓고 바로 금지.
#: ★ 2026-10-01: 가운뎃점(·)·모점(、)·「그리고」가 없어 「scenekit.py · API.md · test_tidy.py 는 고치지 마」에서 마지막 하나만
#:   지켰다(코딩 사다리 L8 을 넣다가 시험이 잡음). 한국어로 파일을 늘어놓을 때 흔히 쓰는 기호다.
_NAMED_FILE_PROHIBITION = re.compile(
    r"((?:[\w./\\-]*[\w-]\.[A-Za-z][A-Za-z0-9]{0,5}\s*(?:와|과|,|·|、|및|하고|그리고|랑|이랑)?\s*)+)"
    r"(?:은|는|도|만은)?\s*(?:절대로?\s*|하나도\s*)?(?:고치|수정하|바꾸|건드리|손대|변경하|편집하)(?:지는|지도|지)\s*(?:마|말)"
)
#: 시험 파일 경로 —`tests/`·`__tests__/` 안이거나 `test_*`·`*_test.*`·`*.test.*`·`*.spec.*`·`conftest.py`.
_TEST_PATH = re.compile(
    r"(?i)(?:^|[\\/])(?:tests?|__tests__|specs?)[\\/]"
    r"|(?:^|[\\/])(?:test_[^\\/]*|[^\\/]*_test\.[a-z0-9]+|[^\\/]*\.(?:test|spec)\.[a-z0-9]+|conftest\.py)$"
)


#: 「언제 것」을 물었는가. 날짜 말이 **분명할 때만** — 「최근·최신」은 「최근 3년」처럼 기간이 제각각이라 넣지 않는다.
_RECENCY_DAY = re.compile(r"오늘|금일|실시간|(?i:\btoday(?:'s)?\b)")
_RECENCY_WEEK = re.compile(r"이번\s*주|금주|최근\s*(?:일주일|1주일|한\s*주)|(?i:\b(?:this|past)\s+week\b)")
#: **소식**을 물었는가. 「오늘자 기준 깃허브 기술 트렌드」·「오늘자로 지원 가능한 공고」의 「오늘」은 「지금 시점」이지
#: 「오늘 게시된 글」이 아니다 — 거기에 날짜 경고를 달면 멀쩡한 2주 전 기술 글을 버리게 만든다(역검증 09-29).
_NEWS_NOUN = re.compile(r"뉴스|이슈|소식|기사|속보|헤드라인|시사|(?i:\b(?:news|headlines?)\b)")


def requested_news_recency(prompt: str) -> str:
    """사용자가 **오늘 소식**(`day`)·**이번 주 소식**(`week`)을 물었나. 아니면 빈 문자열.

    ★ 2026-09-29 실사용: 「오늘자 국내 이슈나 뉴스 한번 조사해주라」 — 모델이 「오늘」을 검색어에만 넣고 네이버 뉴스 탭을
    관련도순·기간 없이 검색해 12일 전 기사를 열었다. 이 값으로 도구가 뉴스 검색에 기간을 붙이고, 읽은 기사의 게시일을 센다.

    역검증(09-29, 과거 고유 요청 1,217개): 날짜 말만 보면 101개가 걸리는데 대부분 「오늘 날씨」·「오늘 기분」·「오늘자 기준
    기술 트렌드」·「오늘자로 지원 가능한 공고」다. 그래서 **소식 낱말이 같이 있을 때만** 값을 낸다.
    """
    compact = " ".join((prompt or "").split())
    if not _NEWS_NOUN.search(compact):
        return ""
    if _RECENCY_DAY.search(compact):
        return "day"
    if _RECENCY_WEEK.search(compact):
        return "week"
    return ""


#: 「A 말고(빼고 · 외에 · 제외하고) 다른(나머지) …」 — A 는 지킬 파일이 아니라 **고칠 파일**이다(2026-10-15 · 551f1fa5).
_EXCEPTED_FILE = re.compile(
    r"([\w./\-]*[\w-]\.[A-Za-z][A-Za-z0-9]{0,5})\s*(?:을|를)?\s*(?:말고|빼고|외(?:에|의|에는)|제외하고|만\s*빼고)\s*(?:다른|나머지)"
)


def _excepted_names(text: str) -> set[str]:
    """「ollama.py 말고 다른 파일은 고치지 마」의 ollama.py 처럼 **예외로 댄** 파일 이름들."""
    return {match.group(1).replace("\\", "/").rsplit("/", 1)[-1].lower() for match in _EXCEPTED_FILE.finditer(text)}


def protected_write_targets(prompt: str) -> dict[str, Any]:
    """지목 금지(「테스트 파일은 고치지 마」)가 **무엇을** 지키라는 것인지.

    ★ 2026-09-28 헤르메스 대결 T7 O3: 「테스트 파일은 고치지 마」 판에서 모델이 테스트 파일의 빈 줄을 지우고
    답에 「들여쓰기만 수정(pytest 구문 오류 해결)」 — pytest 를 한 번도 안 돌린 채 지어낸 이유였다. 금지는 `clause` 로
    **모델에게 말만** 했고 도구는 막지 않았다(사고 R — 안내문은 안 먹는다). 판을 통째로 묶지 않고(09-03 SWE 사고) 그 대상만 막는다.

    돌려주는 값: `{"tests": bool, "names": [파일 이름…], "quote": 금지 문장}` — 대상을 못 알아보면 빈 dict(막지 않는다).
    """
    compact = " ".join((prompt or "").split())
    scope = write_prohibition_scope(compact)
    if scope == "job":
        return {}          # 판 전체가 읽기 전용 — 이미 모든 쓰기 손을 막는다
    # ★ 파일 이름을 **직접 댄** 금지(T7 L7 「test_report.py 와 orders.csv 는 고치지 마」)는 `_WRITE_PROHIBITION` 이
    #   못 읽어 scope 가 "" 였다 — 모델에게 전달조차 안 됐다. 이름을 댔으니 대상이 분명하다.
    named = [
        name
        for found in _NAMED_FILE_PROHIBITION.finditer(compact)
        for name in _TARGET_FILE_NAME.findall(found.group(1))
    ]
    if scope != "clause":
        excepted = _excepted_names(compact)
        named = [name for name in named if name.replace("\\", "/").rsplit("/", 1)[-1].lower() not in excepted]
        return {"tests": False, "names": named, "quote": _NAMED_FILE_PROHIBITION.search(compact).group(0)[:200]} if named else {}
    if _en.english_dominant(compact):
        found = _EN_CLAUSE_SENTENCE.search(compact)
        sentence = found.group(0).strip() if found else ""
    else:
        match = _WRITE_PROHIBITION.search(compact)
        if not match:
            return {}
        starts = [m.end() for m in _SENTENCE_START.finditer(compact[: match.start()])]
        sentence = compact[(starts[-1] if starts else 0): match.end()].strip()
    if not sentence:
        return {}
    tests = bool(_TEST_TARGET_WORD.search(sentence))
    names = [name for name in _TARGET_FILE_NAME.findall(sentence) if not re.fullmatch(r"\d+(?:\.\d+)*", name)]
    names = list(dict.fromkeys([*names, *named]))
    # ★ 2026-10-15 551f1fa5: 「ollama.py 말고 다른 파일은 고치지 마」가 ollama.py 를 지켰다 — 예외로 댄 이름은 고칠 파일이다.
    excepted = _excepted_names(compact)
    names = [name for name in names if name.replace("\\", "/").rsplit("/", 1)[-1].lower() not in excepted]
    if not tests and not names:
        return {}
    return {"tests": tests, "names": names, "quote": sentence[:200]}


def protected_path_hit(path: str, targets: dict[str, Any]) -> str:
    """이 경로가 지목 금지 대상이면 그 이유(짧은 말), 아니면 빈 문자열."""
    if not path or not targets:
        return ""
    normalized = path.replace("\\", "/")
    base = normalized.rsplit("/", 1)[-1].lower()
    for name in targets.get("names") or []:
        wanted = str(name).replace("\\", "/").lower()
        if normalized.lower().endswith("/" + wanted) or normalized.lower() == wanted or base == wanted.rsplit("/", 1)[-1]:
            return f"사용자가 지목한 파일({name})"
    if targets.get("tests") and _TEST_PATH.search(normalized):
        return "사용자가 고치지 말라고 한 테스트 파일"
    return ""


def forbids_writes(prompt: str) -> bool:
    """사용자가 **이 판을** 명시적으로 조사·분석 전용으로 묶었는가.

    대상 하나를 지목한 금지(「테스트 파일은 수정하지 마」)는 여기서 False 다.
    그 금지가 사라지는 것은 아니고 `write_prohibition_scope` 가 "clause" 로
    들고 있다 — 에이전트가 모델에게 그 대상을 따로 일러 준다.
    """
    return write_prohibition_scope(prompt) == "job"


#: 「전부·빠짐없이·끝까지」 — 목록을 다 읽으라는 말. 이 말이 있으면 새 글을 읽는 것도 산출물 단계의 진척이다
#: (`agent._progress_matches_remaining_plan`). 2026-10-03 레딧 기록 판: 「올린 글을 전부」인데 계획은 「근거 1개」뿐이라
#: 6개 중 2개를 읽은 시점에 「진척 없음」으로 파일 쓰기로 넘어갔다.
_EVERYTHING = re.compile(r"(?<![이예])전부(?!터)|빠짐\s*없이|끝까지|하나도\s*빠(?:짐|뜨리지|트리지)|모든\s*(?:글|게시|항목|공고|댓글|상품|파일)|(?i:\b(?:all|every|entire)\b)"
    # ★ 2026-10-03: 「내 레딧 글들」 — 내 것의 여러 개는 내 것 전부다(155bd3e1 판이 글 둘 중 하나만 채우고 접혔다).
    r"|(?:^|\s)내\s+(?:\S+\s+){0,2}?(?:글|게시글|게시물|포스트|댓글|파일|메일|문서)들"
    r"|(?i:\bmy\s+(?:\w+\s+)?(?:posts|comments|files|emails|documents)\b)")


#: 최상급·범위 조건 — 사이트의 정렬·필터가 정답 경로인 요청(2026-10-03 TED 판). 「가장 중요한 이유」 같은 말은 안 건다.
#: 정렬이 필요한 말 — 「가장 많이 본 · 가장 싼 · 최신 · ~순」(2026-10-03 TED 판).
_SORT_ASK = re.compile(
    r"(?i:\b(?:most\s+(?:viewed|popular|watched|played|liked|reviewed|replies|commented|downloaded|followed|recent)"
    r"|cheapest|least\s+expensive|lowest[\s-]+(?:price|priced|cost)|highest[\s-]+(?:rated|priced)|top[\s-]+rated"
    r"|best[\s-]+(?:selling|rated|sellers?)|newest|sort(?:ed)?\s+by"
    # A 판에서 빠졌던 말(2026-10-03): 「the largest hunting land」·「the most experienced agent」 — 크기·경력 순 정렬이다.
    r"|(?:the\s+)?(?:largest|biggest|smallest|oldest|longest|shortest)|most\s+experienced)\b)"
    # 한국어는 **사이트 정렬이 분명한 말만**. 「가장 비싼(모은 것끼리 비교)」·「최저가(상품의 값 칸)」는 운영 판 역검증에서
    #   YES24·교보·크몽·에누리 판을 헛걸었다(2026-10-03, 9판 중 4).
    r"|(?:가장|제일)\s*많이\s*(?:본|조회된|팔린|읽은|시청한)"
    r"|인기순|최신순|조회\s*수?\s*순|평점\s*순|(?:낮은|높은)?\s*가격\s*순|순으로\s*정렬"
)
#: 범위·조건 필터가 필요한 말 — 「12~18분 사이 · 반경 20마일 · $40 이하 · ~로 거르기」.
_RANGE_ASK = re.compile(
    r"(?i:\b(?:filter(?:ed)?\s+(?:by|for)|between\s+\$?\d[\d,.]*\s*(?:and|to|-)\s*\$?\d[\d,.]*"
    r"|within\s+\d+\s*(?:miles?|mi|km|kilometers?)|under\s+\$\s?\d|price\s+range"
    # A 판에서 빠졌던 말: 「a maximum price of $1500」·「within $1000 to $2000」·「33 to 49inch」.
    r"|max(?:imum)?\s+(?:price|budget|rent)|within\s+\$\s?\d|\$\s?\d[\d,]*\s*(?:-|to)\s*\$?\s?\d"
    r"|\d+\s*(?:to|-)\s*\d+\s*(?:inch|inches|in\b|\"|cm|years?|hours?|minutes?|mins?|lbs?|pounds?|kg|gb|tb|hz))"
    r"(?![A-Za-z]))"
)
#: 둘 중 하나라도 있으면 「사이트 정렬·필터로 보기」 칸이 선다.
_SITE_FILTER_ASK = re.compile(_SORT_ASK.pattern + "|" + _RANGE_ASK.pattern)


#: 정렬 방향 — 요청 쪽 낱말과 주소·클릭 쪽 낱말을 같은 표로 읽는다.
_SORT_DIRECTIONS = (
    ("popular", re.compile(r"(?i:most\s+(?:viewed|popular|watched|played|liked|replies|commented|downloaded|followed)"
                           r"|best[\s-]+sell|popular|view|plays?\b|trend|\bhot\b|\btop\b(?![\s-]+rated))"
                           r"|많이\s*(?:본|조회|팔린|읽은|시청)|인기|조회")),
    ("price_low", re.compile(r"(?i:cheapest|least\s+expensive|lowest[\s-]+(?:price|priced|cost)|price[_\s-]*(?:asc|low)"
                             r"|low[_\s-]*price|\bcheap|low\s+to\s+high|\blowest\b)|낮은\s*가격")),
    ("rating", re.compile(r"(?i:highest[\s-]+rated|top[\s-]+rated|best[\s-]+rated|rating|rated|review|score|stars?)|평점|별점")),
    ("newest", re.compile(r"(?i:newest|most\s+recent|latest|\bnew\b|date|recent)|최신")),
)


def _sort_direction(text: str) -> str:
    for name, rx in _SORT_DIRECTIONS:
        if rx.search(text or ""):
            return name
    return "any"


def site_filter_needs(prompt: str) -> tuple[str, ...]:
    """이 요청이 사이트에서 걸어야 하는 것 — ("sort:popular",) · ("range",) · ("sort:price_low", "range") …

    ★ 2026-10-03 TED 2차: 「최신순 + 주제=로봇」이 「많이 본 순 + 12~18분」을 채웠다. 정렬은 **방향까지**, 필터는 범위(길이·가격·거리·날짜)와
      분류(주제·브랜드)를 가른다. 「filter by …」처럼 무엇인지 모르는 거르기는 「filter」(범위든 분류든 하나면 된다).
    """
    compact = " ".join((prompt or "").split())
    needs: list[str] = []
    sort_match = _SORT_ASK.search(compact)
    if sort_match:
        needs.append("sort:" + _sort_direction(sort_match.group(0)))
    ranges = [m.group(0) for m in _RANGE_ASK.finditer(compact)]
    if any(not re.match(r"(?i)filter", r) for r in ranges):  # 감지와 같은 표 — 두 자리가 어긋나지 않게(사고 Q)
        needs.append("range")
    elif re.search(r"(?i:\bfilter(?:ed)?\s+(?:by|for)\b)", compact):
        needs.append("filter")
    return tuple(needs)


#: 정렬·필터 칸을 누른 걸음 — 주소가 안 바뀌는 사이트(Ajax 정렬)도 있다.
_FILTER_CLICK_WORDS = re.compile(
    r"순$|순\b|정렬|필터|(?i:\bsort|\bfilter|most\s+(?:viewed|popular)|lowest|highest|newest|top[\s-]rated|popular"
    r"|low\s+to\s+high|high\s+to\s+low)")
#: 필터가 아닌 주소 값 — 검색어·쪽·화면 상태·추적값.
_NOT_A_FILTER_KEY = re.compile(
    r"(?i)^(?:q|query|search|keyword|keywords|k|s|ss|term|text|page|p|pg|offset|start|limit|per_?page|lang|hl|locale|ref|referral"
    r"|referrer|src|source|utm_\w+|gclid|fbclid|session\w*|sid|tab|view|mode|layout|types?|include\w*|ssps|_\w*)$")


def asks_for_site_filter(prompt: str) -> bool:
    return bool(_SITE_FILTER_ASK.search(" ".join((prompt or "").split())))


#: 주소에 정렬·필터 값이 들어갔는가 — 검색어(q)·쪽(page)·추적값은 필터가 아니다.
_FILTER_QUERY_KEY = re.compile(
    r"^(?:sort|sortby|sort_by|sorting|order|orderby|order_by|filter|filters|f|facet|facets|duration|length|price"
    r"|min|max|range|rating|ratings|stars|distance|radius|within|topic|topics|category|categories|cat|genre|year"
    r"|years|date|dates|from|to|depart|return|condition|brand|brands|size|bed|beds|bedrooms|bath|baths|bathrooms"
    r"|period|time|views|level|availability|in_stock|status|location|zip|postal|guests|adults)(?:\[\d*\])?$"
    r"|^(?:min|max|price)[_-]?\w*$|^\w*[_-](?:min|max)$",
    re.I,
)
_FILTER_PATH_PART = frozenset({
    "top", "popular", "trending", "bestsellers", "best-sellers", "most-viewed", "most-popular", "top-rated",
    "toprated", "newest", "hot", "best",
    # ★ 2026-10-05 AWS 영어 윈도우: 「the 3 newest arXiv cs.CL papers」 판이 `arxiv.org/list/cs.CL/recent` 를 읽고도
    #   「최신순 정렬 안 걸림」이었다. 방향 표(`_SORT_DIRECTIONS`)는 recent·latest·new 를 최신으로 읽는데 이 표에만 없었다.
    "recent", "latest", "most-recent", "new",
})


_SORT_QUERY_KEY = re.compile(r"^(?:sort|sortby|sort_by|sorting|order|orderby|order_by)(?:\[\d*\])?$", re.I)


#: 범위 값 — 길이·가격·거리·날짜·크기·인원. 그 밖의 필터 값(주제·브랜드·종류)은 분류다.
_RANGE_QUERY_KEY = re.compile(
    r"^(?:duration|length|price|min|max|range|rating|ratings|stars|distance|radius|within|year|years|date|dates|from|to"
    r"|depart|return|size|bed|beds|bedrooms|bath|baths|bathrooms|period|time|views|guests|adults|zip|postal)(?:\[\d*\])?$"
    r"|^(?:min|max|price)[_-]?\w*$|^\w*[_-](?:min|max)$",
    re.I,
)


#: 값 이름 속 범위 낱말 — `filter_width.filter`·`screenSize`·`priceMax` 같은 합성 이름(B 판 us-appliance).
_RANGE_KEY_WORD = re.compile(
    r"(?i)width|height|depth|size|weight|length|duration|price|budget|min|max|range|distance|radius|capacity|screen|inch"
    r"|year|date|rating|stars|bed|bath|guest|adult|mileage|miles")


def url_filter_kinds(url: str) -> set[str]:
    """주소에 걸린 것 — {"sort:popular"} · {"range"} · {"category"} (`?sort=popular&duration=12-18&topics[0]=robots`)."""
    parsed = urlparse(str(url or ""))
    if parsed.scheme not in {"http", "https"}:
        return set()
    from urllib.parse import parse_qsl

    kinds: set[str] = set()
    # ★ 2026-10-03 B 판 us-appliance: 정렬·필터가 `#` 뒤에 있었다(`#query=c130&sort_method=pageviews…&filter_width.filter=36`).
    #   쪽 안 앱(단일 페이지)이 흔히 그렇게 한다 — `#` 뒤가 「이름=값」 꼴이면 같이 읽는다.
    pairs = parse_qsl(parsed.query, keep_blank_values=False)
    if "=" in parsed.fragment:
        pairs += parse_qsl(parsed.fragment, keep_blank_values=False)
    for key, value in pairs:
        if not value or _NOT_A_FILTER_KEY.match(key):
            continue
        if _SORT_QUERY_KEY.match(key) or re.match(r"(?i)(?:sort|order)[_.-]?\w*$", key):
            kinds.add("sort:" + _sort_direction(value))
        elif _RANGE_QUERY_KEY.match(key) or _RANGE_KEY_WORD.search(key):
            kinds.add("range")
        else:
            # ★ 2026-10-03 A 판: Petfinder `age=young`·`daysOnPetfinder=1 day`, apartments.com `so=4` 를 처음 보는 이름이라
            #   통째로 버려 「필터 안 씀」 헛판정 — 검색어·쪽·추적값이 아닌 값은 필터다(분류). 이름이 「최근·며칠」이면 최신 쪽.
            kinds.add("category")
            if re.search(r"(?i)days?|added|recent|newest|posted", key):
                kinds.add("sort:newest")
    for part in parsed.path.split("/"):
        if part and part.casefold() in _FILTER_PATH_PART:
            kinds.add("sort:" + _sort_direction(part.replace("-", " ")))
        elif re.match(r"(?i)sort[:=_-]", part or ""):
            # Rotten Tomatoes `/browse/tv_series_browse/sort:popular` — 정렬이 경로 조각에 들어간다(B 판).
            kinds.add("sort:" + _sort_direction(part.split(":", 1)[-1].split("=", 1)[-1]))
        elif re.search(r"(?i)(?:under|over|below|above|max|min|from|price|up-to)[-_]?\$?\d|\d+[-_]to[-_]\d+|\d+[-_](?:beds?|baths?|bedrooms?|bathrooms?)\b",
                       part or ""):
            # apartments.com `/corning-ca/under-1500/` — 가격·방 조건이 경로에 들어간다(A 판 공식 조건). 모델이 대조표와 말다툼했다.
            kinds.add("range")
    return kinds


def filter_needs_met(needs: set[str], applied: set[str]) -> list[str]:
    """아직 안 걸린 것. 「sort:any」는 어떤 정렬이든, 「filter」는 범위든 분류든 하나면 된다."""
    missing: list[str] = []
    for need in sorted(needs):
        if need.startswith("sort:"):
            direction = need.split(":", 1)[1]
            ok = any(a.startswith("sort:") and (direction == "any" or a == need) for a in applied)
        elif need == "filter":
            ok = bool(applied & {"range", "category"})
        else:
            ok = need in applied
        if not ok:
            missing.append(need)
    return missing


def filtered_url(url: str) -> bool:
    """사이트가 정렬·필터해 준 목록 주소인가 — `?sort=popular&duration=12-18` · `/top` · `?pricemax=…`."""
    parsed = urlparse(str(url or ""))
    if parsed.scheme not in {"http", "https"}:
        return False
    from urllib.parse import parse_qsl

    for key, value in parse_qsl(parsed.query, keep_blank_values=False):
        if value and _FILTER_QUERY_KEY.match(key):
            return True
    return any(part.casefold() in _FILTER_PATH_PART for part in parsed.path.split("/") if part)


def asks_for_everything(prompt: str) -> bool:
    return bool(_EVERYTHING.search(" ".join((prompt or "").split())))


#: 받는 곳이 **화면의 폼·입력칸**인 꼴 — 「폼에 넣어」·「칸에 채워」·「입력란에 붙여」.
_INTO_PAGE_FORM = re.compile(
    r"(?:폼|입력\s*(?:칸|란|창)|(?<![가-힣])칸|필드)\s*(?:에다가|에다|에|으로|로)?\s*(?:만\s*)?(?:넣|채워|채우|붙여|입력)"
)
#: 그래도 파일을 **만들라**고 한 말이 따로 있으면 산출물이다(「폼에 넣고, 워드로도 만들어줘」).
_MAKES_A_FILE = re.compile(r"(?:만들(?:어|고|어서|자)|생성해|작성해|파일로\s*저장)")


#: 결과 파일을 남기는 앱을 **이름만** 댄 꼴 — 「메모장으로 … 저장해 줘」 · 「VS Code 로 … 한 줄 더해 줘」.
_APP_NAMED_FOR_RESULT = re.compile(
    r"(메모장|그림판|VS\s*Code|비주얼\s*스튜디오\s*코드|한글\s*프로그램)\s*(?:으로|로|에서|를\s*써서|를\s*열어)", re.I,
)
#: 창 · 화면을 직접 다루라는 말 — 이게 있으면 창으로 한다.
_ASKS_FOR_THE_WINDOW = re.compile(r"창|직접|화면|메뉴|클릭|눌러|단추|버튼|마우스|키보드")


def app_named_only(prompt: str) -> str:
    """앱 이름은 댔지만 창을 직접 다루라고는 안 한 요청이면 그 앱 이름, 아니면 빈 글.

    ★ 2026-10-17 대표 결정: 「앱 이름만 댔으면 결과 파일을 코드로 바로 쓰고, 창으로 직접이라 했을 때만 창」.
      평소 말투로 시킨 측정(apptool-*)에서 엑셀 · 워드 · 파워포인트는 스스로 파일 도구로 16~20초에 끝냈지만, 「메모장으로 … 저장」 ·
      「VS Code 로 … 한 줄 더해」 · 「그림판으로 … JPEG 로」는 창을 몰아 90~160초가 걸렸다. 사람이 앱 이름을 댄 것은 대개 결과의
      모양을 말한 것이지 그 창을 손으로 돌리라는 뜻이 아니다.
      운영 요청 2,296개 대조: 걸리는 것은 그 측정 요청 10개뿐(창 · 메뉴를 말한 31개는 빠진다). 엑셀 · 워드 · 파워포인트는 넣지 않았다 —
      「엑셀로 정리해 줘」 같은 조사 요청이 많고, 그 판들은 이미 파일 도구로 간다. 「한글로」는 언어일 수 있어 「한글 프로그램」만 본다.
    """
    compact = " ".join((prompt or "").split())
    if _ASKS_FOR_THE_WINDOW.search(compact):
        return ""
    match = _APP_NAMED_FOR_RESULT.search(compact)
    return match.group(1) if match else ""


#: 작업 관리자가 보여 주는 값을 물은 꼴 — 「작업 관리자 기준으로」 · 「작업 관리자에서 메모리」.
_TASK_MANAGER_VALUES = re.compile(r"작업\s*관리자|task\s*manager", re.I)


def asks_task_manager_values(prompt: str) -> bool:
    """작업 관리자에 보이는 값을 물었는가 — 그 값은 작업 관리자 창을 읽어야 같은 숫자가 나온다.

    ★ 2026-10-18 평소 말투 측정(작업 관리자 3판): 창을 읽은 판은 56초에 작업 관리자와 같은 값(앱별 개인 작업 집합 —
      10-17 에 화면과 대조해 맞춘 읽기)을 냈다. 나머지 두 판은 `tasklist` · `Get-Process` 를 돌렸다 — 그 명령들은 **작업 집합**을
      준다(작업 관리자가 보여 주는 값과 다르다). 프로세스 전체 목록이 통째로 모델에 들어가 읽기에만 31초가 들었고 판은 88 · 164초.
    """
    return bool(_TASK_MANAGER_VALUES.search(" ".join((prompt or "").split())))


#: 이름 안에 「파일」이 든 앱 — 산출물 낱말 판정 전에 지운다(`wants_an_artifact`).
_APP_NAMES_WITH_FILE = re.compile(r"파일\s*탐색기|\bfile\s+explorer\b", re.I)


# ★ **파일을 만들 필요가 없다고 한 말은 산출물 요구가 아니다** (2026-10-19 판 029fae29 · 클로드 코드가 맡긴 과제).
#   「…FAQ 문서(docs/faq.md)를 찾아 읽고 … 원문에 적힌 그대로 옮기고 … **파일은 만들지 않아도 되고** 답으로만 알려 줘」에
#   「실제 산출물 파일 만들기」가 서고 손자리가 시키지 않은 .md 를 만들었다. 쓰기 **금지**(「만들지 마」)는 `forbids_writes` 가 보지만
#   「않아도 · 필요 없 · 없이」는 금지가 아니라 **안 시킨다는 말**이라 어디에도 안 걸렸다 — 거두는 말은 판을 읽기 전용으로 묶지
#   않고, 거둔 구절만 걷어 낸 나머지로 판정한다(「엑셀은 안 만들어도 되고 워드 파일에 옮겨 줘」는 여전히 산출물이다).
_FILE_DECLINED = re.compile(
    r"(?<![가-힣])(?:파일|문서|산출물|엑셀|워드|한글\s*파일|보고서|마크다운)(?:\s*(?:파일|문서))?"
    r"\s*(?:은|는|을|를|로|으로|까지는|까지|도)?\s*(?:따로|별도로|굳이|새로|꼭)?\s*"
    r"(?:(?:만들|작성하|저장하|생성하|쓰|남기)지\s*(?:는|도)?\s*않(?:아도|으셔도)"
    r"|(?:만들|작성할|저장할|생성할|쓸|남길)\s*필요(?:는|가|까지는|도)?\s*없"
    r"|안\s*(?:만들어|작성해|저장해|써|남겨)\s*(?:도|줘도)"
    r"|필요\s*(?:는|가|도)?\s*없"
    r"|없이)"
)


def _without_declined_files(compact: str) -> str:
    """「파일은 만들지 않아도 되고」 · 「no need to create a file」 같은 거두는 구절을 걷어 낸 글."""
    pattern = _en.DECLINED_FILE if _en.english_dominant(compact) else _FILE_DECLINED
    return pattern.sub(" ", compact)


#: 옮기는 동사 — 「옮기고 · 담아 · 달아 · 넣어」. 만드는 동사(만들 · 작성 · 저장 · 써 …)와 달리 **받는 곳**을 가진다.
_TRANSFER_VERB = re.compile(r"옮기|옮겨|담아|달아|넣어")
#: 대상 낱말 뒤가 이 꼴이면 그 대상이 **받는 곳**이다 — 「한글 파일 하나에」 · 「워드 문서로」 · 「엑셀 시트에다」.
_TARGET_AS_DESTINATION = re.compile(
    r"(?:\s*(?:파일|문서|시트|형식|형태|양식|하나|한\s*개|1\s*개|초안))*\s*(?:에다가|에다|에|으로|로|안에|속에)(?![가-힣])"
)


def _names_a_work_file(compact: str) -> bool:
    """읽을 재료가 아니라 **손댈 파일**로 이름을 댄 파일이 있는가 — 「site_atlas.py 의 SEED 에 여섯을 넣어라」의 site_atlas.py.

    주소 안의 이름(`https://…/index.html`) · 소재(「llama.cpp 저장소」) · 읽는 말이 붙은 이름(「faq.md 를 찾아 읽고」)은 아니다.
    """
    for match in _NAMED_FILE_SUFFIX.finditer(compact):
        token = compact[max(compact.rfind(" ", 0, match.start()), 0):match.end()]
        if "://" in token or "www." in token.lower():
            continue
        if _CODE_TARGET_IS_TOPIC.match(compact, match.end()) or _format_is_source(compact, match.end(), named_file=True):
            continue
        return True
    return False


def _verb_meets_its_target(compact: str, *, real_files_only: bool = False) -> bool:
    """대상 낱말 ∧ 산출 동사 — 옮기는 동사뿐이면 대상이 **받는 곳** 자리에 있어야 한다.

    ★ 2026-10-19 (판 029fae29 · 10-19 arXiv 판 302bf59e). 「FAQ **문서**를 찾아 읽고 … 원문에 적힌 그대로 **옮기고**」 ·
      「논문 PDF **파일**을 내려받아서 … 초록의 첫 문장을 그대로 **옮겨 줘**」 — 둘 다 「옮기다」가 **답에 그대로 옮겨 적으라**는 말이고
      문서 · 파일은 **읽을 재료**였다. 글 전체에서 대상 낱말 하나와 동사 하나를 아무렇게나 짝지어 산출물 요구가 섰고, 두 판 다
      시키지 않은 .md 를 만들었다. 「옮기다 · 담다 · 달다 · 넣다」는 **어디로**를 가지는 동사다 — 「한글 파일**에** 본문을 옮기고」 ·
      「엑셀**에** 넣어줘」 · 「워드 파일**로** 옮기고」. 그래서 이 동사들만 있을 때는 대상이 받는 곳(에 · 로) 자리에 있어야 짝이다.
      낱말을 빼지 않는다(「옮기」를 빼면 「한글 파일에 옮기고 각주 달아줘」가 다시 요구 0개가 된다 — 2026-09-01 사고).

    `real_files_only` — 파일을 만들 필요 없다고 거둔 판에서는 「표 · 문서 · 초안」처럼 **답 안에서도 되는 대상**은 세지 않는다
    (`_TIDY_VERB` 가 이미 지키는 규율: 「표로 정리해줘」는 답에 표를 그리면 되는 일이다).
    """
    targets = list((_REAL_FILE_TARGET if real_files_only else _ARTIFACT_REQUEST_TARGET).finditer(compact))
    verbs = list(_ARTIFACT_REQUEST_VERB.finditer(compact))
    if not targets or not verbs:
        return False
    if any(not _TRANSFER_VERB.match(verb.group(0)) for verb in verbs):
        return True
    # 받는 곳은 둘 중 하나다: 대상 낱말이 「에 · 로」 자리에 있거나, 손댈 파일의 이름을 댔다(「ollama.py 의 … 바꾼 자리에 주석을 달아」 —
    # 운영 요청에 대 보니 저장소 코드 과제 넷이 이 꼴이었다).
    return _names_a_work_file(compact) or any(
        target.group(0).endswith("로") or _TARGET_AS_DESTINATION.match(compact, target.end())
        for target in targets
    )


def wants_an_artifact(prompt: str) -> bool:
    """말이 아니라 파일·코드·문서 같은 산출물을 요구했는가."""
    # ★ 「**파일** 탐색기」는 앱 이름이다 (2026-10-17 앱 묶음 · 탐색기 b09c6f60). 「파일 탐색기로 … 새 폴더를 만들어 줘」가 「파일 + 만들어」로
    #   산출물 요구가 되어, 폴더를 창으로 제대로 만든 판이 시키지 않은 결과 문서(.md)를 하나 더 만들었다.
    compact = _APP_NAMES_WITH_FILE.sub("탐색기", " ".join(prompt.split()))
    if forbids_writes(compact):
        return False
    # 「파일은 만들지 않아도 되고」 — 거둔 구절은 시킨 말이 아니다(`_FILE_DECLINED`). 거둔 판에서는 남은 글에 **파일 · 형식 이름**이
    # 있어야 산출물이다 — 「표로 정리해서 줘. 파일은 만들 필요 없어」의 표는 답에 그리는 표다.
    stripped = " ".join(_without_declined_files(compact).split())
    declined = stripped != compact
    compact = stripped
    if wants_code(compact):
        return True
    if _en.english_dominant(compact):
        # 한국어와 같은 세 길: 강한 동사 ∧ 대상 · 「정리」류 ∧ 「문서로」 · 「정리」류 ∧ 파일 낱말.
        verb_en = _en.first_unnegated(_en.ARTIFACT_REQUEST_VERB, compact)
        tidy_en = _en.first_unnegated(_en.TIDY_VERB, compact)
        if verb_en and _en.verb_meets_its_target(
            compact, real_files_only=declined, work_file_named=_names_a_work_file(compact)
        ):
            return True
        if tidy_en and not declined and _en.DOC_FORMAT_REQUEST.search(compact):
            return True
        return bool(tidy_en and _en.REAL_FILE_TARGET.search(compact))
    # ★ **화면의 폼·칸에 넣는 것은 파일 만들기가 아니다** (2026-10-07 대표 실사용). 「자기소개서 **초안** … 그거 이제
    #   **폼 넣어주라**」가 산출물 요구가 됐다 — 「초안」(문서 낱말) + 「넣어」(옮기는 동사). 되받기를 받은 모델은
    #   「워드 파일부터 만들겠습니다」로 빠졌고 폼은 못 채웠다. 받는 곳이 화면의 폼이면 그건 화면에서 하는 동작이다.
    #   운영 요청 1,643개에서 이 꼴은 5개였고, 산출물 판정이 바뀌는 것은 그 한 판뿐이다.
    into_form = bool(_INTO_PAGE_FORM.search(compact)) and not _MAKES_A_FILE.search(compact)
    if not into_form and _verb_meets_its_target(compact, real_files_only=declined):
        return True
    if _TIDY_VERB.search(compact) and not declined and _DOC_FORMAT_REQUEST.search(compact):
        return True
    # ★ 문서 파일 이름이 **받는 곳**(「…xlsx 에 / 로」)이면 동사가 무엇이든 산출물이다 (2026-10-02 밤 레딧 기록 예약).
    #   「reddit_반응기록.xlsx 에 … 기록해줘」가 요구 0개(근거 1개 읽기만)로 돌아, 글 6개를 다 읽고 「이제 엑셀에 기록하겠습니다」로
    #   끝났다 — 계획은 다 채웠다고 보고 글만 쓰는 정리로 넘겼다. 「기록·적어·넣어」를 낱말표에 하나씩 더하지 않고 **대상의 모양**으로 본다.
    #   역검증: 지난 고유 요청 1,440개 중 새로 걸리는 4개 전부 실제 파일 요청(서치어드바이저 결과 .md ×2 · 레딧 기록 ×2), 오탐 0.
    if _NAMED_DOC_DESTINATION.search(compact):
        return True
    # 형식 이름이 받는 꼴이면 동사가 무엇이든 산출물이다 (2026-10-13 · `_FORMAT_AS_DESTINATION` 의 설명).
    if not into_form and _FORMAT_AS_DESTINATION.search(compact):
        return True
    return bool(_TIDY_VERB.search(compact) and _REAL_FILE_TARGET.search(compact))


#: 언어 이름 **뒤에 붙으면 그건 소재다.** 「파이썬 책」의 대상은 책이지 코드가 아니다.
_CODE_TARGET_IS_TOPIC = re.compile(
    # 닫는 따옴표가 끼어도 소재다 — 「'클린 코드' 책」(2026-09-14).
    r"\s*[\"'”’」』]?\s*(책|도서|서적|교재|입문서|강의|강좌|수업|튜토리얼|자료|공고|채용|시장"
    r"|개발자|엔지니어|전문가|학원|과외|가격|후기|리뷰|순위|목록"
    # ★ 「llama.cpp 저장소 문서에서 … 워드 파일로 만들어 줘」가 코딩 과제가 됐다 (2026-09-28, MCP 위임).
    #   `.cpp` 가 파일 대상에, `만들어` 가 동사에 걸려 워드를 다 만든 판이 `coding_verification_missing` 으로 실패했다.
    #   저장소·README·위키 앞의 이름은 **읽을 곳**이다. 「문서」는 넣지 않는다 — 「app.py 문서화」는 코딩이다.
    r"|저장소|레포|리포지토리|(?:repo|readme|github|wiki)(?![A-Za-z])|깃허브|위키)",
    re.I,
)


def wants_code(prompt: str) -> bool:
    """설명 말고 실행할 코드를 요구했는가.

    ## 소재로 쓰인 언어 이름은 대상이 아니다 (2026-09-07)

    산출물 벤치에서 「YES24에서 **파이썬 책** 3권의 쪽수와 ISBN 을 조사해서 …
    마크다운 문서로 **만들어줘**」가 코딩 과제로 분류됐다. `파이썬` 이 대상에,
    `만들어` 가 동사에 걸린 것이다. 판은 `.md` 파일을 정상으로 만들어 냈는데도
    코드 검증이 없다는 이유로 `coding_verification_missing` 으로 **실패 처리**됐다.

    낱말이 아니라 **그 낱말이 무엇을 꾸미는가**를 본다 — 「셋을 세라」의 셋이
    무엇의 셋인지 따졌던 것과 같은 규율이다. 언어 이름이 전부 소재 자리에만
    있으면 코딩이 아니다.
    """
    compact = " ".join(prompt.split())
    english = _en.english_dominant(compact)
    verb, topic = (_en.CODE_REQUEST_VERB, _en.CODE_TARGET_IS_TOPIC) if english else (_CODE_REQUEST_VERB, _CODE_TARGET_IS_TOPIC)
    # 영어는 「Don't fix anything」처럼 동사 앞에 부정이 온다 — 부정된 동사는 시킨 게 아니다.
    if not (_en.first_unnegated(verb, compact) if english else verb.search(compact)):
        return False
    targets = list(_CODE_TARGET.finditer(compact))
    if english:
        targets += list(_en.CODE_TARGET_WORDS.finditer(compact))
    for match in targets:
        if topic.match(compact, match.end()):
            continue
        if _is_quoted(compact, match):
            continue
        if _is_identifier_code(compact, match):
            continue
        return True
    return False


#: 「코드」가 **프로그램이 아니라 부호**인 말 — 「QR 코드」·「바코드」·「쿠폰 코드」·「오류 코드」.
#:
#: ★ 2026-09-30 QR 판: 「robometafit.com 주소를 담은 QR 코드 PNG 를 만들어」가 코딩 과제가 됐다. 코드 모델로 돌았고,
#:   새 파일 쓰기가 「먼저 읽어라」(code_recon_required)에 네 번 막혔고, 테스트 검증이 없다고
#:   `coding_verification_missing` 으로 죽었다. 09-07 「파이썬 책」·09-14 「'클린 코드' 책」·09-28 「llama.cpp 저장소」에 이은
#:   같은 병의 네 번째 — 낱말이 아니라 **그 낱말이 무엇인가**를 본다. 「파이썬 코드 짜서 QR 만들어」는 여전히 코딩이다
#:   (`파이썬`·`짜서` 가 따로 걸린다).
_IDENTIFIER_CODE = re.compile(
    r"(?:QR|qr|큐알|바|우편\s*번호|인증|쿠폰|할인|추천|초대|색상|컬러|국가|지역|상품|학과|종목|오류|에러"
    r"|보안|확인|입장|결제|주문|이벤트|프로모션|초대|멤버십|회원)\s*$"
)


def _is_identifier_code(compact: str, match: "re.Match[str]") -> bool:
    """이 「코드」가 부호(QR·바·쿠폰·오류 …)를 가리키는가."""
    if match.group(0) != "코드":
        return False
    return bool(_IDENTIFIER_CODE.search(compact[max(0, match.start() - 12): match.start()]))


#: 따옴표 종류. 한글 프롬프트는 「」·""·''·""·'' 를 섞어 쓴다.
_QUOTE_OPEN = "\"'“‘「『"
_QUOTE_CLOSE = "\"'”’」』"


def _is_quoted(compact: str, match: "re.Match[str]") -> bool:
    """이 낱말이 **따옴표 안에** 있는가 — 그러면 쓰는 게 아니라 **말하는** 것이다.

    ## 왜 필요했나 (2026-09-07, 같은 날 세 번째)

    위임 벤치에서 클로드가 자비스에게 이렇게 시켰다:

        검색창에 "파이썬"을 입력하고 검색을 실행해줘

    `파이썬` 은 **검색어**인데 코드 대상으로 잡혔고 `실행해` 가 동사에 걸려
    코딩 과제가 됐다 — `coding_verification_missing` 으로 판이 죽었다. 소재 뒤에
    오는 명사(`책`·`강의`)로만 걸렀더니 이 어법이 그대로 샜다.

    따옴표를 씌운 낱말은 **그것을 가리키는 것**이지 그것으로 만들라는 말이 아니다.
    한 판에 따옴표 밖 대상이 하나라도 있으면 여전히 코딩이다
    (「"파이썬" 스크립트 짜줘」의 `스크립트`).
    """
    before = compact[match.start() - 1] if match.start() else ""
    after = compact[match.end()] if match.end() < len(compact) else ""
    if before in _QUOTE_OPEN and after in _QUOTE_CLOSE:
        return True
    # ★ 따옴표 **안 어딘가에** 있어도 가리키는 말이다 (2026-09-14).
    #
    #   「알라딘과 YES24에서 '클린 코드' 책의 판매가를 비교해줘」가 코딩 과제가 됐다.
    #   `코드` 바로 앞이 따옴표가 아니라 「클린 」이라 이 검사를 빠져나갔고, 모델은 코드부터
    #   짜려 했고, 판은 `coding_verification_missing` 으로 죽었다. 책·강의 제목은 여러 낱말이다.
    #   같은 줄에서 여는 따옴표가 앞에, 닫는 따옴표가 뒤에 있고 그 사이에 따옴표가 없으면 안이다.
    #   (짧은 창만 본다 — 문장 끝과 끝의 따옴표를 한 쌍으로 묶으면 본문 전체가 「안」이 된다.)
    window = 40
    left = compact[max(0, match.start() - window): match.start()]
    right = compact[match.end(): match.end() + window]
    open_at = max(left.rfind(q) for q in _QUOTE_OPEN)
    close_at = min((i for i in (right.find(q) for q in _QUOTE_CLOSE) if i >= 0), default=-1)
    if open_at < 0 or close_at < 0:
        return False
    inside = left[open_at + 1:] + right[:close_at]
    return not any(q in inside for q in _QUOTE_OPEN + _QUOTE_CLOSE)


#: 폴더를 고른 Code 모드에서만 쓰는 **넓힌 변경 동사**. 「고치고」·「추가해」·「바꿔」·「삭제해」…
_CODE_ACTION_EXTRA_VERB = re.compile(
    r"(고치(?:고|면|자|세요|는\s*중)?(?![^\s]*지\s*마)|고칠\s*것"
    r"|추가(?:해|하고|하세요|하라)|넣어|끼워|붙여"
    r"|바꿔|바꾸(?:고|어|세요)|변경(?:해|하고|하세요)"
    r"|삭제(?:해|하고)|지워|제거(?:해|하고)|없애"
    r"|통과(?:시켜|하게|되게)|개선해|최적화해|옮겨|나눠|합쳐|적용해)",
    re.I,
)

#: 설명·질문 표지. 있으면 넓힌 동사는 **행동이 아니라 말하는 것**이다.
_CODE_ACTION_EXPLAIN_MARK = re.compile(
    r"(설명|알려\s*줘|알려주|가르쳐|뭐야|뭔가요|무슨\s*(?:뜻|일|역할)|왜\s|어떻게\s*(?:돼|되|동작|작동|하)"
    r"|분석|리뷰|요약|보여\s*줘|\?)",
)


def wants_code_action(prompt: str) -> bool:
    """Code 작업공간이 이미 선택됐을 때 변경·실행 동사만 판정한다.

    일반 Cowork에서는 대상 언어나 파일이 함께 있어야 ``wants_code``가 참이다.
    Code 모드에서는 선택한 저장소 자체가 대상이므로 "이 버그 고쳐줘"처럼 코드라는
    단어를 생략한 요청도 실제 변경 계약으로 잡아야 한다.
    """
    compact = " ".join(prompt.split())
    if _CODE_REQUEST_VERB.search(compact):
        return True
    # ★ 낱말표 구멍 (2026-09-21, 노트북 4B 사다리). 위 목록은 「고쳐·수정해·만들어」 같은 몇 활용형만 알아서
    #   폴더를 고른 Code 모드에서도 「고치고」·「추가해」·「바꿔」 가 코딩으로 안 잡혔다 — **샌드박스가 있는 이 PC 에서도**
    #   검증 게이트가 꺼져 있었다(사다리 여섯 중 L3·L4·L5). 넓히되, 설명·질문 표지가 있으면 넓힌 동사는 무시한다 —
    #   「이 함수를 어떻게 고치는지 설명해줘」에 검증 게이트를 걸면 `coding_verification_missing` 으로 죽는다.
    #   기존 동사(위)는 설명 표지와 함께여도 그대로 참이다(예전과 같다).
    if _CODE_ACTION_EXPLAIN_MARK.search(compact):
        return False
    return bool(_CODE_ACTION_EXTRA_VERB.search(compact))


def wants_a_computation(prompt: str) -> bool:
    """**셈**을 시켰는가. 계획 층이 `sandbox_run` 을 요구로 넣을지 정한다.

    산출물 요청("표로 정리해줘")은 여기 안 걸린다 — 그건 계산이 아니다.
    동사를 짝으로 요구하는 이유: "소수가 뭐야"는 설명 요청이다.
    """
    compact = " ".join(prompt.split())
    english = _en.english_dominant(compact)
    # ★ 따옴표 안에 옮겨 적은 **구절**은 다룰 내용이지 시키는 말이 아니다 (2026-10-18 한글 서식 측정).
    #   「「간편한 작업을 위한 몇 가지 팁」 줄을 16pt … 바꾸고 저장해 줘」에서 따옴표 안의 「몇 가지」와 밖의 「저장해」가
    #   짝지어 「직접 계산해 보기」 요구가 섰고, 되받기 세 번에 모델이 시키지 않은 확인 문서와 PDF 를 만들었다.
    #   글 길이는 그대로 두므로(빈칸으로 바꿈) 거리 판정은 원래 글 기준 그대로다.
    targets = list((_en.ARITHMETIC_TARGET if english else _ARITHMETIC_TARGET).finditer(_without_quoted_sentences(compact)))
    verbs = list((_en.COMPUTE_VERB if english else _COMPUTE_VERB).finditer(compact))
    # 대상과 동사가 같은 절에 있어야 한다. ``페이지를 만들어줘 … 확률을 표시해``의
    # 멀리 떨어진 '만들어'와 '확률'을 짝지으면 UI 작업에 sandbox 계산 요구가 생긴다.
    # 영어는 같은 뜻이 더 길다(「몇 시간…답해라」 ↔ 「How many hours … Answer with the number」) — 창을 넓힌다.
    reach = 60 if english else 40
    arrange = _en.ARRANGE_VERB if english else _ARRANGE_VERB
    return any(
        max(target.start(), verb.start()) - min(target.end(), verb.end()) <= reach
        # 만드는 동사 + 늘어놓은 칸 이름 = 화면의 값을 옮겨 담는 일이다(셈 아님) — `_a_listed_column`.
        and not (arrange.fullmatch(verb.group(0).strip())
                 and _a_listed_column(compact, target.start(), target.end(), english=english))
        for target in targets
        for verb in verbs
    )


def wants_the_sandbox(prompt: str) -> bool:
    """상자를 **쥐어줘야** 하는가. 라우터가 쓴다(산출물이거나 셈이거나)."""
    compact = " ".join(prompt.split())
    if _en.english_dominant(compact):
        return wants_an_artifact(compact) or bool(
            _en.COMPUTE_VERB.search(compact)
            and (_en.ARTIFACT_TARGET.search(compact) or _en.ARITHMETIC_TARGET.search(compact))
        )
    return wants_an_artifact(compact) or bool(
        _COMPUTE_VERB.search(compact)
        and (_ARTIFACT_TARGET.search(compact) or _ARITHMETIC_TARGET.search(compact))
    )


# ── 이어지는 지시 (2026-09-02) ────────────────────────────────────────────────
#
# ★ 왜 필요한가. 라이브에서 이렇게 났다:
#
#     01:03  "…공고 추천해서 조사해줘…"   사람인 직행 → 사이트내검색 253건 → 상세 1건
#     01:05  "어엉 너가 제안한대로 전부다 진행해줘"
#            → 요구 0개 · 네이버 검색 → 네이버 카페(마감 지난 공고) → 구글 검색 → …
#                                                              15걸음 186초
#
# 앞 판이 사람인에서 253건을 이미 찾아 놨는데 두 번째 판이 그걸 버리고 검색엔진부터
# 다시 갔다. 「제안한대로」에는 **사이트도 개수도 조건도 없으니** 계획 층이 눈을
# 감았고, 사이트 직행 갈래도 안 걸렸다. 사용자가 본 「삽질」의 정체다.
#
# 고치는 방식: 새 프롬프트가 **그 자체로는 아무 요구도 못 내는 이어짓기**면,
# 바로 앞 사람 말을 앞에 붙여 다시 뽑는다. 새 장치를 만들지 않는다 — 기존
# `extract` 가 그대로 일한다.
#
# ⚠ 순위(검색 결과)는 **안 물려준다.** 그건 세운 사이트의 것이고, 판을 넘기면
#   딴 사이트 목록 앞에서 엉뚱한 것을 연다(2026-08-14 에 겪었다). 여기서 물려주는
#   것은 **무엇을 시켰는가**뿐이다.

#: 그 자체로는 내용이 없는 말. 앞 말에 기대는 지시다.
_CONTINUATION = re.compile(
    r"(제안한\s*대로|말한\s*대로|얘기한\s*대로|그대로|시킨\s*대로|아까\s*(?:그|말한)"
    r"|이어서|계속\s*(?:해|진행)|마저|나머지도?|전부\s*다?\s*(?:진행|해)"
    r"|진행해|ㅇㅇ)"
)
#: **따로 선** 「해줘」 — 「그렇게 해줘」·「그걸로 해줘」·「해줘」. 이때의 「하다」는 앞에서 말한 일을 가리키는 대용 동사다.
#:
#: ★ 2026-10-09. 예전에는 「…해줘」로 끝나는 60자 이하 문장이 **전부** 이어짓기 후보였다(`해\s*줘\s*$`). 「파이썬 책 2권
#:   추천해줘」·「이 함수가 왜 느린지 설명해줘」·「이 글을 5문장으로 요약해줘」처럼 그 자체로 할 일이 다 있는 말에서
#:   요구가 하나도 안 뽑히면, 바로 앞 사람 말(「사람인에서 공고 3개 열어 읽어줘」)이 통째로 붙어 그 요구가 되살아났다.
#:   요청 읽기 실험(bench/intent_reading_1009)에서 정규식이 틀린 17문장 중 7문장이 이것이었다.
#:   「추천해줘」의 「해줘」는 낱말(추천)에 붙어 있다 — 할 일이 그 낱말이다. 「그렇게 해줘」의 「해줘」는 따로 서 있다 — 할 일이 앞 말에 있다.
#:   이 고침의 앞 절반(꼬리 규칙을 뺀 것)은 손자리가 우리 숙제로 한 것이고(원인을 스스로 찾았다), 그 고침이 깨뜨린
#:   「그렇게 해줘」·「그걸로 해줘」·「해줘」·「바로 해줘」를 이 규칙으로 되살렸다.
_BARE_DO_IT = re.compile(r"(?:^|\s)해\s*줘(?:요|라)?[\s.!~?ㅎㅋ]*$")

#: 맞장구로 **시작하는** 말 — 「응 그렇게 해」·「좋아 그걸로」. 글 머리에서만 본다.
#:
#: ★ 2026-10-08 숙소 조사 판 뒤의 감탄 한마디: 「와 … 왜이렇게 **좋아** 사람 …」의 `좋아` 가 글 한가운데서 걸려
#:   앞 요청이 통째로 붙었다. 모델은 「칭찬일 뿐 볼 것이 없다」고 세 번 판단했는데 하네스가 「안 보고 답하려 한다」며
#:   두 번 되받아 숙소 둘을 다시 열었다(4걸음·122초·부분 완료). 글 가운데의 「좋아·그래·맞아·응」은 맞장구가 아니라
#:   그 문장의 서술어다. 그리고 맞장구 뒤가 길면(시키는 말 없이) 그것은 지시가 아니라 소감이다.
_ASSENT_LEAD = re.compile(r"^(?:응|그래|좋아|맞아)(?=\s|$)")
#: 맞장구 뒤에 이만큼까지는 「응 그걸로」 같은 짧은 지목으로 본다. 넘으면 시키는 말이 있어야 한다.
_ASSENT_TAIL_MAX_CHARS = 12

#: 이어짓기로 볼 수 있는 길이 상한. 길면 그 안에 이미 할 말이 다 들어 있다.
_CONTINUATION_MAX_CHARS = 60


def is_continuation(prompt: str) -> bool:
    """이 말이 **앞 말에 기대는 지시**인가."""
    compact = " ".join((prompt or "").split())
    if not compact or len(compact) > _CONTINUATION_MAX_CHARS:
        return False
    if _en.english_dominant(compact):
        return _en.is_continuation(compact)
    if _CONTINUATION.search(compact + " ") or _BARE_DO_IT.search(compact):
        return True
    lead = _ASSENT_LEAD.match(compact)
    if not lead:
        return False
    tail = compact[lead.end() :].strip()
    return len(tail) <= _ASSENT_TAIL_MAX_CHARS or bool(_REQUEST_ENDING.search(tail))


def planning_prompt(prompt: str, history: list[dict[str, Any]] | None = None) -> str:
    """계획을 뽑을 때 쓸 **실제 문장**.

    이어짓기이고 그 자체로 요구가 안 나오면, 바로 앞 사람 말을 앞에 붙인다.
    그 말에서도 요구가 안 나오면 원래 문장을 그대로 돌려준다 — 없는 맥락을
    지어내지 않는다.
    """
    compact = " ".join((prompt or "").split())
    if not is_continuation(compact) or extract(compact):
        return compact
    for item in reversed(history or []):
        if str(item.get("role") or "") != "user":
            continue
        previous = " ".join(str(item.get("content") or "").split())
        if not previous or previous == compact:
            continue
        joined = f"{previous} {compact}"
        if extract(joined):
            return joined
        # 그 앞 말도 이어짓기였으면 한 칸 더 거슬러 올라간다.
        if not is_continuation(previous):
            break
    return compact
